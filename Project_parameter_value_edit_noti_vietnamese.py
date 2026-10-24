"""Safely upsert project_parameter from the UpdateしたいValue Excel sheet.

The primary key uses project_id, phase_id, variation_id and the Excel
se-parameter_id. All six mapped cells are authoritative; blank cells overwrite
existing database values with SQL NULL. The default mode is a dry run.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import os

    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402

# PyCharm Run-button configuration.
RUN_PROJECT_ID = 1020
RUN_PHASE_ID = 35
RUN_EXCEL_PATH = Path(
    "/Users/TrongKieu/Downloads/"
    "SE_List_P61QR_HEV_US GCC_4WD_UC-Lot_中間確認会UPDATE_VALUE.xlsx"
    # "/Users/TrongKieu/Downloads/SE_List_P61QR_HEV_US_4WD_UC-Lot_Kickoff_value_UPDATE (1).xlsx"
)
# RUN_SHEET_NAME = "UpdateしたいValue"
RUN_SHEET_NAME = "VALUE"
RUN_EXECUTE = False
RUN_ASSUME_YES = False

HEADER_SCAN_LIMIT = 20
PREVIEW_ROW_LIMIT = 30
ADVISORY_LOCK_TIMEOUT_SECONDS = 10

# Fixed, reviewed mapping; identifiers never come from runtime user input.
EXCEL_TO_DB_COLUMNS = {
    "値": "value",
    "ステータス": "approval_status",
    "メモ": "user_memo",
    "担当者": "z_note",
    "更新者": "employee_number",
    "更新日": "update_day",
}
DATA_COLUMNS = tuple(EXCEL_TO_DB_COLUMNS.values())
SE_PARAMETER_HEADERS = {"separameterid", "seprameterid"}


@dataclass(frozen=True)
class ExcelParameterRow:
    excel_row: int
    se_parameter_id: int
    values: tuple[tuple[str, object], ...]


@dataclass(frozen=True)
class ExcelReadResult:
    variation_name: str
    rows: tuple[ExcelParameterRow, ...]
    header_row: int
    sheet_name: str


@dataclass(frozen=True)
class PlanRow:
    excel_row: int
    se_parameter_id: int
    parameter_name: str
    old_values: tuple[tuple[str, object], ...]
    new_values: tuple[tuple[str, object], ...]
    action: str


@dataclass(frozen=True)
class IdComparison:
    db_only: tuple[int, ...]
    excel_only: tuple[int, ...]
    missing_master: tuple[int, ...] = ()
    invalid_statuses: tuple[object, ...] = ()
    missing_employees: tuple[str, ...] = ()


def normalize_positive_id(raw_value: object, *, label: str) -> int:
    if isinstance(raw_value, bool):
        raise ValueError(f"{label} must be a positive integer, not boolean.")
    if isinstance(raw_value, float) and not raw_value.is_integer():
        raise ValueError(f"{label} must be an integer: {raw_value!r}")
    try:
        value = int(raw_value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {label}: {raw_value!r}") from exc
    if value <= 0:
        raise ValueError(f"{label} must be greater than zero: {value}")
    return value


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    return "".join(char for char in text if char not in " _-\n\r\t")


def normalize_data_header(value: object) -> str:
    """Normalize generated headers such as 値2 and 更新日56."""
    return re.sub(r"\d+$", "", normalize_header(value))


def normalize_variation_name(value: object) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()


def normalize_text(value: object, *, strip: bool = False) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        if not value.strip():
            return None
        return value.strip() if strip else value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value)


def normalize_status(value: object, *, row: int) -> int | str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool) or (isinstance(value, float) and not value.is_integer()):
        raise ValueError(f"approval_status at Excel row {row} must be an integer.")
    normalized_value = value
    if isinstance(value, str):
        display_match = re.fullmatch(r"\s*(\d+)\s*[.．].*", value, flags=re.DOTALL)
        if display_match:
            normalized_value = display_match.group(1)
    try:
        status = int(normalized_value)
    except (TypeError, ValueError):
        if isinstance(value, str):
            return " ".join(value.split())
        raise ValueError(f"approval_status at Excel row {row} is invalid: {value!r}")
    if status < 0:
        raise ValueError(f"approval_status at Excel row {row} cannot be negative.")
    return status


def normalize_update_day(value: object, *, row: int) -> dt.datetime | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None, microsecond=0)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min)
    if isinstance(value, str):
        text = value.strip().lstrip(",").strip()
        try:
            return dt.datetime.fromisoformat(text).replace(tzinfo=None, microsecond=0)
        except ValueError as exc:
            raise ValueError(
                f"update_day at Excel row {row} is invalid: {value!r}. "
                "Use YYYY-MM-DD or YYYY-MM-DD HH:MM:SS."
            ) from exc
    raise ValueError(f"update_day at Excel row {row} is invalid: {value!r}")


def normalize_field(column: str, value: object, *, row: int) -> object:
    if column not in DATA_COLUMNS:
        raise ValueError(f"Unexpected project_parameter column: {column!r}")
    if column == "approval_status":
        return normalize_status(value, row=row)
    if column == "update_day":
        return normalize_update_day(value, row=row)
    if column == "employee_number":
        return normalize_text(value, strip=True)
    return normalize_text(value)


def find_excel_layout(
    worksheet,
) -> tuple[int, int, list[tuple[str, dict[str, int]]]]:
    expected_headers = [normalize_data_header(name) for name in EXCEL_TO_DB_COLUMNS]
    max_row = min(worksheet.max_row, HEADER_SCAN_LIMIT)
    for row_number in range(1, max_row + 1):
        headers = {
            column: normalize_data_header(cell.value)
            for column, cell in enumerate(worksheet[row_number], start=1)
        }
        se_columns = [
            column
            for column, header in headers.items()
            if header in SE_PARAMETER_HEADERS
        ]
        if len(se_columns) != 1:
            continue
        blocks: list[tuple[str, dict[str, int]]] = []
        block_width = len(expected_headers)
        for start_column in range(1, worksheet.max_column - block_width + 2):
            actual = [
                headers.get(start_column + offset, "") for offset in range(block_width)
            ]
            if actual != expected_headers:
                continue
            variation_cell = worksheet.cell(row_number - 1, start_column)
            variation_name = normalize_text(variation_cell.value, strip=True)
            if not variation_name:
                raise ValueError(
                    f"Missing variation name above Excel column {start_column}."
                )
            field_columns = {
                db_column: start_column + offset
                for offset, db_column in enumerate(DATA_COLUMNS)
            }
            blocks.append((variation_name, field_columns))
        if blocks:
            return row_number, se_columns[0], blocks
    required = ["se-parameter_id", *EXCEL_TO_DB_COLUMNS]
    raise ValueError(
        f"Could not find a complete six-column variation block in the first "
        f"{max_row} rows: {required}"
    )


def read_excel_variations(
    excel_path: Path | str,
    sheet_name: str,
) -> list[ExcelReadResult]:
    path = Path(excel_path).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"Excel file not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Excel input must be an .xlsx or .xlsm file.")
    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    value_workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name in formula_workbook.sheetnames:
            actual_sheet_name = sheet_name
        elif len(formula_workbook.sheetnames) == 1:
            actual_sheet_name = formula_workbook.sheetnames[0]
        else:
            raise ValueError(
                f"Sheet {sheet_name!r} not found; "
                f"available={formula_workbook.sheetnames}"
            )
        formula_sheet = formula_workbook[actual_sheet_name]
        value_sheet = value_workbook[actual_sheet_name]
        header_row, se_column, blocks = find_excel_layout(formula_sheet)
        seen_variations: dict[str, str] = {}
        for variation_name, columns in blocks:
            key = normalize_variation_name(variation_name)
            if key in seen_variations:
                raise ValueError(
                    "Duplicate variation blocks after normalization: "
                    f"{seen_variations[key]!r} and {variation_name!r}."
                )
            seen_variations[key] = variation_name

        results: list[ExcelReadResult] = []
        for variation_name, field_columns in blocks:
            rows: list[ExcelParameterRow] = []
            seen: dict[int, int] = {}
            formula_rows = formula_sheet.iter_rows(min_row=header_row + 1)
            value_rows = value_sheet.iter_rows(min_row=header_row + 1)
            for row_number, (formula_row, value_row) in enumerate(
                zip(formula_rows, value_rows), start=header_row + 1
            ):
                se_cell = formula_row[se_column - 1]
                formula_cells = {
                    column: formula_row[column_number - 1]
                    for column, column_number in field_columns.items()
                }
                value_cells = {
                    column: value_row[column_number - 1]
                    for column, column_number in field_columns.items()
                }
                if se_cell.data_type == "f":
                    raise ValueError(
                        f"Formula is not allowed in se-parameter_id at row {row_number}."
                    )
                if se_cell.value in (None, ""):
                    if any(
                        cell.value not in (None, "") for cell in formula_cells.values()
                    ):
                        raise ValueError(f"Excel row {row_number} has data but no ID.")
                    continue
                se_id = normalize_positive_id(
                    se_cell.value, label=f"se-parameter_id at Excel row {row_number}"
                )
                if se_id in seen:
                    raise ValueError(
                        f"Duplicate se-parameter_id {se_id} in variation "
                        f"{variation_name!r} at rows {seen[se_id]} and {row_number}."
                    )
                seen[se_id] = row_number
                for column, formula_cell in formula_cells.items():
                    if (
                        formula_cell.data_type == "f"
                        and value_cells[column].value is None
                    ):
                        raise ValueError(
                            f"Formula in {variation_name!r}, Excel row {row_number}, "
                            f"column {column!r} has no cached value. Recalculate and "
                            "save the workbook in Excel before running this tool."
                        )
                values = tuple(
                    (
                        column,
                        normalize_field(
                            column, value_cells[column].value, row=row_number
                        ),
                    )
                    for column in DATA_COLUMNS
                )
                rows.append(ExcelParameterRow(row_number, se_id, values))
            if not rows:
                raise ValueError(
                    f"No rows with se-parameter_id found for {variation_name!r}."
                )
            results.append(
                ExcelReadResult(
                    variation_name, tuple(rows), header_row, actual_sheet_name
                )
            )
    finally:
        formula_workbook.close()
        value_workbook.close()
    return results


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Cannot create placeholders for an empty sequence.")
    return ", ".join(["%s"] * len(values))


def resolve_approval_status_rows(
    cursor, rows: Sequence[ExcelParameterRow]
) -> tuple[ExcelParameterRow, ...]:
    """Resolve status labels to approval_status_table.state_key when possible."""
    cursor.execute(
        "SELECT state_key, state_name FROM approval_status_table ORDER BY state_key"
    )
    by_name: dict[str, list[int]] = {}
    for raw_key, raw_name in cursor.fetchall():
        name_key = normalize_variation_name(raw_name)
        if name_key:
            by_name.setdefault(name_key, []).append(int(raw_key))

    resolved_rows: list[ExcelParameterRow] = []
    for row in rows:
        values = dict(row.values)
        status = values["approval_status"]
        if isinstance(status, str):
            matches = by_name.get(normalize_variation_name(status), [])
            if len(matches) > 1:
                raise ValueError(
                    f"approval status label {status!r} at Excel row "
                    f"{row.excel_row} matches multiple state keys: {matches}"
                )
            if len(matches) == 1:
                values["approval_status"] = matches[0]
        resolved_rows.append(
            ExcelParameterRow(
                row.excel_row,
                row.se_parameter_id,
                tuple((column, values[column]) for column in DATA_COLUMNS),
            )
        )
    return tuple(resolved_rows)


def resolve_variation_ids(cursor, variation_names: Sequence[str]) -> dict[str, int]:
    """Resolve Excel variation labels without relying on hard-coded IDs."""
    if not variation_names:
        raise ValueError("No variation names were detected in Excel.")
    cursor.execute("SELECT id, variation FROM variation ORDER BY id")
    by_name: dict[str, list[int]] = {}
    for raw_id, raw_name in cursor.fetchall():
        key = normalize_variation_name(raw_name)
        if key:
            by_name.setdefault(key, []).append(int(raw_id))

    resolved: dict[str, int] = {}
    missing: list[str] = []
    ambiguous: dict[str, list[int]] = {}
    for excel_name in variation_names:
        matches = by_name.get(normalize_variation_name(excel_name), [])
        if not matches:
            missing.append(excel_name)
        elif len(matches) > 1:
            ambiguous[excel_name] = matches
        else:
            resolved[excel_name] = matches[0]
    if missing or ambiguous:
        raise ValueError(
            "Could not resolve Excel variation names safely; "
            f"missing_in_variation_table={missing}, ambiguous={ambiguous}"
        )
    return resolved


def fetch_target_label(cursor, table: str, object_id: int, column: str) -> str:
    if (table, column) not in {
        ("project_info", "project_code"),
        ("phase", "phase"),
        ("variation", "variation"),
    }:
        raise ValueError("Unexpected target lookup.")
    cursor.execute(f"SELECT {column} FROM {table} WHERE id = %s", (object_id,))
    rows = cursor.fetchall()
    if len(rows) != 1:
        raise ValueError(f"{table} ID {object_id} was not found or is ambiguous.")
    return str(rows[0][0])


def fetch_parameter_names(cursor, ids: Sequence[int]) -> dict[int, str]:
    cursor.execute(
        f"SELECT id, parameter_name_1, parameter_name_2 FROM se_parameter "
        f"WHERE id IN ({make_placeholders(ids)})",
        tuple(ids),
    )
    names = {
        int(item_id): " / ".join(
            str(value) for value in (name_1, name_2) if value not in (None, "")
        )
        for item_id, name_1, name_2 in cursor.fetchall()
    }
    return names


def validate_references(
    cursor, rows: Sequence[ExcelParameterRow]
) -> tuple[tuple[object, ...], tuple[str, ...]]:
    raw_statuses = {
        dict(row.values)["approval_status"]
        for row in rows
        if dict(row.values)["approval_status"] is not None
    }
    statuses = sorted(
        {
            int(value)
            for value in raw_statuses
            if isinstance(value, int) and not isinstance(value, bool)
        }
    )
    unresolved_statuses = sorted(
        (value for value in raw_statuses if isinstance(value, str)), key=str
    )
    employees = sorted(
        {
            str(dict(row.values)["employee_number"])
            for row in rows
            if dict(row.values)["employee_number"] is not None
        }
    )
    if statuses:
        cursor.execute(
            f"SELECT state_key FROM approval_status_table "
            f"WHERE state_key IN ({make_placeholders(statuses)})",
            tuple(statuses),
        )
        found = {int(row[0]) for row in cursor.fetchall()}
        invalid_statuses = tuple(
            [item for item in statuses if item not in found] + unresolved_statuses
        )
    else:
        invalid_statuses = tuple(unresolved_statuses)
    if employees:
        cursor.execute(
            f"SELECT employee_number FROM user_table "
            f"WHERE employee_number IN ({make_placeholders(employees)})",
            tuple(employees),
        )
        found = {str(row[0]) for row in cursor.fetchall()}
        missing_employees = tuple(item for item in employees if item not in found)
    else:
        missing_employees = ()
    return invalid_statuses, missing_employees


def normalize_database_values(values: Sequence[object]) -> dict[str, object]:
    return {
        column: normalize_field(column, value, row=0)
        for column, value in zip(DATA_COLUMNS, values)
    }


def fetch_existing_rows(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    ids: Sequence[int],
    *,
    lock_rows: bool = False,
) -> dict[int, dict[str, object]]:
    columns = ", ".join(f"`{column}`" for column in DATA_COLUMNS)
    lock = " FOR UPDATE" if lock_rows else ""
    cursor.execute(
        f"""
        SELECT se_parameter_id, {columns}
        FROM project_parameter
        WHERE project_id = %s AND phase_id = %s AND variation_id = %s
          AND se_parameter_id IN ({make_placeholders(ids)})
        ORDER BY se_parameter_id{lock}
        """,
        (project_id, phase_id, variation_id, *ids),
    )
    return {
        int(row[0]): normalize_database_values(row[1:]) for row in cursor.fetchall()
    }


def compare_se_parameter_ids(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    excel_ids: Sequence[int],
    *,
    lock_rows: bool = False,
) -> IdComparison:
    """Compare Excel IDs with DB IDs for the exact three-key target."""
    lock = " FOR UPDATE" if lock_rows else ""
    cursor.execute(
        f"""
        SELECT se_parameter_id
        FROM project_parameter
        WHERE project_id = %s AND phase_id = %s AND variation_id = %s
        ORDER BY se_parameter_id{lock}
        """,
        (project_id, phase_id, variation_id),
    )
    db_ids = {int(row[0]) for row in cursor.fetchall()}
    excel_id_set = set(excel_ids)
    return IdComparison(
        db_only=tuple(sorted(db_ids - excel_id_set)),
        excel_only=tuple(sorted(excel_id_set - db_ids)),
    )


def build_plan(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    rows: Sequence[ExcelParameterRow],
    *,
    lock_rows: bool = False,
) -> tuple[dict[str, str], IdComparison, list[PlanRow]]:
    target = {
        "project": fetch_target_label(
            cursor, "project_info", project_id, "project_code"
        ),
        "phase": fetch_target_label(cursor, "phase", phase_id, "phase"),
        "variation": fetch_target_label(cursor, "variation", variation_id, "variation"),
    }
    ids = [row.se_parameter_id for row in rows]
    names = fetch_parameter_names(cursor, ids)
    missing_master = tuple(sorted(set(ids) - set(names)))
    invalid_statuses, missing_employees = validate_references(cursor, rows)
    base_comparison = compare_se_parameter_ids(
        cursor,
        project_id,
        phase_id,
        variation_id,
        ids,
        lock_rows=lock_rows,
    )
    comparison = IdComparison(
        db_only=base_comparison.db_only,
        excel_only=base_comparison.excel_only,
        missing_master=missing_master,
        invalid_statuses=invalid_statuses,
        missing_employees=missing_employees,
    )
    existing = fetch_existing_rows(
        cursor, project_id, phase_id, variation_id, ids, lock_rows=lock_rows
    )
    plan: list[PlanRow] = []
    for row in rows:
        new = dict(row.values)
        old = existing.get(row.se_parameter_id)
        has_invalid_reference = (
            row.se_parameter_id in comparison.missing_master
            or new["approval_status"] in comparison.invalid_statuses
            or new["employee_number"] in comparison.missing_employees
        )
        if has_invalid_reference:
            action = "BLOCKED_REFERENCE"
            old = old or {column: None for column in DATA_COLUMNS}
        elif old is None:
            action = "INSERT"
            old = {column: None for column in DATA_COLUMNS}
        else:
            action = "UNCHANGED" if old == new else "UPDATE"
        plan.append(
            PlanRow(
                row.excel_row,
                row.se_parameter_id,
                names.get(row.se_parameter_id, "<missing in se_parameter master>"),
                tuple((column, old[column]) for column in DATA_COLUMNS),
                tuple((column, new[column]) for column in DATA_COLUMNS),
                action,
            )
        )
    return target, comparison, plan


def summarize_plan(plan: Iterable[PlanRow]) -> dict[str, int]:
    counts = {"INSERT": 0, "UPDATE": 0, "UNCHANGED": 0, "BLOCKED_REFERENCE": 0}
    for row in plan:
        counts[row.action] += 1
    return counts


def changed_fields(row: PlanRow) -> list[str]:
    if row.action == "BLOCKED_REFERENCE":
        return []
    old = dict(row.old_values)
    return [
        column
        for column, value in row.new_values
        if row.action == "INSERT" or old[column] != value
    ]


def compact(value: object, limit: int = 60) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    text = str(value).replace("\r", "\\r").replace("\n", "\\n")
    return repr(text if len(text) <= limit else text[: limit - 3] + "...")


def print_plan(
    db_connection,
    *,
    project_id: int,
    phase_id: int,
    variation_id: int,
    excel_path: Path,
    sheet_name: str,
    read_result: ExcelReadResult,
    target: Mapping[str, str],
    comparison: IdComparison,
    plan: Sequence[PlanRow],
) -> None:
    settings = db_connection.settings_dict
    counts = summarize_plan(plan)
    field_counts = {column: 0 for column in DATA_COLUMNS}
    for row in plan:
        for column in changed_fields(row):
            field_counts[column] += 1
    print(f"Database: {settings.get('NAME')}")
    print(f"Host: {settings.get('HOST')}:{settings.get('PORT') or 3306}")
    print("Affected table: project_parameter")
    print(f"Excel: {excel_path}")
    print(f"Sheet: {read_result.sheet_name!r}; header row: {read_result.header_row}")
    print(
        f"Target: project={project_id} ({target['project']}), "
        f"phase={phase_id} ({target['phase']}), "
        f"variation={variation_id} ({target['variation']})"
    )
    print(f"Excel rows selected: {len(read_result.rows)}")
    print(
        "DB_ONLY se_parameter_id "
        f"({len(comparison.db_only)}; Có trong cơ sở dữ liệu, nhưng không có trong Excel.):"
    )
    print(f"  {list(comparison.db_only)}")
    print(
        "EXCEL_ONLY se_parameter_id "
        f"({len(comparison.excel_only)}; Có trong Excel, nhưng không có trong cơ sở dữ liệu.):"
    )
    print(f"  {list(comparison.excel_only)}")
    print(
        "MISSING_IN_SE_PARAMETER_MASTER "
        f"({len(comparison.missing_master)}; cannot be inserted):"
    )
    print(f"  {list(comparison.missing_master)}")
    print(
        "INVALID_APPROVAL_STATUS "
        f"({len(comparison.invalid_statuses)}; Không tồn tại trong approval_status_table):"
    )
    print(f"  {list(comparison.invalid_statuses)}")
    print(
        "MISSING_EMPLOYEE_NUMBER "
        f"({len(comparison.missing_employees)}; không tồn tại trong user_table):"
    )
    print(f"  {list(comparison.missing_employees)}")
    print(
        f"Estimated rows: INSERT={counts['INSERT']}, UPDATE={counts['UPDATE']}, "
        f"UNCHANGED={counts['UNCHANGED']}, "
        f"BLOCKED_REFERENCE={counts['BLOCKED_REFERENCE']}"
    )
    print("Changed fields (NULL overwrites included):")
    for column in DATA_COLUMNS:
        print(f"  {column}: {field_counts[column]}")
    print("Planned changes:")
    changes = [row for row in plan if row.action != "UNCHANGED"]
    for row in changes[:PREVIEW_ROW_LIMIT]:
        old = dict(row.old_values)
        new = dict(row.new_values)
        detail = ", ".join(
            f"{column}: {compact(old[column])} -> {compact(new[column])}"
            for column in changed_fields(row)
        )
        print(
            f"- Excel row {row.excel_row}, se_parameter_id={row.se_parameter_id} "
            f"({row.parameter_name}): {row.action}; {detail}"
        )
    if len(changes) > PREVIEW_ROW_LIMIT:
        print(f"  ... and {len(changes) - PREVIEW_ROW_LIMIT} more changed rows")


def plan_fingerprint(plan: Sequence[PlanRow]) -> tuple[tuple[object, ...], ...]:
    return tuple(
        (row.excel_row, row.se_parameter_id, row.old_values, row.new_values, row.action)
        for row in plan
    )


def acquire_lock(cursor, database_name: str) -> str:
    name = f"spdm_project_parameter_excel:{database_name}"[:64]
    cursor.execute("SELECT GET_LOCK(%s, %s)", (name, ADVISORY_LOCK_TIMEOUT_SECONDS))
    result = cursor.fetchone()
    if not result or result[0] != 1:
        raise RuntimeError(f"Could not acquire maintenance lock {name!r}.")
    return name


def apply_plan(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    plan: Sequence[PlanRow],
) -> dict[str, int]:
    counts = {"INSERT": 0, "UPDATE": 0, "UNCHANGED": 0}
    columns = ", ".join(f"`{column}`" for column in DATA_COLUMNS)
    insert_sql = f"""
        INSERT INTO project_parameter
          (project_id, se_parameter_id, phase_id, variation_id, {columns})
        VALUES (%s, %s, %s, %s, {make_placeholders(DATA_COLUMNS)})
    """
    assignments = ", ".join(f"`{column}` = %s" for column in DATA_COLUMNS)
    update_sql = f"""
        UPDATE project_parameter SET {assignments}
        WHERE project_id = %s AND se_parameter_id = %s
          AND phase_id = %s AND variation_id = %s
    """
    for row in plan:
        if row.action == "UNCHANGED":
            counts["UNCHANGED"] += 1
            continue
        new_values = tuple(value for _, value in row.new_values)
        if row.action == "INSERT":
            params = (
                project_id,
                row.se_parameter_id,
                phase_id,
                variation_id,
                *new_values,
            )
            cursor.execute(insert_sql, params)
        elif row.action == "UPDATE":
            params = (
                *new_values,
                project_id,
                row.se_parameter_id,
                phase_id,
                variation_id,
            )
            cursor.execute(update_sql, params)
        else:
            raise RuntimeError(f"Unexpected action: {row.action!r}")
        allowed_rowcounts = {1} if row.action == "INSERT" else {0, 1}
        if cursor.rowcount not in allowed_rowcounts:
            raise RuntimeError(
                f"Unexpected {row.action} count for se_parameter_id "
                f"{row.se_parameter_id}: {cursor.rowcount}; rolling back."
            )
        counts[row.action] += 1
    return counts


def verify_rows(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    rows: Sequence[ExcelParameterRow],
) -> None:
    ids = [row.se_parameter_id for row in rows]
    actual = fetch_existing_rows(cursor, project_id, phase_id, variation_id, ids)
    expected = {row.se_parameter_id: dict(row.values) for row in rows}
    mismatches = {
        item_id: (values, actual.get(item_id))
        for item_id, values in expected.items()
        if actual.get(item_id) != values
    }
    if mismatches:
        raise RuntimeError(
            f"Final verification failed; sample={dict(list(mismatches.items())[:10])}. "
            "Transaction will roll back."
        )


def upsert_project_parameters(
    project_id: object,
    phase_id: object,
    excel_path: Path | str,
    *,
    sheet_name: str = RUN_SHEET_NAME,
    execute: bool = False,
    assume_yes: bool = False,
    db_connection=connection,
) -> list[PlanRow]:
    project_id = normalize_positive_id(project_id, label="project_id")
    phase_id = normalize_positive_id(phase_id, label="phase_id")
    excel_path = Path(excel_path).expanduser().resolve()
    source_read_results = read_excel_variations(excel_path, sheet_name)
    read_results: list[ExcelReadResult] = []
    previews: list[
        tuple[int, ExcelReadResult, dict[str, str], IdComparison, list[PlanRow]]
    ] = []
    with db_connection.cursor() as cursor:
        for source_result in source_read_results:
            read_results.append(
                ExcelReadResult(
                    source_result.variation_name,
                    resolve_approval_status_rows(cursor, source_result.rows),
                    source_result.header_row,
                    source_result.sheet_name,
                )
            )
        variation_ids_by_name = resolve_variation_ids(
            cursor, [result.variation_name for result in read_results]
        )
        for read_result in read_results:
            variation_id = variation_ids_by_name[read_result.variation_name]
            target, comparison, plan = build_plan(
                cursor,
                project_id,
                phase_id,
                variation_id,
                read_result.rows,
            )
            if normalize_variation_name(
                target["variation"]
            ) != normalize_variation_name(read_result.variation_name):
                raise ValueError(
                    f"variation_id {variation_id} is "
                    f"{target['variation']!r} in DB, but its Excel block is "
                    f"{read_result.variation_name!r}."
                )
            previews.append((variation_id, read_result, target, comparison, plan))

    for variation_id, read_result, target, comparison, plan in previews:
        print("\n" + "=" * 80)
        print_plan(
            db_connection,
            project_id=project_id,
            phase_id=phase_id,
            variation_id=variation_id,
            excel_path=excel_path,
            sheet_name=sheet_name,
            read_result=read_result,
            target=target,
            comparison=comparison,
            plan=plan,
        )
    has_reference_errors = any(
        comparison.missing_master
        or comparison.invalid_statuses
        or comparison.missing_employees
        for _, _, _, comparison, _ in previews
    )
    if has_reference_errors:
        print(
            "\nBLOCKED: invalid foreign-key values were found. "
            "Dry-run completed, but execute mode is not allowed."
        )
    if not execute:
        print("\nDRY RUN: no data was changed. Set RUN_EXECUTE=True to apply.")
        return [row for _, _, _, _, plan in previews for row in plan]
    if has_reference_errors:
        raise ValueError(
            "Cannot execute while master se_parameter, approval_status, or employee "
            "references are missing."
        )

    database_name = db_connection.settings_dict.get("NAME") or "<unknown>"
    variation_id_text = ",".join(
        str(variation_id) for variation_id, _, _, _, _ in previews
    )
    confirmation = (
        f"UPSERT_PROJECT_PARAMETER {database_name}:"
        f"{project_id}@{phase_id}@[{variation_id_text}]"
    )
    if not assume_yes:
        answer = input(f"\nType exactly '{confirmation}' to continue: ").strip()
        if answer != confirmation:
            print("Cancelled. No data was changed.")
            return [row for _, _, _, _, plan in previews for row in plan]

    lock_name: str | None = None
    try:
        with db_connection.cursor() as cursor:
            lock_name = acquire_lock(cursor, str(database_name))
        with transaction.atomic(using=db_connection.alias):
            with db_connection.cursor() as cursor:
                current_variation_ids = resolve_variation_ids(
                    cursor, [result.variation_name for result in read_results]
                )
                if current_variation_ids != variation_ids_by_name:
                    raise RuntimeError(
                        "Variation mapping changed after preview; transaction will "
                        "roll back."
                    )
                current_read_results = [
                    ExcelReadResult(
                        source_result.variation_name,
                        resolve_approval_status_rows(cursor, source_result.rows),
                        source_result.header_row,
                        source_result.sheet_name,
                    )
                    for source_result in source_read_results
                ]
                if current_read_results != read_results:
                    raise RuntimeError(
                        "Approval-status mapping changed after preview; transaction "
                        "will roll back."
                    )
                total = {"INSERT": 0, "UPDATE": 0, "UNCHANGED": 0}
                for variation_id, read_result, _, comparison, plan in previews:
                    _, current_comparison, current_plan = build_plan(
                        cursor,
                        project_id,
                        phase_id,
                        variation_id,
                        read_result.rows,
                        lock_rows=True,
                    )
                    if current_comparison != comparison or plan_fingerprint(
                        current_plan
                    ) != plan_fingerprint(plan):
                        raise RuntimeError(
                            "Database changed after preview; transaction will roll back."
                        )
                    actual = apply_plan(
                        cursor,
                        project_id,
                        phase_id,
                        variation_id,
                        current_plan,
                    )
                    for action in total:
                        total[action] += actual[action]
                    verify_rows(
                        cursor,
                        project_id,
                        phase_id,
                        variation_id,
                        read_result.rows,
                    )
        print("\nproject_parameter upsert completed successfully:")
        print(
            f"  INSERT={total['INSERT']}, UPDATE={total['UPDATE']}, "
            f"UNCHANGED={total['UNCHANGED']}"
        )
        return [row for _, _, _, _, plan in previews for row in plan]
    finally:
        if lock_name is not None:
            with db_connection.cursor() as cursor:
                cursor.execute("SELECT RELEASE_LOCK(%s)", (lock_name,))


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Upsert project_parameter from Excel; blank mapped cells become SQL NULL."
        )
    )
    parser.add_argument("--project-id", type=int, default=RUN_PROJECT_ID)
    parser.add_argument("--phase-id", type=int, default=RUN_PHASE_ID)
    parser.add_argument("--excel", type=Path, default=RUN_EXCEL_PATH)
    parser.add_argument("--sheet", default=RUN_SHEET_NAME)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--yes", action="store_true")
    return parser.parse_args(argv)


def main() -> None:
    args = parse_args()
    execute = args.execute or RUN_EXECUTE
    assume_yes = args.yes or RUN_ASSUME_YES
    if assume_yes and not execute:
        raise ValueError("RUN_ASSUME_YES/--yes requires execute mode.")
    upsert_project_parameters(
        args.project_id,
        args.phase_id,
        args.excel,
        sheet_name=args.sheet,
        execute=execute,
        assume_yes=assume_yes,
    )


if __name__ == "__main__":
    main()
