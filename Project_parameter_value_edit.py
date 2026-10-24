"""Safely sync project_parameter from the UpdateしたいValue Excel sheet.

The primary key uses project_id, phase_id, variation_id and the Excel
se-parameter_id. All six mapped cells are authoritative; blank cells overwrite
existing database values with SQL NULL. DB-only rows are report-only deletion
candidates unless deletion is explicitly enabled. The default mode is a dry run.
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402

# PyCharm Run-button configuration.
RUN_PROJECT_ID = 1020
RUN_PHASE_ID = 35
RUN_EXCEL_PATH = Path(
    # "/Users/TrongKieu/Downloads/SE_List_P61QR_HEV_US GCC_4WD_UC-Lot_中間確認会UPDATE_VALUE_EDITED.xlsx"
    "/Users/TrongKieu/Downloads/SE_List_P61QR_HEV_US GCC_4WD_UC-Lot_中間確認会UPDATE_VALUE - Copy.xlsx"
)
RUN_SHEET_NAME = "VALUE"

# RUN_PROJECT_ID = 1019
# RUN_PHASE_ID = 33
# RUN_EXCEL_PATH = Path(
#     "/Users/TrongKieu/Downloads/SE_List_P61QR_HEV_US_4WD_UC-Lot_Kickoff_value_UPDATE (1).xlsx"
# )
# RUN_SHEET_NAME = "UpdateしたいValue"
RUN_REPORT_PATH: Path | None = None


RUN_EXECUTE = False
RUN_ASSUME_YES = False
RUN_DELETE_DB_ONLY = False

HEADER_SCAN_LIMIT = 20
PREVIEW_ROW_LIMIT = 30
ADVISORY_LOCK_TIMEOUT_SECONDS = 10

# Fixed, reviewed mapping; identifiers never come from runtime user input.
EXCEL_TO_DB_COLUMNS = {
    "値": "value",
    "ステータス": "approval_status",
    "メモ": "user_memo",
    "担当者": "z_note",
    "更新者": "employee_number",
    "更新日": "update_day",
}
DATA_COLUMNS = tuple(EXCEL_TO_DB_COLUMNS.values())
SE_PARAMETER_HEADERS = {"separameterid", "seprameterid"}


@dataclass(frozen=True)
class ExcelParameterRow:
    excel_row: int
    se_parameter_id: int
    values: tuple[tuple[str, object], ...]


@dataclass(frozen=True)
class ExcelReadResult:
    variation_name: str
    rows: tuple[ExcelParameterRow, ...]
    header_row: int
    sheet_name: str


@dataclass(frozen=True)
class PlanRow:
    excel_row: int
    se_parameter_id: int
    parameter_name: str
    old_values: tuple[tuple[str, object], ...]
    new_values: tuple[tuple[str, object], ...]
    action: str


@dataclass(frozen=True)
class IdComparison:
    db_only: tuple[int, ...]
    excel_only: tuple[int, ...]
    missing_master: tuple[int, ...] = ()
    invalid_statuses: tuple[object, ...] = ()
    missing_employees: tuple[str, ...] = ()


@dataclass(frozen=True)
class DatabaseOnlyRow:
    se_parameter_id: int
    parameter_name: str
    values: tuple[tuple[str, object], ...]


@dataclass(frozen=True)
class VariationPreview:
    variation_id: int
    read_result: ExcelReadResult
    target: Mapping[str, str]
    comparison: IdComparison
    plan: tuple[PlanRow, ...]
    database_only_rows: tuple[DatabaseOnlyRow, ...]


def normalize_positive_id(raw_value: object, *, label: str) -> int:
    if isinstance(raw_value, bool):
        raise ValueError(f"{label} must be a positive integer, not boolean.")
    if isinstance(raw_value, float) and not raw_value.is_integer():
        raise ValueError(f"{label} must be an integer: {raw_value!r}")
    try:
        value = int(raw_value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {label}: {raw_value!r}") from exc
    if value <= 0:
        raise ValueError(f"{label} must be greater than zero: {value}")
    return value


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    return "".join(char for char in text if char not in " _-\n\r\t")


def normalize_data_header(value: object) -> str:
    """Normalize generated headers such as 値2 and 更新日56."""
    return re.sub(r"\d+$", "", normalize_header(value))


def normalize_variation_name(value: object) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()


def normalize_text(value: object, *, strip: bool = False) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        if not value.strip():
            return None
        return value.strip() if strip else value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value)


def normalize_status(value: object, *, row: int) -> int | str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool) or (isinstance(value, float) and not value.is_integer()):
        raise ValueError(f"approval_status at Excel row {row} must be an integer.")
    normalized_value = value
    if isinstance(value, str):
        display_match = re.fullmatch(r"\s*(\d+)\s*[.．].*", value, flags=re.DOTALL)
        if display_match:
            normalized_value = display_match.group(1)
    try:
        status = int(normalized_value)
    except (TypeError, ValueError):
        if isinstance(value, str):
            return " ".join(value.split())
        raise ValueError(f"approval_status at Excel row {row} is invalid: {value!r}")
    if status < 0:
        raise ValueError(f"approval_status at Excel row {row} cannot be negative.")
    return status


def normalize_update_day(value: object, *, row: int) -> dt.datetime | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None, microsecond=0)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min)
    if isinstance(value, str):
        text = value.strip().lstrip(",").strip()
        try:
            return dt.datetime.fromisoformat(text).replace(tzinfo=None, microsecond=0)
        except ValueError as exc:
            raise ValueError(
                f"update_day at Excel row {row} is invalid: {value!r}. "
                "Use YYYY-MM-DD or YYYY-MM-DD HH:MM:SS."
            ) from exc
    raise ValueError(f"update_day at Excel row {row} is invalid: {value!r}")


def normalize_field(column: str, value: object, *, row: int) -> object:
    if column not in DATA_COLUMNS:
        raise ValueError(f"Unexpected project_parameter column: {column!r}")
    if column == "approval_status":
        return normalize_status(value, row=row)
    if column == "update_day":
        return normalize_update_day(value, row=row)
    if column == "employee_number":
        return normalize_text(value, strip=True)
    return normalize_text(value)


def find_excel_layout(
    worksheet,
) -> tuple[int, int, list[tuple[str, dict[str, int]]]]:
    expected_headers = [normalize_data_header(name) for name in EXCEL_TO_DB_COLUMNS]
    max_row = min(worksheet.max_row, HEADER_SCAN_LIMIT)
    for row_number in range(1, max_row + 1):
        headers = {
            column: normalize_data_header(cell.value)
            for column, cell in enumerate(worksheet[row_number], start=1)
        }
        se_columns = [
            column
            for column, header in headers.items()
            if header in SE_PARAMETER_HEADERS
        ]
        if len(se_columns) != 1:
            continue
        blocks: list[tuple[str, dict[str, int]]] = []
        block_width = len(expected_headers)
        for start_column in range(1, worksheet.max_column - block_width + 2):
            actual = [
                headers.get(start_column + offset, "") for offset in range(block_width)
            ]
            if actual != expected_headers:
                continue
            variation_cell = worksheet.cell(row_number - 1, start_column)
            variation_name = normalize_text(variation_cell.value, strip=True)
            if not variation_name:
                raise ValueError(
                    f"Missing variation name above Excel column {start_column}."
                )
            field_columns = {
                db_column: start_column + offset
                for offset, db_column in enumerate(DATA_COLUMNS)
            }
            blocks.append((variation_name, field_columns))
        if blocks:
            return row_number, se_columns[0], blocks
    required = ["se-parameter_id", *EXCEL_TO_DB_COLUMNS]
    raise ValueError(
        f"Could not find a complete six-column variation block in the first "
        f"{max_row} rows: {required}"
    )


def read_excel_variations(
    excel_path: Path | str,
    sheet_name: str,
) -> list[ExcelReadResult]:
    path = Path(excel_path).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"Excel file not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Excel input must be an .xlsx or .xlsm file.")
    formula_workbook = load_workbook(path, read_only=True, data_only=False)
    value_workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name in formula_workbook.sheetnames:
            actual_sheet_name = sheet_name
        elif len(formula_workbook.sheetnames) == 1:
            actual_sheet_name = formula_workbook.sheetnames[0]
        else:
            raise ValueError(
                f"Sheet {sheet_name!r} not found; "
                f"available={formula_workbook.sheetnames}"
            )
        formula_sheet = formula_workbook[actual_sheet_name]
        value_sheet = value_workbook[actual_sheet_name]
        header_row, se_column, blocks = find_excel_layout(formula_sheet)
        seen_variations: dict[str, str] = {}
        for variation_name, columns in blocks:
            key = normalize_variation_name(variation_name)
            if key in seen_variations:
                raise ValueError(
                    "Duplicate variation blocks after normalization: "
                    f"{seen_variations[key]!r} and {variation_name!r}."
                )
            seen_variations[key] = variation_name

        results: list[ExcelReadResult] = []
        for variation_name, field_columns in blocks:
            rows: list[ExcelParameterRow] = []
            seen: dict[int, int] = {}
            formula_rows = formula_sheet.iter_rows(min_row=header_row + 1)
            value_rows = value_sheet.iter_rows(min_row=header_row + 1)
            for row_number, (formula_row, value_row) in enumerate(
                zip(formula_rows, value_rows), start=header_row + 1
            ):
                se_cell = formula_row[se_column - 1]
                formula_cells = {
                    column: formula_row[column_number - 1]
                    for column, column_number in field_columns.items()
                }
                value_cells = {
                    column: value_row[column_number - 1]
                    for column, column_number in field_columns.items()
                }
                if se_cell.data_type == "f":
                    raise ValueError(
                        f"Formula is not allowed in se-parameter_id at row {row_number}."
                    )
                if se_cell.value in (None, ""):
                    if any(
                        cell.value not in (None, "") for cell in formula_cells.values()
                    ):
                        raise ValueError(f"Excel row {row_number} has data but no ID.")
                    continue
                se_id = normalize_positive_id(
                    se_cell.value, label=f"se-parameter_id at Excel row {row_number}"
                )
                if se_id in seen:
                    raise ValueError(
                        f"Duplicate se-parameter_id {se_id} in variation "
                        f"{variation_name!r} at rows {seen[se_id]} and {row_number}."
                    )
                seen[se_id] = row_number
                for column, formula_cell in formula_cells.items():
                    if (
                        formula_cell.data_type == "f"
                        and value_cells[column].value is None
                    ):
                        raise ValueError(
                            f"Formula in {variation_name!r}, Excel row {row_number}, "
                            f"column {column!r} has no cached value. Recalculate and "
                            "save the workbook in Excel before running this tool."
                        )
                values = tuple(
                    (
                        column,
                        normalize_field(
                            column, value_cells[column].value, row=row_number
                        ),
                    )
                    for column in DATA_COLUMNS
                )
                rows.append(ExcelParameterRow(row_number, se_id, values))
            if not rows:
                raise ValueError(
                    f"No rows with se-parameter_id found for {variation_name!r}."
                )
            results.append(
                ExcelReadResult(
                    variation_name, tuple(rows), header_row, actual_sheet_name
                )
            )
    finally:
        formula_workbook.close()
        value_workbook.close()
    return results


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Cannot create placeholders for an empty sequence.")
    return ", ".join(["%s"] * len(values))


def resolve_approval_status_rows(
    cursor, rows: Sequence[ExcelParameterRow]
) -> tuple[ExcelParameterRow, ...]:
    """Resolve status labels to approval_status_table.state_key when possible."""
    cursor.execute(
        "SELECT state_key, state_name FROM approval_status_table ORDER BY state_key"
    )
    by_name: dict[str, list[int]] = {}
    for raw_key, raw_name in cursor.fetchall():
        name_key = normalize_variation_name(raw_name)
        if name_key:
            by_name.setdefault(name_key, []).append(int(raw_key))

    resolved_rows: list[ExcelParameterRow] = []
    for row in rows:
        values = dict(row.values)
        status = values["approval_status"]
        if isinstance(status, str):
            matches = by_name.get(normalize_variation_name(status), [])
            if len(matches) > 1:
                raise ValueError(
                    f"approval status label {status!r} at Excel row "
                    f"{row.excel_row} matches multiple state keys: {matches}"
                )
            if len(matches) == 1:
                values["approval_status"] = matches[0]
        resolved_rows.append(
            ExcelParameterRow(
                row.excel_row,
                row.se_parameter_id,
                tuple((column, values[column]) for column in DATA_COLUMNS),
            )
        )
    return tuple(resolved_rows)


def resolve_variation_ids(cursor, variation_names: Sequence[str]) -> dict[str, int]:
    """Resolve Excel variation labels without relying on hard-coded IDs."""
    if not variation_names:
        raise ValueError("No variation names were detected in Excel.")
    cursor.execute("SELECT id, variation FROM variation ORDER BY id")
    by_name: dict[str, list[int]] = {}
    for raw_id, raw_name in cursor.fetchall():
        key = normalize_variation_name(raw_name)
        if key:
            by_name.setdefault(key, []).append(int(raw_id))

    resolved: dict[str, int] = {}
    missing: list[str] = []
    ambiguous: dict[str, list[int]] = {}
    for excel_name in variation_names:
        matches = by_name.get(normalize_variation_name(excel_name), [])
        if not matches:
            missing.append(excel_name)
        elif len(matches) > 1:
            ambiguous[excel_name] = matches
        else:
            resolved[excel_name] = matches[0]
    if missing or ambiguous:
        raise ValueError(
            "Could not resolve Excel variation names safely; "
            f"missing_in_variation_table={missing}, ambiguous={ambiguous}"
        )
    return resolved


def fetch_target_label(cursor, table: str, object_id: int, column: str) -> str:
    if (table, column) not in {
        ("project_info", "project_code"),
        ("phase", "phase"),
        ("variation", "variation"),
    }:
        raise ValueError("Unexpected target lookup.")
    cursor.execute(f"SELECT {column} FROM {table} WHERE id = %s", (object_id,))
    rows = cursor.fetchall()
    if len(rows) != 1:
        raise ValueError(f"{table} ID {object_id} was not found or is ambiguous.")
    return str(rows[0][0])


def fetch_parameter_names(cursor, ids: Sequence[int]) -> dict[int, str]:
    cursor.execute(
        f"SELECT id, parameter_name_1, parameter_name_2 FROM se_parameter "
        f"WHERE id IN ({make_placeholders(ids)})",
        tuple(ids),
    )
    names = {
        int(item_id): " / ".join(
            str(value) for value in (name_1, name_2) if value not in (None, "")
        )
        for item_id, name_1, name_2 in cursor.fetchall()
    }
    return names


def validate_references(
    cursor, rows: Sequence[ExcelParameterRow]
) -> tuple[tuple[object, ...], tuple[str, ...]]:
    raw_statuses = {
        dict(row.values)["approval_status"]
        for row in rows
        if dict(row.values)["approval_status"] is not None
    }
    statuses = sorted(
        {
            int(value)
            for value in raw_statuses
            if isinstance(value, int) and not isinstance(value, bool)
        }
    )
    unresolved_statuses = sorted(
        (value for value in raw_statuses if isinstance(value, str)), key=str
    )
    employees = sorted(
        {
            str(dict(row.values)["employee_number"])
            for row in rows
            if dict(row.values)["employee_number"] is not None
        }
    )
    if statuses:
        cursor.execute(
            f"SELECT state_key FROM approval_status_table "
            f"WHERE state_key IN ({make_placeholders(statuses)})",
            tuple(statuses),
        )
        found = {int(row[0]) for row in cursor.fetchall()}
        invalid_statuses = tuple(
            [item for item in statuses if item not in found] + unresolved_statuses
        )
    else:
        invalid_statuses = tuple(unresolved_statuses)
    if employees:
        cursor.execute(
            f"SELECT employee_number FROM user_table "
            f"WHERE employee_number IN ({make_placeholders(employees)})",
            tuple(employees),
        )
        found = {str(row[0]) for row in cursor.fetchall()}
        missing_employees = tuple(item for item in employees if item not in found)
    else:
        missing_employees = ()
    return invalid_statuses, missing_employees


def normalize_database_values(values: Sequence[object]) -> dict[str, object]:
    return {
        column: normalize_field(column, value, row=0)
        for column, value in zip(DATA_COLUMNS, values)
    }


def fetch_existing_rows(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    ids: Sequence[int],
    *,
    lock_rows: bool = False,
) -> dict[int, dict[str, object]]:
    columns = ", ".join(f"`{column}`" for column in DATA_COLUMNS)
    lock = " FOR UPDATE" if lock_rows else ""
    cursor.execute(
        f"""
        SELECT se_parameter_id, {columns}
        FROM project_parameter
        WHERE project_id = %s AND phase_id = %s AND variation_id = %s
          AND se_parameter_id IN ({make_placeholders(ids)})
        ORDER BY se_parameter_id{lock}
        """,
        (project_id, phase_id, variation_id, *ids),
    )
    return {
        int(row[0]): normalize_database_values(row[1:]) for row in cursor.fetchall()
    }


def compare_se_parameter_ids(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    excel_ids: Sequence[int],
    *,
    lock_rows: bool = False,
) -> IdComparison:
    """Compare Excel IDs with DB IDs for the exact three-key target."""
    lock = " FOR UPDATE" if lock_rows else ""
    cursor.execute(
        f"""
        SELECT se_parameter_id
        FROM project_parameter
        WHERE project_id = %s AND phase_id = %s AND variation_id = %s
        ORDER BY se_parameter_id{lock}
        """,
        (project_id, phase_id, variation_id),
    )
    db_ids = {int(row[0]) for row in cursor.fetchall()}
    excel_id_set = set(excel_ids)
    return IdComparison(
        db_only=tuple(sorted(db_ids - excel_id_set)),
        excel_only=tuple(sorted(excel_id_set - db_ids)),
    )


def build_plan(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    rows: Sequence[ExcelParameterRow],
    *,
    lock_rows: bool = False,
) -> tuple[dict[str, str], IdComparison, list[PlanRow]]:
    target = {
        "project": fetch_target_label(
            cursor, "project_info", project_id, "project_code"
        ),
        "phase": fetch_target_label(cursor, "phase", phase_id, "phase"),
        "variation": fetch_target_label(cursor, "variation", variation_id, "variation"),
    }
    ids = [row.se_parameter_id for row in rows]
    names = fetch_parameter_names(cursor, ids)
    missing_master = tuple(sorted(set(ids) - set(names)))
    invalid_statuses, missing_employees = validate_references(cursor, rows)
    base_comparison = compare_se_parameter_ids(
        cursor,
        project_id,
        phase_id,
        variation_id,
        ids,
        lock_rows=lock_rows,
    )
    comparison = IdComparison(
        db_only=base_comparison.db_only,
        excel_only=base_comparison.excel_only,
        missing_master=missing_master,
        invalid_statuses=invalid_statuses,
        missing_employees=missing_employees,
    )
    existing = fetch_existing_rows(
        cursor, project_id, phase_id, variation_id, ids, lock_rows=lock_rows
    )
    plan: list[PlanRow] = []
    for row in rows:
        new = dict(row.values)
        old = existing.get(row.se_parameter_id)
        has_invalid_reference = (
            row.se_parameter_id in comparison.missing_master
            or new["approval_status"] in comparison.invalid_statuses
            or new["employee_number"] in comparison.missing_employees
        )
        if has_invalid_reference:
            action = "BLOCKED_REFERENCE"
            old = old or {column: None for column in DATA_COLUMNS}
        elif old is None:
            action = "INSERT"
            old = {column: None for column in DATA_COLUMNS}
        else:
            action = "UNCHANGED" if old == new else "UPDATE"
        plan.append(
            PlanRow(
                row.excel_row,
                row.se_parameter_id,
                names.get(row.se_parameter_id, "<missing in se_parameter master>"),
                tuple((column, old[column]) for column in DATA_COLUMNS),
                tuple((column, new[column]) for column in DATA_COLUMNS),
                action,
            )
        )
    return target, comparison, plan


def fetch_database_only_rows(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    ids: Sequence[int],
) -> tuple[DatabaseOnlyRow, ...]:
    """Fetch audit details for rows present in DB but absent from Excel."""
    if not ids:
        return ()
    existing = fetch_existing_rows(cursor, project_id, phase_id, variation_id, ids)
    names = fetch_parameter_names(cursor, ids)
    return tuple(
        DatabaseOnlyRow(
            se_parameter_id=item_id,
            parameter_name=names.get(item_id, "<missing in se_parameter master>"),
            values=tuple(
                (column, existing[item_id][column]) for column in DATA_COLUMNS
            ),
        )
        for item_id in ids
        if item_id in existing
    )


def summarize_plan(plan: Iterable[PlanRow]) -> dict[str, int]:
    counts = {"INSERT": 0, "UPDATE": 0, "UNCHANGED": 0, "BLOCKED_REFERENCE": 0}
    for row in plan:
        counts[row.action] += 1
    return counts


def changed_fields(row: PlanRow) -> list[str]:
    if row.action == "BLOCKED_REFERENCE":
        return []
    old = dict(row.old_values)
    return [
        column
        for column, value in row.new_values
        if row.action == "INSERT" or old[column] != value
    ]


def compact(value: object, limit: int = 60) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    text = str(value).replace("\r", "\\r").replace("\n", "\\n")
    return repr(text if len(text) <= limit else text[: limit - 3] + "...")


def print_plan(
    db_connection,
    *,
    project_id: int,
    phase_id: int,
    variation_id: int,
    excel_path: Path,
    sheet_name: str,
    read_result: ExcelReadResult,
    target: Mapping[str, str],
    comparison: IdComparison,
    plan: Sequence[PlanRow],
) -> None:
    settings = db_connection.settings_dict
    counts = summarize_plan(plan)
    field_counts = {column: 0 for column in DATA_COLUMNS}
    for row in plan:
        for column in changed_fields(row):
            field_counts[column] += 1
    print(f"Database: {settings.get('NAME')}")
    print(f"Host: {settings.get('HOST')}:{settings.get('PORT') or 3306}")
    print("Affected table: project_parameter")
    print(f"Excel: {excel_path}")
    print(f"Sheet: {read_result.sheet_name!r}; header row: {read_result.header_row}")
    print(
        f"Target: project={project_id} ({target['project']}), "
        f"phase={phase_id} ({target['phase']}), "
        f"variation={variation_id} ({target['variation']})"
    )
    print(f"Excel rows selected: {len(read_result.rows)}")
    print(
        "DB_ONLY se_parameter_id "
        f"({len(comparison.db_only)}; present in DB, absent from Excel):"
    )
    print(f"  {list(comparison.db_only)}")
    print(
        "EXCEL_ONLY se_parameter_id "
        f"({len(comparison.excel_only)}; present in Excel, absent from DB):"
    )
    print(f"  {list(comparison.excel_only)}")
    print(
        "MISSING_IN_SE_PARAMETER_MASTER "
        f"({len(comparison.missing_master)}; cannot be inserted):"
    )
    print(f"  {list(comparison.missing_master)}")
    print(
        "INVALID_APPROVAL_STATUS "
        f"({len(comparison.invalid_statuses)}; absent from approval_status_table):"
    )
    print(f"  {list(comparison.invalid_statuses)}")
    print(
        "MISSING_EMPLOYEE_NUMBER "
        f"({len(comparison.missing_employees)}; absent from user_table):"
    )
    print(f"  {list(comparison.missing_employees)}")
    print(
        f"Estimated rows: INSERT={counts['INSERT']}, UPDATE={counts['UPDATE']}, "
        f"UNCHANGED={counts['UNCHANGED']}, "
        f"BLOCKED_REFERENCE={counts['BLOCKED_REFERENCE']}"
    )
    print("Changed fields (NULL overwrites included):")
    for column in DATA_COLUMNS:
        print(f"  {column}: {field_counts[column]}")
    print("Planned changes:")
    changes = [row for row in plan if row.action != "UNCHANGED"]
    for row in changes[:PREVIEW_ROW_LIMIT]:
        old = dict(row.old_values)
        new = dict(row.new_values)
        detail = ", ".join(
            f"{column}: {compact(old[column])} -> {compact(new[column])}"
            for column in changed_fields(row)
        )
        print(
            f"- Excel row {row.excel_row}, se_parameter_id={row.se_parameter_id} "
            f"({row.parameter_name}): {row.action}; {detail}"
        )
    if len(changes) > PREVIEW_ROW_LIMIT:
        print(f"  ... and {len(changes) - PREVIEW_ROW_LIMIT} more changed rows")


REPORT_ACTION_FILLS = {
    "INSERT": "D9EAD3",
    "UPDATE": "FFF2CC",
    "UNCHANGED": "E7E6E6",
    "BLOCKED_REFERENCE": "F4CCCC",
    "DELETE_CANDIDATE": "FCE4D6",
}
REPORT_SEVERITY_FILLS = {
    "LỖI": "F4CCCC",
    "CẢNH BÁO": "FFF2CC",
    "THÔNG TIN": "D9EAF7",
}
REPORT_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
REPORT_HEADER_FONT = Font(color="FFFFFF", bold=True)
REPORT_THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)


def resolve_report_path(
    excel_path: Path | str,
    project_id: int,
    phase_id: int,
    report_path: Path | str | None = None,
    *,
    generated_at: dt.datetime | None = None,
) -> Path:
    """Resolve a safe .xlsx report path without overwriting the source file."""
    source = Path(excel_path).expanduser().resolve()
    if report_path is not None:
        output = Path(report_path).expanduser().resolve()
        if output.suffix.lower() != ".xlsx":
            raise ValueError("Report output must be an .xlsx file.")
    else:
        timestamp = (generated_at or dt.datetime.now()).strftime("%Y%m%d_%H%M%S")
        base_name = (
            f"{source.stem}_project_parameter_report_"
            f"P{project_id}_PH{phase_id}_{timestamp}"
        )
        output = source.with_name(f"{base_name}.xlsx")
        counter = 2
        while output.exists():
            output = source.with_name(f"{base_name}_{counter}.xlsx")
            counter += 1
    if os.path.normcase(str(output)) == os.path.normcase(str(source)):
        raise ValueError("Report output cannot overwrite the source Excel file.")
    return output


def report_value(value: object) -> object:
    """Keep dates typed while making SQL NULL explicit for audit purposes."""
    return "NULL" if value is None else value


def _append_report_row(worksheet, values: Sequence[object]) -> None:
    """Append values while keeping leading-equals user text out of formula cells."""
    worksheet.append(values)
    for cell in worksheet[worksheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def format_report_error(exc: Exception, limit: int = 1000) -> str:
    message = " ".join(str(exc).split())
    detail = f"{type(exc).__name__}: {message}"
    return detail if len(detail) <= limit else detail[: limit - 3] + "..."


def _style_header(worksheet, row_number: int) -> None:
    for cell in worksheet[row_number]:
        if cell.value is None:
            continue
        cell.fill = REPORT_HEADER_FILL
        cell.font = REPORT_HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = REPORT_THIN_BORDER
    worksheet.row_dimensions[row_number].height = 34


def _finish_table_sheet(worksheet, *, action_column: int | None = None) -> None:
    _style_header(worksheet, 1)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = REPORT_THIN_BORDER
        if action_column is not None:
            action_cell = row[action_column - 1]
            color = REPORT_ACTION_FILLS.get(str(action_cell.value))
            if color:
                action_cell.fill = PatternFill("solid", fgColor=color)
                action_cell.font = Font(bold=True)
    for column_number, cells in enumerate(worksheet.columns, start=1):
        maximum = max(
            (len(str(cell.value)) for cell in cells if cell.value is not None),
            default=0,
        )
        worksheet.column_dimensions[get_column_letter(column_number)].width = min(
            max(maximum + 2, 11), 48
        )


def _append_summary_sheet(
    worksheet,
    *,
    db_connection,
    project_id: int,
    phase_id: int,
    excel_path: Path,
    execute: bool,
    status: str,
    generated_at: dt.datetime,
    previews: Sequence[VariationPreview],
    delete_db_only: bool,
    deleted_count: int,
    error_detail: str | None,
) -> None:
    settings = db_connection.settings_dict
    target = previews[0].target
    all_plan_rows = [row for preview in previews for row in preview.plan]
    counts = summarize_plan(all_plan_rows)
    field_counts = {column: 0 for column in DATA_COLUMNS}
    for row in all_plan_rows:
        for column in changed_fields(row):
            field_counts[column] += 1
    delete_candidate_count = sum(
        len(preview.database_only_rows) for preview in previews
    )
    issue_count = sum(
        len(preview.comparison.db_only)
        + len(preview.comparison.excel_only)
        + len(preview.comparison.missing_master)
        + sum(
            1
            for row in preview.plan
            if dict(row.new_values)["approval_status"]
            in preview.comparison.invalid_statuses
        )
        + sum(
            1
            for row in preview.plan
            if dict(row.new_values)["employee_number"]
            in preview.comparison.missing_employees
        )
        for preview in previews
    )

    worksheet.append(["Báo cáo cập nhật project_parameter"])
    worksheet.merge_cells("A1:H1")
    worksheet["A1"].fill = REPORT_HEADER_FILL
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    worksheet["A1"].alignment = Alignment(horizontal="center")
    metadata = [
        ("Thời gian tạo", generated_at),
        ("Chế độ", "EXECUTE" if execute else "DRY_RUN"),
        ("Trạng thái", status),
        ("Database", settings.get("NAME")),
        ("Host", f"{settings.get('HOST')}:{settings.get('PORT') or 3306}"),
        ("File Excel nguồn", str(excel_path)),
        ("Sheet nguồn", previews[0].read_result.sheet_name),
        ("Dòng tiêu đề", previews[0].read_result.header_row),
        ("Project", f"{project_id} ({target['project']})"),
        ("Phase", f"{phase_id} ({target['phase']})"),
        ("Số variation", len(previews)),
        ("Tổng dòng Excel", len(all_plan_rows)),
        ("Cho phép xoá DB_ONLY", "CÓ" if delete_db_only else "KHÔNG"),
        ("Số DELETE_CANDIDATE", delete_candidate_count),
        ("Số bản ghi đã xoá", deleted_count),
        ("Tổng vấn đề kiểm tra", issue_count),
    ]
    if error_detail:
        metadata.append(("Lỗi thực thi", error_detail))
    for label, value in metadata:
        _append_report_row(worksheet, [label, value])
    worksheet[4][1].fill = PatternFill(
        "solid",
        fgColor=(
            "D9EAD3"
            if status == "SUCCESS"
            else "F4CCCC"
            if status in {"FAILED", "BLOCKED_VALIDATION"}
            else "FFF2CC"
        ),
    )
    worksheet[4][1].font = Font(bold=True)

    worksheet.append([])
    variation_header_row = worksheet.max_row + 1
    worksheet.append(
        [
            "Variation ID",
            "Variation",
            "Tổng dòng",
            "INSERT",
            "UPDATE",
            "UNCHANGED",
            "BLOCKED_REFERENCE",
        ]
    )
    for preview in previews:
        variation_counts = summarize_plan(preview.plan)
        _append_report_row(
            worksheet,
            [
                preview.variation_id,
                preview.target["variation"],
                len(preview.plan),
                variation_counts["INSERT"],
                variation_counts["UPDATE"],
                variation_counts["UNCHANGED"],
                variation_counts["BLOCKED_REFERENCE"],
            ],
        )
    _style_header(worksheet, variation_header_row)

    worksheet.append([])
    action_header_row = worksheet.max_row + 1
    worksheet.append(["Hành động", "Số dòng"])
    for action in ("INSERT", "UPDATE", "UNCHANGED", "BLOCKED_REFERENCE"):
        worksheet.append([action, counts[action]])
    worksheet.append(["DELETE_CANDIDATE", delete_candidate_count])
    _style_header(worksheet, action_header_row)

    worksheet.append([])
    field_header_row = worksheet.max_row + 1
    worksheet.append(["Trường DB", "Số dòng thay đổi"])
    for column in DATA_COLUMNS:
        worksheet.append([column, field_counts[column]])
    _style_header(worksheet, field_header_row)

    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.border = REPORT_THIN_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 72
    for column in "CDEFGH":
        worksheet.column_dimensions[column].width = 20


def _append_detail_sheet(
    worksheet,
    *,
    project_id: int,
    phase_id: int,
    previews: Sequence[VariationPreview],
) -> None:
    headers = [
        "Project ID",
        "Mã project",
        "Phase ID",
        "Phase",
        "Variation ID",
        "Variation",
        "Sheet nguồn",
        "Dòng Excel",
        "se_parameter_id",
        "Tên parameter",
        "Hành động",
        "Trường thay đổi",
    ]
    for column in DATA_COLUMNS:
        headers.extend([f"Giá trị cũ ({column})", f"Giá trị mới ({column})"])
    worksheet.append(headers)
    for preview in previews:
        for row in preview.plan:
            old = dict(row.old_values)
            new = dict(row.new_values)
            values: list[object] = [
                project_id,
                preview.target["project"],
                phase_id,
                preview.target["phase"],
                preview.variation_id,
                preview.target["variation"],
                preview.read_result.sheet_name,
                row.excel_row,
                row.se_parameter_id,
                row.parameter_name,
                row.action,
                ", ".join(changed_fields(row)),
            ]
            for column in DATA_COLUMNS:
                values.extend([report_value(old[column]), report_value(new[column])])
            _append_report_row(worksheet, values)
        for database_row in preview.database_only_rows:
            old = dict(database_row.values)
            values = [
                project_id,
                preview.target["project"],
                phase_id,
                preview.target["phase"],
                preview.variation_id,
                preview.target["variation"],
                preview.read_result.sheet_name,
                None,
                database_row.se_parameter_id,
                database_row.parameter_name,
                "DELETE_CANDIDATE",
                "Toàn bộ bản ghi",
            ]
            for column in DATA_COLUMNS:
                values.extend([report_value(old[column]), None])
            _append_report_row(worksheet, values)
    _finish_table_sheet(worksheet, action_column=11)
    for column in range(13, worksheet.max_column + 1, 2):
        worksheet.column_dimensions[get_column_letter(column)].width = 32
        worksheet.column_dimensions[get_column_letter(column + 1)].width = 32


def _append_validation_sheet(worksheet, previews: Sequence[VariationPreview]) -> None:
    headers = [
        "Mức độ",
        "Loại kiểm tra",
        "Variation ID",
        "Variation",
        "Dòng Excel",
        "se_parameter_id",
        "Tên parameter",
        "Giá trị liên quan",
        "Chi tiết",
        *[f"Giá trị DB ({column})" for column in DATA_COLUMNS],
    ]
    worksheet.append(headers)

    def append_issue(
        preview: VariationPreview,
        severity: str,
        issue_type: str,
        *,
        excel_row: int | None = None,
        se_parameter_id: int | None = None,
        parameter_name: str | None = None,
        related_value: object = None,
        detail: str,
        database_values: Mapping[str, object] | None = None,
    ) -> None:
        db_values = database_values or {}
        _append_report_row(
            worksheet,
            [
                severity,
                issue_type,
                preview.variation_id,
                preview.target["variation"],
                excel_row,
                se_parameter_id,
                parameter_name,
                related_value,
                detail,
                *[
                    report_value(db_values[column]) if column in db_values else None
                    for column in DATA_COLUMNS
                ],
            ],
        )

    for preview in previews:
        plans_by_id = {row.se_parameter_id: row for row in preview.plan}
        for database_row in preview.database_only_rows:
            append_issue(
                preview,
                "CẢNH BÁO",
                "DB_ONLY",
                se_parameter_id=database_row.se_parameter_id,
                parameter_name=database_row.parameter_name,
                related_value=database_row.se_parameter_id,
                detail="Có trong DB nhưng không có trong file Excel nguồn.",
                database_values=dict(database_row.values),
            )
        for item_id in preview.comparison.excel_only:
            plan_row = plans_by_id[item_id]
            append_issue(
                preview,
                "THÔNG TIN",
                "EXCEL_ONLY",
                excel_row=plan_row.excel_row,
                se_parameter_id=item_id,
                parameter_name=plan_row.parameter_name,
                related_value=item_id,
                detail="Có trong Excel nhưng chưa có trong DB; sẽ INSERT nếu hợp lệ.",
            )
        for row in preview.plan:
            new = dict(row.new_values)
            if row.se_parameter_id in preview.comparison.missing_master:
                append_issue(
                    preview,
                    "LỖI",
                    "MISSING_IN_SE_PARAMETER_MASTER",
                    excel_row=row.excel_row,
                    se_parameter_id=row.se_parameter_id,
                    parameter_name=row.parameter_name,
                    related_value=row.se_parameter_id,
                    detail="Không tồn tại trong bảng master se_parameter.",
                )
            if new["approval_status"] in preview.comparison.invalid_statuses:
                append_issue(
                    preview,
                    "LỖI",
                    "INVALID_APPROVAL_STATUS",
                    excel_row=row.excel_row,
                    se_parameter_id=row.se_parameter_id,
                    parameter_name=row.parameter_name,
                    related_value=report_value(new["approval_status"]),
                    detail="Không tồn tại trong approval_status_table.",
                )
            if new["employee_number"] in preview.comparison.missing_employees:
                append_issue(
                    preview,
                    "LỖI",
                    "MISSING_EMPLOYEE_NUMBER",
                    excel_row=row.excel_row,
                    se_parameter_id=row.se_parameter_id,
                    parameter_name=row.parameter_name,
                    related_value=report_value(new["employee_number"]),
                    detail="Không tồn tại trong user_table.",
                )

    _finish_table_sheet(worksheet)
    for row in worksheet.iter_rows(min_row=2):
        severity_cell = row[0]
        color = REPORT_SEVERITY_FILLS.get(str(severity_cell.value))
        if color:
            severity_cell.fill = PatternFill("solid", fgColor=color)
            severity_cell.font = Font(bold=True)


def build_report_workbook(
    db_connection,
    *,
    project_id: int,
    phase_id: int,
    excel_path: Path,
    execute: bool,
    status: str,
    generated_at: dt.datetime,
    previews: Sequence[VariationPreview],
    delete_db_only: bool = False,
    deleted_count: int = 0,
    error_detail: str | None = None,
) -> Workbook:
    if not previews:
        raise ValueError("Cannot create a report without variation previews.")
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Tổng quan"
    detail_sheet = workbook.create_sheet("Chi tiết")
    validation_sheet = workbook.create_sheet("Kiểm tra")
    _append_summary_sheet(
        summary_sheet,
        db_connection=db_connection,
        project_id=project_id,
        phase_id=phase_id,
        excel_path=excel_path,
        execute=execute,
        status=status,
        generated_at=generated_at,
        previews=previews,
        delete_db_only=delete_db_only,
        deleted_count=deleted_count,
        error_detail=error_detail,
    )
    _append_detail_sheet(
        detail_sheet,
        project_id=project_id,
        phase_id=phase_id,
        previews=previews,
    )
    _append_validation_sheet(validation_sheet, previews)
    return workbook


def save_report_workbook(workbook: Workbook, report_path: Path) -> None:
    """Save through a sibling temporary file, then atomically replace output."""
    report_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{report_path.stem}.",
            suffix=".xlsx",
            dir=report_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, report_path)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def export_report(
    db_connection,
    *,
    report_path: Path,
    project_id: int,
    phase_id: int,
    excel_path: Path,
    execute: bool,
    status: str,
    generated_at: dt.datetime,
    previews: Sequence[VariationPreview],
    delete_db_only: bool = False,
    deleted_count: int = 0,
    error_detail: str | None = None,
) -> None:
    workbook = build_report_workbook(
        db_connection,
        project_id=project_id,
        phase_id=phase_id,
        excel_path=excel_path,
        execute=execute,
        status=status,
        generated_at=generated_at,
        previews=previews,
        delete_db_only=delete_db_only,
        deleted_count=deleted_count,
        error_detail=error_detail,
    )
    try:
        save_report_workbook(workbook, report_path)
    finally:
        workbook.close()


def plan_fingerprint(plan: Sequence[PlanRow]) -> tuple[tuple[object, ...], ...]:
    return tuple(
        (row.excel_row, row.se_parameter_id, row.old_values, row.new_values, row.action)
        for row in plan
    )


def acquire_lock(cursor, database_name: str) -> str:
    name = f"spdm_project_parameter_excel:{database_name}"[:64]
    cursor.execute("SELECT GET_LOCK(%s, %s)", (name, ADVISORY_LOCK_TIMEOUT_SECONDS))
    result = cursor.fetchone()
    if not result or result[0] != 1:
        raise RuntimeError(f"Could not acquire maintenance lock {name!r}.")
    return name


def apply_plan(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    plan: Sequence[PlanRow],
) -> dict[str, int]:
    counts = {"INSERT": 0, "UPDATE": 0, "UNCHANGED": 0}
    columns = ", ".join(f"`{column}`" for column in DATA_COLUMNS)
    insert_sql = f"""
        INSERT INTO project_parameter
          (project_id, se_parameter_id, phase_id, variation_id, {columns})
        VALUES (%s, %s, %s, %s, {make_placeholders(DATA_COLUMNS)})
    """
    assignments = ", ".join(f"`{column}` = %s" for column in DATA_COLUMNS)
    update_sql = f"""
        UPDATE project_parameter SET {assignments}
        WHERE project_id = %s AND se_parameter_id = %s
          AND phase_id = %s AND variation_id = %s
    """
    for row in plan:
        if row.action == "UNCHANGED":
            counts["UNCHANGED"] += 1
            continue
        new_values = tuple(value for _, value in row.new_values)
        if row.action == "INSERT":
            params = (
                project_id,
                row.se_parameter_id,
                phase_id,
                variation_id,
                *new_values,
            )
            cursor.execute(insert_sql, params)
        elif row.action == "UPDATE":
            params = (
                *new_values,
                project_id,
                row.se_parameter_id,
                phase_id,
                variation_id,
            )
            cursor.execute(update_sql, params)
        else:
            raise RuntimeError(f"Unexpected action: {row.action!r}")
        allowed_rowcounts = {1} if row.action == "INSERT" else {0, 1}
        if cursor.rowcount not in allowed_rowcounts:
            raise RuntimeError(
                f"Unexpected {row.action} count for se_parameter_id "
                f"{row.se_parameter_id}: {cursor.rowcount}; rolling back."
            )
        counts[row.action] += 1
    return counts


def delete_database_only_rows(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    ids: Sequence[int],
) -> int:
    """Delete the exact DB_ONLY rows and verify that none remain."""
    if not ids:
        return 0
    placeholders = make_placeholders(ids)
    params = (project_id, phase_id, variation_id, *ids)
    cursor.execute(
        f"""
        DELETE FROM project_parameter
        WHERE project_id = %s AND phase_id = %s AND variation_id = %s
          AND se_parameter_id IN ({placeholders})
        """,
        params,
    )
    if cursor.rowcount != len(ids):
        raise RuntimeError(
            "Unexpected DELETE count for DB_ONLY project_parameter rows: "
            f"expected={len(ids)}, actual={cursor.rowcount}; rolling back."
        )
    cursor.execute(
        f"""
        SELECT COUNT(*)
        FROM project_parameter
        WHERE project_id = %s AND phase_id = %s AND variation_id = %s
          AND se_parameter_id IN ({placeholders})
        """,
        params,
    )
    remaining = cursor.fetchone()
    if not remaining or int(remaining[0]) != 0:
        raise RuntimeError(
            "DB_ONLY deletion verification failed; transaction will roll back."
        )
    return len(ids)


def verify_rows(
    cursor,
    project_id: int,
    phase_id: int,
    variation_id: int,
    rows: Sequence[ExcelParameterRow],
) -> None:
    ids = [row.se_parameter_id for row in rows]
    actual = fetch_existing_rows(cursor, project_id, phase_id, variation_id, ids)
    expected = {row.se_parameter_id: dict(row.values) for row in rows}
    mismatches = {
        item_id: (values, actual.get(item_id))
        for item_id, values in expected.items()
        if actual.get(item_id) != values
    }
    if mismatches:
        raise RuntimeError(
            f"Final verification failed; sample={dict(list(mismatches.items())[:10])}. "
            "Transaction will roll back."
        )


def upsert_project_parameters(
    project_id: object,
    phase_id: object,
    excel_path: Path | str,
    *,
    sheet_name: str = RUN_SHEET_NAME,
    execute: bool = False,
    assume_yes: bool = False,
    delete_db_only: bool = False,
    report_path: Path | str | None = None,
    db_connection=connection,
    excel_reader=None,
) -> list[PlanRow]:
    project_id = normalize_positive_id(project_id, label="project_id")
    phase_id = normalize_positive_id(phase_id, label="phase_id")
    excel_path = Path(excel_path).expanduser().resolve()
    generated_at = dt.datetime.now().replace(microsecond=0)
    resolved_report_path = resolve_report_path(
        excel_path,
        project_id,
        phase_id,
        report_path,
        generated_at=generated_at,
    )
    source_read_results = (excel_reader or read_excel_variations)(
        excel_path, sheet_name
    )
    read_results: list[ExcelReadResult] = []
    previews: list[VariationPreview] = []
    with db_connection.cursor() as cursor:
        for source_result in source_read_results:
            read_results.append(
                ExcelReadResult(
                    source_result.variation_name,
                    resolve_approval_status_rows(cursor, source_result.rows),
                    source_result.header_row,
                    source_result.sheet_name,
                )
            )
        variation_ids_by_name = resolve_variation_ids(
            cursor, [result.variation_name for result in read_results]
        )
        for read_result in read_results:
            variation_id = variation_ids_by_name[read_result.variation_name]
            target, comparison, plan = build_plan(
                cursor,
                project_id,
                phase_id,
                variation_id,
                read_result.rows,
            )
            if normalize_variation_name(
                target["variation"]
            ) != normalize_variation_name(read_result.variation_name):
                raise ValueError(
                    f"variation_id {variation_id} is "
                    f"{target['variation']!r} in DB, but its Excel block is "
                    f"{read_result.variation_name!r}."
                )
            database_only_rows = fetch_database_only_rows(
                cursor,
                project_id,
                phase_id,
                variation_id,
                comparison.db_only,
            )
            previews.append(
                VariationPreview(
                    variation_id=variation_id,
                    read_result=read_result,
                    target=target,
                    comparison=comparison,
                    plan=tuple(plan),
                    database_only_rows=database_only_rows,
                )
            )

    for preview in previews:
        print("\n" + "=" * 80)
        print_plan(
            db_connection,
            project_id=project_id,
            phase_id=phase_id,
            variation_id=preview.variation_id,
            excel_path=excel_path,
            sheet_name=sheet_name,
            read_result=preview.read_result,
            target=preview.target,
            comparison=preview.comparison,
            plan=preview.plan,
        )
    has_reference_errors = any(
        preview.comparison.missing_master
        or preview.comparison.invalid_statuses
        or preview.comparison.missing_employees
        for preview in previews
    )
    delete_candidate_count = sum(
        len(preview.database_only_rows) for preview in previews
    )
    print(
        "DB_ONLY delete candidates: "
        f"{delete_candidate_count}; deletion enabled={delete_db_only}"
    )
    initial_status = (
        "BLOCKED_VALIDATION"
        if has_reference_errors
        else "READY_FOR_CONFIRMATION"
        if execute
        else "DRY_RUN"
    )
    export_report(
        db_connection,
        report_path=resolved_report_path,
        project_id=project_id,
        phase_id=phase_id,
        excel_path=excel_path,
        execute=execute,
        status=initial_status,
        generated_at=generated_at,
        previews=previews,
        delete_db_only=delete_db_only,
    )
    print(f"\nExcel report: {resolved_report_path}")

    if has_reference_errors:
        print(
            "\nBLOCKED: invalid foreign-key values were found. "
            "Dry-run completed, but execute mode is not allowed."
        )
    if not execute:
        print("\nDRY RUN: no data was changed. Set RUN_EXECUTE=True to apply.")
        return [row for preview in previews for row in preview.plan]
    if has_reference_errors:
        raise ValueError(
            "Cannot execute while master se_parameter, approval_status, or employee "
            "references are missing."
        )

    database_name = db_connection.settings_dict.get("NAME") or "<unknown>"
    variation_id_text = ",".join(str(preview.variation_id) for preview in previews)
    operation_name = (
        "UPSERT_AND_DELETE_DB_ONLY"
        if delete_db_only and delete_candidate_count
        else "UPSERT_PROJECT_PARAMETER"
    )
    confirmation = (
        f"{operation_name} {database_name}:"
        f"{project_id}@{phase_id}@[{variation_id_text}]"
    )
    if delete_db_only and delete_candidate_count:
        confirmation += f" DELETE={delete_candidate_count}"
    if not assume_yes:
        answer = input(f"\nType exactly '{confirmation}' to continue: ").strip()
        if answer != confirmation:
            export_report(
                db_connection,
                report_path=resolved_report_path,
                project_id=project_id,
                phase_id=phase_id,
                excel_path=excel_path,
                execute=execute,
                status="CANCELLED",
                generated_at=generated_at,
                previews=previews,
                delete_db_only=delete_db_only,
            )
            print("Cancelled. No data was changed.")
            return [row for preview in previews for row in preview.plan]

    lock_name: str | None = None
    total = {"INSERT": 0, "UPDATE": 0, "UNCHANGED": 0}
    deleted_total = 0
    execution_error: Exception | None = None
    execution_traceback = None
    try:
        with db_connection.cursor() as cursor:
            lock_name = acquire_lock(cursor, str(database_name))
        with transaction.atomic(using=db_connection.alias):
            with db_connection.cursor() as cursor:
                current_variation_ids = resolve_variation_ids(
                    cursor, [result.variation_name for result in read_results]
                )
                if current_variation_ids != variation_ids_by_name:
                    raise RuntimeError(
                        "Variation mapping changed after preview; transaction will "
                        "roll back."
                    )
                current_read_results = [
                    ExcelReadResult(
                        source_result.variation_name,
                        resolve_approval_status_rows(cursor, source_result.rows),
                        source_result.header_row,
                        source_result.sheet_name,
                    )
                    for source_result in source_read_results
                ]
                if current_read_results != read_results:
                    raise RuntimeError(
                        "Approval-status mapping changed after preview; transaction "
                        "will roll back."
                    )
                for preview in previews:
                    _, current_comparison, current_plan = build_plan(
                        cursor,
                        project_id,
                        phase_id,
                        preview.variation_id,
                        preview.read_result.rows,
                        lock_rows=True,
                    )
                    if current_comparison != preview.comparison or plan_fingerprint(
                        current_plan
                    ) != plan_fingerprint(preview.plan):
                        raise RuntimeError(
                            "Database changed after preview; transaction will roll back."
                        )
                    actual = apply_plan(
                        cursor,
                        project_id,
                        phase_id,
                        preview.variation_id,
                        current_plan,
                    )
                    for action in total:
                        total[action] += actual[action]
                    verify_rows(
                        cursor,
                        project_id,
                        phase_id,
                        preview.variation_id,
                        preview.read_result.rows,
                    )
                    if delete_db_only:
                        deleted_total += delete_database_only_rows(
                            cursor,
                            project_id,
                            phase_id,
                            preview.variation_id,
                            preview.comparison.db_only,
                        )
    except Exception as exc:
        execution_error = exc
        execution_traceback = exc.__traceback__
    finally:
        if lock_name is not None:
            try:
                with db_connection.cursor() as cursor:
                    cursor.execute("SELECT RELEASE_LOCK(%s)", (lock_name,))
            except Exception as exc:
                if execution_error is None:
                    execution_error = exc
                    execution_traceback = exc.__traceback__
                else:
                    print(f"WARNING: maintenance lock release failed: {exc}")

    if execution_error is not None:
        try:
            export_report(
                db_connection,
                report_path=resolved_report_path,
                project_id=project_id,
                phase_id=phase_id,
                excel_path=excel_path,
                execute=execute,
                status="FAILED",
                generated_at=generated_at,
                previews=previews,
                delete_db_only=delete_db_only,
                deleted_count=0,
                error_detail=format_report_error(execution_error),
            )
        except Exception as report_error:
            print(
                "WARNING: could not update the Excel report with FAILED status: "
                f"{report_error}"
            )
        raise execution_error.with_traceback(execution_traceback)

    try:
        export_report(
            db_connection,
            report_path=resolved_report_path,
            project_id=project_id,
            phase_id=phase_id,
            excel_path=excel_path,
            execute=execute,
            status="SUCCESS",
            generated_at=generated_at,
            previews=previews,
            delete_db_only=delete_db_only,
            deleted_count=deleted_total,
        )
    except Exception as exc:
        raise RuntimeError(
            "Database update committed, but the final SUCCESS report could not be "
            f"saved to {resolved_report_path}: {exc}"
        ) from exc
    print("\nproject_parameter upsert completed successfully:")
    print(
        f"  INSERT={total['INSERT']}, UPDATE={total['UPDATE']}, "
        f"UNCHANGED={total['UNCHANGED']}, DELETE={deleted_total}"
    )
    print(f"Final Excel report: {resolved_report_path}")
    return [row for preview in previews for row in preview.plan]


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Sync project_parameter from Excel; blank mapped cells become SQL NULL "
            "and DB-only deletion is opt-in."
        )
    )
    parser.add_argument("--project-id", type=int, default=RUN_PROJECT_ID)
    parser.add_argument("--phase-id", type=int, default=RUN_PHASE_ID)
    parser.add_argument("--excel", type=Path, default=RUN_EXCEL_PATH)
    parser.add_argument("--sheet", default=RUN_SHEET_NAME)
    parser.add_argument(
        "--report",
        type=Path,
        default=RUN_REPORT_PATH,
        help="Optional .xlsx report path; defaults next to the source workbook.",
    )
    parser.add_argument("--execute", action="store_true")
    parser.add_argument(
        "--delete-db-only",
        action="store_true",
        help=(
            "Delete project_parameter rows that exist in DB but are absent from "
            "Excel; only takes effect with --execute/RUN_EXECUTE."
        ),
    )
    parser.add_argument("--yes", action="store_true")
    return parser.parse_args(argv)


def main() -> None:
    args = parse_args()
    execute = args.execute or RUN_EXECUTE
    assume_yes = args.yes or RUN_ASSUME_YES
    delete_db_only = args.delete_db_only or RUN_DELETE_DB_ONLY
    if assume_yes and not execute:
        raise ValueError("RUN_ASSUME_YES/--yes requires execute mode.")
    upsert_project_parameters(
        args.project_id,
        args.phase_id,
        args.excel,
        sheet_name=args.sheet,
        execute=execute,
        assume_yes=assume_yes,
        delete_db_only=delete_db_only,
        report_path=args.report,
    )


if __name__ == "__main__":
    main()
