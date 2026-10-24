"""Thay toàn bộ R-list của nhiều project/phase bằng danh sách R từ nguồn.

Mỗi đích giữ nguyên tập variation hiện có. Script xóa toàn bộ dòng R của
project/phase đích rồi thêm mọi r_parameter_id nguồn vào từng variation đó.
Mặc định chỉ preview và xuất Excel; thực thi cần RUN_EXECUTE=True cùng xác nhận.

Các bảng được xóa theo scope đích:
* detail_data_rlist
* temp_data
* project_summary với category='R'
* project_r_parameter

Hai bảng con có FK ON DELETE CASCADE từ project_r_parameter cũng được snapshot
trước khi xóa: project_parameter_attachments và
project_r_parameter_manager_approval_history. Script không xóa file/S3; nếu có
attachment thì mặc định chặn thực thi trừ khi RUN_ALLOW_EXTERNAL_ORPHANS=True.

setup_r_relation thuộc setup dùng chung nên không bị xóa. Script chỉ thêm quan
hệ (setup đích, r_parameter_id nguồn, usecase_id=999999) còn thiếu.
"""

from __future__ import annotations

import tempfile
from datetime import datetime
from pathlib import Path


# Cấu hình chạy bằng PyCharm. Luôn xem PREVIEW trước khi bật RUN_EXECUTE.
RUN_SOURCE = (28, 27)
RUN_DESTINATIONS = [
    (1028, 27),
    (1029, 27),
    (1030, 27),
    (1031, 27),
    (1032, 27),
    (1033, 27),
    (1034, 27),
]
RUN_EXECUTE = False
RUN_ASSUME_YES = False
RUN_ALLOW_EXTERNAL_ORPHANS = False
RUN_REPORT_DIRECTORY = Path(__file__).resolve().parent / "output"
RUN_REPORT_ROW_LIMIT = 100_000

DEFAULT_USECASE_ID = 999999
DELETE_TABLES = (
    "detail_data_rlist",
    "temp_data",
    "project_summary",
    "project_r_parameter",
)
CASCADE_TABLES = (
    "project_parameter_attachments",
    "project_r_parameter_manager_approval_history",
)
EXPECTED_CASCADE_RELATIONS = {
    "project_parameter_attachments": "CASCADE",
    "project_r_parameter_manager_approval_history": "CASCADE",
}


def normalize_scope(value):
    if not isinstance(value, tuple) or len(value) != 2:
        raise ValueError("Mỗi scope phải là tuple (project_id, phase_id).")
    result = []
    for label, raw in zip(("project_id", "phase_id"), value):
        if isinstance(raw, bool) or not isinstance(raw, (int, str)):
            raise ValueError(f"{label} phải là số nguyên dương: {raw!r}")
        try:
            number = int(raw)
        except ValueError as exc:
            raise ValueError(f"{label} không hợp lệ: {raw!r}") from exc
        if number <= 0:
            raise ValueError(f"{label} phải lớn hơn 0: {number}")
        result.append(number)
    return tuple(result)


def fetch_rows(cursor, table, scope, *, lock=False):
    """Snapshot đầy đủ từ một bảng cố định, giới hạn theo project/phase."""
    if table not in {*DELETE_TABLES, *CASCADE_TABLES}:
        raise ValueError(f"Bảng chưa được allowlist: {table}")
    category = " AND category = 'R'" if table == "project_summary" else ""
    lock_clause = " FOR UPDATE" if lock else ""
    cursor.execute(
        f"SELECT * FROM {table} WHERE project_id = %s AND phase_id = %s"
        f"{category}{lock_clause}",
        scope,
    )
    columns = tuple(item[0] for item in cursor.description)
    return columns, tuple(sorted(cursor.fetchall(), key=repr))


def validate_cascade_schema(cursor):
    """Chặn chạy nếu quan hệ FK thực tế khác giả định cascade đã review."""
    cursor.execute(
        """
        SELECT kcu.TABLE_NAME, rc.DELETE_RULE
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        JOIN information_schema.REFERENTIAL_CONSTRAINTS AS rc
          ON rc.CONSTRAINT_SCHEMA = kcu.CONSTRAINT_SCHEMA
         AND rc.CONSTRAINT_NAME = kcu.CONSTRAINT_NAME
         AND rc.TABLE_NAME = kcu.TABLE_NAME
        WHERE kcu.CONSTRAINT_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_NAME = 'project_r_parameter'
        GROUP BY kcu.TABLE_NAME, rc.DELETE_RULE
        ORDER BY kcu.TABLE_NAME
        """
    )
    actual = dict(cursor.fetchall())
    if actual != EXPECTED_CASCADE_RELATIONS:
        raise RuntimeError(
            "FK con của project_r_parameter khác schema đã review: "
            f"expected={EXPECTED_CASCADE_RELATIONS}, actual={actual}"
        )


def build_plan(cursor, source, destinations, *, lock=False, row_limit=RUN_REPORT_ROW_LIMIT):
    suffix = " FOR UPDATE" if lock else ""
    scopes = sorted(set([source, *destinations]))
    for table, ids in (
        ("project_info", sorted({scope[0] for scope in scopes})),
        ("phase", sorted({scope[1] for scope in scopes})),
    ):
        for record_id in ids:
            cursor.execute(f"SELECT id FROM {table} WHERE id = %s{suffix}", [record_id])
            if cursor.fetchone() is None:
                raise ValueError(f"Không tồn tại {table}.id={record_id}.")

    validate_cascade_schema(cursor)
    cursor.execute(
        "SELECT DISTINCT r_parameter_id FROM project_r_parameter "
        f"WHERE project_id = %s AND phase_id = %s ORDER BY r_parameter_id{suffix}",
        source,
    )
    source_ids = tuple(row[0] for row in cursor.fetchall())
    if not source_ids:
        raise ValueError(f"Nguồn {source} không có r_parameter_id; từ chối xóa đích.")

    snapshots = {}
    variations = {}
    setup_ids = set()
    for destination in destinations:
        cursor.execute(
            "SELECT variation_id FROM project_parameter "
            f"WHERE project_id = %s AND phase_id = %s{suffix}",
            destination,
        )
        variation_ids = {row[0] for row in cursor.fetchall()}
        cursor.execute(
            "SELECT variation_id FROM project_r_parameter "
            f"WHERE project_id = %s AND phase_id = %s{suffix}",
            destination,
        )
        variation_ids.update(row[0] for row in cursor.fetchall())
        if not variation_ids:
            raise ValueError(f"Đích {destination} chưa có variation; không tự tạo variation.")
        variations[destination] = tuple(sorted(variation_ids))

        for table in (*DELETE_TABLES, *CASCADE_TABLES):
            columns, rows = fetch_rows(cursor, table, destination, lock=lock)
            if len(rows) > row_limit:
                raise ValueError(
                    f"{table} tại {destination} có {len(rows)} dòng, vượt "
                    f"RUN_REPORT_ROW_LIMIT={row_limit}; chưa được phép xóa."
                )
            snapshots[(destination, table)] = (columns, rows)

        cursor.execute("SELECT setup_id FROM project_info WHERE id = %s", [destination[0]])
        setup_ids.add(cursor.fetchone()[0])

    project_rows = {
        destination: tuple(
            (destination[0], parameter_id, destination[1], variation_id)
            for parameter_id in source_ids
            for variation_id in variations[destination]
        )
        for destination in destinations
    }
    setup_decisions = {}
    for setup_id in sorted(setup_ids):
        cursor.execute(f"SELECT id FROM setup WHERE id = %s{suffix}", [setup_id])
        if cursor.fetchone() is None:
            raise ValueError(f"Không tồn tại setup.id={setup_id}.")
        cursor.execute(
            "SELECT r_parameter_id, usecase_id FROM setup_r_relation "
            f"WHERE setup_id = %s{suffix}",
            [setup_id],
        )
        existing = set(cursor.fetchall())
        for parameter_id in source_ids:
            key = (setup_id, parameter_id, DEFAULT_USECASE_ID)
            setup_decisions[key] = key[1:] not in existing
    if any(setup_decisions.values()):
        cursor.execute(f"SELECT id FROM usecase WHERE id = %s{suffix}", [DEFAULT_USECASE_ID])
        if cursor.fetchone() is None:
            raise ValueError(f"Thiếu usecase mặc định {DEFAULT_USECASE_ID}; không tự tạo master.")

    return {
        "source_ids": source_ids,
        "variations": variations,
        "snapshots": snapshots,
        "project_rows": project_rows,
        "setup_decisions": setup_decisions,
    }


def write_report(path, source, destinations, plan, database, host, status, error=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tong_hop"
    attachment_count = sum(
        len(plan["snapshots"][(scope, "project_parameter_attachments")][1])
        for scope in destinations
    )
    for row in (
        ("Thông tin", "Giá trị"),
        ("Database", database),
        ("Host", host),
        ("Nguồn", str(source)),
        ("Đích", str(destinations)),
        ("Thời điểm", datetime.now().astimezone().isoformat()),
        ("Trạng thái", status),
        ("Lỗi", error),
        ("Số R ID nguồn", len(plan["source_ids"])),
        ("Attachment metadata sẽ cascade", attachment_count),
        ("Lưu ý S3", "Script không xóa file/S3."),
    ):
        summary.append(row)
    summary.append([])
    summary.append([
        "Project đích", "Phase đích", "Variation", "Xóa project_r_parameter",
        "Xóa detail_data_rlist", "Xóa temp_data", "Xóa project_summary R",
        "Thêm project_r_parameter", "Thêm detail_data_rlist",
    ])
    for scope in destinations:
        def count(table):
            return len(plan["snapshots"][(scope, table)][1])

        insert_count = len(plan["project_rows"][scope])
        summary.append([
            *scope,
            ", ".join(map(str, plan["variations"][scope])),
            count("project_r_parameter"),
            count("detail_data_rlist"),
            count("temp_data"),
            count("project_summary"),
            insert_count,
            insert_count,
        ])

    sheets = [summary]
    for table in (*DELETE_TABLES, *CASCADE_TABLES):
        sheet = workbook.create_sheet(table[:31])
        first_columns = plan["snapshots"][(destinations[0], table)][0]
        sheet.append([*first_columns, "action", "result"])
        for scope in destinations:
            columns, rows = plan["snapshots"][(scope, table)]
            if columns != first_columns:
                raise RuntimeError(f"Metadata cột của {table} thay đổi trong lúc lập report.")
            action = "CASCADE_DELETE" if table in CASCADE_TABLES else "DELETE"
            result = "DELETED" if status == "COMMITTED" else status
            for row in rows:
                sheet.append([*row, action, result])
        sheets.append(sheet)

    for name in ("insert_project_r", "insert_detail_r"):
        sheet = workbook.create_sheet(name)
        sheet.append([
            "project_id", "r_parameter_id", "phase_id", "variation_id",
            "action", "result",
        ])
        for scope in destinations:
            for row in plan["project_rows"][scope]:
                sheet.append([*row, "INSERT", "INSERTED" if status == "COMMITTED" else status])
        sheets.append(sheet)

    setup_sheet = workbook.create_sheet("setup_r_relation")
    setup_sheet.append([
        "setup_id", "r_parameter_id", "usecase_id", "action", "result",
    ])
    for key, adding in plan["setup_decisions"].items():
        setup_sheet.append([
            *key,
            "INSERT" if adding else "SKIP_EXISTING",
            ("INSERTED" if status == "COMMITTED" else status) if adding else "SKIPPED",
        ])
    sheets.append(setup_sheet)

    if any(sheet.max_row > 1_048_576 for sheet in sheets):
        raise ValueError("Báo cáo vượt giới hạn dòng Excel; hãy chia nhỏ danh sách đích.")
    for sheet in sheets:
        header_row = 13 if sheet == summary else 1
        sheet.freeze_panes = "A14" if sheet == summary else "A2"
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row in ({1, 13} if sheet == summary else {1}):
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="24476A")
            sheet.row_dimensions[row[0].row].height = 30
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 24

    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as stream:
        temporary = Path(stream.name)
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
        workbook.close()
    print(f"Báo cáo Excel ({status}): {path}")


def delete_destination_rows(cursor, destination, snapshots):
    deleted = 0
    for table in DELETE_TABLES:
        category = " AND category = 'R'" if table == "project_summary" else ""
        cursor.execute(
            f"DELETE FROM {table} WHERE project_id = %s AND phase_id = %s{category}",
            destination,
        )
        expected = len(snapshots[(destination, table)][1])
        if cursor.rowcount != expected:
            raise RuntimeError(
                f"Số dòng xóa không khớp tại {table} {destination}: "
                f"expected={expected}, actual={cursor.rowcount}"
            )
        deleted += cursor.rowcount
    for table in CASCADE_TABLES:
        cursor.execute(
            f"SELECT COUNT(*) FROM {table} WHERE project_id = %s AND phase_id = %s",
            destination,
        )
        if cursor.fetchone()[0]:
            raise RuntimeError(f"Cascade vẫn còn dữ liệu tại {table} {destination}.")
    return deleted


def insert_replacement_rows(cursor, plan):
    inserted = 0
    for rows in plan["project_rows"].values():
        cursor.executemany(
            "INSERT INTO project_r_parameter "
            "(project_id, r_parameter_id, phase_id, variation_id) "
            "VALUES (%s, %s, %s, %s)",
            rows,
        )
        if cursor.rowcount != len(rows):
            raise RuntimeError("Số dòng insert project_r_parameter không khớp.")
        inserted += cursor.rowcount
        cursor.executemany(
            "INSERT INTO detail_data_rlist "
            "(project_id, r_parameter_id, phase_id, variation_id) "
            "VALUES (%s, %s, %s, %s)",
            rows,
        )
        if cursor.rowcount != len(rows):
            raise RuntimeError("Số dòng insert detail_data_rlist không khớp.")
        inserted += cursor.rowcount

    setup_rows = [key for key, adding in plan["setup_decisions"].items() if adding]
    if setup_rows:
        from utils.id_sequence import next_ids

        first_id = next_ids("setup_r_relation", n=len(setup_rows))
        values = [(first_id + index, *row) for index, row in enumerate(setup_rows)]
        cursor.executemany(
            "INSERT INTO setup_r_relation (id, setup_id, r_parameter_id, usecase_id) "
            "VALUES (%s, %s, %s, %s)",
            values,
        )
        if cursor.rowcount != len(values):
            raise RuntimeError("Số dòng insert setup_r_relation không khớp.")
        inserted += cursor.rowcount
    return inserted


def verify_result(cursor, destinations, plan):
    for destination in destinations:
        expected = set(plan["project_rows"][destination])
        for table in ("project_r_parameter", "detail_data_rlist"):
            cursor.execute(
                f"SELECT project_id, r_parameter_id, phase_id, variation_id FROM {table} "
                "WHERE project_id = %s AND phase_id = %s",
                destination,
            )
            rows = cursor.fetchall()
            actual = set(rows)
            if actual != expected or len(rows) != len(expected):
                raise RuntimeError(f"Kết quả {table} không đúng tại {destination}.")
        for table, category in (("temp_data", ""), ("project_summary", " AND category='R'")):
            cursor.execute(
                f"SELECT COUNT(*) FROM {table} WHERE project_id = %s AND phase_id = %s{category}",
                destination,
            )
            if cursor.fetchone()[0]:
                raise RuntimeError(f"Vẫn còn dữ liệu cũ tại {table} {destination}.")
    for key in plan["setup_decisions"]:
        cursor.execute(
            "SELECT COUNT(*) FROM setup_r_relation "
            "WHERE setup_id = %s AND r_parameter_id = %s AND usecase_id = %s",
            key,
        )
        if cursor.fetchone()[0] != 1:
            raise RuntimeError(f"Quan hệ setup không duy nhất hoặc bị thiếu: {key}")


def replace_r_parameters(
    source,
    destinations,
    *,
    execute=False,
    assume_yes=False,
    allow_external_orphans=False,
    report_directory=RUN_REPORT_DIRECTORY,
    report_row_limit=RUN_REPORT_ROW_LIMIT,
):
    source = normalize_scope(source)
    if not isinstance(destinations, list) or not destinations:
        raise ValueError("destinations phải là list không rỗng các tuple.")
    destinations = list(dict.fromkeys(normalize_scope(item) for item in destinations))
    if source in destinations:
        raise ValueError("Nguồn không được đồng thời là đích.")
    if not all(isinstance(value, bool) for value in (execute, assume_yes, allow_external_orphans)):
        raise ValueError("Các cờ execute/assume_yes/allow_external_orphans phải là bool.")
    if isinstance(report_row_limit, bool) or not isinstance(report_row_limit, int) or report_row_limit <= 0:
        raise ValueError("report_row_limit phải là số nguyên dương.")

    from django.db import connection, transaction

    if connection.vendor != "mysql":
        raise ValueError("Script chỉ hỗ trợ MySQL.")
    settings = connection.settings_dict
    database = str(settings.get("NAME") or "<unknown>")
    host = str(settings.get("HOST") or "localhost")
    with connection.cursor() as cursor:
        preview = build_plan(
            cursor, source, destinations, row_limit=report_row_limit,
        )

    report_directory = Path(report_directory).expanduser().resolve()
    report_directory.mkdir(parents=True, exist_ok=True)
    report_path = report_directory / (
        f"Project_copy_r_parameter_v2_p{source[0]}_ph{source[1]}_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
    )

    def report(status, error=""):
        write_report(
            report_path, source, destinations, preview, database, host, status, error,
        )

    report("PREVIEW")
    print(f"Database: {database}; host: {host}; nguồn: {source}; đích: {destinations}")
    print(f"Nguồn có {len(preview['source_ids'])} r_parameter_id: {preview['source_ids']}")
    for destination in destinations:
        counts = {
            table: len(preview["snapshots"][(destination, table)][1])
            for table in (*DELETE_TABLES, *CASCADE_TABLES)
        }
        print(
            f"Đích {destination}; variations={preview['variations'][destination]}; "
            f"xóa={counts}; thêm mỗi bảng R={len(preview['project_rows'][destination])}"
        )

    attachment_count = sum(
        len(preview["snapshots"][(scope, "project_parameter_attachments")][1])
        for scope in destinations
    )
    if attachment_count and execute and not allow_external_orphans:
        raise RuntimeError(
            f"Có {attachment_count} attachment metadata sẽ bị cascade nhưng file/S3 "
            "không được xóa. Xem report; chỉ bật RUN_ALLOW_EXTERNAL_ORPHANS=True "
            "khi đã chấp nhận orphan bên ngoài."
        )
    if not execute:
        print("DRY-RUN: không ghi dữ liệu.")
        return 0

    confirmation = f"REPLACE R {database} TARGETS {destinations}"
    if assume_yes:
        print(f"Bỏ qua xác nhận do assume_yes=True: {confirmation}")
    elif input(f"Nhập chính xác '{confirmation}' để thực thi: ") != confirmation:
        report("CANCELLED")
        print("Đã hủy; không ghi dữ liệu.")
        return 0

    deleted = inserted = 0
    try:
        with transaction.atomic(), connection.cursor() as cursor:
            locked = build_plan(
                cursor, source, destinations, lock=True, row_limit=report_row_limit,
            )
            if locked != preview:
                raise RuntimeError("Dữ liệu đã thay đổi sau preview; hãy chạy lại.")
            for destination in destinations:
                deleted += delete_destination_rows(
                    cursor, destination, preview["snapshots"],
                )
            inserted = insert_replacement_rows(cursor, preview)
            verify_result(cursor, destinations, preview)
    except Exception as exc:
        try:
            report("TRANSACTION_ERROR", str(exc))
        except Exception as report_error:
            print(f"Không cập nhật được report lỗi: {report_error}; preview: {report_path}")
        raise

    print(f"Đã commit: xóa {deleted} dòng trực tiếp, thêm {inserted} dòng.")
    try:
        report("COMMITTED")
    except Exception as exc:
        raise RuntimeError(
            f"DB ĐÃ COMMIT nhưng không cập nhật được Excel; file vẫn là preview: {report_path}"
        ) from exc
    return deleted + inserted


if __name__ == "__main__":
    import os
    import sys

    import django

    project_root = Path(__file__).resolve().parents[2]
    if str(project_root) not in sys.path:
        sys.path.insert(0, str(project_root))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()
    replace_r_parameters(
        RUN_SOURCE,
        RUN_DESTINATIONS,
        execute=RUN_EXECUTE,
        assume_yes=RUN_ASSUME_YES,
        allow_external_orphans=RUN_ALLOW_EXTERNAL_ORPHANS,
        report_directory=RUN_REPORT_DIRECTORY,
        report_row_limit=RUN_REPORT_ROW_LIMIT,
    )
