"""Preview or update ``se_parameter.wp_id`` from an Excel master list safely.

The configured workbook uses column B for ``se_parameter.id`` and column C
for the human-readable ``wp.wp`` value.  The tool resolves each WP name to its
unique ``wp.id`` before updating ``se_parameter.wp_id``.  Every run writes an
Excel audit report, including dry runs, blocked runs, and successful updates.

For PyCharm's Run button, edit the RUN_* constants below.  Mutation is disabled
by default and must be enabled explicitly with RUN_EXECUTE or ``--execute``.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402

# PyCharm Run-button configuration. Always preview before enabling execution.
RUN_EXCEL_PATH = Path("/Users/TrongKieu/Downloads/MASTER_SEリスト項目_UPDATE.xlsx")
RUN_SHEET_NAME = "追加項目"
RUN_REPORT_PATH: Path | None = None
RUN_EXECUTE = False
RUN_ASSUME_YES = False

DEFAULT_REPORT_DIRECTORY = PROJECT_ROOT / "outputs" / "se_master_update_reports"
ID_COLUMN = 2
WP_COLUMN = 3
EXPECTED_ID_HEADER = "ID"
EXPECTED_WP_HEADER = "WP"
PREVIEW_ROW_LIMIT = 50


@dataclass
class ReportRow:
    excel_row: int
    se_parameter_id: int | None
    requested_wp: str | None
    resolved_wp_id: int | None = None
    parameter_name_1: str | None = None
    parameter_name_2: str | None = None
    old_wp_id: int | None = None
    old_wp: str | None = None
    action: str = "PENDING"
    result: str = "PENDING"
    message: str = ""


@dataclass(frozen=True)
class RunResult:
    status: str
    report_path: Path
    updated_count: int
    rows: tuple[ReportRow, ...]


class ToolRunError(RuntimeError):
    """Raised after a failure report has been written."""

    def __init__(self, message: str, report_path: Path) -> None:
        super().__init__(message)
        self.report_path = report_path


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip().casefold()


def normalize_wp_name(value: object) -> str | None:
    """Normalize a WP name for matching without discarding meaningful text."""
    if value is None:
        return None
    normalized = " ".join(unicodedata.normalize("NFKC", str(value)).split())
    return normalized or None


def normalize_positive_id(value: object, *, label: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{label} must be a positive integer, not boolean.")
    if isinstance(value, float) and not value.is_integer():
        raise ValueError(f"{label} must be an integer: {value!r}")
    try:
        normalized = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {label}: {value!r}") from exc
    if normalized <= 0:
        raise ValueError(f"{label} must be greater than zero: {normalized}")
    return normalized


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Cannot create SQL placeholders for an empty sequence.")
    return ", ".join(["%s"] * len(values))


def safe_excel_value(value: object) -> object:
    """Prevent text values in the generated report from becoming formulas."""
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.astimezone().replace(tzinfo=None)
        return value
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def make_report_path(report_path: Path | str | None = None) -> Path:
    if report_path is not None:
        path = Path(report_path).expanduser().resolve()
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Report path must use the .xlsx extension.")
        path.parent.mkdir(parents=True, exist_ok=True)
        return path

    DEFAULT_REPORT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
    return DEFAULT_REPORT_DIRECTORY / f"se_master_update_{timestamp}.xlsx"


def read_excel_rows(
    excel_path: Path | str,
    sheet_name: str,
) -> list[ReportRow]:
    """Read and validate column B IDs and column C WP names."""
    path = Path(excel_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Excel file was not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Excel input must be an .xlsx or .xlsm file.")

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet {sheet_name!r} was not found; available={workbook.sheetnames}"
            )
        worksheet = workbook[sheet_name]
        id_header = worksheet.cell(1, ID_COLUMN).value
        wp_header = worksheet.cell(1, WP_COLUMN).value
        if normalize_header(id_header) != normalize_header(EXPECTED_ID_HEADER):
            raise ValueError(
                f"Expected {EXPECTED_ID_HEADER!r} in B1, found {id_header!r}."
            )
        if normalize_header(wp_header) != normalize_header(EXPECTED_WP_HEADER):
            raise ValueError(
                f"Expected {EXPECTED_WP_HEADER!r} in C1, found {wp_header!r}."
            )

        rows: list[ReportRow] = []
        first_by_id: dict[int, ReportRow] = {}
        for row_number in range(2, worksheet.max_row + 1):
            id_cell = worksheet.cell(row_number, ID_COLUMN)
            wp_cell = worksheet.cell(row_number, WP_COLUMN)
            raw_id = id_cell.value
            raw_wp = wp_cell.value
            if raw_id in (None, "") and raw_wp in (None, ""):
                continue

            report_row = ReportRow(
                excel_row=row_number,
                se_parameter_id=None,
                requested_wp=normalize_wp_name(raw_wp),
            )
            rows.append(report_row)
            try:
                if id_cell.data_type == "f":
                    raise ValueError("Formula is not allowed in the ID column.")
                if wp_cell.data_type == "f":
                    raise ValueError("Formula is not allowed in the WP column.")
                if raw_id in (None, ""):
                    raise ValueError("ID is blank while WP has a value.")
                report_row.se_parameter_id = normalize_positive_id(
                    raw_id, label=f"ID at Excel row {row_number}"
                )
                if report_row.requested_wp is None:
                    raise ValueError("WP is blank.")
            except ValueError as exc:
                report_row.action = "ERROR"
                report_row.result = "ERROR"
                report_row.message = str(exc)
                continue

            first_row = first_by_id.get(report_row.se_parameter_id)
            if first_row is None:
                first_by_id[report_row.se_parameter_id] = report_row
                continue
            if first_row.requested_wp == report_row.requested_wp:
                report_row.action = "DUPLICATE"
                report_row.result = "SKIPPED_DUPLICATE"
                report_row.message = (
                    f"Same ID and WP already appeared at Excel row "
                    f"{first_row.excel_row}."
                )
            else:
                message = (
                    f"Conflicting WP values for ID {report_row.se_parameter_id} "
                    f"at Excel rows {first_row.excel_row} and {row_number}."
                )
                first_row.action = "ERROR"
                first_row.result = "ERROR"
                first_row.message = message
                report_row.action = "ERROR"
                report_row.result = "ERROR"
                report_row.message = message

        if not rows:
            raise ValueError("No input rows were found in columns B and C.")
        return rows
    finally:
        workbook.close()


def database_identity(db_connection=connection) -> tuple[str, str, str]:
    settings = db_connection.settings_dict
    return (
        str(settings.get("NAME") or "<unknown>"),
        str(settings.get("HOST") or "localhost"),
        str(settings.get("PORT") or "3306"),
    )


def active_input_rows(rows: Sequence[ReportRow]) -> list[ReportRow]:
    return [row for row in rows if row.result == "PENDING"]


def fetch_se_parameter_rows(
    cursor: Any,
    se_parameter_ids: Sequence[int],
) -> dict[int, tuple[str | None, str | None, int | None, str | None]]:
    placeholders = make_placeholders(se_parameter_ids)
    cursor.execute(
        f"""
        SELECT sp.id, sp.parameter_name_1, sp.parameter_name_2, sp.wp_id, wp.wp
        FROM se_parameter AS sp
        LEFT JOIN wp ON wp.id = sp.wp_id
        WHERE sp.id IN ({placeholders})
        ORDER BY sp.id
        """,
        list(se_parameter_ids),
    )
    return {
        int(se_id): (parameter_name_1, parameter_name_2, wp_id, wp_name)
        for se_id, parameter_name_1, parameter_name_2, wp_id, wp_name in cursor.fetchall()
    }


def load_wp_catalog(cursor: Any) -> dict[str, list[tuple[int, str]]]:
    cursor.execute("SELECT id, wp FROM wp ORDER BY id")
    catalog: dict[str, list[tuple[int, str]]] = {}
    for raw_id, raw_name in cursor.fetchall():
        normalized_name = normalize_wp_name(raw_name)
        if normalized_name is None:
            continue
        catalog.setdefault(normalized_name, []).append((int(raw_id), str(raw_name)))
    return catalog


def build_plan(cursor: Any, rows: Sequence[ReportRow]) -> None:
    """Resolve input values and attach database state to every report row."""
    pending_rows = active_input_rows(rows)
    if not pending_rows:
        return
    requested_ids = [
        row.se_parameter_id for row in pending_rows if row.se_parameter_id is not None
    ]
    master_rows = fetch_se_parameter_rows(cursor, requested_ids)
    wp_catalog = load_wp_catalog(cursor)

    for row in pending_rows:
        assert row.se_parameter_id is not None
        master = master_rows.get(row.se_parameter_id)
        if master is None:
            row.action = "ERROR"
            row.result = "ERROR"
            row.message = "se_parameter.id was not found in the database."
            continue

        (
            row.parameter_name_1,
            row.parameter_name_2,
            row.old_wp_id,
            row.old_wp,
        ) = master
        wp_matches = wp_catalog.get(row.requested_wp or "", [])
        if not wp_matches:
            row.action = "ERROR"
            row.result = "ERROR"
            row.message = "WP name was not found in wp.wp."
            continue
        if len(wp_matches) > 1:
            row.action = "ERROR"
            row.result = "ERROR"
            row.message = (
                "WP name is ambiguous after normalization; matching wp.id values: "
                + ", ".join(str(wp_id) for wp_id, _ in wp_matches)
            )
            continue

        row.resolved_wp_id = wp_matches[0][0]
        if row.old_wp_id == row.resolved_wp_id:
            row.action = "NO_CHANGE"
            row.result = "NO_CHANGE"
            row.message = "se_parameter.wp_id already has the requested value."
        else:
            row.action = "UPDATE"
            row.result = "READY"
            row.message = (
                f"Will change wp_id from {row.old_wp_id!r} " f"to {row.resolved_wp_id}."
            )


def print_plan(
    rows: Sequence[ReportRow],
    *,
    database: str,
    host: str,
    port: str,
    excel_path: Path,
    sheet_name: str,
    report_path: Path,
    execute: bool,
) -> None:
    ready_rows = [row for row in rows if row.result == "READY"]
    error_rows = [row for row in rows if row.result == "ERROR"]
    no_change_rows = [row for row in rows if row.result == "NO_CHANGE"]
    print(f"Database: {database}")
    print(f"Host: {host}:{port}")
    print(f"Input Excel: {excel_path}")
    print(f"Input sheet: {sheet_name}")
    print("Target: se_parameter.wp_id (resolved from wp.wp)")
    print(f"Mode: {'EXECUTE' if execute else 'PREVIEW'}")
    print(f"Report: {report_path}")
    print(f"Rows to update: {len(ready_rows)}")
    print(f"Rows already correct: {len(no_change_rows)}")
    print(f"Rows with errors: {len(error_rows)}")
    print("Selected records:")
    display_rows = [
        row
        for row in rows
        if row.se_parameter_id is not None and row.result not in {"SKIPPED_DUPLICATE"}
    ]
    for row in display_rows[:PREVIEW_ROW_LIMIT]:
        print(
            f"- Excel row {row.excel_row}: id={row.se_parameter_id}; "
            f"old_wp={row.old_wp!r} ({row.old_wp_id!r}); "
            f"new_wp={row.requested_wp!r} ({row.resolved_wp_id!r}); "
            f"status={row.result}"
        )
    if len(display_rows) > PREVIEW_ROW_LIMIT:
        print(f"- ... {len(display_rows) - PREVIEW_ROW_LIMIT} more rows")


def confirmation_text(database: str, rows: Sequence[ReportRow]) -> str:
    ids = [
        str(row.se_parameter_id)
        for row in rows
        if row.result == "READY" and row.se_parameter_id is not None
    ]
    return f"UPDATE {database} SE_PARAMETER IDS {','.join(ids)}"


def confirm_execution(
    database: str,
    rows: Sequence[ReportRow],
    *,
    assume_yes: bool,
) -> bool:
    required = confirmation_text(database, rows)
    if assume_yes:
        print("Confirmation bypass explicitly enabled with --yes/RUN_ASSUME_YES.")
        return True
    print("Database mutation requires exact confirmation.")
    print(f"Type exactly: {required}")
    return input("> ").strip() == required


def execute_updates(
    rows: Sequence[ReportRow],
    *,
    db_connection=connection,
) -> int:
    update_rows = [row for row in rows if row.result == "READY"]
    if not update_rows:
        return 0

    ids = [
        int(row.se_parameter_id)
        for row in update_rows
        if row.se_parameter_id is not None
    ]
    if len(ids) != len(update_rows):
        raise RuntimeError("An update row is missing se_parameter_id.")
    placeholders = make_placeholders(ids)
    alias = getattr(db_connection, "alias", "default")

    with transaction.atomic(using=alias):
        with db_connection.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT id, wp_id
                FROM se_parameter
                WHERE id IN ({placeholders})
                ORDER BY id
                FOR UPDATE
                """,
                ids,
            )
            locked_state = {int(se_id): wp_id for se_id, wp_id in cursor.fetchall()}
            expected_ids = set(ids)
            if set(locked_state) != expected_ids:
                missing = sorted(expected_ids - set(locked_state))
                raise RuntimeError(
                    f"se_parameter rows disappeared before update: {missing}"
                )

            wp_ids = sorted(
                {
                    row.resolved_wp_id
                    for row in update_rows
                    if row.resolved_wp_id is not None
                }
            )
            wp_placeholders = make_placeholders(wp_ids)
            cursor.execute(
                f"SELECT id FROM wp WHERE id IN ({wp_placeholders}) FOR UPDATE",
                wp_ids,
            )
            existing_wp_ids = {int(item[0]) for item in cursor.fetchall()}
            if existing_wp_ids != set(wp_ids):
                missing_wp_ids = sorted(set(wp_ids) - existing_wp_ids)
                raise RuntimeError(
                    f"Resolved wp rows disappeared before update: {missing_wp_ids}"
                )

            for row in update_rows:
                assert row.se_parameter_id is not None
                assert row.resolved_wp_id is not None
                if locked_state[row.se_parameter_id] != row.old_wp_id:
                    raise RuntimeError(
                        f"se_parameter.id={row.se_parameter_id} changed after "
                        "preview; aborting the whole transaction."
                    )
                cursor.execute(
                    """
                    UPDATE se_parameter
                    SET wp_id = %s
                    WHERE id = %s AND wp_id <=> %s
                    """,
                    [row.resolved_wp_id, row.se_parameter_id, row.old_wp_id],
                )
                if cursor.rowcount != 1:
                    raise RuntimeError(
                        f"Unexpected update count for se_parameter.id="
                        f"{row.se_parameter_id}: {cursor.rowcount}"
                    )

            cursor.execute(
                f"SELECT id, wp_id FROM se_parameter "
                f"WHERE id IN ({placeholders}) ORDER BY id",
                ids,
            )
            actual = {int(se_id): wp_id for se_id, wp_id in cursor.fetchall()}
            expected = {row.se_parameter_id: row.resolved_wp_id for row in update_rows}
            if actual != expected:
                raise RuntimeError(
                    "Post-update verification failed; transaction will roll back. "
                    f"expected={expected}, actual={actual}"
                )
    return len(update_rows)


def _style_report_sheet(sheet: Any, *, header_row: int = 1) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="1F1F1F")
    thin_gray = Side(style="thin", color="D9E2F3")
    for cell in sheet[header_row]:
        if cell.value is not None:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=thin_gray)
    for column_index in range(1, sheet.max_column + 1):
        values = [
            "" if cell.value is None else str(cell.value)
            for cell in list(sheet.columns)[column_index - 1]
        ]
        width = min(max(max((len(value) for value in values), default=0) + 2, 10), 45)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.sheet_view.showGridLines = False


def write_report(
    report_path: Path,
    rows: Sequence[ReportRow],
    *,
    status: str,
    database: str,
    host: str,
    port: str,
    excel_path: Path,
    sheet_name: str,
    execute_requested: bool,
    assume_yes: bool,
    started_at: datetime,
    completed_at: datetime,
    error: str | None = None,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    results = workbook.create_sheet("Results")

    summary.append(["SE master WP update report"])
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color="1F1F1F")
    summary.append([])
    summary.append(["Item", "Value"])
    summary_items: list[tuple[str, object]] = [
        ("Status", status),
        ("Mode", "EXECUTE" if execute_requested else "PREVIEW"),
        ("Database", database),
        ("Host", f"{host}:{port}"),
        ("Target", "se_parameter.wp_id"),
        ("WP lookup", "wp.wp -> wp.id"),
        ("Input file", str(excel_path)),
        ("Input sheet", sheet_name),
        ("Report created", completed_at),
        ("Started", started_at),
        ("Non-interactive confirmation", assume_yes),
        ("Input rows", len(rows)),
        ("Rows ready/would update", sum(row.action == "UPDATE" for row in rows)),
        ("Rows updated", sum(row.result == "UPDATED" for row in rows)),
        ("Rows unchanged", sum(row.result == "NO_CHANGE" for row in rows)),
        (
            "Duplicate rows skipped",
            sum(row.result == "SKIPPED_DUPLICATE" for row in rows),
        ),
        ("Rows with errors", sum(row.result == "ERROR" for row in rows)),
        ("Error", error or ""),
    ]
    for item, value in summary_items:
        summary.append([item, safe_excel_value(value)])
    _style_report_sheet(summary, header_row=3)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 90
    for cell in (summary["B12"], summary["B13"]):
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    status_cell = summary["B4"]
    status_colors = {
        "COMMITTED": "C6EFCE",
        "PREVIEW": "FFF2CC",
        "NO_CHANGES": "D9EAD3",
        "READY_TO_EXECUTE": "FCE5CD",
        "CANCELLED": "E7E6E6",
        "BLOCKED": "F4CCCC",
        "ROLLED_BACK": "F4CCCC",
    }
    status_cell.fill = PatternFill("solid", fgColor=status_colors.get(status, "D9E2F3"))
    status_cell.font = Font(name="Arial", size=10, bold=True)

    result_headers = [
        "excel_row",
        "se_parameter_id",
        "requested_wp",
        "resolved_wp_id",
        "parameter_name_1",
        "parameter_name_2",
        "old_wp_id",
        "old_wp",
        "action",
        "result",
        "message",
    ]
    results.append(result_headers)
    for row in rows:
        results.append(
            [
                safe_excel_value(row.excel_row),
                safe_excel_value(row.se_parameter_id),
                safe_excel_value(row.requested_wp),
                safe_excel_value(row.resolved_wp_id),
                safe_excel_value(row.parameter_name_1),
                safe_excel_value(row.parameter_name_2),
                safe_excel_value(row.old_wp_id),
                safe_excel_value(row.old_wp),
                safe_excel_value(row.action),
                safe_excel_value(row.result),
                safe_excel_value(row.message),
            ]
        )
    _style_report_sheet(results)
    results.freeze_panes = "A2"
    results.auto_filter.ref = results.dimensions
    results.column_dimensions["A"].width = 12
    results.column_dimensions["B"].width = 18
    results.column_dimensions["C"].width = 24
    results.column_dimensions["D"].width = 18
    results.column_dimensions["E"].width = 24
    results.column_dimensions["F"].width = 40
    results.column_dimensions["G"].width = 14
    results.column_dimensions["H"].width = 24
    results.column_dimensions["I"].width = 16
    results.column_dimensions["J"].width = 22
    results.column_dimensions["K"].width = 65
    for row_number in range(2, results.max_row + 1):
        result_cell = results.cell(row_number, 10)
        if result_cell.value in {"ERROR", "BLOCKED", "ROLLED_BACK"}:
            result_cell.fill = PatternFill("solid", fgColor="F4CCCC")
        elif result_cell.value == "UPDATED":
            result_cell.fill = PatternFill("solid", fgColor="C6EFCE")
        elif result_cell.value in {"WOULD_UPDATE", "READY"}:
            result_cell.fill = PatternFill("solid", fgColor="FFF2CC")

    report_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{report_path.stem}_",
            suffix=".xlsx",
            dir=report_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        temporary_path.replace(report_path)
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def update_se_parameter_wp(
    excel_path: Path | str,
    *,
    sheet_name: str = RUN_SHEET_NAME,
    execute: bool = False,
    assume_yes: bool = False,
    report_path: Path | str | None = None,
    db_connection=connection,
) -> RunResult:
    """Build the plan, optionally execute it atomically, and write a report."""
    if assume_yes and not execute:
        raise ValueError("--yes/RUN_ASSUME_YES requires execute mode.")

    input_path = Path(excel_path).expanduser().resolve()
    output_path = make_report_path(report_path)
    started_at = datetime.now().astimezone()
    rows: list[ReportRow] = []
    database, host, port = database_identity(db_connection)

    try:
        rows = read_excel_rows(input_path, sheet_name)
        with db_connection.cursor() as cursor:
            build_plan(cursor, rows)
    except Exception as exc:
        completed_at = datetime.now().astimezone()
        write_report(
            output_path,
            rows,
            status="BLOCKED",
            database=database,
            host=host,
            port=port,
            excel_path=input_path,
            sheet_name=sheet_name,
            execute_requested=execute,
            assume_yes=assume_yes,
            started_at=started_at,
            completed_at=completed_at,
            error=f"{type(exc).__name__}: {exc}",
        )
        raise ToolRunError(str(exc), output_path) from exc

    print_plan(
        rows,
        database=database,
        host=host,
        port=port,
        excel_path=input_path,
        sheet_name=sheet_name,
        report_path=output_path,
        execute=execute,
    )

    # Verify reporting before any mutation. If the final report rewrite fails
    # unexpectedly after commit, this pre-execution audit file still remains.
    if execute and not any(row.result == "ERROR" for row in rows):
        write_report(
            output_path,
            rows,
            status="READY_TO_EXECUTE",
            database=database,
            host=host,
            port=port,
            excel_path=input_path,
            sheet_name=sheet_name,
            execute_requested=execute,
            assume_yes=assume_yes,
            started_at=started_at,
            completed_at=datetime.now().astimezone(),
        )

    if any(row.result == "ERROR" for row in rows):
        status = "BLOCKED"
        error = "Input/lookup errors exist; no database rows were updated."
        for row in rows:
            if row.result == "READY":
                row.result = "BLOCKED"
                row.message = "Not executed because another input row has an error."
    elif not execute:
        status = "PREVIEW"
        error = None
        for row in rows:
            if row.result == "READY":
                row.result = "WOULD_UPDATE"
                row.message = (
                    f"Preview only: would change wp_id from {row.old_wp_id!r} "
                    f"to {row.resolved_wp_id}."
                )
    else:
        ready_rows = [row for row in rows if row.result == "READY"]
        if not ready_rows:
            status = "NO_CHANGES"
            error = None
        elif not confirm_execution(database, rows, assume_yes=assume_yes):
            status = "CANCELLED"
            error = "Exact confirmation was not provided; no rows were updated."
            for row in ready_rows:
                row.result = "CANCELLED"
                row.message = "Database update was cancelled before mutation."
        else:
            try:
                updated_count = execute_updates(rows, db_connection=db_connection)
            except Exception as exc:
                status = "ROLLED_BACK"
                error = f"{type(exc).__name__}: {exc}"
                for row in ready_rows:
                    row.result = "ROLLED_BACK"
                    row.message = "Transaction failed; all changes were rolled back."
            else:
                status = "COMMITTED"
                error = None
                for row in ready_rows:
                    row.result = "UPDATED"
                    row.message = (
                        f"Updated wp_id from {row.old_wp_id!r} "
                        f"to {row.resolved_wp_id}."
                    )
                if updated_count != len(ready_rows):
                    raise RuntimeError(
                        "Internal count mismatch after a committed transaction."
                    )

    completed_at = datetime.now().astimezone()
    write_report(
        output_path,
        rows,
        status=status,
        database=database,
        host=host,
        port=port,
        excel_path=input_path,
        sheet_name=sheet_name,
        execute_requested=execute,
        assume_yes=assume_yes,
        started_at=started_at,
        completed_at=completed_at,
        error=error,
    )
    updated_count = sum(row.result == "UPDATED" for row in rows)
    print(f"Run status: {status}")
    print(f"Excel report: {output_path}")
    print(f"Rows updated: {updated_count}")

    result = RunResult(status, output_path, updated_count, tuple(rows))
    if status == "ROLLED_BACK":
        raise ToolRunError(error or "Database update failed.", output_path)
    return result


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Preview or update se_parameter.wp_id using IDs from Excel column B "
            "and WP names from column C. Every run writes an Excel report."
        )
    )
    parser.add_argument("--excel", type=Path, default=RUN_EXCEL_PATH)
    parser.add_argument("--sheet", default=RUN_SHEET_NAME)
    parser.add_argument("--report", type=Path, default=RUN_REPORT_PATH)
    parser.add_argument(
        "--execute",
        action="store_true",
        help="Apply updates. Without this flag the tool is preview-only.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        help="Bypass the exact interactive confirmation (audited opt-in).",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    execute = args.execute or RUN_EXECUTE
    assume_yes = args.yes or RUN_ASSUME_YES
    try:
        result = update_se_parameter_wp(
            args.excel,
            sheet_name=args.sheet,
            execute=execute,
            assume_yes=assume_yes,
            report_path=args.report,
        )
    except ToolRunError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        print(f"Excel report: {exc.report_path}", file=sys.stderr)
        return 1
    return 1 if result.status == "BLOCKED" else 0


if __name__ == "__main__":
    raise SystemExit(main())
