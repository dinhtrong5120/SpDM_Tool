"""Safely update L-to-SE links and map names from an Excel workbook.

The script is preview-only by default. It updates l_parameter links from the
``id`` and ``related_se_parameter_id`` columns, then synchronizes prj_rfl.map_name
for every matching project/RFL/phase. It does not clear values for blank cells.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Sequence


PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


# =========================================================
# PYCHARM / COMMAND-LINE CONFIGURATION
# =========================================================

RUN_INPUT_PATH = Path("/Users/TrongKieu/Downloads/l_parameter_追加紐付け.xlsx")
RUN_EXECUTE = False
RUN_ASSUME_YES = False
RUN_OUTPUT_PATH: str | None = None
RUN_REPORT_ROW_LIMIT = 50_000

REQUIRED_HEADERS = ("id", "related_se_parameter_id")
REPORT_ROW_LIMIT = 1_000_000
EXCEL_ILLEGAL_CHARACTER_PATTERN = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")
EXCEL_CELL_LIMIT = 32_767


@dataclass(frozen=True)
class InputRecord:
    excel_row: int
    raw_l_id: Any
    raw_se_id: Any
    l_id: int | None
    se_id: int | None
    error: str = ""


@dataclass(frozen=True)
class RflLink:
    l_id: int
    se_id: int
    project_id: int
    phase_id: int
    rfl_id: int
    map_name: str | None


@dataclass(frozen=True)
class ProjectParameterValue:
    project_id: int
    phase_id: int
    se_id: int
    variation_id: int
    value: Any


@dataclass
class RflUpdate:
    project_id: int
    phase_id: int
    rfl_id: int
    current_map_name: str | None
    desired_map_name: str | None
    l_ids: list[int] = field(default_factory=list)
    se_ids: list[int] = field(default_factory=list)
    chosen_variations: list[int] = field(default_factory=list)
    source_values: list[str] = field(default_factory=list)
    status: str = ""
    reason: str = ""

    @property
    def key(self) -> tuple[int, int, int]:
        return self.project_id, self.phase_id, self.rfl_id


@dataclass
class CopyPlan:
    records: list[InputRecord]
    l_existing: dict[int, int | None]
    se_units: dict[int, str | None]
    links: list[RflLink]
    project_parameter_values: list[ProjectParameterValue]
    rfl_updates: list[RflUpdate]
    blockers: list[str]
    warnings: list[str]
    fingerprint: str = ""

    @property
    def l_updates(self) -> list[tuple[int, int | None, int]]:
        updates = []
        for record in self.records:
            if record.l_id is None or record.se_id is None:
                continue
            current = self.l_existing.get(record.l_id)
            if current != record.se_id:
                updates.append((record.l_id, current, record.se_id))
        return updates

    @property
    def map_updates(self) -> list[RflUpdate]:
        return [
            row
            for row in self.rfl_updates
            if row.status == "UPDATE" and row.current_map_name != row.desired_map_name
        ]


def placeholders(values: Sequence[Any]) -> str:
    if not values:
        raise ValueError("Cannot create SQL placeholders for an empty list.")
    return ", ".join(["%s"] * len(values))


def normalize_id(value: Any, *, label: str) -> int:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{label} must be a positive integer.")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{label} must be a finite positive integer.")
    try:
        decimal_value = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, AttributeError) as exc:
        raise ValueError(f"{label} is not a valid integer: {value!r}") from exc
    if not decimal_value.is_finite():
        raise ValueError(f"{label} must be finite: {value!r}")
    if decimal_value != decimal_value.to_integral_value():
        raise ValueError(f"{label} must be an integer: {value!r}")
    result = int(decimal_value)
    if result <= 0:
        raise ValueError(f"{label} must be greater than zero: {result}")
    return result


def safe_excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, datetime, date)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bytes):
        text = value.hex()
    else:
        text = str(value)
    text = EXCEL_ILLEGAL_CHARACTER_PATTERN.sub("�", text)
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    if len(text) > EXCEL_CELL_LIMIT:
        text = text[: EXCEL_CELL_LIMIT - 16] + "...[TRUNCATED]"
    return text


def default_output_path(input_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    stem = input_path.stem or "link_se_l"
    return (
        Path(__file__).resolve().parent
        / "output"
        / f"update_l_se_{stem}_{timestamp}.xlsx"
    )


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Preview/update l_parameter related SE links and RFL map names."
    )
    parser.add_argument("--input", default=str(RUN_INPUT_PATH), help="Input .xlsx file")
    parser.add_argument(
        "--execute",
        action="store_true",
        default=RUN_EXECUTE,
        help="Apply planned updates after confirmation.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        default=RUN_ASSUME_YES,
        help="Skip interactive confirmation (not recommended).",
    )
    parser.add_argument("--output", default=RUN_OUTPUT_PATH, help="Report .xlsx path")
    parser.add_argument(
        "--report-row-limit",
        type=int,
        default=RUN_REPORT_ROW_LIMIT,
        help="Maximum detail rows exported per report sheet.",
    )
    args = parser.parse_args(argv)
    if args.report_row_limit < 1 or args.report_row_limit > REPORT_ROW_LIMIT:
        raise ValueError(f"--report-row-limit must be 1..{REPORT_ROW_LIMIT}.")
    input_path = Path(args.input).expanduser().resolve()
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("Input file must have the .xlsx extension.")
    args.input_path = input_path
    if args.output:
        output_path = Path(args.output).expanduser()
        if not output_path.is_absolute():
            output_path = Path.cwd() / output_path
        output_path = output_path.resolve()
        if output_path.suffix.lower() != ".xlsx":
            raise ValueError("Report path must have the .xlsx extension.")
    else:
        output_path = default_output_path(input_path)
    args.output_path = output_path
    return args


def read_input_workbook(path: Path) -> tuple[list[InputRecord], list[str]]:
    if not path.is_file():
        raise FileNotFoundError(f"Input workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        header_cells = next(worksheet.iter_rows(min_row=1, max_row=1))
        headers = [cell.value for cell in header_cells]
        missing_headers = [
            name for name in REQUIRED_HEADERS if headers.count(name) == 0
        ]
        duplicate_headers = [
            name for name in REQUIRED_HEADERS if headers.count(name) > 1
        ]
        if missing_headers or duplicate_headers:
            problems = []
            if missing_headers:
                problems.append(f"missing headers: {missing_headers}")
            if duplicate_headers:
                problems.append(f"duplicate headers: {duplicate_headers}")
            raise ValueError("Invalid workbook header row: " + "; ".join(problems))
        id_index = headers.index("id")
        se_index = headers.index("related_se_parameter_id")

        records: list[InputRecord] = []
        errors: list[str] = []
        seen_l_ids: dict[int, int] = {}
        for row_number, cells in enumerate(worksheet.iter_rows(min_row=2), start=2):
            id_cell = cells[id_index] if id_index < len(cells) else None
            se_cell = cells[se_index] if se_index < len(cells) else None
            raw_l_id = id_cell.value if id_cell is not None else None
            raw_se_id = se_cell.value if se_cell is not None else None
            if raw_l_id in (None, "") and raw_se_id in (None, ""):
                continue

            row_errors: list[str] = []
            if id_cell is not None and id_cell.data_type == "f":
                row_errors.append("id cell must contain a value, not a formula")
            if se_cell is not None and se_cell.data_type == "f":
                row_errors.append(
                    "related_se_parameter_id cell must contain a value, not a formula"
                )

            try:
                l_id = normalize_id(raw_l_id, label=f"row {row_number} id")
            except ValueError as exc:
                l_id = None
                row_errors.append(str(exc))

            if raw_se_id in (None, ""):
                se_id = None
            else:
                try:
                    se_id = normalize_id(
                        raw_se_id, label=f"row {row_number} related_se_parameter_id"
                    )
                except ValueError as exc:
                    se_id = None
                    row_errors.append(str(exc))

            if l_id is not None:
                if l_id in seen_l_ids:
                    row_errors.append(
                        f"duplicate l_parameter id {l_id}; first appears on "
                        f"row {seen_l_ids[l_id]}"
                    )
                else:
                    seen_l_ids[l_id] = row_number

            record = InputRecord(
                excel_row=row_number,
                raw_l_id=raw_l_id,
                raw_se_id=raw_se_id,
                l_id=l_id,
                se_id=se_id,
                error="; ".join(row_errors),
            )
            records.append(record)
            if record.error:
                errors.append(f"Excel row {row_number}: {record.error}")
        if not records:
            raise ValueError("Workbook has no data rows.")
        return records, errors
    finally:
        workbook.close()


def database_identity() -> tuple[str, str]:
    settings = connection.settings_dict
    return str(settings.get("NAME") or ""), str(settings.get("HOST") or "")


def fetch_l_parameters(
    cursor, l_ids: Sequence[int], *, for_update: bool
) -> dict[int, int | None]:
    if not l_ids:
        return {}
    lock_sql = " FOR UPDATE" if for_update else ""
    cursor.execute(
        f"SELECT id, related_se_parameter_id FROM l_parameter "
        f"WHERE id IN ({placeholders(l_ids)}) ORDER BY id{lock_sql}",
        list(l_ids),
    )
    return {
        int(row[0]): (int(row[1]) if row[1] is not None else None)
        for row in cursor.fetchall()
    }


def fetch_se_units(
    cursor, se_ids: Sequence[int], *, for_update: bool
) -> dict[int, str | None]:
    if not se_ids:
        return {}
    lock_sql = " FOR UPDATE" if for_update else ""
    cursor.execute(
        f"""
        SELECT se_parameter.id, units.unit
        FROM se_parameter
        LEFT JOIN units ON units.id = se_parameter.unit_id
        WHERE se_parameter.id IN ({placeholders(se_ids)})
        ORDER BY se_parameter.id{lock_sql}
        """,
        list(se_ids),
    )
    return {
        int(row[0]): (str(row[1]) if row[1] is not None else None)
        for row in cursor.fetchall()
    }


def fetch_rfl_links(
    cursor, map_l_to_se: dict[int, int], *, for_update: bool
) -> list[RflLink]:
    l_ids = sorted(map_l_to_se)
    if not l_ids:
        return []
    lock_sql = " FOR UPDATE" if for_update else ""
    cursor.execute(
        f"""
        SELECT
            setup_l_relation.l_parameter_id,
            prj_rfl.project_info_id,
            prj_rfl.phase_id,
            prj_rfl.rfl_id,
            prj_rfl.map_name
        FROM prj_rfl
        JOIN rfl ON rfl.id = prj_rfl.rfl_id
        JOIN setup_l_relation ON setup_l_relation.id = rfl.logic_s_id
        WHERE setup_l_relation.l_parameter_id IN ({placeholders(l_ids)})
        ORDER BY setup_l_relation.l_parameter_id,
                 prj_rfl.project_info_id,
                 prj_rfl.phase_id,
                 prj_rfl.rfl_id{lock_sql}
        """,
        l_ids,
    )
    links = []
    for row in cursor.fetchall():
        l_id = int(row[0])
        map_name = str(row[4]) if row[4] is not None else None
        links.append(
            RflLink(
                l_id=l_id,
                se_id=map_l_to_se[l_id],
                project_id=int(row[1]),
                phase_id=int(row[2]),
                rfl_id=int(row[3]),
                map_name=map_name,
            )
        )
    return links


def fetch_project_parameter_values(
    cursor,
    scopes: Sequence[tuple[int, int, int]],
    *,
    for_update: bool,
) -> list[ProjectParameterValue]:
    unique_scopes = sorted(set(scopes))
    if not unique_scopes:
        return []
    values: list[ProjectParameterValue] = []
    # Keep each statement safely below common MySQL parameter-count limits.
    for offset in range(0, len(unique_scopes), 250):
        chunk = unique_scopes[offset : offset + 250]
        where_sql = " OR ".join(
            "(project_id = %s AND phase_id = %s AND se_parameter_id = %s)"
            for _ in chunk
        )
        params = [part for scope in chunk for part in scope]
        lock_sql = " FOR UPDATE" if for_update else ""
        cursor.execute(
            f"""
            SELECT project_id, phase_id, se_parameter_id, variation_id, value
            FROM project_parameter
            WHERE {where_sql}
            ORDER BY project_id, phase_id, se_parameter_id, variation_id{lock_sql}
            """,
            params,
        )
        values.extend(
            ProjectParameterValue(
                project_id=int(row[0]),
                phase_id=int(row[1]),
                se_id=int(row[2]),
                variation_id=int(row[3]),
                value=row[4],
            )
            for row in cursor.fetchall()
        )
    return values


def text_value(value: Any) -> str | None:
    if value is None:
        return None
    value_as_text = str(value)
    return None if not value_as_text.strip() else value_as_text


def build_rfl_updates(
    links: Sequence[RflLink],
    project_parameter_values: Sequence[ProjectParameterValue],
) -> list[RflUpdate]:
    values_by_scope: dict[tuple[int, int, int], list[ProjectParameterValue]] = (
        defaultdict(list)
    )
    for row in project_parameter_values:
        values_by_scope[(row.project_id, row.phase_id, row.se_id)].append(row)

    updates_by_key: dict[tuple[int, int, int], RflUpdate] = {}
    source_scopes_by_key: dict[
        tuple[int, int, int], list[tuple[RflLink, ProjectParameterValue | None]]
    ] = defaultdict(list)
    for link in links:
        scope = (link.project_id, link.phase_id, link.se_id)
        candidates = sorted(
            values_by_scope.get(scope, []), key=lambda row: row.variation_id
        )
        chosen = next(
            (row for row in candidates if text_value(row.value) is not None), None
        )
        source_scopes_by_key[(link.project_id, link.phase_id, link.rfl_id)].append(
            (link, chosen)
        )

    for key, sources in sorted(source_scopes_by_key.items()):
        link = sources[0][0]
        update = RflUpdate(
            project_id=key[0],
            phase_id=key[1],
            rfl_id=key[2],
            current_map_name=link.map_name,
            desired_map_name=None,
            l_ids=sorted({item.l_id for item, _ in sources}),
            se_ids=sorted({item.se_id for item, _ in sources}),
        )
        unresolved_sources = [item.l_id for item, chosen in sources if chosen is None]
        selected = [chosen for _, chosen in sources if chosen is not None]
        if not selected:
            update.status = "SKIP_NO_VALUE"
            update.reason = (
                "Không cập nhật vì không có project_parameter.value không rỗng "
                "cho các SE liên kết từ l_parameter id: "
                f"{update.l_ids}."
            )
        else:
            desired_values = sorted({str(text_value(row.value)) for row in selected})
            update.source_values = desired_values
            update.chosen_variations = sorted({row.variation_id for row in selected})
            if len(desired_values) != 1:
                update.status = "SKIP_CONFLICT"
                update.reason = (
                    "Không cập nhật vì nhiều nguồn nhắm cùng một prj_rfl nhưng "
                    f"có giá trị khác nhau: {desired_values!r}."
                )
            elif len(desired_values[0]) > 255:
                update.status = "BLOCK_VALUE_TOO_LONG"
                update.reason = (
                    "Giá trị project_parameter.value vượt quá giới hạn 255 ký tự "
                    "của prj_rfl.map_name."
                )
            else:
                update.desired_map_name = desired_values[0]
                if update.current_map_name == update.desired_map_name:
                    update.status = "SKIP_UNCHANGED"
                    update.reason = "map_name đã bằng giá trị được chọn."
                else:
                    update.status = "UPDATE"
                    update.reason = (
                        "Cập nhật map_name bằng value không rỗng đầu tiên theo "
                        "variation_id tăng dần."
                    )
                if unresolved_sources:
                    update.reason += (
                        " Một số nguồn không có value, nhưng các nguồn có value "
                        "đều thống nhất."
                    )
        updates_by_key[key] = update
    return list(updates_by_key.values())


def row_as_dict(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): row_as_dict(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [row_as_dict(item) for item in value]
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bytes):
        return value.hex()
    return value


def calculate_fingerprint(plan: CopyPlan) -> str:
    payload = {
        "records": [
            [record.excel_row, record.l_id, record.se_id, record.error]
            for record in plan.records
        ],
        "l_existing": sorted(plan.l_existing.items()),
        "se_units": sorted(plan.se_units.items()),
        "links": [
            [
                row.l_id,
                row.se_id,
                row.project_id,
                row.phase_id,
                row.rfl_id,
                row.map_name,
            ]
            for row in plan.links
        ],
        "project_parameter_values": [
            [
                row.project_id,
                row.phase_id,
                row.se_id,
                row.variation_id,
                row_as_dict(row.value),
            ]
            for row in plan.project_parameter_values
        ],
    }
    serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(serialized).hexdigest()


def build_plan(
    cursor,
    records: list[InputRecord],
    input_errors: list[str],
    *,
    for_update: bool,
) -> CopyPlan:
    blockers = list(input_errors)
    valid_l_ids = sorted({record.l_id for record in records if record.l_id is not None})
    requested_records = [
        record
        for record in records
        if record.l_id is not None and record.se_id is not None
    ]
    requested_se_ids = sorted(
        {record.se_id for record in requested_records if record.se_id}
    )

    l_existing = fetch_l_parameters(cursor, valid_l_ids, for_update=for_update)
    missing_l_ids = sorted(set(valid_l_ids) - set(l_existing))
    if missing_l_ids:
        blockers.append(f"l_parameter.id không tồn tại: {missing_l_ids}")

    se_units = fetch_se_units(cursor, requested_se_ids, for_update=for_update)
    missing_se_ids = sorted(set(requested_se_ids) - set(se_units))
    if missing_se_ids:
        blockers.append(
            f"related_se_parameter_id không tồn tại trong se_parameter: {missing_se_ids}"
        )

    map_l_to_se = {
        record.l_id: record.se_id
        for record in requested_records
        if record.l_id in l_existing
        and record.se_id in se_units
        and se_units[record.se_id] == "Map"
    }
    links = fetch_rfl_links(cursor, map_l_to_se, for_update=for_update)
    scopes = [(link.project_id, link.phase_id, link.se_id) for link in links]
    pp_values = fetch_project_parameter_values(cursor, scopes, for_update=for_update)
    rfl_updates = build_rfl_updates(links, pp_values)
    for row in rfl_updates:
        if row.status == "BLOCK_VALUE_TOO_LONG":
            blockers.append(
                f"prj_rfl ({row.project_id}, {row.phase_id}, {row.rfl_id}): "
                f"{row.reason}"
            )
    warnings = []
    for record in records:
        if (
            record.l_id is not None
            and record.se_id is not None
            and record.se_id in se_units
        ):
            if se_units[record.se_id] != "Map":
                warnings.append(
                    f"Excel row {record.excel_row}: SE {record.se_id} có unit "
                    f"{se_units[record.se_id]!r}, không phải 'Map'; chỉ cập nhật "
                    "l_parameter.related_se_parameter_id."
                )
    for record in requested_records:
        if record.l_id not in map_l_to_se:
            continue
        if not any(link.l_id == record.l_id for link in links):
            warnings.append(
                f"l_parameter {record.l_id} có unit Map nhưng không có liên kết "
                "RFL trong bảng nền."
            )

    plan = CopyPlan(
        records=records,
        l_existing=l_existing,
        se_units=se_units,
        links=links,
        project_parameter_values=pp_values,
        rfl_updates=rfl_updates,
        blockers=blockers,
        warnings=warnings,
    )
    plan.fingerprint = calculate_fingerprint(plan)
    return plan


def apply_plan(cursor, plan: CopyPlan) -> tuple[int, int]:
    l_updated = 0
    for l_id, expected_old, new_se_id in plan.l_updates:
        cursor.execute(
            """
            UPDATE l_parameter
            SET related_se_parameter_id = %s
            WHERE id = %s AND related_se_parameter_id <=> %s
            """,
            [new_se_id, l_id, expected_old],
        )
        if cursor.rowcount != 1:
            raise RuntimeError(
                f"l_parameter {l_id}: expected one updated row, got {cursor.rowcount}."
            )
        l_updated += 1

    map_updated = 0
    for item in plan.map_updates:
        cursor.execute(
            """
            UPDATE prj_rfl
            SET map_name = %s
            WHERE project_info_id = %s
              AND phase_id = %s
              AND rfl_id = %s
              AND map_name <=> %s
            """,
            [
                item.desired_map_name,
                item.project_id,
                item.phase_id,
                item.rfl_id,
                item.current_map_name,
            ],
        )
        if cursor.rowcount != 1:
            raise RuntimeError(
                f"prj_rfl {item.key}: expected one updated row, got {cursor.rowcount}."
            )
        map_updated += 1

    expected_l = {l_id: new_id for l_id, _, new_id in plan.l_updates}
    if expected_l:
        cursor.execute(
            f"SELECT id, related_se_parameter_id FROM l_parameter "
            f"WHERE id IN ({placeholders(list(expected_l))}) ORDER BY id",
            list(expected_l),
        )
        actual_l = {int(row[0]): int(row[1]) for row in cursor.fetchall()}
        if actual_l != expected_l:
            raise RuntimeError("Post-update verification failed for l_parameter.")

    if plan.map_updates:
        expected_map = {item.key: item.desired_map_name for item in plan.map_updates}
        actual_map: dict[tuple[int, int, int], str | None] = {}
        for offset in range(0, len(plan.map_updates), 250):
            chunk = plan.map_updates[offset : offset + 250]
            conditions = " OR ".join(
                "(project_info_id = %s AND phase_id = %s AND rfl_id = %s)"
                for _ in chunk
            )
            keys = [part for item in chunk for part in item.key]
            cursor.execute(
                f"SELECT project_info_id, phase_id, rfl_id, map_name FROM prj_rfl "
                f"WHERE {conditions}",
                keys,
            )
            actual_map.update(
                {
                    (int(row[0]), int(row[1]), int(row[2])): row[3]
                    for row in cursor.fetchall()
                }
            )
        if actual_map != expected_map:
            raise RuntimeError("Post-update verification failed for prj_rfl.map_name.")
    return l_updated, map_updated


def set_sheet_style(worksheet, header_row: int, *, freeze: str = "A2") -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = freeze
    worksheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )
    for cell in worksheet[header_row]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F1F1F")
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for column_index in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column_index)
        max_length = max(
            (
                len(str(cell.value))
                for cell in worksheet[letter]
                if cell.value is not None
            ),
            default=10,
        )
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 60)


def format_status_column(worksheet, column_index: int) -> None:
    fills = {
        "UPDATE": PatternFill("solid", fgColor="E2F0D9"),
        "UNCHANGED": PatternFill("solid", fgColor="E7E6E6"),
        "SKIP_EMPTY_RELATION": PatternFill("solid", fgColor="DDEBF7"),
        "SKIP_NO_VALUE": PatternFill("solid", fgColor="FFF2CC"),
        "SKIP_CONFLICT": PatternFill("solid", fgColor="FCE4D6"),
        "SKIP_UNCHANGED": PatternFill("solid", fgColor="E7E6E6"),
        "SKIP_NO_RFL": PatternFill("solid", fgColor="DDEBF7"),
    }
    blocked_fill = PatternFill("solid", fgColor="F4CCCC")
    for row_index in range(2, worksheet.max_row + 1):
        cell = worksheet.cell(row=row_index, column=column_index)
        action = str(cell.value or "")
        if action.startswith("BLOCK"):
            cell.fill = blocked_fill
            cell.font = Font(name="Arial", size=10, bold=True, color="C00000")
        elif action in fills:
            cell.fill = fills[action]
            cell.font = Font(name="Arial", size=10, bold=True, color="1F1F1F")


def limited_rows(worksheet, rows: Sequence[Sequence[Any]], limit: int) -> None:
    for row in rows[:limit]:
        worksheet.append([safe_excel_value(value) for value in row])
    if len(rows) > limit:
        worksheet.append([])
        worksheet.append(
            [
                f"Truncated: exported {limit} of {len(rows)} detail rows; summary totals remain exact."
            ]
        )


def report_input_rows(plan: CopyPlan) -> list[list[Any]]:
    rows = []
    for record in plan.records:
        if record.error:
            status, reason = "BLOCKED", record.error
        elif record.se_id is None:
            status, reason = (
                "SKIP_EMPTY_RELATION",
                "Ô trống: giữ nguyên related_se_parameter_id hiện tại.",
            )
        elif record.l_id not in plan.l_existing:
            status, reason = "BLOCKED", "Không tìm thấy l_parameter.id trong database."
        elif record.se_id not in plan.se_units:
            status, reason = (
                "BLOCKED",
                "Không tìm thấy related_se_parameter_id trong se_parameter.",
            )
        elif plan.l_existing[record.l_id] == record.se_id:
            status, reason = (
                "UNCHANGED",
                "related_se_parameter_id đã đúng theo workbook.",
            )
        else:
            status, reason = "UPDATE", "Cập nhật related_se_parameter_id theo workbook."
        unit = plan.se_units.get(record.se_id) if record.se_id is not None else None
        if record.se_id is not None and record.se_id in plan.se_units:
            if unit != "Map":
                reason += f" Unit {unit!r} không phải 'Map'; không đồng bộ prj_rfl."
            elif record.l_id is not None and not any(
                link.l_id == record.l_id for link in plan.links
            ):
                reason += " SE unit là Map nhưng không tìm thấy liên kết RFL."
        rows.append(
            [
                record.excel_row,
                record.raw_l_id,
                record.raw_se_id,
                record.l_id,
                record.se_id,
                plan.l_existing.get(record.l_id) if record.l_id is not None else None,
                plan.se_units.get(record.se_id) if record.se_id is not None else None,
                status,
                reason,
            ]
        )
    return rows


def report_variation_rows(plan: CopyPlan) -> list[list[Any]]:
    rows = []
    links_by_scope: dict[tuple[int, int, int], list[RflLink]] = defaultdict(list)
    for link in plan.links:
        links_by_scope[(link.project_id, link.phase_id, link.se_id)].append(link)
    chosen_by_scope: dict[tuple[int, int, int], int] = {}
    for scope, candidates in links_by_scope.items():
        ordered = sorted(
            [
                row
                for row in plan.project_parameter_values
                if (row.project_id, row.phase_id, row.se_id) == scope
            ],
            key=lambda row: row.variation_id,
        )
        selected = next(
            (row for row in ordered if text_value(row.value) is not None), None
        )
        if selected:
            chosen_by_scope[scope] = selected.variation_id
        for row in ordered:
            matching_links = [
                link
                for link in candidates
                if link.project_id == scope[0] and link.phase_id == scope[1]
            ]
            rows.append(
                [
                    scope[0],
                    scope[1],
                    scope[2],
                    row.variation_id,
                    row.value,
                    "YES" if chosen_by_scope.get(scope) == row.variation_id else "NO",
                    ", ".join(str(link.l_id) for link in matching_links),
                ]
            )
    return rows


def report_link_rows(plan: CopyPlan) -> list[list[Any]]:
    update_by_key = {row.key: row for row in plan.rfl_updates}
    rows = []
    for item in plan.rfl_updates:
        rows.append(
            [
                item.project_id,
                item.phase_id,
                item.rfl_id,
                ", ".join(map(str, item.l_ids)),
                ", ".join(map(str, item.se_ids)),
                ", ".join(map(str, item.chosen_variations)),
                ", ".join(item.source_values),
                item.current_map_name,
                item.desired_map_name,
                item.status,
                item.reason,
            ]
        )
    # Keep the lookup intentional: it also guards against accidental duplicate keys.
    if len(update_by_key) != len(plan.rfl_updates):
        raise RuntimeError("Duplicate prj_rfl keys exist in the report plan.")
    linked_l_ids = {link.l_id for link in plan.links}
    for record in plan.records:
        if (
            record.l_id is None
            or record.se_id is None
            or plan.se_units.get(record.se_id) != "Map"
            or record.l_id in linked_l_ids
        ):
            continue
        rows.append(
            [
                None,
                None,
                None,
                str(record.l_id),
                str(record.se_id),
                "",
                "",
                "",
                "",
                "SKIP_NO_RFL",
                "SE có unit Map nhưng không có liên kết RFL trong các bảng nền.",
            ]
        )
    return rows


def write_report(
    path: Path,
    *,
    input_path: Path,
    status: str,
    message: str,
    database_name: str,
    database_host: str,
    records: list[InputRecord],
    plan: CopyPlan | None,
    input_errors: list[str],
    execute_requested: bool,
    assume_yes: bool,
    output_limits: int,
    l_updated: int | None = None,
    map_updated: int | None = None,
    allow_overwrite: bool = False,
) -> None:
    if path.exists() and not allow_overwrite:
        raise FileExistsError(
            f"Report already exists and will not be overwritten: {path}"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary.append(["L–SE link and RFL map update report"])
    summary.append(["Field", "Value"])
    summary_values = [
        ("status", status),
        ("generated_at", datetime.now()),
        ("database", database_name),
        ("host", database_host),
        ("input_file", str(input_path)),
        ("execute_requested", execute_requested),
        ("assume_yes", assume_yes),
        ("input_rows", len(records)),
        (
            "rows_with_related_se_id",
            sum(record.se_id is not None for record in records),
        ),
        ("input_errors", "\n".join(input_errors)),
        ("message", message),
    ]
    if plan is not None:
        summary_values.extend(
            [
                ("l_parameter_rows_to_update", len(plan.l_updates)),
                ("l_parameter_rows_updated", l_updated),
                ("map_unit_rfl_links_found", len(plan.links)),
                ("prj_rfl_rows_to_update", len(plan.map_updates)),
                ("prj_rfl_rows_updated", map_updated),
                (
                    "prj_rfl_rows_skipped_or_unchanged",
                    sum(row.status != "UPDATE" for row in plan.rfl_updates),
                ),
                ("blockers", "\n".join(plan.blockers)),
                ("warnings", "\n".join(plan.warnings)),
                ("plan_fingerprint", plan.fingerprint),
            ]
        )
    for key, value in summary_values:
        summary.append([safe_excel_value(key), safe_excel_value(value)])
    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 90
    for cell in summary[2]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for cell in summary[1]:
        cell.font = Font(name="Arial", size=14, bold=True)
    for cell in summary["B"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    summary.freeze_panes = "A3"

    input_sheet = workbook.create_sheet("Input rows")
    input_sheet.append(
        [
            "excel_row",
            "source_id",
            "source_related_se_id",
            "l_parameter_id",
            "requested_se_id",
            "current_se_id",
            "se_unit",
            "action",
            "reason",
        ]
    )
    limited_rows(input_sheet, report_input_rows(plan) if plan else [], output_limits)
    set_sheet_style(input_sheet, 1, freeze="A2")
    format_status_column(input_sheet, 8)

    variation_sheet = workbook.create_sheet("Variation values")
    variation_sheet.append(
        [
            "project_id",
            "phase_id",
            "se_parameter_id",
            "variation_id",
            "project_parameter.value",
            "selected",
            "l_parameter_ids",
        ]
    )
    limited_rows(
        variation_sheet, report_variation_rows(plan) if plan else [], output_limits
    )
    set_sheet_style(variation_sheet, 1, freeze="A2")

    rfl_sheet = workbook.create_sheet("RFL map updates")
    rfl_sheet.append(
        [
            "project_id",
            "phase_id",
            "rfl_id",
            "l_parameter_ids",
            "se_parameter_ids",
            "selected_variation_ids",
            "source_values",
            "map_name_before",
            "map_name_after",
            "action",
            "reason",
        ]
    )
    limited_rows(rfl_sheet, report_link_rows(plan) if plan else [], output_limits)
    set_sheet_style(rfl_sheet, 1, freeze="A2")
    format_status_column(rfl_sheet, 10)

    temp_path = path.with_name(f".{path.name}.tmp.xlsx")
    workbook.save(temp_path)
    temp_path.replace(path)


def confirmation_text(database_name: str, plan: CopyPlan) -> str:
    input_l_ids = sorted(
        {
            record.l_id
            for record in plan.records
            if record.l_id is not None and record.se_id is not None
        }
    )
    rfl_keys = sorted(item.key for item in plan.map_updates)
    return (
        f"UPDATE L-SE AND RFL MAP {database_name} "
        f"L_UPDATES {len(plan.l_updates)} RFL_UPDATES {len(plan.map_updates)} "
        f"L_IDS {','.join(map(str, input_l_ids))} "
        f"RFL_KEYS {';'.join(':'.join(map(str, key)) for key in rfl_keys)} "
        "ALL_PROJECT_PHASES"
    )


def print_plan(database_name: str, database_host: str, plan: CopyPlan) -> None:
    print("=" * 88)
    print("PREVIEW L-SE LINK / RFL MAP UPDATE")
    print(f"Database       : {database_name}")
    print(f"Host           : {database_host}")
    print(f"Input rows     : {len(plan.records)}")
    print(f"L links changed: {len(plan.l_updates)}")
    print(f"RFL links found: {len(plan.links)}")
    print(f"Map rows change: {len(plan.map_updates)}")
    print("Affected tables : l_parameter, prj_rfl")
    print(
        "L parameter IDs : "
        + ", ".join(
            str(l_id)
            for l_id in sorted(
                {
                    record.l_id
                    for record in plan.records
                    if record.l_id is not None and record.se_id is not None
                }
            )
        )
    )
    print(
        "prj_rfl keys    : "
        + "; ".join(
            ":".join(map(str, item.key))
            for item in sorted(plan.map_updates, key=lambda row: row.key)
        )
    )
    print(f"Blockers       : {len(plan.blockers)}")
    print(f"Warnings       : {len(plan.warnings)}")
    for warning in plan.warnings:
        print(f"WARNING: {warning}")
    for blocker in plan.blockers:
        print(f"BLOCKER: {blocker}")
    print("=" * 88)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    database_name, database_host = database_identity()
    try:
        records, input_errors = read_input_workbook(args.input_path)
    except Exception as exc:
        message = f"{type(exc).__name__}: {exc}"
        write_report(
            args.output_path,
            input_path=args.input_path,
            status="INPUT_ERROR",
            message=message,
            database_name=database_name,
            database_host=database_host,
            records=[],
            plan=None,
            input_errors=[message],
            execute_requested=args.execute,
            assume_yes=args.yes,
            output_limits=args.report_row_limit,
        )
        print(message)
        print(f"Excel report: {args.output_path}")
        return 1

    plan: CopyPlan | None = None
    report_written = False
    committed = False
    l_updated: int | None = None
    map_updated: int | None = None
    try:
        with connection.cursor() as cursor:
            plan = build_plan(cursor, records, input_errors, for_update=False)
        print_plan(database_name, database_host, plan)
        write_report(
            args.output_path,
            input_path=args.input_path,
            status="PREVIEW",
            message="Preview completed; database has not been changed.",
            database_name=database_name,
            database_host=database_host,
            records=records,
            plan=plan,
            input_errors=input_errors,
            execute_requested=args.execute,
            assume_yes=args.yes,
            output_limits=args.report_row_limit,
        )
        report_written = True
        print(f"Excel report   : {args.output_path}")

        if not args.execute:
            print("Preview only. No database updates were applied.")
            return 0
        if plan.blockers:
            print("Execution blocked. Resolve blockers listed in the report.")
            write_report(
                args.output_path,
                input_path=args.input_path,
                status="BLOCKED",
                message="Execution blocked by preflight validation.",
                database_name=database_name,
                database_host=database_host,
                records=records,
                plan=plan,
                input_errors=input_errors,
                execute_requested=args.execute,
                assume_yes=args.yes,
                output_limits=args.report_row_limit,
                allow_overwrite=True,
            )
            return 1

        if not plan.l_updates and not plan.map_updates:
            write_report(
                args.output_path,
                input_path=args.input_path,
                status="NO_OP",
                message="No database changes are required.",
                database_name=database_name,
                database_host=database_host,
                records=records,
                plan=plan,
                input_errors=input_errors,
                execute_requested=args.execute,
                assume_yes=args.yes,
                output_limits=args.report_row_limit,
                l_updated=0,
                map_updated=0,
                allow_overwrite=True,
            )
            print("No-op. No database updates were required.")
            return 0

        confirmation = confirmation_text(database_name, plan)
        if not args.yes:
            print("Type the exact confirmation to continue:")
            print(confirmation)
            if input("> ").strip() != confirmation:
                write_report(
                    args.output_path,
                    input_path=args.input_path,
                    status="CANCELLED",
                    message="Confirmation was cancelled or did not match.",
                    database_name=database_name,
                    database_host=database_host,
                    records=records,
                    plan=plan,
                    input_errors=input_errors,
                    execute_requested=args.execute,
                    assume_yes=args.yes,
                    output_limits=args.report_row_limit,
                    allow_overwrite=True,
                )
                print("Cancelled. No database updates were applied.")
                return 0

        with transaction.atomic():
            with connection.cursor() as cursor:
                locked_plan = build_plan(cursor, records, input_errors, for_update=True)
                if locked_plan.blockers:
                    raise RuntimeError(
                        "Transactional preflight blocked: "
                        + " | ".join(locked_plan.blockers)
                    )
                if locked_plan.fingerprint != plan.fingerprint:
                    raise RuntimeError(
                        "Relevant database data changed after preview. "
                        "Transaction rolled back; run preview again."
                    )
                l_updated, map_updated = apply_plan(cursor, locked_plan)
        committed = True

        write_report(
            args.output_path,
            input_path=args.input_path,
            status="UPDATED",
            message="Both table updates committed and post-update checks passed.",
            database_name=database_name,
            database_host=database_host,
            records=records,
            plan=plan,
            input_errors=input_errors,
            execute_requested=args.execute,
            assume_yes=args.yes,
            output_limits=args.report_row_limit,
            l_updated=l_updated,
            map_updated=map_updated,
            allow_overwrite=True,
        )
        print(f"Updated l_parameter rows : {l_updated}")
        print(f"Updated prj_rfl rows     : {map_updated}")
        print(f"Excel report             : {args.output_path}")
        return 0
    except Exception as exc:
        status = "REPORT_FAILED_AFTER_COMMIT" if committed else "FAILED"
        message = f"{type(exc).__name__}: {exc}"
        if (report_written or not args.output_path.exists()) and not committed:
            try:
                write_report(
                    args.output_path,
                    input_path=args.input_path,
                    status=status,
                    message=message,
                    database_name=database_name,
                    database_host=database_host,
                    records=records,
                    plan=plan,
                    input_errors=input_errors,
                    execute_requested=args.execute,
                    assume_yes=args.yes,
                    output_limits=args.report_row_limit,
                    allow_overwrite=True,
                )
            except Exception as report_error:
                print(f"Unable to update failure report: {report_error}")
        print(message)
        if committed:
            print(
                "Database updates were committed, but the final report could not be saved."
            )
        return 1
    finally:
        connection.close()


if __name__ == "__main__":
    raise SystemExit(main())
