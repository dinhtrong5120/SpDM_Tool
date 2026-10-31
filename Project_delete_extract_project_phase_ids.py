"""Resolve SPDM project/phase IDs from an Excel project list.

This utility is read-only: it queries the configured Django database and writes
the match result to a separate Excel workbook.
"""

from __future__ import annotations

import argparse
import os
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RUN_INPUT_PATH = Path("/Users/TrongKieu/Downloads/CLW_消去するProject.xlsx")
RUN_OUTPUT_PATH = (
    PROJECT_ROOT
    / "outputs"
    / "clw_project_phase_ids"
    / "CLW_project_phase_ids.xlsx"
)
RUN_EXPECTED_PROJECT_COUNT = 19

SOURCE_HEADERS = (
    "Project",
    "Destination",
    "Drivetrain",
    "Powertrain",
    "Body type",
    "Lot",
    "Phase",
)
OUTPUT_HEADERS = (*SOURCE_HEADERS, "project_id", "phase_id", "match_status")

PROJECT_PHASE_SQL = """
    SELECT DISTINCT
        pjf.id AS project_id,
        ph.id AS phase_id
    FROM project_info AS pjf
    JOIN setup AS stp ON stp.id = pjf.setup_id
    JOIN architecture AS archi ON archi.id = stp.architecture_id
    JOIN destination AS dest ON dest.id = stp.destination_id
    JOIN drivetrain AS dt ON dt.id = stp.drivetrain_id
    JOIN body_type AS bt ON bt.id = stp.body_type_id
    JOIN lot AS lt ON lt.id = pjf.lot_id
    JOIN (
        SELECT project_id, phase_id FROM project_parameter
        UNION
        SELECT project_id, phase_id FROM project_r_parameter
    ) AS project_phase
      ON project_phase.project_id = pjf.id
    JOIN phase AS ph ON ph.id = project_phase.phase_id
    WHERE pjf.project_code = %s
      AND dest.destination = %s
      AND dt.drivetrain = %s
      AND archi.architecturename = %s
      AND bt.body_type = %s
      AND lt.lot = %s
      AND ph.phase = %s
    ORDER BY pjf.id, ph.id
"""


@dataclass(frozen=True)
class ProjectKey:
    project: str
    destination: str
    drivetrain: str
    powertrain: str
    body_type: str
    lot: str
    phase: str

    def as_query_params(self) -> tuple[str, ...]:
        return (
            self.project,
            self.destination,
            self.drivetrain,
            self.powertrain,
            self.body_type,
            self.lot,
            self.phase,
        )

    def as_output_values(self) -> tuple[str, ...]:
        return self.as_query_params()


@dataclass(frozen=True)
class MatchResult:
    key: ProjectKey
    project_id: int | None
    phase_id: int | None
    status: str


def _clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_project_keys(input_path: Path, expected_count: int | None) -> list[ProjectKey]:
    if not input_path.is_file():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    header_index = None
    header_positions: dict[str, int] = {}
    for index, row in enumerate(rows):
        positions = {
            _clean_text(value): column_index
            for column_index, value in enumerate(row)
            if _clean_text(value)
        }
        if all(header in positions for header in SOURCE_HEADERS):
            header_index = index
            header_positions = positions
            break

    if header_index is None:
        raise ValueError(
            "Could not find the required header row: " + ", ".join(SOURCE_HEADERS)
        )

    keys: list[ProjectKey] = []
    for excel_row_number, row in enumerate(rows[header_index + 1 :], header_index + 2):
        values = {
            header: _clean_text(row[header_positions[header]])
            if header_positions[header] < len(row)
            else ""
            for header in SOURCE_HEADERS
        }
        if not any(values.values()):
            continue

        missing = [header for header, value in values.items() if not value]
        if missing:
            raise ValueError(
                f"Row {excel_row_number} is missing required values: {', '.join(missing)}"
            )

        keys.append(
            ProjectKey(
                project=values["Project"],
                destination=values["Destination"],
                drivetrain=values["Drivetrain"],
                powertrain=values["Powertrain"],
                body_type=values["Body type"],
                lot=values["Lot"],
                phase=values["Phase"],
            )
        )

    if expected_count is not None and len(keys) != expected_count:
        raise ValueError(
            f"Expected {expected_count} project rows but found {len(keys)} in {input_path}"
        )
    return keys


def resolve_project_keys(cursor, keys: Iterable[ProjectKey]) -> list[MatchResult]:
    results: list[MatchResult] = []
    for key in keys:
        cursor.execute(PROJECT_PHASE_SQL, key.as_query_params())
        matches = list(cursor.fetchall())
        if len(matches) == 1:
            project_id, phase_id = matches[0]
            results.append(
                MatchResult(key, int(project_id), int(phase_id), "MATCHED")
            )
        elif not matches:
            results.append(MatchResult(key, None, None, "NOT_FOUND"))
        else:
            results.append(MatchResult(key, None, None, f"AMBIGUOUS ({len(matches)} matches)"))
    return results


def write_output(output_path: Path, results: Sequence[MatchResult]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Project Phase IDs"
    worksheet.append(OUTPUT_HEADERS)

    for result in results:
        worksheet.append(
            (
                *result.key.as_output_values(),
                result.project_id,
                result.phase_id,
                result.status,
            )
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")

    widths = (30, 20, 14, 18, 18, 12, 20, 12, 12, 24)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    workbook.save(output_path)


def initialize_django() -> None:
    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")

    import django

    django.setup()


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Read project attributes from Excel and export project_id/phase_id matches."
    )
    parser.add_argument("--input", type=Path, default=RUN_INPUT_PATH)
    parser.add_argument("--output", type=Path, default=RUN_OUTPUT_PATH)
    parser.add_argument(
        "--expected-count",
        type=int,
        default=RUN_EXPECTED_PROJECT_COUNT,
        help="Fail before querying when the input row count differs; use 0 to disable.",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    expected_count = args.expected_count or None
    keys = read_project_keys(args.input, expected_count)

    initialize_django()
    from django.db import connection

    db_settings = connection.settings_dict
    print(
        "Read-only lookup against "
        f"database={db_settings.get('NAME')} host={db_settings.get('HOST')}"
    )
    with connection.cursor() as cursor:
        results = resolve_project_keys(cursor, keys)

    write_output(args.output, results)
    matched_count = sum(result.status == "MATCHED" for result in results)
    print(f"Matched {matched_count}/{len(results)} project rows")
    print(f"Output: {args.output}")

    unresolved = [result for result in results if result.status != "MATCHED"]
    if unresolved:
        print("Unresolved rows:")
        for result in unresolved:
            print(f"- {result.key.project}: {result.status}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
