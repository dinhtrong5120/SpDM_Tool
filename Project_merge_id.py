"""Preview or safely merge one ``project_info.id`` into another.

The default mode is read-only and always writes an Excel audit report.  A real
merge requires ``--execute`` plus an exact database-specific confirmation.
All database mutations run in one transaction: direct/legacy references are
updated first, the source project is verified unused, then its root row is
deleted.

PyCharm: review the ``RUN_*`` constants, run once with ``RUN_EXECUTE=False``,
inspect the Excel report, then explicitly enable execution if the plan is safe.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Sequence

# ---------------------------------------------------------------------------
# PyCharm Run-button configuration
# ---------------------------------------------------------------------------

RUN_SOURCE_PROJECT_ID = 1014
RUN_TARGET_PROJECT_ID = 1000
RUN_EXECUTE = False
RUN_ASSUME_YES = False
RUN_OUTPUT_PATH: str | None = None
RUN_REPORT_ROW_LIMIT = 50_000
# Set this only when the SPDM backend is not already the PyCharm content root.
RUN_PROJECT_ROOT: str | None = None

PROJECT_TABLE = "project_info"
PROJECT_KEY = "id"
OUTPUT_DIR = Path(__file__).resolve().parent / "output"

# Reviewed legacy project columns used by the existing SPDM maintenance tools.
# Never add a column only because its name looks relevant; verify its semantics.
LEGACY_PROJECT_COLUMNS: dict[str, tuple[str, ...]] = {
    "column_resized_tbl": ("project_info_id",),
    "detail_data": ("project_id",),
    "detail_data_design": ("project_id",),
    "detail_data_rfl": ("project_id",),
    "detail_data_rlist": ("project_id",),
    "detail_data_to": ("project_id",),
    "input_surrogate_model_parameter": ("project_id",),
    "map_folder_url": ("project_id",),
    "operational_state": ("project_id",),
    "parameter_map_relation": ("project_id",),
    "prj_rfl": ("project_info_id",),
    "project_parameter": ("project_id",),
    "project_parameter_attachments": ("project_id",),
    "project_r_parameter": ("project_id",),
    "project_r_parameter_history": ("project_id",),
    "project_r_parameter_manager_approval_history": ("project_id",),
    "project_statement": ("project_id",),
    "project_se_confirmation_history": ("project_id",),
    "project_se_input_completion": ("project_id",),
    "project_se_notification": ("project_id",),
    "project_se_notification_target": ("project_id",),
    "project_summary": ("project_id",),
    "rfl_order": ("project_id",),
    "rfl_to_description": ("project_info_id",),
    "rfl_to_statement": ("project_info_id",),
    "sampling_data_path": ("project_id",),
    "studies_comparison_dashboard_settings": ("project_id",),
    "study_group": ("project_id",),
    "surrogate_model_path": ("project_id",),
    "temp_data": ("project_id",),
    "temp_project_parameter": ("project_id",),
    "test_project_senario_parameter": ("project_id",),
    "test_senario_parameter_map_relation": ("project_id",),
    "wp_statement": ("project_id",),
}

# Reviewed backup/recovery snapshots. Their historical project IDs are kept.
IGNORED_SNAPSHOT_PROJECT_COLUMNS: dict[str, tuple[str, ...]] = {
    "bk_prj_rfl_20260901": ("project_info_id",),
    "bk_project_parameter_20260901": ("project_id",),
    "bk_project_r_parameter_20260901": ("project_id",),
    "zz_oos_pp": ("project_id",),
    "zz_oos_prp": ("project_id",),
    "zz_oos_rfl": ("project_info_id",),
}

CANDIDATE_COLUMN_NAMES = ("project_id", "project_info_id")
IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_$]*$")
EXCEL_ILLEGAL_CHARACTER_PATTERN = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")
EXCEL_CELL_LIMIT = 32_767


@dataclass(frozen=True)
class Reference:
    table: str
    column: str
    sources: tuple[str, ...]
    source_rows: int
    target_rows: int


@dataclass(frozen=True)
class UniqueCheck:
    table: str
    index: str
    columns: tuple[str, ...]
    conflict_groups: int | None
    note: str = ""


@dataclass(frozen=True)
class Candidate:
    table: str
    column: str
    source_rows: int


@dataclass
class TableSnapshot:
    table: str
    matched_columns: tuple[str, ...]
    total_rows: int
    rows: list[dict[str, Any]] = field(default_factory=list)

    @property
    def truncated(self) -> bool:
        return self.total_rows > len(self.rows)


@dataclass
class MergePlan:
    source_id: int
    target_id: int
    source_project: dict[str, Any]
    target_project: dict[str, Any]
    project_differences: dict[str, tuple[Any, Any]]
    references: tuple[Reference, ...]
    unique_checks: tuple[UniqueCheck, ...]
    unreviewed_candidates: tuple[Candidate, ...]
    ignored_snapshots: tuple[Candidate, ...] = ()
    snapshots: tuple[TableSnapshot, ...] = ()

    @property
    def blockers(self) -> tuple[str, ...]:
        blockers: list[str] = []
        for check in self.unique_checks:
            if check.conflict_groups is None:
                blockers.append(
                    f"{check.table}.{check.index}: {check.note or 'không thể kiểm tra unique index'}"
                )
            elif check.conflict_groups:
                blockers.append(
                    f"{check.table}.{check.index}: {check.conflict_groups} nhóm sẽ trùng unique key."
                )
        for candidate in self.unreviewed_candidates:
            if candidate.source_rows:
                blockers.append(
                    f"{candidate.table}.{candidate.column}: còn {candidate.source_rows} dòng "
                    "dùng source ID nhưng quan hệ chưa được review."
                )
        return tuple(blockers)

    @property
    def source_reference_rows(self) -> int:
        return sum(reference.source_rows for reference in self.references)

    def fingerprint(self) -> tuple[Any, ...]:
        return (
            tuple(sorted(self.source_project.items())),
            tuple(sorted(self.target_project.items())),
            tuple(
                (r.table, r.column, r.sources, r.source_rows, r.target_rows)
                for r in self.references
            ),
            tuple(
                (c.table, c.index, c.columns, c.conflict_groups, c.note)
                for c in self.unique_checks
            ),
            tuple(
                (c.table, c.column, c.source_rows) for c in self.unreviewed_candidates
            ),
            tuple((c.table, c.column, c.source_rows) for c in self.ignored_snapshots),
        )


def quote_identifier(identifier: str) -> str:
    if not IDENTIFIER_PATTERN.fullmatch(identifier):
        raise ValueError(f"Tên bảng/cột không hợp lệ: {identifier!r}")
    return f"`{identifier}`"


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Không thể tạo SQL placeholders cho danh sách rỗng.")
    return ", ".join(["%s"] * len(values))


def normalize_positive_id(raw_id: Any, *, label: str) -> int:
    if isinstance(raw_id, bool):
        raise ValueError(f"{label} phải là số nguyên dương, không phải boolean.")
    try:
        value = int(raw_id)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} không hợp lệ: {raw_id!r}") from exc
    if value <= 0:
        raise ValueError(f"{label} phải lớn hơn 0: {value}")
    return value


def validate_ids(source_id: Any, target_id: Any) -> tuple[int, int]:
    source = normalize_positive_id(source_id, label="source_project_id")
    target = normalize_positive_id(target_id, label="target_project_id")
    if source == target:
        raise ValueError("Source project ID và target project ID phải khác nhau.")
    return source, target


def normalize_report_row_limit(raw_limit: Any) -> int:
    limit = normalize_positive_id(raw_limit, label="report_row_limit")
    if limit > 1_000_000:
        raise ValueError("report_row_limit không được lớn hơn 1.000.000.")
    return limit


def resolve_output_path(raw_path: str | None, source_id: int, target_id: int) -> Path:
    if raw_path:
        path = Path(raw_path).expanduser()
        if not path.is_absolute():
            path = Path.cwd() / path
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Đường dẫn báo cáo phải có phần mở rộng .xlsx.")
        return path.resolve()
    timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
    return OUTPUT_DIR / f"project_id_merge_{source_id}_to_{target_id}_{timestamp}.xlsx"


def safe_excel_value(value: Any) -> Any:
    if value is None or isinstance(
        value, (bool, int, float, Decimal, date, datetime, time)
    ):
        return value
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    elif not isinstance(value, str):
        value = str(value)
    value = EXCEL_ILLEGAL_CHARACTER_PATTERN.sub("", value)
    if value.startswith(("=", "+", "-", "@")):
        value = f"'{value}"
    return value[:EXCEL_CELL_LIMIT]


def database_identity(cursor, db_connection) -> tuple[str, str]:
    cursor.execute("SELECT DATABASE(), @@hostname")
    database, server_host = cursor.fetchone()
    configured_host = db_connection.settings_dict.get("HOST") or server_host
    configured_port = db_connection.settings_dict.get("PORT") or 3306
    return str(database or "<unknown>"), f"{configured_host}:{configured_port}"


def table_columns(cursor, table: str) -> list[str]:
    cursor.execute(
        """
        SELECT COLUMN_NAME
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s
        ORDER BY ORDINAL_POSITION
        """,
        [table],
    )
    return [str(row[0]) for row in cursor.fetchall()]


def fetch_project(cursor, project_id: int, *, lock: bool = False) -> dict[str, Any]:
    columns = table_columns(cursor, PROJECT_TABLE)
    if PROJECT_KEY not in columns:
        raise RuntimeError(f"Không tìm thấy {PROJECT_TABLE}.{PROJECT_KEY}.")
    selected = ", ".join(quote_identifier(column) for column in columns)
    lock_sql = " FOR UPDATE" if lock else ""
    cursor.execute(
        f"SELECT {selected} FROM {quote_identifier(PROJECT_TABLE)} "
        f"WHERE {quote_identifier(PROJECT_KEY)} = %s{lock_sql}",
        [project_id],
    )
    row = cursor.fetchone()
    if row is None:
        raise ValueError(f"Không tìm thấy {PROJECT_TABLE}.id={project_id}.")
    return dict(zip(columns, row, strict=True))


def compare_projects(
    source: dict[str, Any], target: dict[str, Any]
) -> dict[str, tuple[Any, Any]]:
    return {
        column: (source.get(column), target.get(column))
        for column in source
        if column != PROJECT_KEY and source.get(column) != target.get(column)
    }


def discover_references(cursor) -> dict[tuple[str, str], set[str]]:
    cursor.execute(
        """
        SELECT kcu.TABLE_NAME, kcu.COLUMN_NAME
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        JOIN information_schema.TABLES AS tables_metadata
          ON tables_metadata.TABLE_SCHEMA = kcu.TABLE_SCHEMA
         AND tables_metadata.TABLE_NAME = kcu.TABLE_NAME
         AND tables_metadata.TABLE_TYPE = 'BASE TABLE'
        WHERE kcu.TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_NAME = %s
          AND kcu.REFERENCED_COLUMN_NAME = %s
        ORDER BY kcu.TABLE_NAME, kcu.COLUMN_NAME
        """,
        [PROJECT_TABLE, PROJECT_KEY],
    )
    references: dict[tuple[str, str], set[str]] = defaultdict(set)
    for table, column in cursor.fetchall():
        references[(str(table), str(column))].add("foreign_key")

    legacy_pairs = [
        (table, column)
        for table, columns in LEGACY_PROJECT_COLUMNS.items()
        for column in columns
    ]
    if legacy_pairs:
        table_names = sorted({table for table, _ in legacy_pairs})
        cursor.execute(
            f"""
            SELECT c.TABLE_NAME, c.COLUMN_NAME
            FROM information_schema.COLUMNS AS c
            JOIN information_schema.TABLES AS tables_metadata
              ON tables_metadata.TABLE_SCHEMA = c.TABLE_SCHEMA
             AND tables_metadata.TABLE_NAME = c.TABLE_NAME
             AND tables_metadata.TABLE_TYPE = 'BASE TABLE'
            WHERE c.TABLE_SCHEMA = DATABASE()
              AND c.TABLE_NAME IN ({make_placeholders(table_names)})
            """,
            table_names,
        )
        existing = {(str(table), str(column)) for table, column in cursor.fetchall()}
        for pair in legacy_pairs:
            if pair in existing:
                references[pair].add("legacy_allowlist")
    return references


def count_value(cursor, table: str, column: str, value: int) -> int:
    cursor.execute(
        f"SELECT COUNT(*) FROM {quote_identifier(table)} "
        f"WHERE {quote_identifier(column)} = %s",
        [value],
    )
    return int(cursor.fetchone()[0])


def build_references(
    cursor,
    discovered: dict[tuple[str, str], set[str]],
    source_id: int,
    target_id: int,
) -> tuple[Reference, ...]:
    return tuple(
        Reference(
            table=table,
            column=column,
            sources=tuple(sorted(sources)),
            source_rows=count_value(cursor, table, column, source_id),
            target_rows=count_value(cursor, table, column, target_id),
        )
        for (table, column), sources in sorted(discovered.items())
    )


def discover_candidates(
    cursor,
    reviewed_pairs: set[tuple[str, str]],
    ignored_snapshot_pairs: set[tuple[str, str]],
    source_id: int,
) -> tuple[tuple[Candidate, ...], tuple[Candidate, ...]]:
    cursor.execute(
        f"""
        SELECT c.TABLE_NAME, c.COLUMN_NAME
        FROM information_schema.COLUMNS AS c
        JOIN information_schema.TABLES AS tables_metadata
          ON tables_metadata.TABLE_SCHEMA = c.TABLE_SCHEMA
         AND tables_metadata.TABLE_NAME = c.TABLE_NAME
         AND tables_metadata.TABLE_TYPE = 'BASE TABLE'
        WHERE c.TABLE_SCHEMA = DATABASE()
          AND c.COLUMN_NAME IN ({make_placeholders(CANDIDATE_COLUMN_NAMES)})
          AND c.TABLE_NAME <> %s
        ORDER BY c.TABLE_NAME, c.COLUMN_NAME
        """,
        [*CANDIDATE_COLUMN_NAMES, PROJECT_TABLE],
    )
    unreviewed: list[Candidate] = []
    ignored_snapshots: list[Candidate] = []
    for table, column in cursor.fetchall():
        pair = (str(table), str(column))
        if pair not in reviewed_pairs:
            candidate = Candidate(
                pair[0], pair[1], count_value(cursor, *pair, source_id)
            )
            if pair in ignored_snapshot_pairs:
                ignored_snapshots.append(candidate)
            else:
                unreviewed.append(candidate)
    return tuple(unreviewed), tuple(ignored_snapshots)


def unique_indexes(
    cursor, table: str
) -> dict[str, list[tuple[str | None, int | None]]]:
    cursor.execute(
        """
        SELECT INDEX_NAME, COLUMN_NAME, SUB_PART
        FROM information_schema.STATISTICS
        WHERE TABLE_SCHEMA = DATABASE()
          AND TABLE_NAME = %s
          AND NON_UNIQUE = 0
        ORDER BY INDEX_NAME, SEQ_IN_INDEX
        """,
        [table],
    )
    indexes: dict[str, list[tuple[str | None, int | None]]] = defaultdict(list)
    for index_name, column_name, sub_part in cursor.fetchall():
        indexes[str(index_name)].append(
            (
                str(column_name) if column_name is not None else None,
                int(sub_part) if sub_part is not None else None,
            )
        )
    return dict(indexes)


def count_unique_conflict_groups(
    cursor,
    table: str,
    index_columns: Sequence[str],
    project_columns: set[str],
    source_id: int,
    target_id: int,
) -> int:
    select_parts: list[str] = []
    params: list[int] = []
    for position, column in enumerate(index_columns, start=1):
        identifier = quote_identifier(column)
        if column in project_columns:
            select_parts.append(
                f"IF({identifier} = %s, %s, {identifier}) AS k{position}"
            )
            params.extend([source_id, target_id])
        else:
            select_parts.append(f"{identifier} AS k{position}")
    source_predicate = " OR ".join(
        f"{quote_identifier(column)} = %s" for column in sorted(project_columns)
    )
    params.extend([source_id] * len(project_columns))
    key_names = [f"k{position}" for position in range(1, len(index_columns) + 1)]
    not_null = " AND ".join(f"{name} IS NOT NULL" for name in key_names)
    grouped_columns = ", ".join(key_names)
    cursor.execute(
        f"""
        SELECT COUNT(*)
        FROM (
            SELECT {grouped_columns}
            FROM (
                SELECT {", ".join(select_parts)},
                       IF({source_predicate}, 1, 0) AS source_marker
                FROM {quote_identifier(table)}
            ) AS mapped_rows
            WHERE {not_null}
            GROUP BY {grouped_columns}
            HAVING COUNT(*) > 1 AND SUM(source_marker) > 0
        ) AS conflicting_groups
        """,
        params,
    )
    return int(cursor.fetchone()[0])


def build_unique_checks(
    cursor,
    references: Sequence[Reference],
    source_id: int,
    target_id: int,
) -> tuple[UniqueCheck, ...]:
    project_columns_by_table: dict[str, set[str]] = defaultdict(set)
    source_rows_by_table: dict[str, int] = defaultdict(int)
    for reference in references:
        project_columns_by_table[reference.table].add(reference.column)
        source_rows_by_table[reference.table] += reference.source_rows

    checks: list[UniqueCheck] = []
    for table, project_columns in sorted(project_columns_by_table.items()):
        if not source_rows_by_table[table]:
            continue
        for index_name, metadata in sorted(unique_indexes(cursor, table).items()):
            columns = tuple(column for column, _ in metadata if column is not None)
            if not project_columns.intersection(columns):
                continue
            if any(column is None for column, _ in metadata):
                checks.append(
                    UniqueCheck(
                        table,
                        index_name,
                        columns,
                        None,
                        "Functional unique index cần review thủ công.",
                    )
                )
                continue
            if any(sub_part is not None for _, sub_part in metadata):
                checks.append(
                    UniqueCheck(
                        table,
                        index_name,
                        columns,
                        None,
                        "Prefix unique index cần review thủ công.",
                    )
                )
                continue
            conflicts = count_unique_conflict_groups(
                cursor,
                table,
                columns,
                project_columns,
                source_id,
                target_id,
            )
            checks.append(UniqueCheck(table, index_name, columns, conflicts))
    return tuple(checks)


def snapshot_tables(
    cursor,
    references: Sequence[Reference],
    source_id: int,
    row_limit: int,
) -> tuple[TableSnapshot, ...]:
    columns_by_table: dict[str, set[str]] = defaultdict(set)
    for reference in references:
        if reference.source_rows:
            columns_by_table[reference.table].add(reference.column)

    snapshots: list[TableSnapshot] = []
    remaining = row_limit
    for table, matched_columns in sorted(columns_by_table.items()):
        predicate = " OR ".join(
            f"{quote_identifier(column)} = %s" for column in sorted(matched_columns)
        )
        params = [source_id] * len(matched_columns)
        cursor.execute(
            f"SELECT COUNT(*) FROM {quote_identifier(table)} WHERE {predicate}",
            params,
        )
        total_rows = int(cursor.fetchone()[0])
        rows: list[dict[str, Any]] = []
        if remaining > 0:
            columns = table_columns(cursor, table)
            selected = ", ".join(quote_identifier(column) for column in columns)
            cursor.execute(
                f"SELECT {selected} FROM {quote_identifier(table)} "
                f"WHERE {predicate} LIMIT %s",
                [*params, remaining],
            )
            rows = [dict(zip(columns, row, strict=True)) for row in cursor.fetchall()]
            remaining -= len(rows)
        snapshots.append(
            TableSnapshot(table, tuple(sorted(matched_columns)), total_rows, rows)
        )
    return tuple(snapshots)


def build_plan(
    cursor,
    source_id: int,
    target_id: int,
    *,
    row_limit: int,
    lock_projects: bool = False,
    include_snapshots: bool = True,
) -> MergePlan:
    source_project = fetch_project(cursor, source_id, lock=lock_projects)
    target_project = fetch_project(cursor, target_id, lock=lock_projects)
    discovered = discover_references(cursor)
    references = build_references(cursor, discovered, source_id, target_id)
    ignored_snapshot_pairs = {
        (table, column)
        for table, columns in IGNORED_SNAPSHOT_PROJECT_COLUMNS.items()
        for column in columns
    }
    unreviewed_candidates, ignored_snapshots = discover_candidates(
        cursor, set(discovered), ignored_snapshot_pairs, source_id
    )
    plan = MergePlan(
        source_id=source_id,
        target_id=target_id,
        source_project=source_project,
        target_project=target_project,
        project_differences=compare_projects(source_project, target_project),
        references=references,
        unique_checks=build_unique_checks(cursor, references, source_id, target_id),
        unreviewed_candidates=unreviewed_candidates,
        ignored_snapshots=ignored_snapshots,
    )
    if include_snapshots:
        plan.snapshots = snapshot_tables(cursor, references, source_id, row_limit)
    return plan


def update_table_references(
    cursor,
    table: str,
    columns: Sequence[str],
    source_id: int,
    target_id: int,
) -> int:
    assignments = ", ".join(
        f"{quote_identifier(column)} = IF({quote_identifier(column)} = %s, %s, "
        f"{quote_identifier(column)})"
        for column in columns
    )
    predicate = " OR ".join(f"{quote_identifier(column)} = %s" for column in columns)
    count_params = [source_id] * len(columns)
    expected = count_table_reference_rows(cursor, table, columns, source_id, lock=True)
    update_params: list[int] = []
    for _ in columns:
        update_params.extend([source_id, target_id])
    update_params.extend(count_params)
    cursor.execute(
        f"UPDATE {quote_identifier(table)} SET {assignments} WHERE {predicate}",
        update_params,
    )
    actual = int(cursor.rowcount)
    if actual != expected:
        raise RuntimeError(
            f"{table}: số dòng update không đúng; expected={expected}, actual={actual}."
        )
    return actual


def count_table_reference_rows(
    cursor,
    table: str,
    columns: Sequence[str],
    source_id: int,
    *,
    lock: bool = False,
) -> int:
    predicate = " OR ".join(f"{quote_identifier(column)} = %s" for column in columns)
    lock_sql = " FOR UPDATE" if lock else ""
    cursor.execute(
        f"SELECT COUNT(*) FROM {quote_identifier(table)} WHERE {predicate}{lock_sql}",
        [source_id] * len(columns),
    )
    return int(cursor.fetchone()[0])


def discover_cascade_update_dependencies(
    cursor, tables: Sequence[str]
) -> tuple[tuple[str, str], ...]:
    """Return parent/child table pairs whose project key updates cascade."""
    ordered_tables = sorted(set(tables))
    if not ordered_tables:
        return ()
    placeholders = make_placeholders(ordered_tables)
    cursor.execute(
        f"""
        SELECT DISTINCT kcu.REFERENCED_TABLE_NAME, kcu.TABLE_NAME
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        JOIN information_schema.REFERENTIAL_CONSTRAINTS AS rc
          ON rc.CONSTRAINT_SCHEMA = kcu.CONSTRAINT_SCHEMA
         AND rc.TABLE_NAME = kcu.TABLE_NAME
         AND rc.CONSTRAINT_NAME = kcu.CONSTRAINT_NAME
        WHERE kcu.TABLE_SCHEMA = DATABASE()
          AND kcu.TABLE_NAME IN ({placeholders})
          AND kcu.REFERENCED_TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_NAME IN ({placeholders})
          AND rc.UPDATE_RULE = 'CASCADE'
        ORDER BY kcu.REFERENCED_TABLE_NAME, kcu.TABLE_NAME
        """,
        [*ordered_tables, *ordered_tables],
    )
    return tuple((str(parent), str(child)) for parent, child in cursor.fetchall())


def order_tables_parent_first(
    tables: Iterable[str], dependencies: Iterable[tuple[str, str]]
) -> tuple[str, ...]:
    """Topologically order tables so a cascading parent is updated first."""
    remaining = set(tables)
    dependencies = tuple(
        (parent, child)
        for parent, child in dependencies
        if parent in remaining and child in remaining and parent != child
    )
    ordered: list[str] = []
    while remaining:
        children_with_pending_parent = {
            child
            for parent, child in dependencies
            if parent in remaining and child in remaining
        }
        ready = sorted(remaining - children_with_pending_parent)
        if not ready:
            raise RuntimeError(
                "Không thể xác định thứ tự update do vòng lặp foreign key: "
                + ", ".join(sorted(remaining))
            )
        ordered.extend(ready)
        remaining.difference_update(ready)
    return tuple(ordered)


def execute_plan(cursor, plan: MergePlan) -> dict[str, int]:
    columns_by_table: dict[str, set[str]] = defaultdict(set)
    for reference in plan.references:
        columns_by_table[reference.table].add(reference.column)

    dependencies = discover_cascade_update_dependencies(cursor, tuple(columns_by_table))
    table_order = order_tables_parent_first(columns_by_table, dependencies)
    planned_updates = {
        table: count_table_reference_rows(
            cursor, table, sorted(columns), plan.source_id
        )
        for table, columns in columns_by_table.items()
    }

    updated: dict[str, int] = {}
    for table in table_order:
        columns = columns_by_table[table]
        update_table_references(
            cursor,
            table,
            sorted(columns),
            plan.source_id,
            plan.target_id,
        )
        updated[table] = planned_updates[table]

    remaining: dict[str, int] = {}
    for table, columns in sorted(columns_by_table.items()):
        for column in sorted(columns):
            count = count_value(cursor, table, column, plan.source_id)
            if count:
                remaining[f"{table}.{column}"] = count
    for candidate in plan.unreviewed_candidates:
        count = count_value(cursor, candidate.table, candidate.column, plan.source_id)
        if count:
            remaining[f"{candidate.table}.{candidate.column}"] = count
    if remaining:
        raise RuntimeError(f"Vẫn còn tham chiếu source ID; rollback: {remaining}")

    cursor.execute(
        f"DELETE FROM {quote_identifier(PROJECT_TABLE)} "
        f"WHERE {quote_identifier(PROJECT_KEY)} = %s",
        [plan.source_id],
    )
    if cursor.rowcount != 1:
        raise RuntimeError(
            f"Phải xóa đúng 1 dòng {PROJECT_TABLE}; actual={cursor.rowcount}."
        )
    cursor.execute(
        f"SELECT COUNT(*) FROM {quote_identifier(PROJECT_TABLE)} "
        f"WHERE {quote_identifier(PROJECT_KEY)} IN (%s, %s)",
        [plan.source_id, plan.target_id],
    )
    if int(cursor.fetchone()[0]) != 1:
        raise RuntimeError("Xác minh project_info sau merge thất bại; rollback.")
    return updated


def write_report(
    path: Path,
    plan: MergePlan,
    *,
    status: str,
    database: str,
    host: str,
    execute_requested: bool,
    updated_counts: dict[str, int] | None = None,
    error: str | None = None,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Excel reporting requires openpyxl. Install project requirements first."
        ) from exc

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    projects = workbook.create_sheet("Projects")
    references = workbook.create_sheet("References")
    unique_checks = workbook.create_sheet("Unique checks")
    unreviewed = workbook.create_sheet("Unreviewed")
    ignored_snapshots = workbook.create_sheet("Ignored snapshots")
    source_rows = workbook.create_sheet("Source rows")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", color="FFFFFF", bold=True)
    title_font = Font(name="Arial", size=16, bold=True)
    body_font = Font(name="Arial", size=10)
    status_fill = {
        "PREVIEW": "FFF2CC",
        "READY_TO_EXECUTE": "FCE5CD",
        "COMMITTED": "C6EFCE",
        "CANCELLED": "E7E6E6",
        "BLOCKED": "F4CCCC",
        "ROLLED_BACK": "F4CCCC",
    }.get(status, "E7E6E6")

    def append(sheet, values: Iterable[Any]) -> None:
        sheet.append([safe_excel_value(value) for value in values])

    def finish(sheet, *, freeze: str = "A2") -> None:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = freeze
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_index in range(1, sheet.max_column + 1):
            letter = get_column_letter(column_index)
            width = max(
                (len(str(cell.value or "")) for cell in sheet[letter]), default=0
            )
            sheet.column_dimensions[letter].width = min(max(width + 2, 12), 80)

    generated_at = datetime.now().astimezone()
    summary.merge_cells("A1:D1")
    summary["A1"] = "Project ID merge report"
    summary["A1"].font = title_font
    summary.append([])
    summary.append(["Report field", "Value"])
    metadata = [
        ("Status", status),
        ("Mode", "EXECUTE" if execute_requested else "PREVIEW"),
        ("Database", database),
        ("Host", host),
        ("Generated at", generated_at.replace(tzinfo=None)),
        ("Timezone", str(generated_at.tzinfo)),
        ("Source project ID", plan.source_id),
        ("Target project ID", plan.target_id),
        ("Reviewed reference columns", len(plan.references)),
        ("Source reference rows by column", plan.source_reference_rows),
        ("Project differences (information only)", len(plan.project_differences)),
        ("Ignored snapshot columns", len(plan.ignored_snapshots)),
        (
            "Ignored snapshot source rows",
            sum(candidate.source_rows for candidate in plan.ignored_snapshots),
        ),
        ("Blocking issues", len(plan.blockers)),
        ("Captured source rows", sum(len(item.rows) for item in plan.snapshots)),
        (
            "External resources",
            "Database rows only; S3/filesystem paths are not renamed by this tool.",
        ),
    ]
    if updated_counts is not None:
        metadata.append(("Updated rows", sum(updated_counts.values())))
        metadata.append(
            ("Deleted source project rows", 1 if status == "COMMITTED" else 0)
        )
    if error:
        metadata.append(("Error", error))
    for key, value in metadata:
        append(summary, [key, value])
    if plan.blockers:
        summary.append([])
        append(summary, ["Blocking issue", "Detail"])
        for position, blocker in enumerate(plan.blockers, start=1):
            append(summary, [position, blocker])
    summary["B4"].fill = PatternFill("solid", fgColor=status_fill)
    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "A4"
    for cell in summary[3]:
        cell.fill = header_fill
        cell.font = header_font
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 90
    for row in summary.iter_rows(min_row=4):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    append(projects, ["Column", "Source value", "Target value", "Match"])
    for column in plan.source_project:
        source_value = plan.source_project.get(column)
        target_value = plan.target_project.get(column)
        append(
            projects,
            [
                column,
                source_value,
                target_value,
                "YES" if source_value == target_value else "NO",
            ],
        )

    append(
        references,
        [
            "Table",
            "Column",
            "Relationship source",
            "Source rows",
            "Target rows",
            "Planned action",
            "Result",
        ],
    )
    updated_counts = updated_counts or {}
    for reference in plan.references:
        result = (
            "UPDATED" if status == "COMMITTED" and reference.source_rows else status
        )
        append(
            references,
            [
                reference.table,
                reference.column,
                ", ".join(reference.sources),
                reference.source_rows,
                reference.target_rows,
                f"UPDATE {plan.source_id} TO {plan.target_id}",
                result,
            ],
        )

    append(
        unique_checks,
        ["Table", "Unique index", "Columns", "Conflict groups", "Note"],
    )
    for check in plan.unique_checks:
        append(
            unique_checks,
            [
                check.table,
                check.index,
                ", ".join(check.columns),
                check.conflict_groups
                if check.conflict_groups is not None
                else "NOT CHECKED",
                check.note,
            ],
        )

    append(unreviewed, ["Table", "Column", "Source rows", "Execution status"])
    for candidate in plan.unreviewed_candidates:
        append(
            unreviewed,
            [
                candidate.table,
                candidate.column,
                candidate.source_rows,
                "BLOCKED" if candidate.source_rows else "NO SOURCE ROWS",
            ],
        )

    append(
        ignored_snapshots,
        ["Table", "Column", "Source rows", "Planned action", "Reason"],
    )
    for candidate in plan.ignored_snapshots:
        append(
            ignored_snapshots,
            [
                candidate.table,
                candidate.column,
                candidate.source_rows,
                "KEEP UNCHANGED",
                "Reviewed backup/recovery snapshot",
            ],
        )

    append(
        source_rows,
        ["Table", "Matched columns", "Row before merge", "Planned action", "Result"],
    )
    for snapshot in plan.snapshots:
        result = "UPDATED" if status == "COMMITTED" else status
        for row in snapshot.rows:
            append(
                source_rows,
                [
                    snapshot.table,
                    ", ".join(snapshot.matched_columns),
                    row,
                    f"REPLACE {plan.source_id} WITH {plan.target_id}",
                    result,
                ],
            )
        if snapshot.truncated:
            append(
                source_rows,
                [
                    snapshot.table,
                    ", ".join(snapshot.matched_columns),
                    f"TRUNCATED: captured {len(snapshot.rows)} of {snapshot.total_rows} rows",
                    "REVIEW SUMMARY COUNTS",
                    status,
                ],
            )

    for sheet in (
        projects,
        references,
        unique_checks,
        unreviewed,
        ignored_snapshots,
        source_rows,
    ):
        finish(sheet)
    source_rows.column_dimensions["C"].width = 100

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def print_plan(plan: MergePlan, database: str, host: str, report_path: Path) -> None:
    print(f"Database: {database}")
    print(f"Host: {host}")
    print(f"Merge: {plan.source_id} -> {plan.target_id}")
    print(f"Reviewed reference columns: {len(plan.references)}")
    for reference in plan.references:
        if reference.source_rows or reference.target_rows:
            print(
                f"- {reference.table}.{reference.column}: "
                f"source={reference.source_rows}, target={reference.target_rows}, "
                f"via={','.join(reference.sources)}"
            )
    if plan.ignored_snapshots:
        print("Ignored snapshots (kept unchanged):")
        for candidate in plan.ignored_snapshots:
            if candidate.source_rows:
                print(
                    f"- {candidate.table}.{candidate.column}: "
                    f"source={candidate.source_rows}, action=KEEP UNCHANGED"
                )
    if plan.blockers:
        print("Blocking issues:")
        for blocker in plan.blockers:
            print(f"- {blocker}")
    print(f"Excel report: {report_path}")


def initialize_django():
    # When copied under <backend>/utils/Tools*, parents[2] is the Django root.
    # A PyCharm content root or SPDM_PROJECT_ROOT may also provide the project.
    configured_root = RUN_PROJECT_ROOT or os.environ.get("SPDM_PROJECT_ROOT")
    candidates = [Path(configured_root).expanduser()] if configured_root else []
    candidates.extend([Path.cwd(), Path(__file__).resolve().parents[2]])
    for candidate in candidates:
        if (candidate / "backend" / "settings.py").is_file():
            if str(candidate) not in sys.path:
                sys.path.insert(0, str(candidate))
            break
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    import django

    django.setup()
    from django.db import connection, transaction

    return connection, transaction


def run_self_checks() -> None:
    assert validate_ids("1014", 1000) == (1014, 1000)
    assert make_placeholders([1, 2, 3]) == "%s, %s, %s"
    assert quote_identifier("project_info_id") == "`project_info_id`"
    assert safe_excel_value("=danger") == "'=danger"
    assert order_tables_parent_first(
        {"project_parameter_attachments", "project_r_parameter", "temp_data"},
        {("project_r_parameter", "project_parameter_attachments")},
    ) == ("project_r_parameter", "temp_data", "project_parameter_attachments")
    try:
        order_tables_parent_first(
            {"parent", "child"}, {("parent", "child"), ("child", "parent")}
        )
    except RuntimeError:
        pass
    else:  # pragma: no cover - executable assertion
        raise AssertionError("Cyclic update dependencies were accepted")
    matching_plan = MergePlan(
        source_id=1014,
        target_id=1000,
        source_project={"id": 1014, "project_code": "P62C"},
        target_project={"id": 1000, "project_code": "P62C"},
        project_differences={},
        references=(),
        unique_checks=(),
        unreviewed_candidates=(),
    )
    assert not matching_plan.blockers
    blocked_plan = MergePlan(
        source_id=1014,
        target_id=1000,
        source_project={"id": 1014, "project_code": "P62C"},
        target_project={"id": 1000, "project_code": "OTHER"},
        project_differences={"project_code": ("P62C", "OTHER")},
        references=(),
        unique_checks=(UniqueCheck("child", "uniq", ("project_id",), 1),),
        unreviewed_candidates=(Candidate("unknown", "project_id", 2),),
        ignored_snapshots=(Candidate("backup", "project_id", 3),),
    )
    assert len(blocked_plan.blockers) == 2
    for invalid in (0, -1, True, "x"):
        try:
            normalize_positive_id(invalid, label="id")
        except ValueError:
            pass
        else:  # pragma: no cover - executable assertion
            raise AssertionError(f"Expected invalid ID: {invalid!r}")
    try:
        quote_identifier("bad-name")
    except ValueError:
        pass
    else:  # pragma: no cover - executable assertion
        raise AssertionError("Unsafe identifier was accepted")
    print("Self-checks passed; no database connection was opened.")


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Preview or atomically replace every reviewed project ID reference, "
            "delete the duplicate source project, and write an Excel report."
        )
    )
    parser.add_argument("--source-id", type=int, default=RUN_SOURCE_PROJECT_ID)
    parser.add_argument("--target-id", type=int, default=RUN_TARGET_PROJECT_ID)
    parser.add_argument(
        "--execute",
        action="store_true",
        default=RUN_EXECUTE,
        help="Execute after exact confirmation; default is preview only.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        default=RUN_ASSUME_YES,
        help="Skip interactive confirmation (only for controlled automation).",
    )
    parser.add_argument("--output", default=RUN_OUTPUT_PATH)
    parser.add_argument("--report-row-limit", type=int, default=RUN_REPORT_ROW_LIMIT)
    parser.add_argument(
        "--self-test",
        action="store_true",
        help="Run pure helper checks without Django or a database connection.",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    if args.self_test:
        run_self_checks()
        return 0

    source_id, target_id = validate_ids(args.source_id, args.target_id)
    row_limit = normalize_report_row_limit(args.report_row_limit)
    if args.yes and not args.execute:
        raise ValueError("--yes/RUN_ASSUME_YES chỉ hợp lệ khi bật execute.")
    report_path = resolve_output_path(args.output, source_id, target_id)
    connection, transaction = initialize_django()

    with connection.cursor() as cursor:
        database, host = database_identity(cursor, connection)
        plan = build_plan(cursor, source_id, target_id, row_limit=row_limit)

    initial_status = "BLOCKED" if plan.blockers else "PREVIEW"
    write_report(
        report_path,
        plan,
        status=initial_status,
        database=database,
        host=host,
        execute_requested=args.execute,
    )
    print_plan(plan, database, host, report_path)

    if plan.blockers:
        print("BLOCKED: không có dữ liệu nào bị thay đổi.")
        return 2
    if not args.execute:
        print(
            "PREVIEW: không có dữ liệu nào bị thay đổi. Review Excel trước khi execute."
        )
        return 0

    confirmation = f"MERGE_PROJECT_ID {database}:{source_id}->{target_id}"
    if not args.yes:
        print(f"Nhập chính xác '{confirmation}' để tiếp tục:")
        if input("> ").strip() != confirmation:
            write_report(
                report_path,
                plan,
                status="CANCELLED",
                database=database,
                host=host,
                execute_requested=True,
            )
            print("CANCELLED: không có dữ liệu nào bị thay đổi.")
            return 1

    # Preserve an audited pre-mutation file even if a later report save fails.
    write_report(
        report_path,
        plan,
        status="READY_TO_EXECUTE",
        database=database,
        host=host,
        execute_requested=True,
    )
    try:
        with transaction.atomic(using=connection.alias):
            with connection.cursor() as cursor:
                current_identity = database_identity(cursor, connection)
                if current_identity != (database, host):
                    raise RuntimeError("Database/host đã đổi sau preview.")
                locked_plan = build_plan(
                    cursor,
                    source_id,
                    target_id,
                    row_limit=0,
                    lock_projects=True,
                    include_snapshots=False,
                )
                if locked_plan.blockers:
                    raise RuntimeError(
                        "Plan mới bị block: " + " | ".join(locked_plan.blockers)
                    )
                if locked_plan.fingerprint() != plan.fingerprint():
                    raise RuntimeError(
                        "Dữ liệu/schema đã thay đổi sau preview; hãy chạy preview lại."
                    )
                updated_counts = execute_plan(cursor, locked_plan)
    except Exception as exc:
        write_report(
            report_path,
            plan,
            status="ROLLED_BACK",
            database=database,
            host=host,
            execute_requested=True,
            error=str(exc),
        )
        print(f"ROLLED_BACK: {exc}")
        print(f"Excel report: {report_path}")
        return 3

    write_report(
        report_path,
        plan,
        status="COMMITTED",
        database=database,
        host=host,
        execute_requested=True,
        updated_counts=updated_counts,
    )
    print(f"COMMITTED: đã chuyển {source_id} -> {target_id} và xóa project nguồn.")
    print(f"Excel report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
