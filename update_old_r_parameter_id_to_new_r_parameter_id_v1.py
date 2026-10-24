"""Preview or execute multiple R parameter ID remap groups atomically.

Examples::

    [1, 2, 3] -> 4
    [5, 6] -> 7

The sibling ``update_old_r_parameter_id_to_new_r_parameter_id`` module owns
the schema-aware plan for one group. This module coordinates many disjoint
groups in one preview, one confirmation, and one database transaction. Every
run also writes a detailed Excel audit report without querying the database a
second time merely for reporting.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Sequence


PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import os

    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402

from utils.Tools_new import (  # noqa: E402
    update_old_r_parameter_id_to_new_r_parameter_id as remap_core,
)

# import pandas as pd  # noqa: E402
#
# file_path = "/Users/TrongKieu/Downloads/R_List_Master_重複統合.xlsx"
# df = pd.read_excel(
#     file_path, usecols=["変更前 r_parameter_id", "変更後 r_parameter_id"]
# )
# RUN_REMAP_GROUPS = [
#     ([int(old_id)], int(new_id))
#     for old_id, new_id in zip(df["変更前 r_parameter_id"], df["変更後 r_parameter_id"])
# ]
# print(RUN_REMAP_GROUPS)
# PyCharm Run-button configuration. Preview before enabling execution.
RUN_REMAP_GROUPS: list[tuple[list[int], int]] = [
    ([290145], 9999),
    ([10041], 10016),
]
RUN_EXECUTE = False
RUN_ASSUME_YES = False
# Hard-delete source r_parameter rows only after every reference is remapped
# and verified inside the same transaction.
RUN_DELETE_REPLACED_PARAMETERS = False
# Separate opt-in because hard deletion is more destructive than ID remapping.
RUN_ASSUME_DELETE_YES = False
# ``None`` writes reports under ``<project>/outputs/r_parameter_remap_reports``.
RUN_REPORT_DIRECTORY: str | None = None

DEFAULT_REPORT_DIRECTORY = PROJECT_ROOT / "outputs" / "r_parameter_remap_reports"
EXCEL_CELL_TEXT_LIMIT = 32_700


@dataclass(frozen=True)
class RemapGroup:
    old_ids: tuple[int, ...]
    new_id: int

    @property
    def mappings(self) -> tuple[tuple[int, int], ...]:
        return tuple((old_id, self.new_id) for old_id in self.old_ids)

    def display(self) -> str:
        return f"[{','.join(map(str, self.old_ids))}]->{self.new_id}"


@dataclass(frozen=True)
class BatchPlan:
    groups: tuple[RemapGroup, ...]
    group_plans: tuple[remap_core.MigrationPlan, ...]
    master_list_changes: tuple[remap_core.MasterListChange, ...]
    conflicts: tuple[str, ...]

    @property
    def mutation_count(self) -> int:
        return sum(plan.mutation_count for plan in self.group_plans) + len(
            self.master_list_changes
        )

    @property
    def rfl_id_map(self) -> dict[int, int]:
        return {
            old_rfl_id: new_rfl_id
            for plan in self.group_plans
            for old_rfl_id, new_rfl_id in plan.rfl_merges
        }


def replaced_parameter_ids(plan: BatchPlan) -> tuple[int, ...]:
    """Return unique source R parameter IDs in stable batch order."""
    return tuple(
        dict.fromkeys(old_id for group in plan.groups for old_id in group.old_ids)
    )


def estimated_batch_mutations(
    plan: BatchPlan, *, delete_replaced_parameters: bool
) -> int:
    delete_count = (
        len(replaced_parameter_ids(plan)) if delete_replaced_parameters else 0
    )
    return plan.mutation_count + delete_count


def make_report_path(report_directory: str | Path | None = None) -> Path:
    """Reserve a collision-resistant path without touching database state."""
    output_directory = (
        Path(report_directory).expanduser()
        if report_directory is not None
        else DEFAULT_REPORT_DIRECTORY
    )
    output_directory.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
    return output_directory / f"r_parameter_remap_{timestamp}.xlsx"


def safe_excel_value(value: Any) -> Any:
    """Return a typed Excel value while preventing formula injection."""
    if value is None or isinstance(
        value, (bool, int, float, Decimal, date, datetime, time)
    ):
        return value
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    elif not isinstance(value, str):
        value = str(value)
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def split_excel_text(value: str) -> tuple[str, ...]:
    """Preserve long JSON values by splitting them across Excel rows."""
    if not value:
        return ("",)
    return tuple(
        value[offset : offset + EXCEL_CELL_TEXT_LIMIT]
        for offset in range(0, len(value), EXCEL_CELL_TEXT_LIMIT)
    )


def operation_summary_rows(
    plan: BatchPlan, *, delete_replaced_parameters: bool
) -> list[tuple[str, int]]:
    group_plans = plan.group_plans
    parameter_delete_count = (
        len(replaced_parameter_ids(plan)) if delete_replaced_parameters else 0
    )
    return [
        (
            "Direct R parameter reference updates",
            sum(
                reference_plan.row_count
                for group_plan in group_plans
                for reference_plan in group_plan.direct_reference_plans
            ),
        ),
        (
            "setup_r_relation updates",
            sum(len(group_plan.setup_relation_updates) for group_plan in group_plans),
        ),
        (
            "setup_r_relation merges",
            sum(len(group_plan.setup_relation_merges) for group_plan in group_plans),
        ),
        ("rfl rewires", sum(len(group_plan.rfl_updates) for group_plan in group_plans)),
        ("rfl merges", sum(len(group_plan.rfl_merges) for group_plan in group_plans)),
        (
            "RFL reference updates",
            sum(
                reference_plan.row_count
                for group_plan in group_plans
                for reference_plan in group_plan.rfl_reference_plans
            ),
        ),
        (
            "prj_rfl duplicate deletes",
            sum(len(group_plan.prj_rfl_deletes) for group_plan in group_plans),
        ),
        (
            "prj_rfl updates",
            sum(len(group_plan.prj_rfl_updates) for group_plan in group_plans),
        ),
        (
            "setup relation reference updates",
            sum(
                reference_plan.row_count
                for group_plan in group_plans
                for reference_plan in group_plan.setup_reference_plans
            ),
        ),
        (
            "l_r_relation duplicate deletes",
            sum(len(group_plan.l_r_relation_deletes) for group_plan in group_plans),
        ),
        (
            "l_r_relation updates",
            sum(len(group_plan.l_r_relation_updates) for group_plan in group_plans),
        ),
        ("master_list_order JSON updates", len(plan.master_list_changes)),
        ("r_parameter source deletes", parameter_delete_count),
        (
            "Total estimated mutations",
            estimated_batch_mutations(
                plan, delete_replaced_parameters=delete_replaced_parameters
            ),
        ),
    ]


def relation_change_rows(
    plan: BatchPlan, *, delete_replaced_parameters: bool
) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    for group_index, (group, group_plan) in enumerate(
        zip(plan.groups, plan.group_plans, strict=True), start=1
    ):
        rows.extend(
            (
                group_index,
                "setup_r_relation",
                "UPDATE_R_PARAMETER_ID",
                relation_id,
                None,
                group_plan.new_id,
                "",
            )
            for relation_id in group_plan.setup_relation_updates
        )
        rows.extend(
            (
                group_index,
                "setup_r_relation",
                "MERGE",
                old_relation_id,
                old_relation_id,
                new_relation_id,
                "Source relation is merged into target relation",
            )
            for old_relation_id, new_relation_id in group_plan.setup_relation_merges
        )
        rows.extend(
            (
                group_index,
                "rfl",
                "REWIRE_REQUIREMENT",
                rfl_id,
                old_relation_id,
                new_relation_id,
                "requirement_s_id",
            )
            for rfl_id, old_relation_id, new_relation_id in group_plan.rfl_updates
        )
        rows.extend(
            (
                group_index,
                "rfl",
                "MERGE",
                old_rfl_id,
                old_rfl_id,
                new_rfl_id,
                "Source RFL is merged into target RFL",
            )
            for old_rfl_id, new_rfl_id in group_plan.rfl_merges
        )
        rows.extend(
            (
                group_index,
                "prj_rfl",
                "DELETE_DUPLICATE",
                None,
                rfl_id,
                None,
                f"project_info_id={project_id}; phase_id={phase_id}",
            )
            for project_id, rfl_id, phase_id in group_plan.prj_rfl_deletes
        )
        rows.extend(
            (
                group_index,
                "prj_rfl",
                "UPDATE_RFL_ID",
                None,
                old_rfl_id,
                new_rfl_id,
                f"project_info_id={project_id}; phase_id={phase_id}",
            )
            for project_id, phase_id, old_rfl_id, new_rfl_id in (
                group_plan.prj_rfl_updates
            )
        )
        rows.extend(
            (
                group_index,
                "l_r_relation",
                "DELETE_DUPLICATE",
                relation_row_id,
                None,
                None,
                "",
            )
            for relation_row_id in group_plan.l_r_relation_deletes
        )
        rows.extend(
            (
                group_index,
                "l_r_relation",
                "UPDATE_R_S_ID",
                relation_row_id,
                old_relation_id,
                new_relation_id,
                "",
            )
            for relation_row_id, old_relation_id, new_relation_id in (
                group_plan.l_r_relation_updates
            )
        )
        if delete_replaced_parameters:
            rows.extend(
                (
                    group_index,
                    "r_parameter",
                    "DELETE_REPLACED_PARAMETER",
                    old_id,
                    old_id,
                    group.new_id,
                    "Delete source after all references pass verification",
                )
                for old_id in group.old_ids
            )
    return rows


def write_batch_report(
    report_path: Path,
    plan: BatchPlan,
    *,
    status: str,
    database: str,
    host: str,
    execute_requested: bool,
    delete_replaced_parameters: bool,
    error: str | None = None,
) -> None:
    """Write a styled, filterable Excel audit workbook from an existing plan."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependency is project-pinned
        raise RuntimeError(
            "Excel reporting requires openpyxl. Install project requirements first."
        ) from exc

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    parameters_sheet = workbook.create_sheet("Parameters")
    references_sheet = workbook.create_sheet("References")
    relations_sheet = workbook.create_sheet("Relation Changes")
    master_list_sheet = workbook.create_sheet("Master List")
    conflicts_sheet = workbook.create_sheet("Conflicts")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", color="FFFFFF", bold=True)
    title_font = Font(name="Arial", size=16, bold=True, color="1F1F1F")
    body_font = Font(name="Arial", size=10)
    thin_gray = Side(style="thin", color="D9E2F3")
    table_border = Border(bottom=thin_gray)
    status_fills = {
        "COMMITTED": PatternFill("solid", fgColor="C6EFCE"),
        "PREVIEW": PatternFill("solid", fgColor="FFF2CC"),
        "NO_CHANGES": PatternFill("solid", fgColor="D9EAD3"),
        "READY_TO_EXECUTE": PatternFill("solid", fgColor="FCE5CD"),
        "CANCELLED": PatternFill("solid", fgColor="E7E6E6"),
        "BLOCKED": PatternFill("solid", fgColor="F4CCCC"),
        "ROLLED_BACK": PatternFill("solid", fgColor="F4CCCC"),
    }

    def style_header(sheet: Any, row_number: int) -> None:
        for cell in sheet[row_number]:
            if cell.value is not None:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")

    def append_row(sheet: Any, values: Iterable[Any]) -> None:
        sheet.append([safe_excel_value(value) for value in values])

    def finish_detail_sheet(sheet: Any) -> None:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        style_header(sheet, 1)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = table_border
        for column_index in range(1, sheet.max_column + 1):
            letter = get_column_letter(column_index)
            longest = max(
                (
                    len(str(cell.value))
                    for cell in sheet[letter]
                    if cell.value is not None
                ),
                default=0,
            )
            sheet.column_dimensions[letter].width = min(max(longest + 2, 12), 60)

    generated_at = datetime.now().astimezone()
    summary_sheet.merge_cells("A1:F1")
    summary_sheet["A1"] = "R parameter batch remap report"
    summary_sheet["A1"].font = title_font
    summary_sheet["A1"].alignment = Alignment(vertical="center")
    summary_sheet.row_dimensions[1].height = 24
    summary_sheet.append([])
    summary_sheet.append(["Report field", "Value"])
    metadata = [
        ("Status", status),
        ("Mode", "EXECUTE" if execute_requested else "PREVIEW"),
        (
            "Delete replaced parameters",
            "YES" if delete_replaced_parameters else "NO",
        ),
        ("Database", database),
        ("Host", host),
        ("Generated at", generated_at.replace(tzinfo=None)),
        ("Timezone", str(generated_at.tzinfo)),
        ("Remap groups", len(plan.groups)),
        (
            "Estimated mutations",
            estimated_batch_mutations(
                plan, delete_replaced_parameters=delete_replaced_parameters
            ),
        ),
        ("Blocking conflicts", len(plan.conflicts)),
        ("Error", error or ""),
    ]
    for label, value in metadata:
        append_row(summary_sheet, (label, value))
    summary_sheet["B4"].fill = status_fills.get(
        status, PatternFill("solid", fgColor="FFF2CC")
    )
    summary_sheet["B9"].number_format = "yyyy-mm-dd hh:mm:ss"

    group_header_row = summary_sheet.max_row + 2
    append_row(
        summary_sheet,
        (
            "Group",
            "Old R parameter IDs",
            "New R parameter ID",
            "Selected parameters",
            "Estimated mutations",
            "Conflicts",
        ),
    )
    for group_index, (group, group_plan) in enumerate(
        zip(plan.groups, plan.group_plans, strict=True), start=1
    ):
        append_row(
            summary_sheet,
            (
                group_index,
                ", ".join(map(str, group.old_ids)),
                group.new_id,
                len(group_plan.parameter_rows),
                group_plan.mutation_count
                + (len(group.old_ids) if delete_replaced_parameters else 0),
                len(group_plan.conflicts),
            ),
        )

    operation_header_row = summary_sheet.max_row + 2
    append_row(summary_sheet, ("Operation", "Affected rows"))
    for operation, row_count in operation_summary_rows(
        plan, delete_replaced_parameters=delete_replaced_parameters
    ):
        append_row(summary_sheet, (operation, row_count))

    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.freeze_panes = "A4"
    style_header(summary_sheet, 3)
    style_header(summary_sheet, group_header_row)
    style_header(summary_sheet, operation_header_row)
    summary_sheet.column_dimensions["A"].width = 38
    summary_sheet.column_dimensions["B"].width = 58
    for letter in ("C", "D", "E", "F"):
        summary_sheet.column_dimensions[letter].width = 22
    for row in summary_sheet.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(
                name="Arial",
                size=10,
                bold=cell.font.bold,
                color=cell.font.color,
            )
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    append_row(
        parameters_sheet,
        (
            "Group",
            "Role",
            "Planned action",
            "ID",
            "Performance",
            "Design item 1",
            "Design item 2",
            "Design item 3",
            "Unit",
            "Hierarchy ID",
            "WP ID",
            "Unit ID",
        ),
    )
    for group_index, (group, group_plan) in enumerate(
        zip(plan.groups, plan.group_plans, strict=True), start=1
    ):
        for parameter_row in group_plan.parameter_rows:
            parameter_id, *parameter_values = parameter_row
            is_target = int(parameter_id) == group.new_id
            append_row(
                parameters_sheet,
                (
                    group_index,
                    "TARGET" if is_target else "SOURCE",
                    (
                        "DELETE_AFTER_REMAP"
                        if delete_replaced_parameters and not is_target
                        else "KEEP"
                    ),
                    parameter_id,
                    *parameter_values,
                ),
            )

    append_row(
        references_sheet,
        (
            "Group",
            "Reference scope",
            "Table",
            "Column",
            "Old ID",
            "New ID",
            "Rows",
            "Unique conflicts",
        ),
    )
    for group_index, group_plan in enumerate(plan.group_plans, start=1):
        reference_groups = (
            ("R_PARAMETER", group_plan.direct_reference_plans),
            ("RFL", group_plan.rfl_reference_plans),
            ("SETUP_R_RELATION", group_plan.setup_reference_plans),
        )
        for scope, reference_plans in reference_groups:
            for reference_plan in reference_plans:
                for (old_id, new_id), row_count in zip(
                    reference_plan.mappings,
                    reference_plan.row_counts,
                    strict=True,
                ):
                    append_row(
                        references_sheet,
                        (
                            group_index,
                            scope,
                            reference_plan.reference.table,
                            reference_plan.reference.column,
                            old_id,
                            new_id,
                            row_count,
                            " | ".join(reference_plan.unique_conflicts),
                        ),
                    )

    append_row(
        relations_sheet,
        (
            "Group",
            "Entity",
            "Action",
            "Record ID",
            "Source ID",
            "Target ID",
            "Context",
        ),
    )
    for row in relation_change_rows(
        plan, delete_replaced_parameters=delete_replaced_parameters
    ):
        append_row(relations_sheet, row)

    append_row(
        master_list_sheet,
        (
            "master_list_order ID",
            "Chunk",
            "Chunk count",
            "Old JSON chunk",
            "New JSON chunk",
        ),
    )
    for change in plan.master_list_changes:
        old_chunks = split_excel_text(change.old_json)
        new_chunks = split_excel_text(change.new_json)
        chunk_count = max(len(old_chunks), len(new_chunks))
        for chunk_index in range(chunk_count):
            append_row(
                master_list_sheet,
                (
                    change.row_id,
                    chunk_index + 1,
                    chunk_count,
                    old_chunks[chunk_index] if chunk_index < len(old_chunks) else "",
                    new_chunks[chunk_index] if chunk_index < len(new_chunks) else "",
                ),
            )

    append_row(conflicts_sheet, ("Severity", "Message"))
    if plan.conflicts:
        for conflict in plan.conflicts:
            append_row(conflicts_sheet, ("BLOCKING", conflict))
    else:
        append_row(conflicts_sheet, ("NONE", "No blocking conflicts detected."))

    for sheet in (
        parameters_sheet,
        references_sheet,
        relations_sheet,
        master_list_sheet,
        conflicts_sheet,
    ):
        finish_detail_sheet(sheet)

    report_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = report_path.with_name(f".{report_path.stem}.tmp.xlsx")
    try:
        workbook.save(temporary_path)
        temporary_path.replace(report_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def normalize_remap_groups(
    raw_groups: Iterable[tuple[Iterable[object], object]],
) -> tuple[RemapGroup, ...]:
    """Validate groups, merge repeated targets, and reject mapping chains."""
    old_ids_by_target: dict[int, list[int]] = {}
    source_targets: dict[int, int] = {}

    for group_index, (raw_old_ids, raw_new_id) in enumerate(raw_groups, start=1):
        try:
            old_ids, new_id = remap_core.normalize_id_mapping(raw_old_ids, raw_new_id)
        except ValueError as exc:
            raise ValueError(f"Invalid remap group #{group_index}: {exc}") from exc

        target_old_ids = old_ids_by_target.setdefault(new_id, [])
        for old_id in old_ids:
            previous_target = source_targets.get(old_id)
            if previous_target is not None and previous_target != new_id:
                raise ValueError(
                    f"Old ID {old_id} maps to both {previous_target} and {new_id}."
                )
            source_targets[old_id] = new_id
            if old_id not in target_old_ids:
                target_old_ids.append(old_id)

    if not old_ids_by_target:
        raise ValueError("At least one remap group is required.")

    targets = set(old_ids_by_target)
    chained_ids = sorted(set(source_targets) & targets)
    if chained_ids:
        raise ValueError(
            "Chained or cyclic mappings are not allowed. IDs used as both source "
            f"and target: {chained_ids}"
        )

    return tuple(
        RemapGroup(tuple(old_ids), new_id)
        for new_id, old_ids in old_ids_by_target.items()
    )


def parse_mapping_spec(raw_value: str) -> tuple[tuple[int, ...], int]:
    """Parse ``OLD[,OLD...]:NEW`` used by the repeatable --map option."""
    if raw_value.count(":") != 1:
        raise ValueError(f"Invalid mapping {raw_value!r}; expected OLD[,OLD...]:NEW.")
    raw_old_ids, raw_new_id = raw_value.split(":", maxsplit=1)
    old_ids = remap_core.parse_id_list(raw_old_ids)
    new_id = remap_core.normalize_positive_id(
        raw_new_id.strip(), label="new R parameter ID"
    )
    return old_ids, new_id


def merge_rfl_mappings(
    plans: Sequence[remap_core.MigrationPlan],
) -> tuple[tuple[int, int], ...]:
    merged: dict[int, int] = {}
    for plan in plans:
        for old_rfl_id, new_rfl_id in plan.rfl_merges:
            previous_target = merged.get(old_rfl_id)
            if previous_target is not None and previous_target != new_rfl_id:
                raise RuntimeError(
                    f"RFL ID {old_rfl_id} has conflicting targets "
                    f"{previous_target} and {new_rfl_id}."
                )
            merged[old_rfl_id] = new_rfl_id
    return tuple(merged.items())


def build_batch_plan(
    cursor: Any,
    groups: Sequence[RemapGroup],
    *,
    lock_rows: bool,
) -> BatchPlan:
    """Build all group plans against the same database snapshot."""
    group_plans: list[remap_core.MigrationPlan] = []
    for group in groups:
        plan = remap_core.build_plan(
            cursor,
            group.old_ids,
            group.new_id,
            lock_rows=lock_rows,
        )
        # A batch must update each JSON row once with the complete RFL map.
        group_plans.append(replace(plan, master_list_changes=()))

    rfl_mappings = merge_rfl_mappings(group_plans)
    master_list_changes = remap_core.build_master_list_changes(
        cursor,
        rfl_mappings,
        lock_rows=lock_rows,
    )
    conflicts = tuple(
        f"{group.display()}: {conflict}"
        for group, plan in zip(groups, group_plans, strict=True)
        for conflict in plan.conflicts
    )
    return BatchPlan(
        groups=tuple(groups),
        group_plans=tuple(group_plans),
        master_list_changes=master_list_changes,
        conflicts=conflicts,
    )


def print_batch_plan(
    plan: BatchPlan,
    *,
    database: str,
    host: str,
    delete_replaced_parameters: bool,
) -> None:
    print("R parameter batch remap preview")
    print(f"Database: {database}")
    print(f"Host: {host}")
    print(f"Groups: {len(plan.groups)}")

    for group_index, (group, group_plan) in enumerate(
        zip(plan.groups, plan.group_plans, strict=True), start=1
    ):
        print("\n" + "=" * 80)
        print(f"Group #{group_index}: {group.display()}")
        print("Selected r_parameter records:")
        for row in group_plan.parameter_rows:
            marker = "TARGET" if int(row[0]) == group.new_id else "SOURCE"
            print(f"  [{marker}] {row}")
        print(f"  setup_r_relation updates: {len(group_plan.setup_relation_updates)}")
        print(f"  setup_r_relation merges: {group_plan.setup_relation_merges}")
        print(f"  rfl rewires: {group_plan.rfl_updates}")
        print(f"  rfl merges: {group_plan.rfl_merges}")
        print(f"  prj_rfl duplicate deletes: {group_plan.prj_rfl_deletes}")
        print(f"  l_r_relation duplicate deletes: {group_plan.l_r_relation_deletes}")
        for reference_plan in (
            *group_plan.direct_reference_plans,
            *group_plan.rfl_reference_plans,
            *group_plan.setup_reference_plans,
        ):
            if reference_plan.row_count:
                reference = reference_plan.reference
                print(
                    f"  {reference.table}.{reference.column}: "
                    f"{reference_plan.row_count} row(s)"
                )

    print("\n" + "=" * 80)
    print(
        f"master_list_order rows updated once globally: {len(plan.master_list_changes)}"
    )
    delete_ids = replaced_parameter_ids(plan) if delete_replaced_parameters else ()
    print(f"Delete replaced r_parameter rows: {bool(delete_ids)}")
    print(f"r_parameter source rows to delete: {len(delete_ids)}")
    print(
        "Estimated mutation statements/rows: "
        f"{estimated_batch_mutations(plan, delete_replaced_parameters=delete_replaced_parameters)}"
    )
    if plan.conflicts:
        print("BLOCKING unique-key conflicts:")
        for conflict in plan.conflicts:
            print(f"  - {conflict}")
    elif not estimated_batch_mutations(
        plan, delete_replaced_parameters=delete_replaced_parameters
    ):
        print("No references require changes; this is a successful no-op.")


def require_batch_confirmation(
    database: str,
    groups: Sequence[RemapGroup],
    *,
    delete_replaced_parameters: bool,
) -> None:
    mappings = ";".join(group.display() for group in groups)
    action = "BATCH_REMAP_DELETE_OLD" if delete_replaced_parameters else "BATCH_REMAP"
    expected = f"{action} {database} {mappings}"
    print("\nMutation requested. Type the following text exactly to continue:")
    print(expected)
    if input("> ").strip() != expected:
        raise RuntimeError("Confirmation did not match; nothing was changed.")


def execute_master_list_changes(
    cursor: Any, changes: Sequence[remap_core.MasterListChange]
) -> None:
    for change in changes:
        cursor.execute(
            "UPDATE master_list_order SET list_order = %s WHERE id = %s",
            [change.new_json, change.row_id],
        )
        remap_core.expect_row_count(cursor, 1, "update master_list_order.list_order")


def contains_integer_id(value: Any, ids: set[int]) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int):
        return value in ids
    if isinstance(value, list):
        return any(contains_integer_id(item, ids) for item in value)
    if isinstance(value, dict):
        return any(contains_integer_id(item, ids) for item in value.values())
    return False


def verify_master_list_order(cursor: Any, old_rfl_ids: set[int]) -> None:
    if not old_rfl_ids:
        return
    cursor.execute("SELECT id, list_order FROM master_list_order ORDER BY id")
    remaining_rows: list[int] = []
    for row_id, raw_value in cursor.fetchall():
        parsed = json.loads(raw_value) if isinstance(raw_value, str) else raw_value
        if contains_integer_id(parsed, old_rfl_ids):
            remaining_rows.append(int(row_id))
    if remaining_rows:
        raise RuntimeError(
            "Post-update verification found old RFL IDs in master_list_order "
            f"rows: {remaining_rows}"
        )


def delete_replaced_parameter_rows(cursor: Any, old_ids: Sequence[int]) -> None:
    """Hard-delete verified source parameters and require an exact row count."""
    if not old_ids:
        return
    placeholders = remap_core.make_placeholders(old_ids)
    cursor.execute(
        f"DELETE FROM r_parameter WHERE id IN ({placeholders})",
        list(old_ids),
    )
    remap_core.expect_row_count(
        cursor,
        len(old_ids),
        "delete replaced r_parameter rows",
    )


def verify_replaced_parameter_rows_deleted(cursor: Any, old_ids: Sequence[int]) -> None:
    """Fail the transaction if any requested source parameter still exists."""
    if not old_ids:
        return
    placeholders = remap_core.make_placeholders(old_ids)
    cursor.execute(
        f"SELECT id FROM r_parameter WHERE id IN ({placeholders}) ORDER BY id",
        list(old_ids),
    )
    remaining_ids = [int(row[0]) for row in cursor.fetchall()]
    if remaining_ids:
        raise RuntimeError(
            f"Post-delete verification found replaced r_parameter rows: {remaining_ids}"
        )


def execute_batch_plan(
    cursor: Any,
    plan: BatchPlan,
    *,
    delete_replaced_parameters: bool,
) -> None:
    if plan.conflicts:
        raise RuntimeError("Cannot execute a batch plan with unique-key conflicts.")
    for group_plan in plan.group_plans:
        remap_core.execute_plan(cursor, group_plan)
    execute_master_list_changes(cursor, plan.master_list_changes)
    for group_plan in plan.group_plans:
        remap_core.verify_result(cursor, group_plan)
    verify_master_list_order(cursor, set(plan.rfl_id_map))
    if delete_replaced_parameters:
        old_ids = replaced_parameter_ids(plan)
        delete_replaced_parameter_rows(cursor, old_ids)
        verify_replaced_parameter_rows_deleted(cursor, old_ids)


def remap_r_parameter_id_groups(
    raw_groups: Iterable[tuple[Iterable[object], object]],
    *,
    execute: bool = False,
    assume_yes: bool = False,
    delete_replaced_parameters: bool = False,
    assume_delete_yes: bool = False,
    report_directory: str | Path | None = None,
) -> BatchPlan:
    groups = normalize_remap_groups(raw_groups)
    database, host = remap_core.database_identity()
    report_path = make_report_path(report_directory)

    with connection.cursor() as cursor:
        preview_plan = build_batch_plan(cursor, groups, lock_rows=False)
    print_batch_plan(
        preview_plan,
        database=database,
        host=host,
        delete_replaced_parameters=delete_replaced_parameters,
    )

    if preview_plan.conflicts:
        write_batch_report(
            report_path,
            preview_plan,
            status="BLOCKED",
            database=database,
            host=host,
            execute_requested=execute,
            delete_replaced_parameters=delete_replaced_parameters,
            error="Preview found unique-key conflicts.",
        )
        print(f"\nExcel report: {report_path}")
        raise RuntimeError(
            "Preview found unique-key conflicts. Resolve them before remapping."
        )
    if not execute:
        write_batch_report(
            report_path,
            preview_plan,
            status="PREVIEW",
            database=database,
            host=host,
            execute_requested=False,
            delete_replaced_parameters=delete_replaced_parameters,
        )
        print("\nPREVIEW ONLY: no database rows were changed.")
        print(f"Excel report: {report_path}")
        return preview_plan
    if not estimated_batch_mutations(
        preview_plan, delete_replaced_parameters=delete_replaced_parameters
    ):
        write_batch_report(
            report_path,
            preview_plan,
            status="NO_CHANGES",
            database=database,
            host=host,
            execute_requested=True,
            delete_replaced_parameters=delete_replaced_parameters,
        )
        print("\nNothing to update: all references already use their target IDs.")
        print(f"Excel report: {report_path}")
        return preview_plan

    # Validate report dependencies and destination before asking for confirmation
    # or starting the database transaction.
    write_batch_report(
        report_path,
        preview_plan,
        status="READY_TO_EXECUTE",
        database=database,
        host=host,
        execute_requested=True,
        delete_replaced_parameters=delete_replaced_parameters,
    )
    confirmation_required = not assume_yes or (
        delete_replaced_parameters and not assume_delete_yes
    )
    if confirmation_required:
        try:
            require_batch_confirmation(
                database,
                groups,
                delete_replaced_parameters=delete_replaced_parameters,
            )
        except Exception as exc:
            write_batch_report(
                report_path,
                preview_plan,
                status="CANCELLED",
                database=database,
                host=host,
                execute_requested=True,
                delete_replaced_parameters=delete_replaced_parameters,
                error=str(exc),
            )
            print(f"Excel report: {report_path}")
            raise

    try:
        with transaction.atomic():
            with connection.cursor() as cursor:
                locked_plan = build_batch_plan(cursor, groups, lock_rows=True)
                if locked_plan != preview_plan:
                    raise RuntimeError(
                        "Database state changed after preview; the whole batch was "
                        "rolled back. Run preview again."
                    )
                execute_batch_plan(
                    cursor,
                    locked_plan,
                    delete_replaced_parameters=delete_replaced_parameters,
                )
    except Exception as exc:
        try:
            write_batch_report(
                report_path,
                preview_plan,
                status="ROLLED_BACK",
                database=database,
                host=host,
                execute_requested=True,
                delete_replaced_parameters=delete_replaced_parameters,
                error=str(exc),
            )
        except Exception as report_exc:
            print(f"WARNING: could not finalize rollback report: {report_exc}")
        print(f"Excel report: {report_path}")
        raise

    try:
        write_batch_report(
            report_path,
            preview_plan,
            status="COMMITTED",
            database=database,
            host=host,
            execute_requested=True,
            delete_replaced_parameters=delete_replaced_parameters,
        )
    except Exception as exc:
        raise RuntimeError(
            "Database committed successfully, but the final Excel report could "
            f"not be written. The pre-execution report remains at {report_path}."
        ) from exc

    print(
        f"\nCommitted {len(groups)} R parameter remap group(s) in database {database}."
    )
    if delete_replaced_parameters:
        print(f"Deleted {len(replaced_parameter_ids(preview_plan))} source parameters.")
    print(f"Excel report: {report_path}")
    return preview_plan


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Preview or execute multiple R parameter ID remap groups."
    )
    parser.add_argument(
        "--map",
        dest="mapping_specs",
        action="append",
        metavar="OLD_IDS:NEW_ID",
        help=(
            "Repeatable mapping, for example --map 1,2,3:4 --map 5,6:7. "
            "Defaults to RUN_REMAP_GROUPS when omitted."
        ),
    )
    parser.add_argument(
        "--execute",
        action="store_true",
        help="Apply every mapping atomically; without this flag only preview.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        help=(
            "Skip remap confirmation for audited non-interactive runs; hard "
            "deletion still requires its separate confirmation opt-in."
        ),
    )
    parser.add_argument(
        "--delete-replaced-parameters",
        action="store_true",
        help=(
            "Hard-delete old r_parameter rows after all references are remapped "
            "and verified. Disabled by default."
        ),
    )
    parser.add_argument(
        "--yes-delete-replaced-parameters",
        action="store_true",
        help=(
            "Skip the separate hard-delete confirmation in an audited "
            "non-interactive run."
        ),
    )
    parser.add_argument(
        "--report-directory",
        metavar="DIRECTORY",
        help=(
            "Directory for the Excel audit report. Defaults to "
            "<project>/outputs/r_parameter_remap_reports."
        ),
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    raw_groups = (
        [parse_mapping_spec(spec) for spec in args.mapping_specs]
        if args.mapping_specs
        else RUN_REMAP_GROUPS
    )
    remap_r_parameter_id_groups(
        raw_groups,
        execute=args.execute or RUN_EXECUTE,
        assume_yes=args.yes or RUN_ASSUME_YES,
        delete_replaced_parameters=(
            args.delete_replaced_parameters or RUN_DELETE_REPLACED_PARAMETERS
        ),
        assume_delete_yes=(
            args.yes_delete_replaced_parameters or RUN_ASSUME_DELETE_YES
        ),
        report_directory=args.report_directory or RUN_REPORT_DIRECTORY,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
