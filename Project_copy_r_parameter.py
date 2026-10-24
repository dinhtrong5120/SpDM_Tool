"""Copy các liên kết r_parameter_id từ một project/phase sang nhiều đích.

Sửa RUN_SOURCE và RUN_DESTINATIONS để chạy bằng PyCharm.
Mặc định chỉ xem trước; RUN_EXECUTE bật ghi dữ liệu sau bước xác nhận.
Lấy DISTINCT r_parameter_id từ project_r_parameter nguồn, thêm vào mọi
variation đích đang xuất hiện trong project_parameter hoặc project_r_parameter.
Thêm liên kết còn thiếu trong project_r_parameter, detail_data_rlist và
setup_r_relation (setup đích, usecase_id=999999 như API thêm R).
Không copy giá trị/phê duyệt; các trường khác dùng default của DB.
Khóa đã tồn tại được bỏ qua riêng từng bảng. Setup dùng chung chỉ thêm một lần.
ID setup_r_relation dùng utils.id_sequence; ID đã cấp không hoàn lại khi rollback.
Đích chưa có variation sẽ báo lỗi; không tự tạo variation.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import tempfile


# Input: (project_id, phase_id). Điền ID thực tế trước khi chạy.
RUN_SOURCE = (28, 27)
RUN_DESTINATIONS = [(1028, 27), (1029, 27), (1030, 27), (1031, 27), (1032, 27),(1033, 27),(1034, 27)]  # Ví dụ: [(1020, 35), (1021, 27)]
RUN_EXECUTE = False
RUN_ASSUME_YES = False
RUN_REPORT_DIRECTORY = Path(__file__).resolve().parent / "output"
DEFAULT_USECASE_ID = 999999
RELATED_COLUMNS = {
    "detail_data_rlist": ("project_id", "r_parameter_id", "phase_id", "variation_id"),
    "setup_r_relation": ("setup_id", "r_parameter_id", "usecase_id"),
}


def write_report(path, source, preview, database, host, status, error=""):
    """Xuất snapshot kế hoạch; chỉ ghi INSERTED sau khi transaction commit."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    source_ids, plan, related = preview
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tong_hop"
    for row in (
        ("Thông tin", "Giá trị"),
        ("Database", database), ("Host", host),
        ("Nguồn (project, phase)", str(source)),
        ("Thời điểm báo cáo", datetime.now().astimezone().isoformat()),
        ("Trạng thái", status), ("Lỗi", error),
        ("Bảng ghi", "project_r_parameter, detail_data_rlist, setup_r_relation"),
        ("Số R ID nguồn", len(source_ids)),
        ("Phạm vi", "Chỉ thêm liên kết ID; trường khác dùng default DB."),
    ):
        summary.append(row)
    summary.append([])
    summary.append(["Project đích", "Phase đích", "Variation", "Dự kiến thêm",
                    "Đã tồn tại", "Đã commit"])
    detail = workbook.create_sheet("Chi_tiet")
    detail.append(["Project nguồn", "Phase nguồn", "Project đích", "Phase đích",
                   "Variation đích", "r_parameter_id", "Hành động", "Kết quả", "Lý do"])
    for destination, (variations, missing) in plan.items():
        missing_keys = set(missing)
        summary.append([
            *destination, ", ".join(map(str, variations)), len(missing),
            len(source_ids) * len(variations) - len(missing),
            len(missing) if status == "COMMITTED" else 0,
        ])
        for parameter_id in source_ids:
            for variation_id in variations:
                key = (destination[0], parameter_id, destination[1], variation_id)
                adding = key in missing_keys
                result = ("INSERTED" if status == "COMMITTED" else status) if adding else "SKIPPED"
                detail.append([
                    *source, *destination, variation_id, parameter_id,
                    "INSERT" if adding else "SKIP_EXISTING", result,
                    "Thiếu khóa ở thời điểm preview." if adding else "Khóa đã tồn tại; giữ nguyên.",
                ])
    extra_sheets = []
    summary.append([])
    summary.append(["Bảng bổ sung", "Dự kiến thêm", "Đã tồn tại", "Đã commit"])
    for table, decisions in related.items():
        missing_count = sum(decisions.values())
        summary.append([table, missing_count, len(decisions) - missing_count,
                        missing_count if status == "COMMITTED" else 0])
        sheet = workbook.create_sheet(table)
        sheet.append([*RELATED_COLUMNS[table], "Hành động", "Kết quả", "Lý do"])
        for key, adding in decisions.items():
            result = ("INSERTED" if status == "COMMITTED" else status) if adding else "SKIPPED"
            sheet.append([
                *key, "INSERT" if adding else "SKIP_EXISTING", result,
                "Thiếu khóa ở thời điểm preview." if adding else "Khóa đã tồn tại; giữ nguyên.",
            ])
        extra_sheets.append((sheet, (1,), (*([24] * len(RELATED_COLUMNS[table])), 24, 26, 55)))
    if any(sheet.max_row > 1_048_576 for sheet in workbook):
        raise ValueError("Báo cáo vượt giới hạn dòng Excel; hãy chia nhỏ danh sách đích.")
    for sheet, header_rows, widths in (
        (summary, (1, 12), (28, 65, 30, 22, 22, 22)),
        (detail, (1,), (20, 20, 20, 20, 22, 22, 24, 26, 55)),
        *extra_sheets,
    ):
        sheet.freeze_panes = "A13" if sheet == summary else "A2"
        last_row = 12 + len(plan) if sheet == summary else sheet.max_row
        sheet.auto_filter.ref = f"A{12 if sheet == summary else 1}:{get_column_letter(sheet.max_column)}{last_row}"
        for row in sheet:
            for cell in row:
                # DB metadata/error messages must remain literal Excel text.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row in header_rows:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="24476A")
            sheet.row_dimensions[row[0].row].height = 32
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    # Thay file bằng bản hoàn chỉnh, tránh làm hỏng preview khi save thất bại.
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as stream:
        temporary = Path(stream.name)
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
        workbook.close()
    print(f"Báo cáo Excel ({status}): {path}")


def normalize_scope(value):
    """Chuẩn hóa đúng hai ID nguyên dương, không làm tròn số thập phân."""
    if not isinstance(value, tuple) or len(value) != 2:
        raise ValueError("Mỗi scope phải là tuple (project_id, phase_id).")
    result = []
    for label, raw in zip(("project_id", "phase_id"), value):
        if isinstance(raw, bool) or not isinstance(raw, (int, str)):
            raise ValueError(f"{label} phải là số nguyên dương: {raw!r}")
        try:
            number = int(raw)
        except ValueError as exc:
            raise ValueError(f"{label} không hợp lệ: {raw!r}") from exc
        if number <= 0:
            raise ValueError(f"{label} phải lớn hơn 0: {number}")
        result.append(number)
    return tuple(result)


def build_plan(cursor, source, destinations, *, lock=False):
    """Đọc và kiểm tra tất cả scope trước khi lập danh sách khóa cần thêm."""
    suffix = " FOR UPDATE" if lock else ""
    scopes = sorted(set([source, *destinations]))
    # Tên bảng là hằng đã kiểm tra trong schema, không nhận từ input.
    for table, ids in (
        ("project_info", sorted({scope[0] for scope in scopes})),
        ("phase", sorted({scope[1] for scope in scopes})),
    ):
        for record_id in ids:
            cursor.execute(f"SELECT id FROM {table} WHERE id = %s{suffix}", [record_id])
            if cursor.fetchone() is None:
                raise ValueError(f"Không tồn tại {table}.id={record_id}.")

    r_rows = {}
    variations = {}
    for scope in scopes:
        cursor.execute(
            "SELECT r_parameter_id, variation_id FROM project_r_parameter "
            f"WHERE project_id = %s AND phase_id = %s{suffix}", scope,
        )
        r_rows[scope] = set(cursor.fetchall())
        cursor.execute(
            "SELECT variation_id FROM project_parameter "
            f"WHERE project_id = %s AND phase_id = %s{suffix}", scope,
        )
        variations[scope] = {row[0] for row in cursor.fetchall()}
        variations[scope].update(row[1] for row in r_rows[scope])
        if not variations[scope]:
            raise ValueError(f"Scope {scope} chưa có variation trong bảng SE hoặc R.")

    source_ids = sorted({row[0] for row in r_rows[source]})
    plan = {}
    related = {table: {} for table in RELATED_COLUMNS}
    setup_ids = set()
    for destination in destinations:
        target_variations = sorted(variations[destination])
        missing = [
            (destination[0], parameter_id, destination[1], variation_id)
            for parameter_id in source_ids
            for variation_id in target_variations
            if (parameter_id, variation_id) not in r_rows[destination]
        ]
        plan[destination] = (target_variations, missing)
        # Không phụ thuộc vào missing của project_r_parameter: bù cả dữ liệu
        # detail thiếu sau một lần chạy utility cũ chỉ insert bảng chính.
        cursor.execute(
            "SELECT r_parameter_id, variation_id FROM detail_data_rlist "
            f"WHERE project_id = %s AND phase_id = %s{suffix}", destination,
        )
        existing_details = set(cursor.fetchall())
        for parameter_id in source_ids:
            for variation_id in target_variations:
                key = (destination[0], parameter_id, destination[1], variation_id)
                related["detail_data_rlist"][key] = (
                    parameter_id, variation_id
                ) not in existing_details
        cursor.execute("SELECT setup_id FROM project_info WHERE id = %s", [destination[0]])
        setup_ids.add(cursor.fetchone()[0])

    for setup_id in sorted(setup_ids):
        # Khóa setup chung trước khi đọc quan hệ để các lần chạy tool được tuần tự.
        cursor.execute(f"SELECT id FROM setup WHERE id = %s{suffix}", [setup_id])
        if cursor.fetchone() is None:
            raise ValueError(f"Không tồn tại setup.id={setup_id}.")
        cursor.execute(
            "SELECT r_parameter_id, usecase_id FROM setup_r_relation "
            f"WHERE setup_id = %s{suffix}", [setup_id],
        )
        existing_relations = set(cursor.fetchall())
        for parameter_id in source_ids:
            key = (setup_id, parameter_id, DEFAULT_USECASE_ID)
            related["setup_r_relation"][key] = key[1:] not in existing_relations
    if any(related["setup_r_relation"].values()):
        cursor.execute(f"SELECT id FROM usecase WHERE id = %s{suffix}", [DEFAULT_USECASE_ID])
        if cursor.fetchone() is None:
            raise ValueError(f"Thiếu usecase mặc định {DEFAULT_USECASE_ID}; không tự tạo master.")
    return source_ids, plan, related


def copy_r_parameters(source, destinations, *, execute=False, assume_yes=False,
                      report_directory=RUN_REPORT_DIRECTORY):
    """Nhận một tuple nguồn và list tuple đích; trả tổng dòng insert cả ba bảng.

    Khi gọi từ Django shell, Django phải được khởi tạo sẵn.
    execute=False luôn chỉ đọc. assume_yes chỉ dùng cho chạy không tương tác
    đã được người vận hành cho phép; mặc định vẫn yêu cầu xác nhận chính xác.
    """
    source = normalize_scope(source)
    if not isinstance(destinations, list) or not destinations:
        raise ValueError("destinations phải là list không rỗng các tuple.")
    destinations = list(dict.fromkeys(normalize_scope(item) for item in destinations))
    if source in destinations:
        raise ValueError("Nguồn không được đồng thời là đích.")
    if not isinstance(execute, bool) or not isinstance(assume_yes, bool):
        raise ValueError("execute và assume_yes phải là bool.")

    from django.db import connection, transaction

    if connection.vendor != "mysql":
        raise ValueError("Script chỉ hỗ trợ MySQL.")
    settings = connection.settings_dict
    database = settings["NAME"]
    print(f"Database: {database}; host: {settings.get('HOST') or 'localhost'}")
    print(f"Nguồn: {source}; đích: {destinations}")
    print("Bảng ghi: project_r_parameter, detail_data_rlist, setup_r_relation; chế độ:",
          "EXECUTE" if execute else "DRY-RUN")
    with connection.cursor() as cursor:
        preview = build_plan(cursor, source, destinations)
    source_ids, plan, related = preview
    report_directory = Path(report_directory).expanduser().resolve()
    report_directory.mkdir(parents=True, exist_ok=True)
    report_path = report_directory / (
        f"Project_copy_r_parameter_p{source[0]}_ph{source[1]}_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
    )

    def report(status, error=""):
        write_report(report_path, source, preview, database,
                     settings.get("HOST") or "localhost", status, error)

    # Phải lưu được preview trước khi cho phép bất kỳ mutation nào.
    report("PREVIEW")
    print(f"Nguồn có {len(source_ids)} r_parameter_id: {source_ids}")
    total = 0
    for destination, (variation_ids, rows) in plan.items():
        existing = len(source_ids) * len(variation_ids) - len(rows)
        print(
            f"Đích {destination}: variations={variation_ids}; "
            f"thêm={len(rows)}; bỏ qua đã tồn tại={existing}"
        )
        print("Khóa sẽ thêm (project_id, r_parameter_id, phase_id, variation_id):", rows)
        total += len(rows)
    for table, decisions in related.items():
        missing_count = sum(decisions.values())
        print(f"{table}: thêm={missing_count}; bỏ qua đã tồn tại={len(decisions) - missing_count}")
        total += missing_count
    if not total or not execute:
        print("Không ghi dữ liệu.")
        return 0

    confirmation = f"COPY {database} TO {destinations}"
    if assume_yes:
        print(f"Bỏ qua xác nhận do assume_yes=True: {confirmation}")
    elif input(f"Nhập chính xác '{confirmation}' để thực thi: ") != confirmation:
        report("CANCELLED")
        print("Đã hủy; không ghi dữ liệu.")
        return 0

    try:
        with transaction.atomic(), connection.cursor() as cursor:
            # Khóa và đọc lại để tránh ghi theo preview đã thay đổi.
            if build_plan(cursor, source, destinations, lock=True) != preview:
                raise RuntimeError("Dữ liệu đã thay đổi sau preview; hãy chạy lại.")
            for destination, (_, rows) in plan.items():
                if not rows:
                    continue
                cursor.executemany(
                    "INSERT INTO project_r_parameter "
                    "(project_id, r_parameter_id, phase_id, variation_id) "
                    "VALUES (%s, %s, %s, %s)", rows,
                )
                if cursor.rowcount != len(rows):
                    raise RuntimeError(f"Số dòng insert không khớp tại {destination}.")
            for table, decisions in related.items():
                rows = [key for key, adding in decisions.items() if adding]
                if not rows:
                    continue
                columns = RELATED_COLUMNS[table]
                if table == "setup_r_relation":
                    from utils.id_sequence import next_ids

                    first_id = next_ids(table, n=len(rows))
                    rows = [(first_id + index, *row) for index, row in enumerate(rows)]
                    columns = ("id", *columns)
                # Identifiers chỉ lấy từ allowlist cố định phía trên.
                placeholders = ", ".join(["%s"] * len(columns))
                cursor.executemany(
                    f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders})", rows,
                )
                if cursor.rowcount != len(rows):
                    raise RuntimeError(f"Số dòng insert không khớp tại {table}.")
            _, after, related_after = build_plan(cursor, source, destinations, lock=True)
            if (any(rows for _, rows in after.values())
                    or any(any(decisions.values()) for decisions in related_after.values())):
                raise RuntimeError("Vẫn còn liên kết thiếu sau insert; rollback toàn bộ.")
    except Exception as exc:
        try:
            report("TRANSACTION_ERROR", str(exc))
        except Exception as report_error:
            print(f"Không cập nhật được báo cáo lỗi: {report_error}; preview: {report_path}")
        raise
    print(f"Đã commit tổng {total} dòng vào ba bảng.")
    try:
        report("COMMITTED")
    except Exception as exc:
        raise RuntimeError(
            f"DB ĐÃ COMMIT {total} dòng nhưng không cập nhật được Excel. "
            f"File vẫn là preview: {report_path}"
        ) from exc
    return total


if __name__ == "__main__":
    import os
    import sys

    import django

    project_root = Path(__file__).resolve().parents[2]
    if str(project_root) not in sys.path:
        sys.path.insert(0, str(project_root))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()
    copy_r_parameters(
        RUN_SOURCE, RUN_DESTINATIONS,
        execute=RUN_EXECUTE, assume_yes=RUN_ASSUME_YES,
        report_directory=RUN_REPORT_DIRECTORY,
    )
