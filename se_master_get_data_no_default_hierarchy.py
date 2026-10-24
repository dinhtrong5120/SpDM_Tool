"""Look up SE parameter IDs across every hierarchy without changing the DB.

Matches Excel columns WP, Category, パラメータ名, and 単位. All hierarchy_id
values are eligible. If multiple matching se_parameter rows exist, the
smallest se_parameter.id is selected and candidate IDs are noted in Excel.
The original workbook is never overwritten.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection  # noqa: E402

from utils.Tools_new.se_master_get_data import (  # noqa: E402
    LookupResult,
    ReferenceCatalog,
    RunResult,
    _set_result_cell,
    database_identity,
    find_header_columns,
    find_or_add_output_column,
    normalize_match_key,
    resolve_se_parameter,
)

# PyCharm Run-button configuration. There is intentionally no hierarchy ID.
RUN_EXCEL_PATH = Path("/Users/TrongKieu/Downloads/seリストへの追加.xlsx")
RUN_SHEET_NAME = "登録する項目"
RUN_OUTPUT_PATH: Path | None = "/Users/TrongKieu/Downloads/untitled folder/seリストへの追加_re.xlsx"

DEFAULT_OUTPUT_DIRECTORY = (
    PROJECT_ROOT / "outputs" / "se_master_get_data_no_default_hierarchy"
)


def make_output_path(
    input_path: Path,
    output_path: Path | str | None = None,
) -> Path:
    if output_path is None:
        timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
        result = DEFAULT_OUTPUT_DIRECTORY / (
            f"se_master_get_data_all_hierarchies_{timestamp}"
            f"{input_path.suffix.lower()}"
        )
    else:
        result = Path(output_path).expanduser().resolve()
    if result.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Output path must use the .xlsx or .xlsm extension.")
    if result == input_path:
        raise ValueError("Output must be a separate file; input will not be overwritten.")
    if result.suffix.lower() != input_path.suffix.lower():
        raise ValueError("Output extension must match the input Excel extension.")
    result.parent.mkdir(parents=True, exist_ok=True)
    return result


def load_reference_catalog_all_hierarchies(cursor: Any) -> ReferenceCatalog:
    cursor.execute("SELECT id, wp FROM wp ORDER BY id")
    wp_ids: dict[str, list[int]] = {}
    for raw_id, raw_wp in cursor.fetchall():
        wp_key = normalize_match_key(raw_wp)
        if wp_key is not None:
            wp_ids.setdefault(wp_key, []).append(int(raw_id))

    cursor.execute(
        """
        SELECT id, wp_id, parameter_name_1, parameter_name_2, parameter_unit
        FROM se_parameter
        ORDER BY id
        """
    )
    se_ids: dict[tuple[object, ...], list[int]] = {}
    for (
        raw_id,
        raw_wp_id,
        raw_category,
        raw_parameter_name,
        raw_parameter_unit,
    ) in cursor.fetchall():
        if raw_wp_id is None:
            continue
        key = (
            int(raw_wp_id),
            normalize_match_key(raw_category),
            normalize_match_key(raw_parameter_name),
            normalize_match_key(raw_parameter_unit),
        )
        se_ids.setdefault(key, []).append(int(raw_id))

    return ReferenceCatalog(
        wp_ids_by_name={
            key: tuple(sorted(set(values))) for key, values in wp_ids.items()
        },
        se_ids_by_key={
            key: tuple(sorted(set(values))) for key, values in se_ids.items()
        },
    )


def enrich_workbook_all_hierarchies(
    input_path: Path | str,
    *,
    sheet_name: str = RUN_SHEET_NAME,
    output_path: Path | str | None = None,
    db_connection=connection,
) -> RunResult:
    resolved_input = Path(input_path).expanduser().resolve()
    if not resolved_input.is_file():
        raise FileNotFoundError(f"Excel file was not found: {resolved_input}")
    if resolved_input.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Input Excel must use the .xlsx or .xlsm extension.")
    resolved_output = make_output_path(resolved_input, output_path)

    workbook = load_workbook(
        resolved_input,
        data_only=False,
        keep_vba=resolved_input.suffix.lower() == ".xlsm",
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet {sheet_name!r} was not found; available={workbook.sheetnames}"
            )
        worksheet = workbook[sheet_name]
        header_columns = find_header_columns(worksheet)
        output_column = find_or_add_output_column(worksheet)

        with db_connection.cursor() as cursor:
            catalog = load_reference_catalog_all_hierarchies(cursor)

        counts: Counter[str] = Counter()
        for row_number in range(2, worksheet.max_row + 1):
            source_cells = {
                header: worksheet.cell(row_number, column)
                for header, column in header_columns.items()
            }
            formula_headers = [
                header for header, cell in source_cells.items() if cell.data_type == "f"
            ]
            if formula_headers:
                result = LookupResult(
                    "INVALID",
                    None,
                    "Formula cells are not supported for lookup fields: "
                    + ", ".join(formula_headers)
                    + ".",
                )
            else:
                result = resolve_se_parameter(
                    catalog,
                    wp=source_cells["WP"].value,
                    category=source_cells["Category"].value,
                    parameter_name=source_cells["パラメータ名"].value,
                    parameter_unit=source_cells["単位"].value,
                )
            counts[result.status] += 1
            _set_result_cell(worksheet.cell(row_number, output_column), result)

        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{resolved_output.stem}_",
                suffix=resolved_output.suffix,
                dir=resolved_output.parent,
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            temporary_path.replace(resolved_output)
        finally:
            if temporary_path is not None and temporary_path.exists():
                temporary_path.unlink()
    finally:
        workbook.close()

    database, host, port = database_identity(db_connection)
    print("SE parameter lookup completed (all hierarchies)")
    print(f"Database (read-only): {database}")
    print(f"Host: {host}:{port}")
    print(f"Input Excel: {resolved_input}")
    print(f"Input sheet: {sheet_name}")
    print(f"Output Excel: {resolved_output}")
    print("Results:")
    for status in ("FOUND", "SELECTED_MIN", "NOT_FOUND", "INVALID", "SKIPPED"):
        print(f"- {status}: {counts[status]}")
    return RunResult(resolved_output, counts)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Look up se_parameter.id from Excel across every hierarchy_id and "
            "write results to a separate workbook. Database is read-only."
        )
    )
    parser.add_argument("--excel", type=Path, default=RUN_EXCEL_PATH)
    parser.add_argument("--sheet", default=RUN_SHEET_NAME)
    parser.add_argument("--output", type=Path, default=RUN_OUTPUT_PATH)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        enrich_workbook_all_hierarchies(
            args.excel,
            sheet_name=args.sheet,
            output_path=args.output,
        )
    except Exception as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
