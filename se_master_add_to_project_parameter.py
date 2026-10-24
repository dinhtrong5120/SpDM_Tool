"""Safely add checked SE IDs from Excel to one project/phase/variation.

Only the ``se_parameter_id check`` column is used for insertion. The original
``Se-paramemter-ID`` column is copied to the report for comparison, not used as
an ID source. Existing project_parameter primary keys are skipped. Preview is
the default; execution requires an explicit opt-in and exact confirmation.
The database mutation is one transaction. Every run writes an Excel report.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402

# PyCharm Run-button configuration. Set the three target IDs before running.
RUN_PROJECT_ID: int | None = 1021
RUN_PHASE_ID: int | None = 27
RUN_VARIATION_ID: int | None = 81
# RUN_EXCEL_PATH = PROJECT_ROOT / "outputs" / "se_master_get_data_no_default_hierarchy" / (
#     "se_master_get_data_all_hierarchies_20260916_065902_208174.xlsx"
# )
RUN_EXCEL_PATH = "/Users/TrongKieu/Downloads/untitled folder/seリストへの追加_re.xlsx"
RUN_SHEET_NAME = "登録する項目"
RUN_REPORT_PATH: Path | None = None
RUN_EXECUTE = False
RUN_ASSUME_YES = False

DEFAULT_APPROVAL_STATUS = 0
DEFAULT_EMPLOYEE_NUMBER = "00000007"
ID_HEADER = "se_parameter_id check"
SOURCE_ID_HEADER = "Se-paramemter-ID"
DEFAULT_REPORT_DIRECTORY = PROJECT_ROOT / "outputs" / "se_master_project_parameter_reports"
QUERY_BATCH_SIZE = 500


@dataclass
class PlanRow:
    excel_row: int
    source_id: object
    se_parameter_id: int | None
    action: str = "PENDING"
    result: str = "PENDING"
    message: str = ""


@dataclass(frozen=True)
class RunResult:
    status: str
    report_path: Path
    inserted_count: int
    rows: list[PlanRow]


class ToolRunError(RuntimeError):
    def __init__(self, message: str, report_path: Path) -> None:
        super().__init__(message)
        self.report_path = report_path


def normalize_positive_id(value: object, *, label: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{label} must be a positive integer, not boolean.")
    if isinstance(value, float) and not value.is_integer():
        raise ValueError(f"{label} must be an integer: {value!r}")
    try:
        normalized = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {label}: {value!r}") from exc
    if normalized <= 0:
        raise ValueError(f"{label} must be greater than zero: {normalized}")
    return normalized


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Cannot build placeholders for an empty list.")
    return ", ".join("%s" for _ in values)


def safe_excel_value(value: object) -> object:
    if value is None or isinstance(value, (bool, int, float, datetime)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def database_identity(db_connection=connection) -> tuple[str, str, str]:
    settings = db_connection.settings_dict
    return (
        str(settings.get("NAME") or "<unknown>"),
        str(settings.get("HOST") or "localhost"),
        str(settings.get("PORT") or "3306"),
    )


def make_report_path(
    excel_path: Path,
    report_path: Path | str | None = None,
) -> Path:
    if report_path is None:
        timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
        path = DEFAULT_REPORT_DIRECTORY / f"se_master_project_add_{timestamp}.xlsx"
    else:
        path = Path(report_path).expanduser().resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Report path must use .xlsx.")
    if path == excel_path:
        raise ValueError("Report cannot overwrite the input Excel file.")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def read_excel_rows(excel_path: Path | str, sheet_name: str) -> list[PlanRow]:
    path = Path(excel_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Excel file was not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Input Excel must be .xlsx or .xlsm.")

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet {sheet_name!r} was not found; available={workbook.sheetnames}"
            )
        sheet = workbook[sheet_name]
        headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
        if headers.count(ID_HEADER) != 1:
            raise ValueError(f"Expected exactly one {ID_HEADER!r} header.")
        id_column = headers.index(ID_HEADER) + 1
        source_column = (
            headers.index(SOURCE_ID_HEADER) + 1 if SOURCE_ID_HEADER in headers else None
        )

        rows: list[PlanRow] = []
        for row_number in range(2, sheet.max_row + 1):
            id_cell = sheet.cell(row_number, id_column)
            source_id = (
                sheet.cell(row_number, source_column).value
                if source_column is not None
                else None
            )
            if id_cell.value in (None, ""):
                rows.append(
                    PlanRow(
                        row_number,
                        source_id,
                        None,
                        "SKIP_NO_ID",
                        "SKIPPED",
                        "The check column is blank; no ID was selected.",
                    )
                )
                continue
            if id_cell.data_type == "f":
                rows.append(
                    PlanRow(
                        row_number,
                        source_id,
                        None,
                        "ERROR",
                        "ERROR",
                        "Formulas are not allowed in the check column.",
                    )
                )
                continue
            try:
                se_parameter_id = normalize_positive_id(
                    id_cell.value, label=f"Excel row {row_number} SE parameter ID"
                )
            except ValueError as exc:
                rows.append(
                    PlanRow(row_number, source_id, None, "ERROR", "ERROR", str(exc))
                )
            else:
                rows.append(PlanRow(row_number, source_id, se_parameter_id))
        if not rows:
            raise ValueError("The input sheet has no data rows.")
        return rows
    finally:
        workbook.close()


def fetch_existing_se_ids(cursor: Any, ids: Sequence[int], *, lock: bool) -> set[int]:
    found: set[int] = set()
    lock_clause = " FOR UPDATE" if lock else ""
    for start in range(0, len(ids), QUERY_BATCH_SIZE):
        batch = list(ids[start : start + QUERY_BATCH_SIZE])
        if not batch:
            continue
        cursor.execute(
            f"SELECT id FROM se_parameter WHERE id IN "
            f"({make_placeholders(batch)}){lock_clause}",
            batch,
        )
        found.update(int(record[0]) for record in cursor.fetchall())
    return found


def validate_targets(
    cursor: Any,
    *,
    project_id: int,
    phase_id: int,
    variation_id: int,
    employee_number: str,
    lock: bool,
) -> None:
    targets = (
        ("project_info", "id", project_id),
        ("phase", "id", phase_id),
        ("variation", "id", variation_id),
        ("user_table", "employee_number", employee_number),
    )
    for table, column, value in targets:
        # Identifiers above are hard-coded, never read from Excel or CLI.
        lock_clause = " FOR UPDATE" if lock else ""
        cursor.execute(
            f"SELECT 1 FROM `{table}` WHERE `{column}` = %s{lock_clause}",
            [value],
        )
        if cursor.fetchone() is None:
            raise ValueError(f"Required {table}.{column}={value!r} does not exist.")


def build_plan(
    cursor: Any,
    rows: list[PlanRow],
    *,
    project_id: int,
    phase_id: int,
    variation_id: int,
    lock: bool = False,
) -> list[PlanRow]:
    validate_targets(
        cursor,
        project_id=project_id,
        phase_id=phase_id,
        variation_id=variation_id,
        employee_number=DEFAULT_EMPLOYEE_NUMBER,
        lock=lock,
    )
    requested_ids = sorted(
        {row.se_parameter_id for row in rows if row.se_parameter_id is not None}
    )
    valid_ids = fetch_existing_se_ids(cursor, requested_ids, lock=lock)

    existing_ids: set[int] = set()
    lock_clause = " FOR UPDATE" if lock else ""
    for start in range(0, len(requested_ids), QUERY_BATCH_SIZE):
        batch = requested_ids[start : start + QUERY_BATCH_SIZE]
        if not batch:
            continue
        cursor.execute(
            "SELECT se_parameter_id FROM project_parameter "
            "WHERE project_id = %s AND phase_id = %s AND variation_id = %s "
            f"AND se_parameter_id IN ({make_placeholders(batch)}){lock_clause}",
            [project_id, phase_id, variation_id, *batch],
        )
        existing_ids.update(int(record[0]) for record in cursor.fetchall())

    seen: dict[int, int] = {}
    for row in rows:
        if row.result != "PENDING":
            continue
        assert row.se_parameter_id is not None
        if row.se_parameter_id not in valid_ids:
            row.action = "ERROR"
            row.result = "ERROR"
            row.message = "se_parameter.id does not exist in the database."
        elif row.se_parameter_id in existing_ids:
            row.action = "SKIP_EXISTING"
            row.result = "EXISTING"
            row.message = "This project/phase/variation/SE ID already exists."
        elif row.se_parameter_id in seen:
            row.action = "SKIP_DUPLICATE_INPUT"
            row.result = "DUPLICATE_INPUT"
            row.message = f"Same SE ID already appears at Excel row {seen[row.se_parameter_id]}."
        else:
            seen[row.se_parameter_id] = row.excel_row
            row.action = "INSERT"
            row.result = "WOULD_INSERT"
            row.message = "Preview only; no database row has been inserted."
    return rows


def plan_signature(rows: Sequence[PlanRow]) -> tuple[tuple[int, int], ...]:
    return tuple(
        (row.excel_row, row.se_parameter_id)
        for row in rows
        if row.result == "WOULD_INSERT" and row.se_parameter_id is not None
    )


def mark_rows(rows: Sequence[PlanRow], *, result: str, message: str) -> None:
    for row in rows:
        if row.result == "WOULD_INSERT":
            row.result = result
            row.message = message


def apply_plan(
    cursor: Any,
    rows: Sequence[PlanRow],
    *,
    project_id: int,
    phase_id: int,
    variation_id: int,
) -> int:
    if any(row.result == "ERROR" for row in rows):
        raise RuntimeError("Plan contains validation errors.")
    inserted_ids: list[int] = []
    for row in rows:
        if row.result != "WOULD_INSERT":
            continue
        assert row.se_parameter_id is not None
        cursor.execute(
            """
            INSERT INTO project_parameter
                (project_id, se_parameter_id, phase_id, variation_id,
                 approval_status, employee_number, update_day)
            VALUES (%s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
            """,
            [
                project_id,
                row.se_parameter_id,
                phase_id,
                variation_id,
                DEFAULT_APPROVAL_STATUS,
                DEFAULT_EMPLOYEE_NUMBER,
            ],
        )
        if cursor.rowcount != 1:
            raise RuntimeError(
                f"Unexpected insert count for se_parameter.id={row.se_parameter_id}: "
                f"{cursor.rowcount}"
            )
        inserted_ids.append(row.se_parameter_id)
        row.result = "INSERTED"
        row.message = "Inserted in the committed transaction."

    for start in range(0, len(inserted_ids), QUERY_BATCH_SIZE):
        batch = inserted_ids[start : start + QUERY_BATCH_SIZE]
        cursor.execute(
            "SELECT COUNT(*) FROM project_parameter "
            "WHERE project_id = %s AND phase_id = %s AND variation_id = %s "
            f"AND se_parameter_id IN ({make_placeholders(batch)})",
            [project_id, phase_id, variation_id, *batch],
        )
        actual_count = int(cursor.fetchone()[0])
        if actual_count != len(batch):
            raise RuntimeError(
                f"Post-insert verification found {actual_count}/{len(batch)} rows."
            )
    return len(inserted_ids)


def confirmation_text(
    database: str,
    rows: Sequence[PlanRow],
    *,
    project_id: int,
    phase_id: int,
    variation_id: int,
) -> str:
    ids = ",".join(
        str(row.se_parameter_id) for row in rows if row.result == "WOULD_INSERT"
    )
    return (
        f"INSERT PROJECT_PARAMETER DB {database} PROJECT {project_id} "
        f"PHASE {phase_id} VARIATION {variation_id} SE IDS {ids or '-'}"
    )


def write_report(
    report_path: Path,
    rows: Sequence[PlanRow],
    *,
    status: str,
    database: str,
    host: str,
    port: str,
    excel_path: Path,
    sheet_name: str,
    project_id: int,
    phase_id: int,
    variation_id: int,
    execute: bool,
    assume_yes: bool,
    error: str | None = None,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    results = workbook.create_sheet("Results")

    summary.append(["Project SE parameter import"])
    summary.append([])
    summary.append(["Item", "Value"])
    summary_items = [
        ("Status", status),
        ("Mode", "EXECUTE" if execute else "PREVIEW"),
        ("Database", database),
        ("Host", f"{host}:{port}"),
        ("Input file", str(excel_path)),
        ("Input sheet", sheet_name),
        ("Project ID", project_id),
        ("Phase ID", phase_id),
        ("Variation ID", variation_id),
        ("Approval status", DEFAULT_APPROVAL_STATUS),
        ("Employee number", DEFAULT_EMPLOYEE_NUMBER),
        ("Non-interactive confirmation", assume_yes),
        ("Input rows", len(rows)),
        ("Rows planned", sum(row.action == "INSERT" for row in rows)),
        ("Rows inserted", sum(row.result == "INSERTED" for row in rows)),
        ("Existing skipped", sum(row.result == "EXISTING" for row in rows)),
        ("Duplicate input skipped", sum(row.result == "DUPLICATE_INPUT" for row in rows)),
        ("Blank ID skipped", sum(row.action == "SKIP_NO_ID" for row in rows)),
        ("Error rows", sum(row.result == "ERROR" for row in rows)),
        ("Error", error or ""),
    ]
    for item, value in summary_items:
        summary.append([item, safe_excel_value(value)])
    summary.column_dimensions["A"].width = 35
    summary.column_dimensions["B"].width = 105

    results.append(
        [
            "excel_row",
            "source_Se-paramemter-ID",
            "se_parameter_id check",
            "project_id",
            "phase_id",
            "variation_id",
            "approval_status",
            "employee_number",
            "action",
            "result",
            "message",
        ]
    )
    for row in rows:
        results.append(
            [
                row.excel_row,
                safe_excel_value(row.source_id),
                row.se_parameter_id,
                project_id,
                phase_id,
                variation_id,
                DEFAULT_APPROVAL_STATUS,
                DEFAULT_EMPLOYEE_NUMBER,
                row.action,
                row.result,
                safe_excel_value(row.message),
            ]
        )
    results.freeze_panes = "A2"
    results.auto_filter.ref = results.dimensions
    for column, width in {
        "A": 12,
        "B": 24,
        "C": 24,
        "D": 14,
        "E": 14,
        "F": 14,
        "G": 18,
        "H": 20,
        "I": 24,
        "J": 22,
        "K": 70,
    }.items():
        results.column_dimensions[column].width = width
    for sheet, header_row in ((summary, 3), (results, 1)):
        sheet.sheet_view.showGridLines = False
        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_number in range(2, results.max_row + 1):
        cell = results.cell(row_number, 10)
        if cell.value in {"ERROR", "BLOCKED", "ROLLED_BACK"}:
            cell.fill = PatternFill("solid", fgColor="F4CCCC")
        elif cell.value == "INSERTED":
            cell.fill = PatternFill("solid", fgColor="C6EFCE")
        elif cell.value == "WOULD_INSERT":
            cell.fill = PatternFill("solid", fgColor="FFF2CC")

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{report_path.stem}_",
            suffix=".xlsx",
            dir=report_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        temporary_path.replace(report_path)
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def import_project_parameters(
    project_id: object,
    phase_id: object,
    variation_id: object,
    excel_path: Path | str,
    *,
    sheet_name: str = RUN_SHEET_NAME,
    report_path: Path | str | None = None,
    execute: bool = False,
    assume_yes: bool = False,
    db_connection=connection,
) -> RunResult:
    normalized_project_id = normalize_positive_id(project_id, label="project_id")
    normalized_phase_id = normalize_positive_id(phase_id, label="phase_id")
    normalized_variation_id = normalize_positive_id(variation_id, label="variation_id")
    if assume_yes and not execute:
        raise ValueError("--yes/RUN_ASSUME_YES requires execution mode.")

    resolved_excel = Path(excel_path).expanduser().resolve()
    resolved_report = make_report_path(resolved_excel, report_path)
    database, host, port = database_identity(db_connection)
    rows: list[PlanRow] = []
    error: str | None = None

    try:
        rows = read_excel_rows(resolved_excel, sheet_name)
        with db_connection.cursor() as cursor:
            build_plan(
                cursor,
                rows,
                project_id=normalized_project_id,
                phase_id=normalized_phase_id,
                variation_id=normalized_variation_id,
            )
    except Exception as exc:
        error = f"{type(exc).__name__}: {exc}"
        write_report(
            resolved_report,
            rows,
            status="BLOCKED",
            database=database,
            host=host,
            port=port,
            excel_path=resolved_excel,
            sheet_name=sheet_name,
            project_id=normalized_project_id,
            phase_id=normalized_phase_id,
            variation_id=normalized_variation_id,
            execute=execute,
            assume_yes=assume_yes,
            error=error,
        )
        raise ToolRunError(error, resolved_report) from exc

    planned_ids = [
        row.se_parameter_id for row in rows if row.result == "WOULD_INSERT"
    ]
    print("Project SE parameter import plan")
    print(f"Database: {database}")
    print(f"Host: {host}:{port}")
    print(f"Target: project_id={normalized_project_id}, phase_id={normalized_phase_id}, "
          f"variation_id={normalized_variation_id}")
    print(f"Input: {resolved_excel} / {sheet_name}")
    print(f"Mode: {'EXECUTE' if execute else 'PREVIEW'}")
    print(f"Target table: project_parameter; planned inserts: {len(planned_ids)}")
    print(f"Selected se_parameter IDs: {planned_ids}")
    print(f"Report: {resolved_report}")

    if any(row.result == "ERROR" for row in rows):
        status = "BLOCKED"
        error = "Input IDs are invalid or missing from se_parameter; no rows inserted."
        mark_rows(rows, result="BLOCKED", message="Not executed because another row has an error.")
    elif not execute:
        status = "PREVIEW"
    elif not planned_ids:
        status = "NO_CHANGES"
    else:
        write_report(
            resolved_report,
            rows,
            status="READY_TO_EXECUTE",
            database=database,
            host=host,
            port=port,
            excel_path=resolved_excel,
            sheet_name=sheet_name,
            project_id=normalized_project_id,
            phase_id=normalized_phase_id,
            variation_id=normalized_variation_id,
            execute=execute,
            assume_yes=assume_yes,
        )
        expected_confirmation = confirmation_text(
            database,
            rows,
            project_id=normalized_project_id,
            phase_id=normalized_phase_id,
            variation_id=normalized_variation_id,
        )
        if assume_yes:
            print("Exact confirmation bypass explicitly enabled.")
            confirmed = True
        else:
            print(f"Type exactly: {expected_confirmation}")
            confirmed = input("> ").strip() == expected_confirmation
        if not confirmed:
            status = "CANCELLED"
            error = "Exact confirmation was not provided; nothing was inserted."
            mark_rows(rows, result="CANCELLED", message="Cancelled before database mutation.")
        else:
            preview_signature = plan_signature(rows)
            try:
                execution_rows = read_excel_rows(resolved_excel, sheet_name)
                alias = getattr(db_connection, "alias", "default")
                with transaction.atomic(using=alias):
                    with db_connection.cursor() as cursor:
                        build_plan(
                            cursor,
                            execution_rows,
                            project_id=normalized_project_id,
                            phase_id=normalized_phase_id,
                            variation_id=normalized_variation_id,
                            lock=True,
                        )
                        if any(row.result == "ERROR" for row in execution_rows):
                            raise RuntimeError("Execution plan contains invalid IDs.")
                        if plan_signature(execution_rows) != preview_signature:
                            raise RuntimeError(
                                "Input or database changed after preview; "
                                "transaction rolled back."
                            )
                        apply_plan(
                            cursor,
                            execution_rows,
                            project_id=normalized_project_id,
                            phase_id=normalized_phase_id,
                            variation_id=normalized_variation_id,
                        )
                rows = execution_rows
                status = "COMMITTED"
            except Exception as exc:
                status = "ROLLED_BACK"
                error = f"{type(exc).__name__}: {exc}"
                mark_rows(rows, result="ROLLED_BACK", message="Transaction rolled back; no rows inserted.")

    write_report(
        resolved_report,
        rows,
        status=status,
        database=database,
        host=host,
        port=port,
        excel_path=resolved_excel,
        sheet_name=sheet_name,
        project_id=normalized_project_id,
        phase_id=normalized_phase_id,
        variation_id=normalized_variation_id,
        execute=execute,
        assume_yes=assume_yes,
        error=error,
    )
    inserted_count = sum(row.result == "INSERTED" for row in rows)
    print(f"Status: {status}; inserted: {inserted_count}; report: {resolved_report}")
    result = RunResult(status, resolved_report, inserted_count, rows)
    if status == "ROLLED_BACK":
        raise ToolRunError(error or "Transaction failed.", resolved_report)
    return result


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Preview or insert checked SE IDs into project_parameter."
    )
    parser.add_argument("--project-id", type=int, default=RUN_PROJECT_ID)
    parser.add_argument("--phase-id", type=int, default=RUN_PHASE_ID)
    parser.add_argument("--variation-id", type=int, default=RUN_VARIATION_ID)
    parser.add_argument("--excel", type=Path, default=RUN_EXCEL_PATH)
    parser.add_argument("--sheet", default=RUN_SHEET_NAME)
    parser.add_argument("--report", type=Path, default=RUN_REPORT_PATH)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--yes", action="store_true")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        result = import_project_parameters(
            args.project_id,
            args.phase_id,
            args.variation_id,
            args.excel,
            sheet_name=args.sheet,
            report_path=args.report,
            execute=args.execute or RUN_EXECUTE,
            assume_yes=args.yes or RUN_ASSUME_YES,
        )
    except ToolRunError as exc:
        print(f"ERROR: {exc}; report: {exc.report_path}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    return 1 if result.status == "BLOCKED" else 0


if __name__ == "__main__":
    raise SystemExit(main())
