"""Preview or safely delete R parameters and export an Excel audit report.

The default mode is read-only. Real deletion requires ``--execute`` and an
exact confirmation containing the database name and requested IDs. Foreign-key
relationships are discovered from ``information_schema`` at runtime; legacy
columns without a declared FK must be reviewed and listed explicitly below.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict, deque
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Sequence

import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import os

    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

# PyCharm Run configuration. Preview first, then review the Excel report.
RUN_R_PARAMETER_IDS = [
    123123367,
    123123368,
    123123369
]
# df = pd.read_excel("/Users/TrongKieu/Downloads/QR0023_未使用R項目.xlsx")
# RUN_R_PARAMETER_IDS = df['r_parameter_id'].tolist()

RUN_EXECUTE = False
RUN_ASSUME_YES = False
RUN_ALLOW_EXTERNAL_ORPHANS = False
RUN_OUTPUT_PATH: str | None = None
RUN_REPORT_ROW_LIMIT = 50_000

ROOT_TABLE = "r_parameter"
ROOT_COLUMN = "id"

# Reviewed legacy relationship without a declared FK in the schema dump.
# Never add every column named r_parameter_id automatically.
LEGACY_RELATIONSHIPS = (
    ("detail_data_rlist", "r_parameter_id", ROOT_TABLE, ROOT_COLUMN),
)

# These DB rows can point to resources outside MySQL. The script never deletes
# external files/S3 objects and blocks execution by default when rows exist.
EXTERNAL_RESOURCE_TABLES = {
    "project_parameter_attachments": (
        "Attachment metadata found; this script does not delete S3 objects."
    ),
}

IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z0-9_$]+$")
EXCEL_ILLEGAL_CHARACTER_PATTERN = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")
EXCEL_CELL_LIMIT = 32_767


@dataclass(frozen=True)
class RunConfig:
    r_parameter_ids: tuple[int, ...]
    execute: bool
    assume_yes: bool
    allow_external_orphans: bool
    output_path: Path
    report_row_limit: int


@dataclass(frozen=True)
class ForeignKeyRelation:
    constraint_name: str
    child_table: str
    child_columns: tuple[str, ...]
    parent_table: str
    parent_columns: tuple[str, ...]
    delete_rule: str
    source: str = "foreign_key"


@dataclass
class DeletionPlan:
    delete_paths: dict[str, list[tuple[ForeignKeyRelation, ...]]]
    set_null_paths: dict[str, list[tuple[ForeignKeyRelation, ...]]]
    delete_order: list[str]
    used_relations: list[ForeignKeyRelation]


@dataclass
class TableSnapshot:
    table_name: str
    action: str
    columns: list[str]
    rows: list[tuple[Any, ...]]
    total_rows: int

    @property
    def truncated(self) -> bool:
        return self.total_rows > len(self.rows)


def quote_identifier(identifier: str) -> str:
    if not IDENTIFIER_PATTERN.fullmatch(identifier):
        raise ValueError(f"Tên bảng/cột không hợp lệ: {identifier!r}")
    return f"`{identifier}`"


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Không thể tạo SQL placeholders cho danh sách rỗng.")
    return ", ".join(["%s"] * len(values))


def normalize_positive_id(raw_id: Any, *, label: str) -> int:
    if isinstance(raw_id, bool):
        raise ValueError(f"{label} phải là số nguyên dương, không phải boolean.")
    try:
        normalized_id = int(raw_id)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} không hợp lệ: {raw_id!r}") from exc
    if normalized_id <= 0:
        raise ValueError(f"{label} phải lớn hơn 0: {normalized_id}")
    return normalized_id


def normalize_ids(raw_ids: Iterable[Any]) -> tuple[int, ...]:
    normalized: list[int] = []
    seen: set[int] = set()
    for raw_id in raw_ids:
        value = normalize_positive_id(raw_id, label="r_parameter_id")
        if value not in seen:
            normalized.append(value)
            seen.add(value)
    if not normalized:
        raise ValueError("R_PARAMETER_IDS không được để trống.")
    return tuple(normalized)


def normalize_report_row_limit(raw_limit: Any) -> int:
    limit = normalize_positive_id(raw_limit, label="report_row_limit")
    if limit > 1_000_000:
        raise ValueError("report_row_limit không được lớn hơn 1.000.000.")
    return limit


def default_output_path(ids: Sequence[int]) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    id_text = "_".join(map(str, ids[:5]))
    if len(ids) > 5:
        id_text += f"_and_{len(ids) - 5}_more"
    filename = f"r_parameter_delete_report_{id_text}_{timestamp}.xlsx"
    return Path(__file__).resolve().parent / "output" / filename


def resolve_output_path(raw_path: str | None, ids: Sequence[int]) -> Path:
    if raw_path:
        path = Path(raw_path).expanduser()
        if not path.is_absolute():
            path = Path.cwd() / path
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Đường dẫn báo cáo phải có phần mở rộng .xlsx.")
        return path.resolve()
    return default_output_path(ids)


def parse_args(argv: Sequence[str] | None = None) -> RunConfig:
    parser = argparse.ArgumentParser(
        description=("Xem trước/xóa R parameter cùng dữ liệu liên quan và xuất Excel.")
    )
    parser.add_argument(
        "--ids", nargs="+", type=int, default=RUN_R_PARAMETER_IDS, metavar="ID"
    )
    parser.add_argument(
        "--execute",
        action="store_true",
        default=RUN_EXECUTE,
        help="Cho phép thực hiện xóa sau bước xác nhận.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        default=RUN_ASSUME_YES,
        help="Bỏ qua xác nhận tương tác (không khuyến nghị).",
    )
    parser.add_argument(
        "--allow-external-orphans",
        action="store_true",
        default=RUN_ALLOW_EXTERNAL_ORPHANS,
        help="Chấp nhận file/S3 không bị xóa khi metadata DB bị xóa.",
    )
    parser.add_argument("--output", default=RUN_OUTPUT_PATH)
    parser.add_argument("--report-row-limit", type=int, default=RUN_REPORT_ROW_LIMIT)
    args = parser.parse_args(argv)

    ids = normalize_ids(args.ids)
    return RunConfig(
        r_parameter_ids=ids,
        execute=bool(args.execute),
        assume_yes=bool(args.yes),
        allow_external_orphans=bool(args.allow_external_orphans),
        output_path=resolve_output_path(args.output, ids),
        report_row_limit=normalize_report_row_limit(args.report_row_limit),
    )


def database_identity() -> tuple[str, str]:
    settings = connection.settings_dict
    name = str(settings.get("NAME") or "")
    host = str(settings.get("HOST") or "")
    port = str(settings.get("PORT") or "")
    return name, f"{host}:{port}" if port else host


def _split_metadata_columns(value: str | None) -> tuple[str, ...]:
    return tuple(value.split("\x1f")) if value else ()


def discover_foreign_keys(cursor) -> list[ForeignKeyRelation]:
    cursor.execute("""
        SELECT
            kcu.CONSTRAINT_NAME,
            kcu.TABLE_NAME,
            GROUP_CONCAT(
                kcu.COLUMN_NAME ORDER BY kcu.ORDINAL_POSITION SEPARATOR '\x1f'
            ),
            kcu.REFERENCED_TABLE_NAME,
            GROUP_CONCAT(
                kcu.REFERENCED_COLUMN_NAME
                ORDER BY kcu.ORDINAL_POSITION SEPARATOR '\x1f'
            ),
            rc.DELETE_RULE
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        JOIN information_schema.REFERENTIAL_CONSTRAINTS AS rc
          ON rc.CONSTRAINT_SCHEMA = kcu.CONSTRAINT_SCHEMA
         AND rc.CONSTRAINT_NAME = kcu.CONSTRAINT_NAME
         AND rc.TABLE_NAME = kcu.TABLE_NAME
        JOIN information_schema.TABLES AS child_metadata
          ON child_metadata.TABLE_SCHEMA = kcu.TABLE_SCHEMA
         AND child_metadata.TABLE_NAME = kcu.TABLE_NAME
         AND child_metadata.TABLE_TYPE = 'BASE TABLE'
        JOIN information_schema.TABLES AS parent_metadata
          ON parent_metadata.TABLE_SCHEMA = kcu.REFERENCED_TABLE_SCHEMA
         AND parent_metadata.TABLE_NAME = kcu.REFERENCED_TABLE_NAME
         AND parent_metadata.TABLE_TYPE = 'BASE TABLE'
        WHERE kcu.TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_NAME IS NOT NULL
        GROUP BY
            kcu.CONSTRAINT_NAME,
            kcu.TABLE_NAME,
            kcu.REFERENCED_TABLE_NAME,
            rc.DELETE_RULE
        ORDER BY kcu.REFERENCED_TABLE_NAME, kcu.TABLE_NAME, kcu.CONSTRAINT_NAME
        """)
    relations: list[ForeignKeyRelation] = []
    for row in cursor.fetchall():
        child_columns = _split_metadata_columns(row[2])
        parent_columns = _split_metadata_columns(row[4])
        if not child_columns or len(child_columns) != len(parent_columns):
            raise RuntimeError(f"Metadata FK không hợp lệ: {row[0]}")
        relations.append(
            ForeignKeyRelation(
                constraint_name=str(row[0]),
                child_table=str(row[1]),
                child_columns=child_columns,
                parent_table=str(row[3]),
                parent_columns=parent_columns,
                delete_rule=str(row[5]).upper(),
            )
        )
    return relations


def fetch_existing_columns(
    cursor, table_columns: Sequence[tuple[str, str]]
) -> set[tuple[str, str]]:
    if not table_columns:
        return set()
    tables = sorted({table for table, _ in table_columns})
    placeholders = make_placeholders(tables)
    cursor.execute(
        f"""
        SELECT c.TABLE_NAME, c.COLUMN_NAME
        FROM information_schema.COLUMNS AS c
        JOIN information_schema.TABLES AS t
          ON t.TABLE_SCHEMA = c.TABLE_SCHEMA
         AND t.TABLE_NAME = c.TABLE_NAME
         AND t.TABLE_TYPE = 'BASE TABLE'
        WHERE c.TABLE_SCHEMA = DATABASE()
          AND c.TABLE_NAME IN ({placeholders})
        """,
        tables,
    )
    return {(str(table), str(column)) for table, column in cursor.fetchall()}


def add_legacy_relationships(
    cursor, relations: Sequence[ForeignKeyRelation]
) -> list[ForeignKeyRelation]:
    combined = list(relations)
    requested_columns = [
        (child_table, child_column)
        for child_table, child_column, _, _ in LEGACY_RELATIONSHIPS
    ]
    existing_columns = fetch_existing_columns(cursor, requested_columns)
    existing_signatures = {
        (
            relation.child_table,
            relation.child_columns,
            relation.parent_table,
            relation.parent_columns,
        )
        for relation in relations
    }

    for child_table, child_column, parent_table, parent_column in LEGACY_RELATIONSHIPS:
        if (child_table, child_column) not in existing_columns:
            continue
        signature = (child_table, (child_column,), parent_table, (parent_column,))
        if signature in existing_signatures:
            continue
        combined.append(
            ForeignKeyRelation(
                constraint_name=f"legacy_{child_table}_{child_column}",
                child_table=child_table,
                child_columns=(child_column,),
                parent_table=parent_table,
                parent_columns=(parent_column,),
                delete_rule="NO ACTION",
                source="legacy_allowlist",
            )
        )
    return combined


def _path_signature(
    path: Sequence[ForeignKeyRelation],
) -> tuple[tuple[str, str], ...]:
    return tuple((item.child_table, item.constraint_name) for item in path)


def build_deletion_plan(
    relations: Sequence[ForeignKeyRelation],
) -> DeletionPlan:
    by_parent: dict[str, list[ForeignKeyRelation]] = defaultdict(list)
    for relation in relations:
        by_parent[relation.parent_table].append(relation)

    delete_paths: dict[str, list[tuple[ForeignKeyRelation, ...]]] = {
        ROOT_TABLE: [tuple()]
    }
    set_null_paths: dict[str, list[tuple[ForeignKeyRelation, ...]]] = defaultdict(list)
    seen_paths: set[tuple[str, tuple[tuple[str, str], ...]]] = {(ROOT_TABLE, tuple())}
    queue: deque[tuple[str, tuple[ForeignKeyRelation, ...], tuple[str, ...]]] = deque(
        [(ROOT_TABLE, tuple(), (ROOT_TABLE,))]
    )

    while queue:
        parent_table, parent_path, tables_in_path = queue.popleft()
        for relation in by_parent.get(parent_table, []):
            if relation.child_table in tables_in_path:
                cycle = " -> ".join((*tables_in_path, relation.child_table))
                raise RuntimeError(f"Phát hiện vòng lặp FK: {cycle}")
            child_path = (*parent_path, relation)
            signature = (relation.child_table, _path_signature(child_path))
            if signature in seen_paths:
                continue
            seen_paths.add(signature)
            if relation.delete_rule == "SET NULL":
                set_null_paths[relation.child_table].append(child_path)
                continue
            delete_paths.setdefault(relation.child_table, []).append(child_path)
            queue.append(
                (
                    relation.child_table,
                    child_path,
                    (*tables_in_path, relation.child_table),
                )
            )

    delete_tables = set(delete_paths)
    children: dict[str, set[str]] = {table: set() for table in delete_tables}
    used_by_key: dict[tuple[str, str], ForeignKeyRelation] = {}
    for table_paths in (*delete_paths.values(), *set_null_paths.values()):
        for path in table_paths:
            for relation in path:
                used_by_key[(relation.child_table, relation.constraint_name)] = relation
                if (
                    relation.parent_table in delete_tables
                    and relation.child_table in delete_tables
                    and relation.parent_table != relation.child_table
                ):
                    children[relation.parent_table].add(relation.child_table)

    ordered: list[str] = []
    visiting: set[str] = set()
    visited: set[str] = set()

    def visit(table_name: str) -> None:
        if table_name in visited:
            return
        if table_name in visiting:
            raise RuntimeError(f"Phát hiện vòng lặp FK tại {table_name}.")
        visiting.add(table_name)
        for child_table in sorted(children.get(table_name, ())):
            visit(child_table)
        visiting.remove(table_name)
        visited.add(table_name)
        ordered.append(table_name)

    visit(ROOT_TABLE)
    for table_name in sorted(delete_tables):
        visit(table_name)

    return DeletionPlan(
        delete_paths=delete_paths,
        set_null_paths=dict(set_null_paths),
        delete_order=ordered,
        used_relations=sorted(
            used_by_key.values(),
            key=lambda item: (
                item.parent_table,
                item.child_table,
                item.constraint_name,
            ),
        ),
    )


def _build_path_predicate(
    path: Sequence[ForeignKeyRelation],
    root_ids: Sequence[int],
    *,
    current_alias: str,
    alias_counter: list[int],
) -> tuple[str, list[int]]:
    if not path:
        placeholders = make_placeholders(root_ids)
        return (
            f"{current_alias}.{quote_identifier(ROOT_COLUMN)} IN ({placeholders})",
            list(root_ids),
        )

    relation = path[-1]
    alias_counter[0] += 1
    parent_alias = f"p{alias_counter[0]}"
    joins = [
        (
            f"{current_alias}.{quote_identifier(child_column)} = "
            f"{parent_alias}.{quote_identifier(parent_column)}"
        )
        for child_column, parent_column in zip(
            relation.child_columns, relation.parent_columns
        )
    ]
    parent_predicate, params = _build_path_predicate(
        path[:-1],
        root_ids,
        current_alias=parent_alias,
        alias_counter=alias_counter,
    )
    predicate = (
        f"EXISTS (SELECT 1 FROM {quote_identifier(relation.parent_table)} "
        f"AS {parent_alias} WHERE {' AND '.join(joins)} "
        f"AND ({parent_predicate}))"
    )
    return predicate, params


def build_table_filter(
    paths: Sequence[Sequence[ForeignKeyRelation]],
    root_ids: Sequence[int],
    *,
    table_alias: str = "target_row",
) -> tuple[str, list[int]]:
    if not paths:
        raise ValueError("Không có đường quan hệ để tạo bộ lọc.")
    predicates: list[str] = []
    params: list[int] = []
    alias_counter = [0]
    for path in paths:
        predicate, path_params = _build_path_predicate(
            path,
            root_ids,
            current_alias=table_alias,
            alias_counter=alias_counter,
        )
        predicates.append(f"({predicate})")
        params.extend(path_params)
    return " OR ".join(predicates), params


def fetch_root_rows(
    cursor, ids: Sequence[int]
) -> tuple[list[str], list[tuple[Any, ...]]]:
    placeholders = make_placeholders(ids)
    cursor.execute(
        f"SELECT * FROM {quote_identifier(ROOT_TABLE)} "
        f"WHERE {quote_identifier(ROOT_COLUMN)} IN ({placeholders}) "
        f"ORDER BY {quote_identifier(ROOT_COLUMN)}",
        list(ids),
    )
    columns = [description[0] for description in cursor.description]
    return columns, list(cursor.fetchall())


def validate_all_ids_exist(
    ids: Sequence[int], root_columns: Sequence[str], root_rows: Sequence[Sequence[Any]]
) -> None:
    try:
        id_index = list(root_columns).index(ROOT_COLUMN)
    except ValueError as exc:
        raise RuntimeError(f"Bảng {ROOT_TABLE} không có cột {ROOT_COLUMN}.") from exc
    found_ids = {int(row[id_index]) for row in root_rows}
    missing_ids = [value for value in ids if value not in found_ids]
    if missing_ids:
        raise ValueError(f"r_parameter.id không tồn tại: {missing_ids}")


def count_rows_for_paths(
    cursor,
    table_name: str,
    paths: Sequence[Sequence[ForeignKeyRelation]],
    ids: Sequence[int],
) -> int:
    where_sql, params = build_table_filter(paths, ids)
    cursor.execute(
        f"SELECT COUNT(*) FROM {quote_identifier(table_name)} AS target_row "
        f"WHERE {where_sql}",
        params,
    )
    return int(cursor.fetchone()[0])


def fetch_snapshot(
    cursor,
    table_name: str,
    action: str,
    paths: Sequence[Sequence[ForeignKeyRelation]],
    ids: Sequence[int],
    row_limit: int,
) -> TableSnapshot:
    where_sql, params = build_table_filter(paths, ids)
    cursor.execute(
        f"SELECT * FROM {quote_identifier(table_name)} AS target_row "
        f"WHERE {where_sql} LIMIT %s",
        [*params, row_limit],
    )
    columns = [description[0] for description in cursor.description]
    rows = list(cursor.fetchall())
    total_rows = count_rows_for_paths(cursor, table_name, paths, ids)
    return TableSnapshot(table_name, action, columns, rows, total_rows)


def collect_snapshots(
    cursor, config: RunConfig, plan: DeletionPlan
) -> list[TableSnapshot]:
    snapshots = [
        fetch_snapshot(
            cursor,
            table_name,
            "DELETE",
            plan.delete_paths[table_name],
            config.r_parameter_ids,
            config.report_row_limit,
        )
        for table_name in plan.delete_order
    ]
    snapshots.extend(
        fetch_snapshot(
            cursor,
            table_name,
            "SET NULL",
            plan.set_null_paths[table_name],
            config.r_parameter_ids,
            config.report_row_limit,
        )
        for table_name in sorted(plan.set_null_paths)
        if table_name not in plan.delete_paths
    )
    return snapshots


def delete_rows(
    config: RunConfig,
    plan: DeletionPlan,
    before_counts: dict[str, int],
) -> tuple[dict[str, int], dict[str, int]]:
    deleted_counts: dict[str, int] = {}
    with transaction.atomic():
        with connection.cursor() as cursor:
            current_counts = {
                table_name: count_rows_for_paths(
                    cursor,
                    table_name,
                    plan.delete_paths[table_name],
                    config.r_parameter_ids,
                )
                for table_name in plan.delete_order
            }
            if current_counts != before_counts:
                raise RuntimeError(
                    "Dữ liệu đã thay đổi sau preview; hủy xóa. "
                    f"Preview={before_counts}, hiện tại={current_counts}."
                )

            for table_name in plan.delete_order:
                where_sql, params = build_table_filter(
                    plan.delete_paths[table_name], config.r_parameter_ids
                )
                cursor.execute(
                    f"DELETE target_row FROM {quote_identifier(table_name)} "
                    "AS target_row "
                    f"WHERE {where_sql}",
                    params,
                )
                deleted_count = int(cursor.rowcount)
                expected_count = before_counts[table_name]
                if deleted_count != expected_count:
                    raise RuntimeError(
                        f"{table_name}: dự kiến {expected_count}, thực tế "
                        f"{deleted_count}. Đã rollback."
                    )
                deleted_counts[table_name] = deleted_count

            placeholders = make_placeholders(config.r_parameter_ids)
            cursor.execute(
                f"SELECT COUNT(*) FROM {quote_identifier(ROOT_TABLE)} "
                f"WHERE {quote_identifier(ROOT_COLUMN)} IN ({placeholders})",
                list(config.r_parameter_ids),
            )
            remaining = int(cursor.fetchone()[0])
            if remaining:
                raise RuntimeError(f"Vẫn còn {remaining} R parameter. Đã rollback.")
            after_counts = {table_name: 0 for table_name in plan.delete_order}
    return deleted_counts, after_counts


def excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, datetime, date)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    text_value = value.hex() if isinstance(value, bytes) else str(value)
    text_value = EXCEL_ILLEGAL_CHARACTER_PATTERN.sub("�", text_value)
    if text_value.startswith(("=", "+", "-", "@")):
        text_value = "'" + text_value
    if len(text_value) > EXCEL_CELL_LIMIT:
        return text_value[: EXCEL_CELL_LIMIT - 16] + "...[TRUNCATED]"
    return text_value


def style_table_sheet(worksheet, *, header_row: int = 1) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    body_font = Font(name="Arial", size=10, color="222222")
    for row in worksheet.iter_rows():
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
    for cell in worksheet[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for column_cells in worksheet.columns:
        max_length = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[letter].width = min(max_length + 2, 60)


def style_summary_sheet(worksheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet["A2"].font = Font(name="Arial", size=14, bold=True)
    for cell in worksheet[4]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=5):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 90


def append_mapping_rows(worksheet, rows: Sequence[dict[str, Any]]) -> None:
    if not rows:
        worksheet.append(["Không có dữ liệu"])
        return
    headers = list(rows[0])
    worksheet.append(headers)
    for row in rows:
        worksheet.append([excel_value(row.get(header)) for header in headers])


def unique_sheet_title(workbook: Workbook, raw_title: str) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", raw_title)[:31] or "Sheet"
    title = base
    suffix = 2
    while title in workbook.sheetnames:
        suffix_text = f"_{suffix}"
        title = f"{base[: 31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    return title


def write_report(
    config: RunConfig,
    *,
    status: str,
    database_name: str,
    database_host: str,
    plan: DeletionPlan | None,
    snapshots: Sequence[TableSnapshot],
    deleted_counts: dict[str, int] | None,
    after_counts: dict[str, int] | None,
    message: str,
    allow_overwrite: bool = False,
) -> None:
    output_path = config.output_path
    if output_path.exists() and not allow_overwrite:
        raise FileExistsError(f"File báo cáo đã tồn tại: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A2"] = "R parameter deletion report"
    summary.append([])
    summary.append(["Field", "Value"])
    summary_rows = [
        ("status", status),
        ("created_at", datetime.now()),
        ("database", database_name),
        ("host", database_host),
        ("r_parameter_ids", ", ".join(map(str, config.r_parameter_ids))),
        ("execute_requested", config.execute),
        ("assume_yes", config.assume_yes),
        ("allow_external_orphans", config.allow_external_orphans),
        ("message", message),
    ]
    for field, value in summary_rows:
        summary.append([field, excel_value(value)])
    summary["B6"].number_format = "yyyy-mm-dd hh:mm:ss"
    style_summary_sheet(summary)

    by_table = {snapshot.table_name: snapshot for snapshot in snapshots}
    count_sheet = workbook.create_sheet("Table counts")
    count_sheet.append(
        [
            "delete_order",
            "table_name",
            "action",
            "before",
            "deleted",
            "after",
            "reported_rows",
            "truncated",
            "note",
        ]
    )
    ordered_tables: list[str] = []
    if plan is not None:
        ordered_tables.extend(plan.delete_order)
        ordered_tables.extend(
            table
            for table in sorted(plan.set_null_paths)
            if table not in plan.delete_paths
        )
    ordered_tables.extend(table for table in by_table if table not in ordered_tables)
    for table_name in ordered_tables:
        snapshot = by_table.get(table_name)
        if snapshot is None:
            continue
        order = (
            plan.delete_order.index(table_name) + 1
            if plan is not None and table_name in plan.delete_order
            else None
        )
        count_sheet.append(
            [
                order,
                table_name,
                snapshot.action,
                snapshot.total_rows,
                None if deleted_counts is None else deleted_counts.get(table_name),
                None if after_counts is None else after_counts.get(table_name),
                len(snapshot.rows),
                snapshot.truncated,
                EXTERNAL_RESOURCE_TABLES.get(table_name, ""),
            ]
        )
    style_table_sheet(count_sheet)

    relation_sheet = workbook.create_sheet("Relationships")
    relation_rows = (
        []
        if plan is None
        else [
            {
                "constraint_name": item.constraint_name,
                "source": item.source,
                "parent_table": item.parent_table,
                "parent_columns": ", ".join(item.parent_columns),
                "child_table": item.child_table,
                "child_columns": ", ".join(item.child_columns),
                "delete_rule": item.delete_rule,
            }
            for item in plan.used_relations
        ]
    )
    append_mapping_rows(relation_sheet, relation_rows)
    style_table_sheet(relation_sheet)

    for snapshot in snapshots:
        worksheet = workbook.create_sheet(
            unique_sheet_title(workbook, f"data_{snapshot.table_name}")
        )
        worksheet.append(snapshot.columns or ["Không có cột"])
        for row in snapshot.rows:
            worksheet.append([excel_value(value) for value in row])
        if not snapshot.rows:
            worksheet.append(["Không có dòng phù hợp"])
        if snapshot.truncated:
            worksheet.append([])
            worksheet.append(
                [
                    f"Chỉ xuất {len(snapshot.rows)}/{snapshot.total_rows} dòng; "
                    "xem Table counts để biết tổng chính xác."
                ]
            )
        style_table_sheet(worksheet)

    temp_path = output_path.with_name(f".{output_path.name}.tmp.xlsx")
    workbook.save(temp_path)
    temp_path.replace(output_path)


def confirmation_text(config: RunConfig, database_name: str) -> str:
    ids = ",".join(map(str, config.r_parameter_ids))
    return f"DELETE {database_name} R_PARAMETER_IDS {ids}"


def print_separator() -> None:
    print("=" * 78)


def main(argv: Sequence[str] | None = None) -> int:
    config = parse_args(argv)
    database_name, database_host = database_identity()
    plan: DeletionPlan | None = None
    snapshots: list[TableSnapshot] = []
    deleted_counts: dict[str, int] | None = None
    after_counts: dict[str, int] | None = None
    report_written = False
    mutation_committed = False

    try:
        with connection.cursor() as cursor:
            root_columns, root_rows = fetch_root_rows(cursor, config.r_parameter_ids)
            validate_all_ids_exist(config.r_parameter_ids, root_columns, root_rows)
            relations = add_legacy_relationships(cursor, discover_foreign_keys(cursor))
            plan = build_deletion_plan(relations)
            snapshots = collect_snapshots(cursor, config, plan)

        before_counts = {
            item.table_name: item.total_rows
            for item in snapshots
            if item.action == "DELETE"
        }
        external_impacts = {
            item.table_name: item.total_rows
            for item in snapshots
            if item.table_name in EXTERNAL_RESOURCE_TABLES and item.total_rows > 0
        }

        print_separator()
        print("XEM TRƯỚC XÓA R PARAMETER")
        print_separator()
        print(f"Database         : {database_name}")
        print(f"Host             : {database_host}")
        print(f"R parameter IDs  : {list(config.r_parameter_ids)}")
        print(f"Execute          : {config.execute}")
        print()
        for order, table_name in enumerate(plan.delete_order, start=1):
            print(
                f"{order:>2}. {table_name:<42} " f"{before_counts[table_name]:>8} dòng"
            )
        for snapshot in snapshots:
            if snapshot.action == "SET NULL":
                print(
                    f"--  {snapshot.table_name:<42} "
                    f"{snapshot.total_rows:>8} dòng SET NULL"
                )
        print(f"Tổng dòng DELETE : {sum(before_counts.values())}")

        preview_message = "Đã xem trước; chưa xóa dữ liệu."
        if external_impacts:
            preview_message += " Tài nguyên ngoài DB: " + ", ".join(
                f"{table}={count}" for table, count in external_impacts.items()
            )
        write_report(
            config,
            status="PREVIEW",
            database_name=database_name,
            database_host=database_host,
            plan=plan,
            snapshots=snapshots,
            deleted_counts=None,
            after_counts=None,
            message=preview_message,
        )
        report_written = True
        print(f"Báo cáo Excel    : {config.output_path}")

        if not config.execute:
            print_separator()
            print("RUN_EXECUTE=False/không có --execute: chưa xóa dữ liệu.")
            return 0

        if external_impacts and not config.allow_external_orphans:
            blocked_message = (
                "Không cho phép execute vì có metadata tài nguyên ngoài DB. "
                "Chỉ dùng --allow-external-orphans sau khi kiểm tra báo cáo."
            )
            write_report(
                config,
                status="BLOCKED",
                database_name=database_name,
                database_host=database_host,
                plan=plan,
                snapshots=snapshots,
                deleted_counts=None,
                after_counts=None,
                message=blocked_message,
                allow_overwrite=True,
            )
            print(blocked_message)
            return 1

        expected_confirmation = confirmation_text(config, database_name)
        if not config.assume_yes:
            print_separator()
            print("Để xác nhận xóa, nhập chính xác nội dung sau:")
            print(expected_confirmation)
            if input("> ").strip() != expected_confirmation:
                write_report(
                    config,
                    status="CANCELLED",
                    database_name=database_name,
                    database_host=database_host,
                    plan=plan,
                    snapshots=snapshots,
                    deleted_counts=None,
                    after_counts=None,
                    message="Người dùng hủy hoặc nhập sai chuỗi xác nhận.",
                    allow_overwrite=True,
                )
                print("Đã hủy. Không có dữ liệu nào bị xóa.")
                return 0

        deleted_counts, after_counts = delete_rows(config, plan, before_counts)
        mutation_committed = True
        write_report(
            config,
            status="DELETED",
            database_name=database_name,
            database_host=database_host,
            plan=plan,
            snapshots=snapshots,
            deleted_counts=deleted_counts,
            after_counts=after_counts,
            message=(
                "Đã commit trong một transaction. Các sheet data_* là snapshot "
                "trước khi xóa."
            ),
            allow_overwrite=True,
        )
        print_separator()
        print(f"Đã xóa và commit {sum(deleted_counts.values())} dòng.")
        print(f"Báo cáo đã cập nhật: {config.output_path}")
        return 0
    except Exception as exc:
        error_message = f"{type(exc).__name__}: {exc}"
        print_separator()
        print(f"ERROR: {error_message}", file=sys.stderr)
        if not mutation_committed:
            print("Không có transaction xóa nào được commit.", file=sys.stderr)
        try:
            write_report(
                config,
                status="ERROR_AFTER_COMMIT" if mutation_committed else "ERROR",
                database_name=database_name,
                database_host=database_host,
                plan=plan,
                snapshots=snapshots,
                deleted_counts=deleted_counts,
                after_counts=after_counts,
                message=error_message,
                allow_overwrite=report_written,
            )
            print(f"Báo cáo lỗi: {config.output_path}", file=sys.stderr)
        except Exception as report_exc:
            print(
                f"Không thể ghi báo cáo lỗi: {type(report_exc).__name__}: "
                f"{report_exc}",
                file=sys.stderr,
            )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
