"""Copy dữ liệu R-list giữa hai project/phase theo cách có thể audit.

Phạm vi copy đã được review:

* ``project_r_parameter``
* ``temp_data``
* ``project_summary`` với ``category = 'R'``
* ``detail_data_rlist``

Script không copy attachment/S3, lịch sử phê duyệt hoặc dữ liệu study/design.
Các bảng liên quan có dữ liệu sẽ được ghi vào báo cáo Excel để người vận hành
xem xét riêng. Báo cáo liệt kê từng record không được copy, business key, dữ
liệu nguồn/đích và lý do chi tiết.

Mặc định chỉ preview và tạo báo cáo. Để chạy bằng nút Run của PyCharm, sửa các
hằng ``RUN_*`` bên dưới. Chỉ đổi ``RUN_EXECUTE`` thành ``True`` sau khi đã kiểm
tra workbook preview.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence


PROJECT_ROOT = Path(__file__).resolve().parents[3]

if __name__ == "__main__":
    import os

    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


# =========================================================
# CẤU HÌNH CHẠY BẰNG PYCHARM
# =========================================================

RUN_SOURCE_PROJECT_ID = 1019
RUN_SOURCE_PHASE_ID = 33
RUN_DESTINATION_PROJECT_ID = 1020
RUN_DESTINATION_PHASE_ID = 35

# False: chỉ xem trước và xuất Excel. True: cho phép đi tới bước xác nhận copy.
RUN_EXECUTE = False

# Nên luôn giữ False. True sẽ bỏ qua xác nhận tương tác chính xác.
RUN_ASSUME_YES = False

# None: tự tạo file có timestamp trong utils/db_scripts/uu3/output/.
RUN_OUTPUT_PATH: str | None = None

# Workbook vẫn ghi tổng chính xác nếu sheet chi tiết bị giới hạn số dòng.
RUN_REPORT_ROW_LIMIT = 50_000


CORE_TABLES = (
    "project_r_parameter",
    "temp_data",
    "project_summary",
    "detail_data_rlist",
)

TABLE_KEY_COLUMNS: dict[str, tuple[str, ...]] = {
    "project_r_parameter": ("r_parameter_id", "variation_id"),
    "project_summary": ("r_parameter_id", "category"),
    # detail_data_rlist là bảng legacy không có PK/unique constraint. Đây là
    # business key đang được các API R-list dùng để xác định một dòng dữ liệu.
    "detail_data_rlist": ("r_parameter_id", "variation_id"),
    # UNIQUE(project_id, r_parameter_id, variation_id, phase_id) được thu gọn
    # theo scope project/phase mà script luôn áp dụng khi đọc và ghi dữ liệu.
    "temp_data": ("r_parameter_id", "variation_id"),
}

REQUIRED_COLUMNS: dict[str, set[str]] = {
    "project_r_parameter": {
        "project_id",
        "r_parameter_id",
        "phase_id",
        "variation_id",
    },
    "project_summary": {
        "project_id",
        "phase_id",
        "r_parameter_id",
        "category",
    },
    "detail_data_rlist": {
        "project_id",
        "r_parameter_id",
        "phase_id",
        "variation_id",
    },
    "temp_data": {
        "project_id",
        "r_parameter_id",
        "phase_id",
        "variation_id",
    },
}

TABLE_COPY_ORDER = (
    "project_r_parameter",
    "temp_data",
    "project_summary",
    "detail_data_rlist",
)

RELATED_TABLE_REASONS = {
    "project_parameter_attachments": (
        "Không copy: bản ghi liên kết file/S3 cần quy trình tài nguyên riêng."
    ),
    "project_r_parameter_manager_approval_history": (
        "Không copy: đây là lịch sử audit/phê duyệt của project nguồn."
    ),
    "detail_data_design": "Không copy: dữ liệu study/design nằm ngoài R-list cốt lõi.",
}

IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z0-9_$]+$")
EXCEL_ILLEGAL_CHARACTER_PATTERN = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")
EXCEL_CELL_LIMIT = 32_767

ACTION_COPY = "COPY"
ACTION_SKIP = "SKIP_EXISTING"
ACTION_SKIP_MISSING_PARENT = "SKIP_MISSING_PARENT"
ACTION_SKIP_OUT_OF_SCOPE = "SKIP_OUT_OF_SCOPE"
ACTION_BLOCK_SOURCE_DUPLICATE = "BLOCK_SOURCE_DUPLICATE"
ACTION_BLOCK_DESTINATION_DUPLICATE = "BLOCK_DESTINATION_DUPLICATE"
ACTION_BLOCK_INVALID_KEY = "BLOCK_INVALID_KEY"


@dataclass(frozen=True)
class RunConfig:
    source_project_id: int
    source_phase_id: int
    destination_project_id: int
    destination_phase_id: int
    execute: bool
    assume_yes: bool
    output_path: Path
    report_row_limit: int


@dataclass(frozen=True)
class RowDecision:
    action: str
    reason: str
    key: tuple[Any, ...]
    source_row: dict[str, Any]
    target_row: dict[str, Any]
    destination_row: dict[str, Any] | None


@dataclass
class TableCopyPlan:
    table_name: str
    columns: tuple[str, ...]
    key_columns: tuple[str, ...]
    source_count: int
    destination_before_count: int
    decisions: list[RowDecision] = field(default_factory=list)

    def action_count(self, action: str) -> int:
        return sum(decision.action == action for decision in self.decisions)

    @property
    def copy_rows(self) -> list[dict[str, Any]]:
        return [
            decision.target_row
            for decision in self.decisions
            if decision.action == ACTION_COPY
        ]

    @property
    def skipped_count(self) -> int:
        return sum(decision.action.startswith("SKIP_") for decision in self.decisions)

    @property
    def blocked_count(self) -> int:
        return sum(decision.action.startswith("BLOCK_") for decision in self.decisions)

    @property
    def not_copied_count(self) -> int:
        return sum(decision.action != ACTION_COPY for decision in self.decisions)


@dataclass(frozen=True)
class RelatedTableImpact:
    table_name: str
    source_rows: int
    action: str
    reason: str
    source_records: list[dict[str, Any]] = field(default_factory=list)
    records_truncated: bool = False


@dataclass
class CopyPlan:
    source_project: dict[str, Any]
    destination_project: dict[str, Any]
    source_phase: dict[str, Any]
    destination_phase: dict[str, Any]
    table_plans: dict[str, TableCopyPlan]
    related_impacts: list[RelatedTableImpact]
    fk_relationships: list[dict[str, Any]]
    warnings: list[str]
    blockers: list[str]
    fingerprint: str = ""

    @property
    def total_copy_rows(self) -> int:
        return sum(len(plan.copy_rows) for plan in self.table_plans.values())

    @property
    def total_skipped_rows(self) -> int:
        return sum(plan.skipped_count for plan in self.table_plans.values())

    @property
    def total_blocked_rows(self) -> int:
        return sum(plan.blocked_count for plan in self.table_plans.values())

    @property
    def total_not_copied_rows(self) -> int:
        core_rows = sum(plan.not_copied_count for plan in self.table_plans.values())
        related_rows = sum(impact.source_rows for impact in self.related_impacts)
        return core_rows + related_rows


def quote_identifier(identifier: str) -> str:
    """Validate rồi quote tên bảng/cột lấy từ allowlist hoặc metadata."""
    if not IDENTIFIER_PATTERN.fullmatch(identifier):
        raise ValueError(f"Tên bảng/cột không hợp lệ: {identifier!r}")
    return f"`{identifier}`"


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Không thể tạo SQL placeholders cho danh sách rỗng.")
    return ", ".join(["%s"] * len(values))


def normalize_positive_id(raw_id: Any, *, label: str) -> int:
    if isinstance(raw_id, bool):
        raise ValueError(f"{label} phải là số nguyên dương, không phải boolean.")
    try:
        normalized_id = int(raw_id)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} không hợp lệ: {raw_id!r}") from exc
    if normalized_id <= 0:
        raise ValueError(f"{label} phải lớn hơn 0: {normalized_id}")
    return normalized_id


def normalize_report_row_limit(raw_limit: Any) -> int:
    limit = normalize_positive_id(raw_limit, label="report_row_limit")
    if limit > 1_000_000:
        raise ValueError("report_row_limit không được lớn hơn 1.000.000.")
    return limit


def default_output_path(config_values: tuple[int, int, int, int]) -> Path:
    source_project_id, source_phase_id, destination_project_id, destination_phase_id = (
        config_values
    )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    filename = (
        "Project_copy_variation_rlist_"
        f"p{source_project_id}_ph{source_phase_id}_to_"
        f"p{destination_project_id}_ph{destination_phase_id}_{timestamp}.xlsx"
    )
    return Path(__file__).resolve().parent / "output" / filename


def resolve_output_path(
    raw_path: str | None,
    *,
    config_values: tuple[int, int, int, int],
) -> Path:
    if raw_path:
        path = Path(raw_path).expanduser()
        if not path.is_absolute():
            path = Path.cwd() / path
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Đường dẫn báo cáo phải có phần mở rộng .xlsx.")
        return path.resolve()
    return default_output_path(config_values)


def parse_args(argv: Sequence[str] | None = None) -> RunConfig:
    parser = argparse.ArgumentParser(
        description=(
            "Preview/copy dữ liệu R-list giữa hai project/phase và xuất báo cáo Excel."
        )
    )
    parser.add_argument("--source-project-id", type=int, default=RUN_SOURCE_PROJECT_ID)
    parser.add_argument("--source-phase-id", type=int, default=RUN_SOURCE_PHASE_ID)
    parser.add_argument(
        "--destination-project-id", type=int, default=RUN_DESTINATION_PROJECT_ID
    )
    parser.add_argument(
        "--destination-phase-id", type=int, default=RUN_DESTINATION_PHASE_ID
    )
    parser.add_argument(
        "--execute",
        action="store_true",
        default=RUN_EXECUTE,
        help="Cho phép insert sau bước xác nhận.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        default=RUN_ASSUME_YES,
        help="Bỏ qua xác nhận tương tác (không khuyến nghị).",
    )
    parser.add_argument("--output", default=RUN_OUTPUT_PATH)
    parser.add_argument(
        "--report-row-limit",
        type=int,
        default=RUN_REPORT_ROW_LIMIT,
    )
    args = parser.parse_args(argv)

    config_values = (
        normalize_positive_id(args.source_project_id, label="source_project_id"),
        normalize_positive_id(args.source_phase_id, label="source_phase_id"),
        normalize_positive_id(
            args.destination_project_id, label="destination_project_id"
        ),
        normalize_positive_id(args.destination_phase_id, label="destination_phase_id"),
    )
    if config_values[:2] == config_values[2:]:
        raise ValueError("Source và destination project/phase không được giống nhau.")

    return RunConfig(
        source_project_id=config_values[0],
        source_phase_id=config_values[1],
        destination_project_id=config_values[2],
        destination_phase_id=config_values[3],
        execute=bool(args.execute),
        assume_yes=bool(args.yes),
        output_path=resolve_output_path(
            args.output,
            config_values=config_values,
        ),
        report_row_limit=normalize_report_row_limit(args.report_row_limit),
    )


def database_identity() -> tuple[str, str]:
    settings_dict = connection.settings_dict
    return str(settings_dict.get("NAME") or ""), str(settings_dict.get("HOST") or "")


def get_insertable_columns(cursor) -> dict[str, tuple[str, ...]]:
    """Lấy cột thật của các bảng từ metadata, bỏ generated/auto increment.

    Vì vậy ``temp_data.id`` không được mang từ source sang destination; MySQL
    tự cấp ID mới, còn toàn bộ cột workflow có thể insert được vẫn được giữ lại.
    """
    placeholders = make_placeholders(CORE_TABLES)
    cursor.execute(
        f"""
        SELECT c.TABLE_NAME, c.COLUMN_NAME, c.EXTRA
        FROM information_schema.COLUMNS AS c
        JOIN information_schema.TABLES AS t
          ON t.TABLE_SCHEMA = c.TABLE_SCHEMA
         AND t.TABLE_NAME = c.TABLE_NAME
         AND t.TABLE_TYPE = 'BASE TABLE'
        WHERE c.TABLE_SCHEMA = DATABASE()
          AND c.TABLE_NAME IN ({placeholders})
        ORDER BY c.TABLE_NAME, c.ORDINAL_POSITION
        """,
        list(CORE_TABLES),
    )

    columns_by_table: dict[str, list[str]] = {table: [] for table in CORE_TABLES}
    all_columns_by_table: dict[str, set[str]] = {table: set() for table in CORE_TABLES}
    for table_name, column_name, extra in cursor.fetchall():
        all_columns_by_table[table_name].add(column_name)
        normalized_extra = str(extra or "").lower()
        if (
            "auto_increment" not in normalized_extra
            and "virtual generated" not in normalized_extra
            and "stored generated" not in normalized_extra
        ):
            columns_by_table[table_name].append(column_name)

    schema_errors = []
    for table_name in CORE_TABLES:
        missing = sorted(
            REQUIRED_COLUMNS[table_name] - all_columns_by_table[table_name]
        )
        if missing:
            schema_errors.append(f"{table_name}: thiếu {', '.join(missing)}")
        if not columns_by_table[table_name]:
            schema_errors.append(f"{table_name}: không tìm thấy bảng/cột insert được")
    if schema_errors:
        raise RuntimeError(
            "Schema database không phù hợp với script: " + "; ".join(schema_errors)
        )

    return {table: tuple(columns) for table, columns in columns_by_table.items()}


def fetch_entities(
    cursor,
    config: RunConfig,
    *,
    for_update: bool,
) -> tuple[dict[int, dict[str, Any]], dict[int, dict[str, Any]]]:
    lock_sql = " FOR UPDATE" if for_update else ""
    project_ids = [config.source_project_id, config.destination_project_id]
    phase_ids = [config.source_phase_id, config.destination_phase_id]

    cursor.execute(
        f"""
        SELECT id, project_code, setup_id
        FROM project_info
        WHERE id IN ({make_placeholders(project_ids)})
        ORDER BY id{lock_sql}
        """,
        project_ids,
    )
    projects = {
        int(row[0]): {"id": int(row[0]), "project_code": row[1], "setup_id": row[2]}
        for row in cursor.fetchall()
    }

    cursor.execute(
        f"""
        SELECT id, phase
        FROM phase
        WHERE id IN ({make_placeholders(phase_ids)})
        ORDER BY id{lock_sql}
        """,
        phase_ids,
    )
    phases = {
        int(row[0]): {"id": int(row[0]), "phase": row[1]} for row in cursor.fetchall()
    }

    missing_projects = [
        project_id for project_id in project_ids if project_id not in projects
    ]
    missing_phases = [phase_id for phase_id in phase_ids if phase_id not in phases]
    if missing_projects or missing_phases:
        problems = []
        if missing_projects:
            problems.append(f"project_id không tồn tại: {missing_projects}")
        if missing_phases:
            problems.append(f"phase_id không tồn tại: {missing_phases}")
        raise ValueError("; ".join(problems))
    return projects, phases


def fetch_scope_rows(
    cursor,
    *,
    table_name: str,
    columns: Sequence[str],
    project_id: int,
    phase_id: int,
    category: str | None = None,
    for_update: bool,
) -> list[dict[str, Any]]:
    if table_name not in CORE_TABLES:
        raise ValueError(f"Bảng không nằm trong allowlist: {table_name}")
    column_sql = ", ".join(quote_identifier(column) for column in columns)
    params: list[Any] = [project_id, phase_id]
    where_parts = ["project_id = %s", "phase_id = %s"]
    if category is not None:
        where_parts.append("category = %s")
        params.append(category)
    order_sql = ", ".join(
        quote_identifier(column) for column in TABLE_KEY_COLUMNS[table_name]
    )
    lock_sql = " FOR UPDATE" if for_update else ""
    cursor.execute(
        f"""
        SELECT {column_sql}
        FROM {quote_identifier(table_name)}
        WHERE {" AND ".join(where_parts)}
        ORDER BY {order_sql}{lock_sql}
        """,
        params,
    )
    return [dict(zip(columns, row, strict=True)) for row in cursor.fetchall()]


def row_key(row: dict[str, Any], key_columns: Sequence[str]) -> tuple[Any, ...]:
    return tuple(row.get(column) for column in key_columns)


def differing_columns(
    left_row: dict[str, Any],
    right_row: dict[str, Any] | None,
) -> tuple[str, ...]:
    if right_row is None:
        return ()
    return tuple(
        column for column in left_row if left_row.get(column) != right_row.get(column)
    )


def make_missing_parent_decisions(
    rows: Sequence[dict[str, Any]],
    *,
    key_columns: Sequence[str],
    destination_project_id: int,
    destination_phase_id: int,
    parent_table: str,
) -> list[RowDecision]:
    decisions = []
    for source_row in rows:
        key = row_key(source_row, key_columns)
        target_row = dict(source_row)
        target_row["project_id"] = destination_project_id
        target_row["phase_id"] = destination_phase_id
        decisions.append(
            RowDecision(
                action=ACTION_SKIP_MISSING_PARENT,
                reason=(
                    f"Không copy vì source không có record {parent_table} tương ứng "
                    f"với business key {key!r}."
                ),
                key=key,
                source_row=dict(source_row),
                target_row=target_row,
                destination_row=None,
            )
        )
    return decisions


def make_out_of_scope_decisions(
    rows: Sequence[dict[str, Any]],
    *,
    key_columns: Sequence[str],
    destination_project_id: int,
    destination_phase_id: int,
    reason: str,
) -> list[RowDecision]:
    decisions = []
    for source_row in rows:
        key = row_key(source_row, key_columns)
        target_row = dict(source_row)
        target_row["project_id"] = destination_project_id
        target_row["phase_id"] = destination_phase_id
        decisions.append(
            RowDecision(
                action=ACTION_SKIP_OUT_OF_SCOPE,
                reason=reason,
                key=key,
                source_row=dict(source_row),
                target_row=target_row,
                destination_row=None,
            )
        )
    return decisions


def classify_rows(
    source_rows: Sequence[dict[str, Any]],
    destination_rows: Sequence[dict[str, Any]],
    *,
    key_columns: Sequence[str],
    destination_project_id: int,
    destination_phase_id: int,
) -> list[RowDecision]:
    """Phân loại COPY/SKIP/BLOCK theo business key, không overwrite."""
    source_counts = Counter(row_key(row, key_columns) for row in source_rows)
    destination_counts = Counter(row_key(row, key_columns) for row in destination_rows)
    destination_by_key: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for destination_row in destination_rows:
        key = row_key(destination_row, key_columns)
        destination_by_key.setdefault(key, []).append(destination_row)
    decisions: list[RowDecision] = []

    for source_row in source_rows:
        key = row_key(source_row, key_columns)
        target_row = dict(source_row)
        target_row["project_id"] = destination_project_id
        target_row["phase_id"] = destination_phase_id
        destination_row = (
            dict(destination_by_key[key][0]) if destination_counts[key] == 1 else None
        )

        null_key_columns = [
            column
            for column, value in zip(key_columns, key, strict=True)
            if value is None
        ]
        if null_key_columns:
            action = ACTION_BLOCK_INVALID_KEY
            reason = (
                "Không copy vì business key chứa NULL tại cột: "
                f"{', '.join(null_key_columns)}."
            )
        elif source_counts[key] > 1:
            action = ACTION_BLOCK_SOURCE_DUPLICATE
            reason = (
                f"Nguồn có {source_counts[key]} dòng cùng business key; "
                "cần làm sạch hoặc xác nhận quy tắc trước khi copy."
            )
        elif destination_counts[key] > 1:
            action = ACTION_BLOCK_DESTINATION_DUPLICATE
            reason = (
                f"Đích có {destination_counts[key]} dòng cùng business key; "
                "không tự chọn dòng để ghi đè."
            )
        elif destination_counts[key] == 1:
            action = ACTION_SKIP
            if destination_row == target_row:
                reason = (
                    "Không copy vì business key và toàn bộ nội dung đã tồn tại "
                    "giống hệt ở destination."
                )
            else:
                changed_columns = differing_columns(target_row, destination_row)
                reason = (
                    "Không copy vì business key đã tồn tại ở destination nhưng "
                    "nội dung khác source tại cột: "
                    f"{', '.join(changed_columns)}. Script giữ nguyên destination "
                    "và không overwrite."
                )
        else:
            action = ACTION_COPY
            reason = "Business key chưa tồn tại ở destination."

        decisions.append(
            RowDecision(
                action=action,
                reason=reason,
                key=key,
                source_row=dict(source_row),
                target_row=target_row,
                destination_row=destination_row,
            )
        )
    return decisions


def fetch_existing_master_ids(
    cursor,
    *,
    table_name: str,
    ids: Sequence[int],
) -> set[int]:
    if table_name not in {"variation", "r_parameter"}:
        raise ValueError(f"Master table không nằm trong allowlist: {table_name}")
    if not ids:
        return set()
    cursor.execute(
        f"SELECT id FROM {quote_identifier(table_name)} "
        f"WHERE id IN ({make_placeholders(ids)})",
        list(ids),
    )
    return {int(row[0]) for row in cursor.fetchall()}


def fetch_destination_setup_gaps(
    cursor,
    *,
    destination_setup_id: int,
    r_parameter_ids: Sequence[int],
) -> list[int]:
    if not r_parameter_ids:
        return []
    cursor.execute(
        f"""
        SELECT DISTINCT r_parameter_id
        FROM setup_r_relation
        WHERE setup_id = %s
          AND r_parameter_id IN ({make_placeholders(r_parameter_ids)})
        """,
        [destination_setup_id, *r_parameter_ids],
    )
    existing = {int(row[0]) for row in cursor.fetchall()}
    return [
        r_parameter_id
        for r_parameter_id in r_parameter_ids
        if r_parameter_id not in existing
    ]


def discover_related_tables(cursor) -> list[str]:
    """Tìm base table có đủ project/r_parameter/phase/variation columns."""
    cursor.execute(
        """
        SELECT c.TABLE_NAME
        FROM information_schema.COLUMNS AS c
        JOIN information_schema.TABLES AS t
          ON t.TABLE_SCHEMA = c.TABLE_SCHEMA
         AND t.TABLE_NAME = c.TABLE_NAME
         AND t.TABLE_TYPE = 'BASE TABLE'
        WHERE c.TABLE_SCHEMA = DATABASE()
          AND c.COLUMN_NAME IN (
              'project_id', 'r_parameter_id', 'phase_id', 'variation_id'
          )
        GROUP BY c.TABLE_NAME
        HAVING COUNT(DISTINCT c.COLUMN_NAME) = 4
        ORDER BY c.TABLE_NAME
        """
    )
    return [str(row[0]) for row in cursor.fetchall()]


def fetch_primary_key_columns(cursor, table_name: str) -> list[str]:
    quote_identifier(table_name)
    cursor.execute(
        """
        SELECT kcu.COLUMN_NAME
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        WHERE kcu.TABLE_SCHEMA = DATABASE()
          AND kcu.TABLE_NAME = %s
          AND kcu.CONSTRAINT_NAME = 'PRIMARY'
        ORDER BY kcu.ORDINAL_POSITION
        """,
        [table_name],
    )
    return [str(row[0]) for row in cursor.fetchall()]


def fetch_related_impacts(
    cursor,
    *,
    config: RunConfig,
) -> list[RelatedTableImpact]:
    impacts: list[RelatedTableImpact] = []
    for table_name in discover_related_tables(cursor):
        if table_name in CORE_TABLES:
            continue
        cursor.execute(
            f"""
            SELECT COUNT(*)
            FROM {quote_identifier(table_name)}
            WHERE project_id = %s
              AND phase_id = %s
            """,
            [config.source_project_id, config.source_phase_id],
        )
        row_count = int(cursor.fetchone()[0])
        identity_columns = [
            "project_id",
            "r_parameter_id",
            "variation_id",
            "phase_id",
        ]
        for primary_key_column in fetch_primary_key_columns(cursor, table_name):
            if primary_key_column not in identity_columns:
                identity_columns.append(primary_key_column)
        identity_column_sql = ", ".join(
            quote_identifier(column) for column in identity_columns
        )
        cursor.execute(
            f"""
            SELECT {identity_column_sql}
            FROM {quote_identifier(table_name)}
            WHERE project_id = %s
              AND phase_id = %s
            ORDER BY r_parameter_id, variation_id
            LIMIT %s
            """,
            [
                config.source_project_id,
                config.source_phase_id,
                config.report_row_limit,
            ],
        )
        source_records = [
            dict(zip(identity_columns, row, strict=True)) for row in cursor.fetchall()
        ]
        impacts.append(
            RelatedTableImpact(
                table_name=table_name,
                source_rows=row_count,
                action="NOT_COPIED",
                reason=RELATED_TABLE_REASONS.get(
                    table_name,
                    "Không copy: bảng nằm ngoài phạm vi R-list cốt lõi đã review.",
                ),
                source_records=source_records,
                records_truncated=row_count > len(source_records),
            )
        )
    return impacts


def fetch_fk_relationships(cursor) -> list[dict[str, Any]]:
    placeholders = make_placeholders(CORE_TABLES)
    params = [*CORE_TABLES, *CORE_TABLES]
    cursor.execute(
        f"""
        SELECT
            kcu.CONSTRAINT_NAME,
            kcu.TABLE_NAME AS child_table,
            GROUP_CONCAT(
                kcu.COLUMN_NAME ORDER BY kcu.ORDINAL_POSITION SEPARATOR ', '
            ) AS child_columns,
            kcu.REFERENCED_TABLE_NAME AS parent_table,
            GROUP_CONCAT(
                kcu.REFERENCED_COLUMN_NAME
                ORDER BY kcu.ORDINAL_POSITION SEPARATOR ', '
            ) AS parent_columns,
            rc.DELETE_RULE,
            rc.UPDATE_RULE
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        JOIN information_schema.REFERENTIAL_CONSTRAINTS AS rc
          ON rc.CONSTRAINT_SCHEMA = kcu.CONSTRAINT_SCHEMA
         AND rc.CONSTRAINT_NAME = kcu.CONSTRAINT_NAME
         AND rc.TABLE_NAME = kcu.TABLE_NAME
        WHERE kcu.TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_SCHEMA = DATABASE()
          AND (
              kcu.TABLE_NAME IN ({placeholders})
              OR kcu.REFERENCED_TABLE_NAME IN ({placeholders})
          )
        GROUP BY
            kcu.CONSTRAINT_NAME,
            kcu.TABLE_NAME,
            kcu.REFERENCED_TABLE_NAME,
            rc.DELETE_RULE,
            rc.UPDATE_RULE
        ORDER BY parent_table, child_table, kcu.CONSTRAINT_NAME
        """,
        params,
    )
    return [
        {
            "constraint_name": row[0],
            "child_table": row[1],
            "child_columns": row[2],
            "parent_table": row[3],
            "parent_columns": row[4],
            "delete_rule": row[5],
            "update_rule": row[6],
        }
        for row in cursor.fetchall()
    ]


def canonical_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    return str(value)


def canonical_row(row: dict[str, Any] | None) -> dict[str, Any] | None:
    if row is None:
        return None
    return {column: canonical_value(value) for column, value in row.items()}


def calculate_plan_fingerprint(table_plans: dict[str, TableCopyPlan]) -> str:
    payload = []
    for table_name in TABLE_COPY_ORDER:
        table_plan = table_plans[table_name]
        payload.append(
            {
                "table": table_name,
                "source_count": table_plan.source_count,
                "destination_before_count": table_plan.destination_before_count,
                "decisions": [
                    {
                        "action": decision.action,
                        "key": [canonical_value(value) for value in decision.key],
                        "target_row": {
                            column: canonical_value(decision.target_row.get(column))
                            for column in table_plan.columns
                        },
                        "destination_row": canonical_row(decision.destination_row),
                    }
                    for decision in table_plan.decisions
                ],
            }
        )
    serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(serialized).hexdigest()


def build_copy_plan(
    cursor,
    config: RunConfig,
    *,
    for_update: bool,
) -> CopyPlan:
    columns_by_table = get_insertable_columns(cursor)
    projects, phases = fetch_entities(cursor, config, for_update=for_update)

    source_project_rows = fetch_scope_rows(
        cursor,
        table_name="project_r_parameter",
        columns=columns_by_table["project_r_parameter"],
        project_id=config.source_project_id,
        phase_id=config.source_phase_id,
        for_update=for_update,
    )
    destination_project_rows = fetch_scope_rows(
        cursor,
        table_name="project_r_parameter",
        columns=columns_by_table["project_r_parameter"],
        project_id=config.destination_project_id,
        phase_id=config.destination_phase_id,
        for_update=for_update,
    )

    source_parent_keys = {
        row_key(row, TABLE_KEY_COLUMNS["project_r_parameter"])
        for row in source_project_rows
    }
    source_r_parameter_ids = sorted(
        {int(row["r_parameter_id"]) for row in source_project_rows}
    )
    source_r_parameter_id_set = set(source_r_parameter_ids)
    source_variation_ids = sorted(
        {int(row["variation_id"]) for row in source_project_rows}
    )

    source_summary_all = fetch_scope_rows(
        cursor,
        table_name="project_summary",
        columns=columns_by_table["project_summary"],
        project_id=config.source_project_id,
        phase_id=config.source_phase_id,
        for_update=for_update,
    )
    destination_summary_rows = fetch_scope_rows(
        cursor,
        table_name="project_summary",
        columns=columns_by_table["project_summary"],
        project_id=config.destination_project_id,
        phase_id=config.destination_phase_id,
        category="R",
        for_update=for_update,
    )
    source_summary_r_rows = [
        row
        for row in source_summary_all
        if str(row.get("category") or "").strip() == "R"
    ]
    out_of_scope_summary_rows = [
        row
        for row in source_summary_all
        if str(row.get("category") or "").strip() != "R"
    ]
    source_summary_rows = [
        row
        for row in source_summary_r_rows
        if row.get("r_parameter_id") in source_r_parameter_id_set
    ]
    orphan_summary_rows = [
        row
        for row in source_summary_r_rows
        if row.get("r_parameter_id") not in source_r_parameter_id_set
    ]
    orphan_summary_count = len(orphan_summary_rows)
    out_of_scope_summary_count = len(out_of_scope_summary_rows)

    source_detail_all = fetch_scope_rows(
        cursor,
        table_name="detail_data_rlist",
        columns=columns_by_table["detail_data_rlist"],
        project_id=config.source_project_id,
        phase_id=config.source_phase_id,
        for_update=for_update,
    )
    destination_detail_rows = fetch_scope_rows(
        cursor,
        table_name="detail_data_rlist",
        columns=columns_by_table["detail_data_rlist"],
        project_id=config.destination_project_id,
        phase_id=config.destination_phase_id,
        for_update=for_update,
    )
    source_detail_rows = [
        row
        for row in source_detail_all
        if row_key(row, TABLE_KEY_COLUMNS["detail_data_rlist"]) in source_parent_keys
    ]
    orphan_detail_rows = [
        row
        for row in source_detail_all
        if row_key(row, TABLE_KEY_COLUMNS["detail_data_rlist"])
        not in source_parent_keys
    ]
    orphan_detail_count = len(orphan_detail_rows)

    source_temp_all = fetch_scope_rows(
        cursor,
        table_name="temp_data",
        columns=columns_by_table["temp_data"],
        project_id=config.source_project_id,
        phase_id=config.source_phase_id,
        for_update=for_update,
    )
    destination_temp_rows = fetch_scope_rows(
        cursor,
        table_name="temp_data",
        columns=columns_by_table["temp_data"],
        project_id=config.destination_project_id,
        phase_id=config.destination_phase_id,
        for_update=for_update,
    )
    source_temp_rows = [
        row
        for row in source_temp_all
        if row_key(row, TABLE_KEY_COLUMNS["temp_data"]) in source_parent_keys
    ]
    orphan_temp_rows = [
        row
        for row in source_temp_all
        if row_key(row, TABLE_KEY_COLUMNS["temp_data"]) not in source_parent_keys
    ]
    orphan_temp_count = len(orphan_temp_rows)

    source_rows_by_table = {
        "project_r_parameter": source_project_rows,
        "temp_data": source_temp_rows,
        "project_summary": source_summary_rows,
        "detail_data_rlist": source_detail_rows,
    }
    destination_rows_by_table = {
        "project_r_parameter": destination_project_rows,
        "temp_data": destination_temp_rows,
        "project_summary": destination_summary_rows,
        "detail_data_rlist": destination_detail_rows,
    }
    source_counts_by_table = {
        "project_r_parameter": len(source_project_rows),
        "temp_data": len(source_temp_all),
        "project_summary": len(source_summary_all),
        "detail_data_rlist": len(source_detail_all),
    }
    missing_parent_rows_by_table = {
        "project_r_parameter": [],
        "temp_data": orphan_temp_rows,
        "project_summary": orphan_summary_rows,
        "detail_data_rlist": orphan_detail_rows,
    }
    out_of_scope_rows_by_table = {
        "project_r_parameter": [],
        "temp_data": [],
        "project_summary": out_of_scope_summary_rows,
        "detail_data_rlist": [],
    }

    table_plans: dict[str, TableCopyPlan] = {}
    for table_name in TABLE_COPY_ORDER:
        source_rows = source_rows_by_table[table_name]
        destination_rows = destination_rows_by_table[table_name]
        decisions = classify_rows(
            source_rows,
            destination_rows,
            key_columns=TABLE_KEY_COLUMNS[table_name],
            destination_project_id=config.destination_project_id,
            destination_phase_id=config.destination_phase_id,
        )
        decisions.extend(
            make_missing_parent_decisions(
                missing_parent_rows_by_table[table_name],
                key_columns=TABLE_KEY_COLUMNS[table_name],
                destination_project_id=config.destination_project_id,
                destination_phase_id=config.destination_phase_id,
                parent_table="project_r_parameter",
            )
        )
        decisions.extend(
            make_out_of_scope_decisions(
                out_of_scope_rows_by_table[table_name],
                key_columns=TABLE_KEY_COLUMNS[table_name],
                destination_project_id=config.destination_project_id,
                destination_phase_id=config.destination_phase_id,
                reason=(
                    "Không copy vì project_summary chỉ thuộc phạm vi R-list khi "
                    "category = 'R'. Record này có category khác 'R'."
                ),
            )
        )
        table_plans[table_name] = TableCopyPlan(
            table_name=table_name,
            columns=columns_by_table[table_name],
            key_columns=TABLE_KEY_COLUMNS[table_name],
            source_count=source_counts_by_table[table_name],
            destination_before_count=len(destination_rows),
            decisions=decisions,
        )

    warnings: list[str] = []
    blockers: list[str] = []
    if not source_project_rows:
        blockers.append("Không có dữ liệu project_r_parameter ở source project/phase.")
    if out_of_scope_summary_count:
        warnings.append(
            f"Có {out_of_scope_summary_count} dòng project_summary có category "
            "khác 'R'; từng record và lý do ngoài phạm vi được ghi trong sheet "
            "Records not copied."
        )
    if orphan_summary_count:
        warnings.append(
            f"Có {orphan_summary_count} dòng project_summary nguồn không có "
            "project_r_parameter tương ứng; từng record và lý do được ghi trong "
            "sheet Records not copied."
        )
    if orphan_detail_count:
        warnings.append(
            f"Có {orphan_detail_count} dòng detail_data_rlist nguồn không có "
            "project_r_parameter tương ứng; từng record và lý do được ghi trong "
            "sheet Records not copied."
        )
    if orphan_temp_count:
        warnings.append(
            f"Có {orphan_temp_count} dòng temp_data nguồn không có "
            "project_r_parameter tương ứng; từng record và lý do được ghi trong "
            "sheet Records not copied."
        )

    existing_r_parameter_ids = fetch_existing_master_ids(
        cursor,
        table_name="r_parameter",
        ids=source_r_parameter_ids,
    )
    missing_r_parameter_ids = [
        value
        for value in source_r_parameter_ids
        if value not in existing_r_parameter_ids
    ]
    if missing_r_parameter_ids:
        blockers.append(
            f"r_parameter_id không tồn tại trong master: {missing_r_parameter_ids}"
        )

    existing_variation_ids = fetch_existing_master_ids(
        cursor,
        table_name="variation",
        ids=source_variation_ids,
    )
    missing_variation_ids = [
        value for value in source_variation_ids if value not in existing_variation_ids
    ]
    if missing_variation_ids:
        blockers.append(
            f"variation_id không tồn tại trong master: {missing_variation_ids}"
        )

    setup_gaps = fetch_destination_setup_gaps(
        cursor,
        destination_setup_id=int(projects[config.destination_project_id]["setup_id"]),
        r_parameter_ids=source_r_parameter_ids,
    )
    if setup_gaps:
        warnings.append(
            "Destination setup chưa có setup_r_relation cho r_parameter_id: "
            f"{setup_gaps}. Script không tự tạo quan hệ setup."
        )

    for table_name, table_plan in table_plans.items():
        if table_plan.blocked_count:
            blockers.append(
                f"{table_name} có {table_plan.blocked_count} dòng bị chặn do "
                "business key NULL hoặc trùng lặp."
            )
        different_existing_count = sum(
            decision.action == ACTION_SKIP
            and decision.destination_row != decision.target_row
            for decision in table_plan.decisions
        )
        if different_existing_count:
            warnings.append(
                f"{table_name} có {different_existing_count} business key đã tồn tại "
                "nhưng nội dung khác source; script giữ nguyên destination."
            )

    related_impacts = fetch_related_impacts(cursor, config=config)
    for impact in related_impacts:
        if impact.records_truncated:
            warnings.append(
                f"{impact.table_name} có {impact.source_rows} record không copy nhưng "
                f"sheet chỉ liệt kê {len(impact.source_records)} record theo "
                "report_row_limit."
            )
    fk_relationships = fetch_fk_relationships(cursor)
    plan = CopyPlan(
        source_project=projects[config.source_project_id],
        destination_project=projects[config.destination_project_id],
        source_phase=phases[config.source_phase_id],
        destination_phase=phases[config.destination_phase_id],
        table_plans=table_plans,
        related_impacts=related_impacts,
        fk_relationships=fk_relationships,
        warnings=warnings,
        blockers=blockers,
    )
    plan.fingerprint = calculate_plan_fingerprint(table_plans)
    return plan


def count_destination_rows(cursor, config: RunConfig, table_name: str) -> int:
    where_parts = ["project_id = %s", "phase_id = %s"]
    params: list[Any] = [
        config.destination_project_id,
        config.destination_phase_id,
    ]
    if table_name == "project_summary":
        where_parts.append("category = %s")
        params.append("R")
    cursor.execute(
        f"SELECT COUNT(*) FROM {quote_identifier(table_name)} "
        f"WHERE {' AND '.join(where_parts)}",
        params,
    )
    return int(cursor.fetchone()[0])


def execute_copy(
    cursor,
    config: RunConfig,
    plan: CopyPlan,
) -> tuple[dict[str, int], dict[str, int]]:
    inserted_counts: dict[str, int] = {}
    destination_after_counts: dict[str, int] = {}

    for table_name in TABLE_COPY_ORDER:
        table_plan = plan.table_plans[table_name]
        rows = table_plan.copy_rows
        if not rows:
            inserted_counts[table_name] = 0
        else:
            column_sql = ", ".join(
                quote_identifier(column) for column in table_plan.columns
            )
            value_sql = make_placeholders(table_plan.columns)
            insert_sql = (
                f"INSERT INTO {quote_identifier(table_name)} ({column_sql}) "
                f"VALUES ({value_sql})"
            )
            values = [
                tuple(row.get(column) for column in table_plan.columns) for row in rows
            ]
            cursor.executemany(insert_sql, values)
            inserted_count = int(cursor.rowcount)
            if inserted_count != len(rows):
                raise RuntimeError(
                    f"{table_name}: dự kiến insert {len(rows)} dòng nhưng database "
                    f"báo {inserted_count} dòng. Transaction đã rollback."
                )
            inserted_counts[table_name] = inserted_count

        after_count = count_destination_rows(cursor, config, table_name)
        expected_count = (
            table_plan.destination_before_count + inserted_counts[table_name]
        )
        if after_count != expected_count:
            raise RuntimeError(
                f"{table_name}: destination dự kiến có {expected_count} dòng nhưng "
                f"đếm được {after_count}. Transaction đã rollback."
            )
        destination_after_counts[table_name] = after_count

    return inserted_counts, destination_after_counts


def excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, datetime, date)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bytes):
        text_value = value.hex()
    else:
        text_value = str(value)
    text_value = EXCEL_ILLEGAL_CHARACTER_PATTERN.sub("�", text_value)
    if text_value.startswith(("=", "+", "-", "@")):
        text_value = "'" + text_value
    if len(text_value) > EXCEL_CELL_LIMIT:
        return text_value[: EXCEL_CELL_LIMIT - 16] + "...[TRUNCATED]"
    return text_value


def apply_sheet_style(
    worksheet,
    *,
    header_row: int,
    freeze_cell: str,
) -> None:
    dark_blue = "1F4E78"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = freeze_cell

    title_cell = worksheet.cell(row=2, column=1)
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1F1F1F")

    for cell in worksheet[header_row]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    max_column = worksheet.max_column
    max_row = worksheet.max_row
    last_column = get_column_letter(max_column)
    worksheet.auto_filter.ref = f"A{header_row}:{last_column}{max_row}"
    if max_row > header_row:
        for row in worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=max_row,
            min_col=1,
            max_col=max_column,
        ):
            for cell in row:
                cell.font = Font(name="Arial", size=10, color="1F1F1F")
                cell.alignment = Alignment(vertical="center")

    for column_index in range(1, max_column + 1):
        column_letter = get_column_letter(column_index)
        max_length = 0
        for cell in worksheet[column_letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 10), 60
        )

    if worksheet.title == "Summary":
        worksheet.column_dimensions["B"].width = 80
        for cell in worksheet["B"]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def decision_report_rows(
    config: RunConfig,
    table_plan: TableCopyPlan,
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "action",
        "copy_result",
        "reason_detail",
        "business_key",
        "different_columns",
        "source_record",
        "planned_destination_record",
        "destination_existing_row",
        "source_project_id",
        "source_phase_id",
        *table_plan.columns,
    ]
    rows = []
    ordered_decisions = [
        decision for decision in table_plan.decisions if decision.action != ACTION_COPY
    ] + [
        decision for decision in table_plan.decisions if decision.action == ACTION_COPY
    ]
    for decision in ordered_decisions[: config.report_row_limit]:
        changed_columns = differing_columns(
            decision.target_row,
            decision.destination_row,
        )
        rows.append(
            [
                decision.action,
                "COPIED" if decision.action == ACTION_COPY else "NOT_COPIED",
                decision.reason,
                repr(decision.key),
                ", ".join(changed_columns),
                excel_value(
                    json.dumps(
                        canonical_row(decision.source_row),
                        ensure_ascii=False,
                        sort_keys=True,
                    )
                ),
                excel_value(
                    json.dumps(
                        canonical_row(decision.target_row),
                        ensure_ascii=False,
                        sort_keys=True,
                    )
                ),
                (
                    ""
                    if decision.destination_row is None
                    else excel_value(
                        json.dumps(
                            canonical_row(decision.destination_row),
                            ensure_ascii=False,
                            sort_keys=True,
                        )
                    )
                ),
                config.source_project_id,
                config.source_phase_id,
                *[
                    excel_value(decision.target_row.get(column))
                    for column in table_plan.columns
                ],
            ]
        )
    return headers, rows


def not_copied_report_rows(plan: CopyPlan) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "table_name",
        "action",
        "reason_detail",
        "business_key",
        "different_columns",
        "source_record",
        "planned_destination_record",
        "destination_existing_record",
    ]
    rows: list[list[Any]] = []
    for table_name in TABLE_COPY_ORDER:
        for decision in plan.table_plans[table_name].decisions:
            if decision.action == ACTION_COPY:
                continue
            changed_columns = differing_columns(
                decision.target_row,
                decision.destination_row,
            )
            rows.append(
                [
                    table_name,
                    decision.action,
                    decision.reason,
                    repr(decision.key),
                    ", ".join(changed_columns),
                    excel_value(
                        json.dumps(
                            canonical_row(decision.source_row),
                            ensure_ascii=False,
                            sort_keys=True,
                        )
                    ),
                    excel_value(
                        json.dumps(
                            canonical_row(decision.target_row),
                            ensure_ascii=False,
                            sort_keys=True,
                        )
                    ),
                    (
                        ""
                        if decision.destination_row is None
                        else excel_value(
                            json.dumps(
                                canonical_row(decision.destination_row),
                                ensure_ascii=False,
                                sort_keys=True,
                            )
                        )
                    ),
                ]
            )

    for impact in plan.related_impacts:
        for source_record in impact.source_records:
            business_key = tuple(source_record.items())
            rows.append(
                [
                    impact.table_name,
                    impact.action,
                    impact.reason,
                    repr(business_key),
                    "",
                    excel_value(
                        json.dumps(
                            canonical_row(source_record),
                            ensure_ascii=False,
                            sort_keys=True,
                        )
                    ),
                    "",
                    "",
                ]
            )
    return headers, rows


def apply_action_formats(
    worksheet,
    *,
    first_data_row: int = 5,
    action_column: int = 1,
) -> None:
    fills = {
        ACTION_COPY: PatternFill("solid", fgColor="E2F0D9"),
        ACTION_SKIP: PatternFill("solid", fgColor="E7E6E6"),
        ACTION_SKIP_MISSING_PARENT: PatternFill("solid", fgColor="FFF2CC"),
        ACTION_SKIP_OUT_OF_SCOPE: PatternFill("solid", fgColor="DDEBF7"),
        "NOT_COPIED": PatternFill("solid", fgColor="E7E6E6"),
    }
    blocked_fill = PatternFill("solid", fgColor="FCE4D6")
    for row_index in range(first_data_row, worksheet.max_row + 1):
        action_cell = worksheet.cell(row=row_index, column=action_column)
        action = str(action_cell.value or "")
        if action.startswith("BLOCK_"):
            action_cell.fill = blocked_fill
            action_cell.font = Font(name="Arial", size=10, bold=True, color="C00000")
        elif action in fills:
            action_cell.fill = fills[action]
            action_cell.font = Font(name="Arial", size=10, bold=True, color="1F1F1F")


def write_report(
    config: RunConfig,
    *,
    status: str,
    message: str,
    database_name: str,
    database_host: str,
    plan: CopyPlan | None,
    inserted_counts: dict[str, int] | None = None,
    destination_after_counts: dict[str, int] | None = None,
    allow_overwrite: bool = False,
) -> None:
    output_path = config.output_path
    if output_path.exists() and not allow_overwrite:
        raise FileExistsError(f"File báo cáo đã tồn tại, không ghi đè: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append([])
    summary.append(["R-list copy report"])
    summary.append([])
    summary.append(["Field", "Value"])
    summary_values = [
        ("status", status),
        ("generated_at", datetime.now()),
        ("database", database_name),
        ("host", database_host),
        ("source_project_id", config.source_project_id),
        ("source_phase_id", config.source_phase_id),
        ("destination_project_id", config.destination_project_id),
        ("destination_phase_id", config.destination_phase_id),
        ("execute_requested", config.execute),
        ("assume_yes", config.assume_yes),
        ("message", message),
    ]
    if plan is not None:
        summary_values.extend(
            [
                ("source_project_code", plan.source_project.get("project_code")),
                (
                    "destination_project_code",
                    plan.destination_project.get("project_code"),
                ),
                ("source_phase", plan.source_phase.get("phase")),
                ("destination_phase", plan.destination_phase.get("phase")),
                ("plan_fingerprint", plan.fingerprint),
                ("rows_to_copy", plan.total_copy_rows),
                ("rows_skipped", plan.total_skipped_rows),
                ("rows_blocked", plan.total_blocked_rows),
                ("records_not_copied", plan.total_not_copied_rows),
                ("blockers", "\n".join(plan.blockers)),
                ("warnings", "\n".join(plan.warnings)),
            ]
        )
    for field_name, value in summary_values:
        summary.append([field_name, excel_value(value)])
    apply_sheet_style(summary, header_row=4, freeze_cell="A5")

    count_sheet = workbook.create_sheet("Table counts")
    count_sheet.append([])
    count_sheet.append(["Copy counts by table"])
    count_sheet.append([])
    count_sheet.append(
        [
            "table_name",
            "source_rows",
            "destination_before",
            "copy",
            "skip_existing",
            "skip_missing_parent",
            "skip_out_of_scope",
            "blocked",
            "not_copied_total",
            "inserted",
            "destination_after",
        ]
    )
    if plan is not None:
        for table_name in TABLE_COPY_ORDER:
            table_plan = plan.table_plans[table_name]
            count_sheet.append(
                [
                    table_name,
                    table_plan.source_count,
                    table_plan.destination_before_count,
                    len(table_plan.copy_rows),
                    table_plan.action_count(ACTION_SKIP),
                    table_plan.action_count(ACTION_SKIP_MISSING_PARENT),
                    table_plan.action_count(ACTION_SKIP_OUT_OF_SCOPE),
                    table_plan.blocked_count,
                    table_plan.not_copied_count,
                    None
                    if inserted_counts is None
                    else inserted_counts.get(table_name, 0),
                    None
                    if destination_after_counts is None
                    else destination_after_counts.get(table_name),
                ]
            )
    apply_sheet_style(count_sheet, header_row=4, freeze_cell="A5")

    if plan is not None:
        not_copied_sheet = workbook.create_sheet("Records not copied")
        not_copied_sheet.append([])
        not_copied_sheet.append(["Records not copied and detailed reasons"])
        not_copied_sheet.append([])
        not_copied_headers, not_copied_rows = not_copied_report_rows(plan)
        not_copied_sheet.append(not_copied_headers)
        for row in not_copied_rows:
            not_copied_sheet.append(row)
        apply_sheet_style(not_copied_sheet, header_row=4, freeze_cell="A5")
        apply_action_formats(not_copied_sheet, action_column=2)

        for table_name in TABLE_COPY_ORDER:
            table_plan = plan.table_plans[table_name]
            worksheet = workbook.create_sheet(table_name[:31])
            worksheet.append([])
            worksheet.append([f"Planned rows: {table_name}"])
            worksheet.append([])
            headers, rows = decision_report_rows(config, table_plan)
            worksheet.append(headers)
            for row in rows:
                worksheet.append(row)
            if len(table_plan.decisions) > len(rows):
                worksheet.append([])
                worksheet.append(
                    [
                        f"Report truncated: exported {len(rows)} of "
                        f"{len(table_plan.decisions)} rows."
                    ]
                )
            apply_sheet_style(worksheet, header_row=4, freeze_cell="A5")
            apply_action_formats(worksheet)

        related_sheet = workbook.create_sheet("Not copied")
        related_sheet.append([])
        related_sheet.append(["Related tables outside copy scope"])
        related_sheet.append([])
        related_sheet.append(
            [
                "table_name",
                "source_rows",
                "records_listed",
                "records_not_listed",
                "action",
                "reason_detail",
            ]
        )
        for impact in plan.related_impacts:
            related_sheet.append(
                [
                    impact.table_name,
                    impact.source_rows,
                    len(impact.source_records),
                    impact.source_rows - len(impact.source_records),
                    impact.action,
                    impact.reason,
                ]
            )
        apply_sheet_style(related_sheet, header_row=4, freeze_cell="A5")

        fk_sheet = workbook.create_sheet("FK relationships")
        fk_sheet.append([])
        fk_sheet.append(["Foreign-key relationships"])
        fk_sheet.append([])
        fk_headers = [
            "constraint_name",
            "child_table",
            "child_columns",
            "parent_table",
            "parent_columns",
            "delete_rule",
            "update_rule",
        ]
        fk_sheet.append(fk_headers)
        for relationship in plan.fk_relationships:
            fk_sheet.append(
                [excel_value(relationship.get(header)) for header in fk_headers]
            )
        apply_sheet_style(fk_sheet, header_row=4, freeze_cell="A5")

    temp_path = output_path.with_name(f".{output_path.name}.tmp.xlsx")
    workbook.save(temp_path)
    temp_path.replace(output_path)


def confirmation_text(config: RunConfig, database_name: str) -> str:
    return (
        f"COPY RLIST {database_name} FROM PROJECT {config.source_project_id} "
        f"PHASE {config.source_phase_id} TO PROJECT "
        f"{config.destination_project_id} PHASE {config.destination_phase_id}"
    )


def print_separator() -> None:
    print("=" * 88)


def print_plan(
    config: RunConfig, database_name: str, database_host: str, plan: CopyPlan
) -> None:
    print_separator()
    print("PREVIEW COPY R-LIST")
    print_separator()
    print(f"Database     : {database_name}")
    print(f"Host         : {database_host}")
    print(
        "Source       : "
        f"project={config.source_project_id} "
        f"({plan.source_project.get('project_code')}), "
        f"phase={config.source_phase_id} ({plan.source_phase.get('phase')})"
    )
    print(
        "Destination  : "
        f"project={config.destination_project_id} "
        f"({plan.destination_project.get('project_code')}), "
        f"phase={config.destination_phase_id} "
        f"({plan.destination_phase.get('phase')})"
    )
    print(f"Execute      : {config.execute}")
    print()
    for table_name in TABLE_COPY_ORDER:
        table_plan = plan.table_plans[table_name]
        print(
            f"{table_name:<28} source={table_plan.source_count:>6}  "
            f"copy={len(table_plan.copy_rows):>6}  "
            f"skip={table_plan.skipped_count:>6}  "
            f"blocked={table_plan.blocked_count:>6}"
        )
    if plan.warnings:
        print("\nWarnings:")
        for warning in plan.warnings:
            print(f"- {warning}")
    if plan.blockers:
        print("\nBlockers:")
        for blocker in plan.blockers:
            print(f"- {blocker}")
    related_with_data = [
        impact for impact in plan.related_impacts if impact.source_rows > 0
    ]
    if related_with_data:
        print("\nRelated tables not copied:")
        for impact in related_with_data:
            print(f"- {impact.table_name}: {impact.source_rows} rows")


def main(argv: Sequence[str] | None = None) -> int:
    config = parse_args(argv)
    database_name, database_host = database_identity()
    plan: CopyPlan | None = None
    inserted_counts: dict[str, int] | None = None
    destination_after_counts: dict[str, int] | None = None
    report_written = False
    mutation_committed = False

    try:
        with connection.cursor() as cursor:
            plan = build_copy_plan(cursor, config, for_update=False)

        print_plan(config, database_name, database_host, plan)
        write_report(
            config,
            status="PREVIEW",
            message="Preview hoàn tất; chưa ghi dữ liệu vào database.",
            database_name=database_name,
            database_host=database_host,
            plan=plan,
        )
        report_written = True
        print(f"\nExcel report : {config.output_path}")

        if not config.execute:
            print_separator()
            print("RUN_EXECUTE=False/không có --execute: chưa copy dữ liệu.")
            return 0

        if plan.blockers:
            message = "Execute bị chặn: " + " | ".join(plan.blockers)
            write_report(
                config,
                status="BLOCKED",
                message=message,
                database_name=database_name,
                database_host=database_host,
                plan=plan,
                allow_overwrite=True,
            )
            print_separator()
            print(message)
            return 1

        if plan.total_copy_rows == 0:
            message = (
                "Không có record hợp lệ mới để copy. Xem sheet Records not copied "
                "để biết từng record và lý do."
            )
            write_report(
                config,
                status="NO_OP",
                message=message,
                database_name=database_name,
                database_host=database_host,
                plan=plan,
                inserted_counts={table: 0 for table in TABLE_COPY_ORDER},
                destination_after_counts={
                    table: plan.table_plans[table].destination_before_count
                    for table in TABLE_COPY_ORDER
                },
                allow_overwrite=True,
            )
            print_separator()
            print(message)
            return 0

        expected_confirmation = confirmation_text(config, database_name)
        if not config.assume_yes:
            print_separator()
            print("Để xác nhận copy, nhập chính xác nội dung sau:")
            print(expected_confirmation)
            actual_confirmation = input("> ").strip()
            if actual_confirmation != expected_confirmation:
                write_report(
                    config,
                    status="CANCELLED",
                    message="Người dùng hủy hoặc nhập sai chuỗi xác nhận.",
                    database_name=database_name,
                    database_host=database_host,
                    plan=plan,
                    allow_overwrite=True,
                )
                print("Đã hủy. Không có dữ liệu nào được copy.")
                return 0

        with transaction.atomic():
            with connection.cursor() as cursor:
                locked_plan = build_copy_plan(cursor, config, for_update=True)
                if locked_plan.blockers:
                    raise RuntimeError(
                        "Preflight trong transaction bị chặn: "
                        + " | ".join(locked_plan.blockers)
                    )
                if locked_plan.fingerprint != plan.fingerprint:
                    raise RuntimeError(
                        "Dữ liệu source/destination đã thay đổi sau preview. "
                        "Transaction đã rollback; hãy chạy lại preview."
                    )
                inserted_counts, destination_after_counts = execute_copy(
                    cursor,
                    config,
                    locked_plan,
                )
        mutation_committed = True

        write_report(
            config,
            status="COPIED",
            message="Copy thành công trong một transaction và đã kiểm tra row count.",
            database_name=database_name,
            database_host=database_host,
            plan=plan,
            inserted_counts=inserted_counts,
            destination_after_counts=destination_after_counts,
            allow_overwrite=True,
        )

        print_separator()
        print("COPY R-LIST THÀNH CÔNG")
        for table_name in TABLE_COPY_ORDER:
            print(f"{table_name:<28}: {inserted_counts[table_name]:>6} rows")
        print(f"Excel report : {config.output_path}")
        return 0

    except KeyboardInterrupt:
        if mutation_committed:
            status = "INTERRUPTED_AFTER_COMMIT"
            message = "Database đã commit; chương trình bị ngắt khi cập nhật báo cáo."
        else:
            status = "INTERRUPTED"
            message = "Chương trình bị ngắt; transaction (nếu có) đã rollback."
        if report_written:
            try:
                write_report(
                    config,
                    status=status,
                    message=message,
                    database_name=database_name,
                    database_host=database_host,
                    plan=plan,
                    inserted_counts=inserted_counts,
                    destination_after_counts=destination_after_counts,
                    allow_overwrite=True,
                )
            except Exception as report_error:
                print(f"Không thể cập nhật báo cáo: {report_error}")
        print(f"\n{message}")
        return 130
    except Exception as error:
        if mutation_committed:
            status = "FAILED_AFTER_COMMIT"
            message = (
                "Database đã commit nhưng cập nhật báo cáo sau copy thất bại. "
                f"{type(error).__name__}: {error}"
            )
        else:
            status = "FAILED"
            message = f"{type(error).__name__}: {error}"
        if report_written or not config.output_path.exists():
            try:
                write_report(
                    config,
                    status=status,
                    message=message,
                    database_name=database_name,
                    database_host=database_host,
                    plan=plan,
                    inserted_counts=inserted_counts,
                    destination_after_counts=destination_after_counts,
                    allow_overwrite=report_written,
                )
            except Exception as report_error:
                print(f"Không thể tạo/cập nhật báo cáo lỗi: {report_error}")
        print_separator()
        print("XẢY RA LỖI")
        print(message)
        return 1
    finally:
        connection.close()


if __name__ == "__main__":
    raise SystemExit(main())
