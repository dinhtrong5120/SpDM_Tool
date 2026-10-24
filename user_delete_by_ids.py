"""Delete selected account rows by username; preserve user_table and business data.

PyCharm: set RUN_USERNAMES (or use RUN_INPUT_EXCEL), review preview, then set
RUN_EXECUTE=True. This script does NOT change an external Cognito account.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")

# Keep this False until the Excel preview has been reviewed.
RUN_USERNAMES: list[str] = ['KNT21617']
# RUN_INPUT_EXCEL = Path("/Users/TrongKieu/Downloads/削除するユーザー.xlsx")
RUN_INPUT_EXCEL = Path("")
RUN_EXECUTE = False
OUTPUT_DIR = Path(__file__).resolve().parent / "output"

# Only these rows and auth_user are deleted. user_table and every other table stay.
DELETE_BY_AUTH_ID = {
    "auth_user_groups": "user_id",
    "auth_user_user_permissions": "user_id",
    "token_blacklist_outstandingtoken": "user_id",
    "user_group": "user_id",
    "user_info": "user_id",
}
DELETE_BY_USERNAME = {
    "bookmark": "employee_number",
    "column_resized_tbl": "employee_number",
    "condition_log": "employee_number",
    "login_log": "employee_number",
}
BLACKLIST_TABLE = "token_blacklist_blacklistedtoken"
DELETE_TABLES = set(DELETE_BY_AUTH_ID) | set(DELETE_BY_USERNAME) | {BLACKLIST_TABLE, "auth_user"}
REDACT_COLUMNS = {"password", "token", "secret", "access_token", "refresh_token"}
REPORT_ROW_LIMIT = 50_000
IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


@dataclass(frozen=True)
class ForeignKey:
    child: str
    column: str
    parent: str
    parent_column: str
    constraint: str
    delete_rule: str


@dataclass
class Target:
    username: str
    auth_id: int | None
    employee_exists: bool
    counts: dict[str, int] = field(default_factory=dict)
    retained: list[tuple[str, str, str, int]] = field(default_factory=list)
    result: str = "PREVIEW_ONLY"

    def fingerprint(self) -> tuple:
        return (self.auth_id, self.employee_exists, tuple(sorted(self.counts.items())),
                tuple(sorted(self.retained)))


def quote_identifier(value: str) -> str:
    if not IDENTIFIER.fullmatch(value):
        raise ValueError(f"Tên bảng/cột không hợp lệ: {value!r}")
    return f"`{value}`"


def normalize_usernames(values: list[str]) -> list[str]:
    if not values:
        raise ValueError("Danh sách username trống.")
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        if not isinstance(value, str) or not value or value != value.strip() or len(value) > 150:
            raise ValueError(f"Username không hợp lệ: {value!r}")
        if value in seen:
            raise ValueError(f"Username trùng: {value!r}")
        seen.add(value)
        result.append(value)
    return result


def input_usernames(cli_values: list[str] | None) -> list[str]:
    if cli_values is not None:
        return normalize_usernames(cli_values)
    if RUN_USERNAMES:
        return normalize_usernames(RUN_USERNAMES)
    import pandas as pd

    frame = pd.read_excel(RUN_INPUT_EXCEL, dtype={"Username": str})
    if "Username" not in frame.columns:
        raise ValueError(f"Thiếu cột Username trong {RUN_INPUT_EXCEL}.")
    return normalize_usernames(frame["Username"].dropna().tolist())


def database_identity(cursor) -> tuple[str, str]:
    cursor.execute("SELECT DATABASE(), @@hostname")
    database, host = cursor.fetchone()
    if not database:
        raise RuntimeError("Không xác định được database hiện tại.")
    return str(database), str(host)


def all_foreign_keys(cursor) -> list[ForeignKey]:
    cursor.execute("""
        SELECT k.TABLE_NAME, k.COLUMN_NAME, k.REFERENCED_TABLE_NAME,
               k.REFERENCED_COLUMN_NAME, k.CONSTRAINT_NAME, rc.DELETE_RULE
        FROM information_schema.KEY_COLUMN_USAGE AS k
        JOIN information_schema.TABLES AS t
          ON t.TABLE_SCHEMA = k.TABLE_SCHEMA AND t.TABLE_NAME = k.TABLE_NAME
         AND t.TABLE_TYPE = 'BASE TABLE'
        JOIN information_schema.REFERENTIAL_CONSTRAINTS AS rc
          ON rc.CONSTRAINT_SCHEMA = k.TABLE_SCHEMA
         AND rc.CONSTRAINT_NAME = k.CONSTRAINT_NAME
         AND rc.TABLE_NAME = k.TABLE_NAME
        WHERE k.TABLE_SCHEMA = DATABASE()
          AND k.REFERENCED_TABLE_SCHEMA = DATABASE()
        ORDER BY k.TABLE_NAME, k.CONSTRAINT_NAME, k.ORDINAL_POSITION
    """)
    rows = [ForeignKey(*row) for row in cursor.fetchall()]
    for fk in rows:
        for name in (fk.child, fk.column, fk.parent, fk.parent_column):
            quote_identifier(name)
    return rows


def validate_schema(cursor, foreign_keys: list[ForeignKey]) -> None:
    names = sorted(DELETE_TABLES | {"user_table"})
    placeholders = ", ".join(["%s"] * len(names))
    cursor.execute(f"""
        SELECT TABLE_NAME FROM information_schema.TABLES
        WHERE TABLE_SCHEMA = DATABASE() AND TABLE_TYPE = 'BASE TABLE'
          AND TABLE_NAME IN ({placeholders})
    """, names)
    missing = set(names) - {row[0] for row in cursor.fetchall()}
    if missing:
        raise RuntimeError(f"Thiếu bảng bắt buộc: {sorted(missing)}")

    expected = {(table, column, "auth_user") for table, column in DELETE_BY_AUTH_ID.items()}
    expected |= {(table, column, "user_table") for table, column in DELETE_BY_USERNAME.items()}
    expected.add((BLACKLIST_TABLE, "token_id", "token_blacklist_outstandingtoken"))
    found = {(fk.child, fk.column, fk.parent) for fk in foreign_keys}
    missing_fk = expected - found
    if missing_fk:
        raise RuntimeError(f"FK khác schema đã kiểm tra: {sorted(missing_fk)}")
    # A new FK into one of the to-be-deleted tables can cause cascading data
    # loss. Refuse it until its ownership and row predicate are reviewed.
    unexpected = [fk for fk in foreign_keys
                  if fk.parent in DELETE_TABLES and fk.child not in DELETE_TABLES]
    if unexpected:
        raise RuntimeError(f"Có FK con ngoài danh sách xoá: {unexpected}")


def retained_columns(cursor, foreign_keys: list[ForeignKey]) -> list[tuple[str, str, str]]:
    """All direct business FKs plus legacy employee/manager columns are KEEP."""
    pairs: set[tuple[str, str, str]] = set()
    for fk in foreign_keys:
        if fk.parent == "user_table" and fk.child not in DELETE_TABLES:
            pairs.add((fk.child, fk.column, "FK:" + fk.delete_rule))
        elif fk.parent == "auth_user" and fk.child not in DELETE_TABLES:
            pairs.add((fk.child, fk.column, "FK:" + fk.delete_rule))
    cursor.execute("""
        SELECT c.TABLE_NAME, c.COLUMN_NAME
        FROM information_schema.COLUMNS AS c
        JOIN information_schema.TABLES AS t
          ON t.TABLE_SCHEMA = c.TABLE_SCHEMA AND t.TABLE_NAME = c.TABLE_NAME
         AND t.TABLE_TYPE = 'BASE TABLE'
        WHERE c.TABLE_SCHEMA = DATABASE()
          AND c.COLUMN_NAME IN ('employee_number', 'manager_number')
          AND c.TABLE_NAME <> 'user_table'
    """)
    existing = {(table, column) for table, column, _ in pairs}
    for table, column in cursor.fetchall():
        if table not in DELETE_TABLES and (table, column) not in existing:
            pairs.add((table, column, "LEGACY_NO_FK"))
    return sorted(pairs)


def exact_count(cursor, table: str, column: str, value: str | int) -> int:
    predicate = (f"BINARY {quote_identifier(column)} = BINARY %s" if isinstance(value, str)
                 else f"{quote_identifier(column)} = %s")
    cursor.execute(f"SELECT COUNT(*) FROM {quote_identifier(table)} WHERE {predicate}", [value])
    return int(cursor.fetchone()[0])


def load_user(cursor, username: str, lock: bool) -> tuple[int | None, bool]:
    suffix = " FOR UPDATE" if lock else ""
    cursor.execute(f"SELECT id FROM auth_user WHERE BINARY username = BINARY %s{suffix}", [username])
    auth_rows = cursor.fetchall()
    cursor.execute(f"SELECT employee_number FROM user_table "
                   f"WHERE BINARY employee_number = BINARY %s{suffix}", [username])
    employee_rows = cursor.fetchall()
    if len(auth_rows) > 1 or len(employee_rows) > 1:
        raise RuntimeError(f"Nhiều record khớp chính xác username {username!r}.")
    return (int(auth_rows[0][0]) if auth_rows else None, bool(employee_rows))


def build_plan(cursor, usernames: list[str], keep_columns: list[tuple[str, str, str]],
               lock: bool = False) -> list[Target]:
    plans: list[Target] = []
    for username in usernames:
        auth_id, employee_exists = load_user(cursor, username, lock)
        target = Target(username, auth_id, employee_exists)
        target.counts["auth_user"] = 1 if auth_id is not None else 0
        for table, column in DELETE_BY_AUTH_ID.items():
            target.counts[table] = exact_count(cursor, table, column, auth_id) if auth_id else 0
        for table, column in DELETE_BY_USERNAME.items():
            target.counts[table] = exact_count(cursor, table, column, username)
        if auth_id is not None:
            cursor.execute("""
                SELECT COUNT(*) FROM token_blacklist_blacklistedtoken AS b
                JOIN token_blacklist_outstandingtoken AS o ON o.id = b.token_id
                WHERE o.user_id = %s
            """, [auth_id])
            target.counts[BLACKLIST_TABLE] = int(cursor.fetchone()[0])
        else:
            target.counts[BLACKLIST_TABLE] = 0
        for table, column, relation in keep_columns:
            count = exact_count(cursor, table, column, username)
            target.retained.append((table, column, relation, count))
        plans.append(target)
    return plans


def safe_record(columns: list[str], row: tuple) -> str:
    data = {column: "[REDACTED]" if column.lower() in REDACT_COLUMNS and value is not None
            else value for column, value in zip(columns, row, strict=True)}
    value = json.dumps(data, ensure_ascii=False, default=str, sort_keys=True)
    return value if len(value) <= 32_000 else value[:31_900] + " ... [TRUNCATED FOR EXCEL]"


def collect_records(cursor, plans: list[Target]) -> dict[str, list[tuple[str, str, str, str]]]:
    records: dict[str, list[tuple[str, str, str, str]]] = {}

    def add(table: str, username: str, match: str, action: str, sql: str,
            params: list, count: int) -> None:
        if not count:
            return
        if count > REPORT_ROW_LIMIT:
            raise RuntimeError(f"{table}/{username} có {count} record, vượt giới hạn báo cáo; chưa xoá.")
        cursor.execute(sql, params)
        columns = [column[0] for column in cursor.description]
        rows = cursor.fetchall()
        if len(rows) != count:
            raise RuntimeError(f"Record {table}/{username} thay đổi khi lập báo cáo; chưa xoá.")
        bucket = records.setdefault(table, [])
        if len(bucket) + len(rows) >= 1_048_576:
            raise RuntimeError(f"Sheet {table} vượt giới hạn Excel; chưa xoá.")
        bucket.extend((username, match, action, safe_record(columns, row)) for row in rows)

    for plan in plans:
        username, auth_id = plan.username, plan.auth_id
        if auth_id is not None:
            add("auth_user", username, "id", "DELETE", "SELECT * FROM auth_user WHERE id = %s",
                [auth_id], 1)
        if plan.employee_exists:
            add("user_table", username, "employee_number", "KEEP",
                "SELECT * FROM user_table WHERE BINARY employee_number = BINARY %s",
                [username], 1)
        for table, column in DELETE_BY_AUTH_ID.items():
            count = plan.counts[table]
            if count:
                add(table, username, column, "DELETE",
                    f"SELECT * FROM {quote_identifier(table)} WHERE {quote_identifier(column)} = %s",
                    [auth_id], count)
        for table, column in DELETE_BY_USERNAME.items():
            count = plan.counts[table]
            if count:
                add(table, username, column, "DELETE",
                    f"SELECT * FROM {quote_identifier(table)} "
                    f"WHERE BINARY {quote_identifier(column)} = BINARY %s", [username], count)
        if plan.counts[BLACKLIST_TABLE]:
            add(BLACKLIST_TABLE, username, "token_id -> user_id", "DELETE",
                """SELECT b.* FROM token_blacklist_blacklistedtoken AS b
                   JOIN token_blacklist_outstandingtoken AS o ON o.id = b.token_id
                   WHERE o.user_id = %s""", [auth_id], plan.counts[BLACKLIST_TABLE])
        for table, column, _, count in plan.retained:
            if count:
                add(table, username, column, "KEEP",
                    f"SELECT * FROM {quote_identifier(table)} "
                    f"WHERE BINARY {quote_identifier(column)} = BINARY %s", [username], count)
    return records


def deletion_order(foreign_keys: list[ForeignKey]) -> list[str]:
    """Order selected tables child before parent, including nested token FK."""
    children: dict[str, set[str]] = {table: set() for table in DELETE_TABLES}
    for fk in foreign_keys:
        if fk.child in DELETE_TABLES and fk.parent in DELETE_TABLES and fk.child != fk.parent:
            children[fk.parent].add(fk.child)
    ordered: list[str] = []
    visiting: set[str] = set()
    visited: set[str] = set()

    def visit(table: str) -> None:
        if table in visiting:
            raise RuntimeError(f"Vòng FK giữa các bảng sẽ xoá tại {table}.")
        if table in visited:
            return
        visiting.add(table)
        for child in sorted(children[table]):
            visit(child)
        visiting.remove(table)
        visited.add(table)
        ordered.append(table)

    for table in sorted(DELETE_TABLES):
        visit(table)
    if ordered[-1] != "auth_user":
        ordered.remove("auth_user")
        ordered.append("auth_user")
    return ordered


def delete_rows(cursor, table: str, column: str, value: str | int, expected: int) -> None:
    if not expected:
        return
    predicate = (f"BINARY {quote_identifier(column)} = BINARY %s" if isinstance(value, str)
                 else f"{quote_identifier(column)} = %s")
    cursor.execute(f"DELETE FROM {quote_identifier(table)} WHERE {predicate}", [value])
    if cursor.rowcount != expected:
        raise RuntimeError(f"Sai số dòng xoá {table}/{value!r}: {cursor.rowcount} != {expected}.")


def execute_plan(cursor, plans: list[Target], order: list[str]) -> None:
    for table in order:
        for target in plans:
            expected = target.counts[table]
            if not expected:
                continue
            if table == BLACKLIST_TABLE:
                cursor.execute("""
                    DELETE b FROM token_blacklist_blacklistedtoken AS b
                    JOIN token_blacklist_outstandingtoken AS o ON o.id = b.token_id
                    WHERE o.user_id = %s
                """, [target.auth_id])
                if cursor.rowcount != expected:
                    raise RuntimeError(f"Sai số dòng xoá {table}/{target.username}.")
            elif table == "auth_user":
                cursor.execute("DELETE FROM auth_user WHERE id = %s AND BINARY username = BINARY %s",
                               [target.auth_id, target.username])
                if cursor.rowcount != 1:
                    raise RuntimeError(f"Không xoá được auth_user/{target.username}.")
            elif table in DELETE_BY_AUTH_ID:
                delete_rows(cursor, table, DELETE_BY_AUTH_ID[table], target.auth_id, expected)
            else:
                delete_rows(cursor, table, DELETE_BY_USERNAME[table], target.username, expected)
    for target in plans:
        auth_id, employee_exists = load_user(cursor, target.username, False)
        if auth_id is not None or employee_exists != target.employee_exists:
            raise RuntimeError(f"Kiểm tra sau xoá thất bại cho {target.username}; rollback.")
        for table, column, relation, count in target.retained:
            if exact_count(cursor, table, column, target.username) != count:
                raise RuntimeError(f"Dữ liệu giữ lại đã đổi: {table}.{column}/{target.username}; rollback.")
        for table, column in DELETE_BY_USERNAME.items():
            if exact_count(cursor, table, column, target.username):
                raise RuntimeError(f"Còn record ở {table}/{target.username}; rollback.")
        target.result = "DELETED_SELECTED_TABLES"


def sheet_name(table: str) -> str:
    reserved = {"Summary", "Users", "Delete_order", "Table_index", "Relations"}
    if len(table) <= 31 and table not in reserved:
        return table
    digest = hashlib.sha1(table.encode()).hexdigest()[:6]
    return f"{table[:24]}_{digest}"


def report_path(custom: str | None) -> Path:
    if custom:
        path = Path(custom).expanduser().resolve()
        if path.suffix.lower() != ".xlsx":
            raise ValueError("--output phải là file .xlsx")
        return path
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    return OUTPUT_DIR / f"user_delete_by_ids_{timestamp}.xlsx"


def write_report(path: Path, database: str, host: str, mode: str, plans: list[Target],
                 records: dict[str, list[tuple[str, str, str, str]]], order: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Database", database])
    summary.append(["Host", host])
    summary.append(["Mode", mode])
    summary.append(["Requested usernames", len(plans)])
    summary.append(["Deleted auth_user", sum(plan.result == "DELETED_SELECTED_TABLES" for plan in plans)])
    summary.append(["Retained user_table", "YES — all project/business references stay in place"])
    summary.append(["External Cognito", "NOT CHANGED — database deletion alone does not disable Cognito login"])

    users = workbook.create_sheet("Users")
    users.append(["username", "auth_user.id", "user_table exists", "planned delete rows",
                  "retained reference rows", "result"])
    for plan in plans:
        users.append([plan.username, plan.auth_id, "yes" if plan.employee_exists else "no",
                      sum(plan.counts.values()), sum(item[3] for item in plan.retained), plan.result])

    sequence = workbook.create_sheet("Delete_order")
    sequence.append(["step", "table_name", "action", "scope", "reason"])
    for number, table in enumerate(order, 1):
        if table == BLACKLIST_TABLE:
            scope = "token_id of outstanding tokens belonging to auth_user.id"
        elif table == "auth_user":
            scope = "exact username + auth_user.id"
        elif table in DELETE_BY_AUTH_ID:
            scope = f"{DELETE_BY_AUTH_ID[table]} = auth_user.id"
        else:
            scope = f"{DELETE_BY_USERNAME[table]} = exact username"
        sequence.append([number, table, "DELETE", scope, "Child before parent; never delete user_table."])

    relations = workbook.create_sheet("Relations")
    relations.append(["username", "table_name", "column", "relationship", "matched rows", "action"])
    for plan in plans:
        for table in order:
            relations.append([plan.username, table, "", "selected", plan.counts[table], "DELETE"])
        relations.append([plan.username, "user_table", "employee_number", "identity anchor",
                          int(plan.employee_exists), "KEEP"])
        for table, column, relation, count in plan.retained:
            relations.append([plan.username, table, column, relation, count, "KEEP"])

    index = workbook.create_sheet("Table_index")
    index.append(["table_name", "sheet_name", "matched rows", "action"])
    for table in sorted(records):
        name = sheet_name(table)
        action = "DELETE" if table in DELETE_TABLES else "KEEP"
        index.append([table, name, len(records[table]), action])
        sheet = workbook.create_sheet(name)
        sheet.append(["username", "matched by", "action", "record before deletion", "result"])
        results = {plan.username: plan.result for plan in plans}
        for username, matched_by, row_action, record in records[table]:
            sheet.append([username, matched_by, row_action, record, results[username]])

    for sheet in workbook:
        if sheet != summary:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183153")
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 14), 80)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Delete only selected account-related tables by username.")
    parser.add_argument("--usernames", nargs="+", help="Username list; otherwise RUN_USERNAMES or input Excel.")
    parser.add_argument("--execute", action="store_true", help="Delete after exact confirmation; default preview.")
    parser.add_argument("--output", help="Report .xlsx path (default timestamped output directory).")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    usernames = input_usernames(args.usernames)
    execute = args.execute or RUN_EXECUTE
    path = report_path(args.output)

    import django

    django.setup()
    from django.db import connection, transaction

    with connection.cursor() as cursor:
        database, host = database_identity(cursor)
        foreign_keys = all_foreign_keys(cursor)
        validate_schema(cursor, foreign_keys)
        keep_columns = retained_columns(cursor, foreign_keys)
        preview = build_plan(cursor, usernames, keep_columns)
        records = collect_records(cursor, preview)
    order = deletion_order(foreign_keys)
    print(f"Database: {database} @ {host}")
    for plan in preview:
        print(f"{plan.username}: auth_user.id={plan.auth_id}; delete_rows={sum(plan.counts.values())}; "
              f"keep_business_rows={sum(item[3] for item in plan.retained)}")
    print("user_table and other business/history tables: KEEP")
    print("Cognito account: NOT CHANGED by this script")
    if not execute:
        write_report(path, database, host, "PREVIEW - NO DATABASE WRITE", preview, records, order)
        print(f"Preview report: {path}")
        return 0
    if any(plan.auth_id is None for plan in preview):
        write_report(path, database, host, "BLOCKED - AUTH_USER NOT FOUND", preview, records, order)
        print(f"Một số username không có auth_user; không xoá bất kỳ dòng nào. Report: {path}")
        return 2
    confirmation = f"DELETE {','.join(usernames)} FROM {database} KEEP user_table"
    print(f"Nhập chính xác để xác nhận: {confirmation}")
    if input("> ").strip() != confirmation:
        write_report(path, database, host, "CANCELLED - NO DATABASE WRITE", preview, records, order)
        return 1
    try:
        with transaction.atomic():
            with connection.cursor() as cursor:
                if database_identity(cursor) != (database, host):
                    raise RuntimeError("Database/host đã đổi sau preview.")
                current_fks = all_foreign_keys(cursor)
                if current_fks != foreign_keys:
                    raise RuntimeError("FK schema đã đổi sau preview.")
                validate_schema(cursor, current_fks)
                if retained_columns(cursor, current_fks) != keep_columns:
                    raise RuntimeError("Danh sách bảng giữ lại đã đổi sau preview.")
                current = build_plan(cursor, usernames, keep_columns, lock=True)
                for before, after in zip(preview, current, strict=True):
                    if before.fingerprint() != after.fingerprint():
                        raise RuntimeError(f"Dữ liệu của {before.username} đã đổi; rollback.")
                execute_plan(cursor, current, order)
    except Exception as exc:
        for plan in preview:
            plan.result = "ROLLED_BACK"
        write_report(path, database, host, f"FAILED/ROLLED BACK: {exc}", preview, records, order)
        raise
    write_report(path, database, host, "EXECUTED - COMMITTED", current, records, order)
    print(f"Report: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
