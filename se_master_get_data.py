"""Look up ``se_parameter.id`` values for rows in an SE master workbook.

The tool is read-only with respect to the database. It matches each Excel row
by ``WP``, ``Category``, ``パラメータ名``, and ``単位`` within one hierarchy,
then writes the selected ID to a new ``se_parameter_id check`` column in a
separate output workbook.

When duplicate WP names exist, the smallest ``wp.id`` is used. When multiple
``se_parameter`` rows match the complete business key, the smallest
``se_parameter.id`` is used and the output cell receives an audit comment.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
import unicodedata
from collections import Counter
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]

if __name__ == "__main__":
    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection  # noqa: E402

# PyCharm Run-button configuration.
RUN_EXCEL_PATH = Path("/Users/TrongKieu/Downloads/seリストへの追加 copy.xlsx")
RUN_SHEET_NAME = "登録する項目"
RUN_HIERARCHY_ID = 1
RUN_OUTPUT_PATH: Path | None = "/Users/TrongKieu/Downloads/untitled folder/seリストへの追加_re.xlsx"

OUTPUT_HEADER = "se_parameter_id check"
REQUIRED_HEADERS = ("WP", "Category", "パラメータ名", "単位")
DEFAULT_OUTPUT_DIRECTORY = PROJECT_ROOT / "outputs" / "se_master_get_data"

FOUND_FILL = PatternFill("solid", fgColor="C6EFCE")
SELECTED_MIN_FILL = PatternFill("solid", fgColor="FFF2CC")
NOT_FOUND_FILL = PatternFill("solid", fgColor="F4CCCC")
SKIPPED_FILL = PatternFill("solid", fgColor="E7E6E6")


@dataclass(frozen=True)
class LookupResult:
    status: str
    se_parameter_id: int | None
    message: str


@dataclass(frozen=True)
class ReferenceCatalog:
    wp_ids_by_name: dict[str, tuple[int, ...]]
    se_ids_by_key: dict[tuple[object, ...], tuple[int, ...]]


@dataclass(frozen=True)
class RunResult:
    output_path: Path
    counts: Counter[str]


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    return "".join(character for character in text if not character.isspace())


def normalize_text(value: object) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def normalize_match_key(value: object) -> str | None:
    cleaned = normalize_text(value)
    if cleaned is None:
        return None
    return unicodedata.normalize("NFKC", cleaned)


def normalize_positive_id(value: object, *, label: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{label} must be a positive integer, not boolean.")
    if isinstance(value, float) and not value.is_integer():
        raise ValueError(f"{label} must be an integer: {value!r}")
    try:
        normalized = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid {label}: {value!r}") from exc
    if normalized <= 0:
        raise ValueError(f"{label} must be greater than zero: {normalized}")
    return normalized


def database_identity(db_connection=connection) -> tuple[str, str, str]:
    settings = db_connection.settings_dict
    return (
        str(settings.get("NAME") or "<unknown>"),
        str(settings.get("HOST") or "localhost"),
        str(settings.get("PORT") or "3306"),
    )


def make_output_path(
    input_path: Path,
    output_path: Path | str | None = None,
) -> Path:
    if output_path is None:
        DEFAULT_OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
        return DEFAULT_OUTPUT_DIRECTORY / (
            f"se_master_get_data_{timestamp}{input_path.suffix.lower()}"
        )

    resolved = Path(output_path).expanduser().resolve()
    if resolved.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Output path must use the .xlsx or .xlsm extension.")
    if resolved == input_path:
        raise ValueError("Output must be a separate file; input will not be overwritten.")
    resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


def find_header_columns(worksheet: Any) -> dict[str, int]:
    normalized_expected = {
        normalize_header(header): header for header in REQUIRED_HEADERS
    }
    columns: dict[str, int] = {}
    for column in range(1, worksheet.max_column + 1):
        normalized = normalize_header(worksheet.cell(1, column).value)
        expected = normalized_expected.get(normalized)
        if expected is not None:
            if expected in columns:
                raise ValueError(
                    f"Header {expected!r} appears more than once in sheet "
                    f"{worksheet.title!r}."
                )
            columns[expected] = column

    missing = [header for header in REQUIRED_HEADERS if header not in columns]
    if missing:
        raise ValueError(
            f"Missing required headers in sheet {worksheet.title!r}: {missing}"
        )
    return columns


def find_or_add_output_column(worksheet: Any) -> int:
    expected = normalize_header(OUTPUT_HEADER)
    for column in range(1, worksheet.max_column + 1):
        if normalize_header(worksheet.cell(1, column).value) == expected:
            return column

    output_column = worksheet.max_column + 1
    source_header = worksheet.cell(1, max(1, output_column - 1))
    target_header = worksheet.cell(1, output_column, OUTPUT_HEADER)
    if source_header.number_format:
        target_header.number_format = source_header.number_format
    target_header.font = copy(source_header.font)
    target_header.fill = copy(source_header.fill)
    target_header.border = copy(source_header.border)
    target_header.alignment = copy(source_header.alignment)
    target_header.protection = copy(source_header.protection)
    worksheet.column_dimensions[get_column_letter(output_column)].width = 24
    return output_column


def load_reference_catalog(cursor: Any, *, hierarchy_id: int) -> ReferenceCatalog:
    normalized_hierarchy_id = normalize_positive_id(
        hierarchy_id, label="hierarchy_id"
    )

    cursor.execute("SELECT id, wp FROM wp ORDER BY id")
    wp_ids: dict[str, list[int]] = {}
    for raw_id, raw_wp in cursor.fetchall():
        wp_key = normalize_match_key(raw_wp)
        if wp_key is None:
            continue
        wp_ids.setdefault(wp_key, []).append(int(raw_id))

    cursor.execute(
        """
        SELECT id, wp_id, parameter_name_1, parameter_name_2, parameter_unit
        FROM se_parameter
        WHERE hierarchy_id = %s
        ORDER BY id
        """,
        [normalized_hierarchy_id],
    )
    se_ids: dict[tuple[object, ...], list[int]] = {}
    for (
        raw_id,
        raw_wp_id,
        raw_category,
        raw_parameter_name,
        raw_parameter_unit,
    ) in cursor.fetchall():
        if raw_wp_id is None:
            continue
        key = (
            int(raw_wp_id),
            normalize_match_key(raw_category),
            normalize_match_key(raw_parameter_name),
            normalize_match_key(raw_parameter_unit),
        )
        se_ids.setdefault(key, []).append(int(raw_id))

    return ReferenceCatalog(
        wp_ids_by_name={
            key: tuple(sorted(set(values))) for key, values in wp_ids.items()
        },
        se_ids_by_key={
            key: tuple(sorted(set(values))) for key, values in se_ids.items()
        },
    )


def resolve_se_parameter(
    catalog: ReferenceCatalog,
    *,
    wp: object,
    category: object,
    parameter_name: object,
    parameter_unit: object,
) -> LookupResult:
    normalized_values = (
        normalize_match_key(wp),
        normalize_match_key(category),
        normalize_match_key(parameter_name),
        normalize_match_key(parameter_unit),
    )
    if all(value is None for value in normalized_values):
        return LookupResult("SKIPPED", None, "Blank input row.")

    wp_key, category_key, parameter_name_key, parameter_unit_key = normalized_values
    if wp_key is None:
        return LookupResult(
            "INVALID",
            None,
            "WP is blank; the row cannot be matched.",
        )

    matching_wp_ids = catalog.wp_ids_by_name.get(wp_key, ())
    if not matching_wp_ids:
        return LookupResult("NOT_FOUND", None, "WP was not found in wp.wp.")

    selected_wp_id = matching_wp_ids[0]
    business_key = (
        selected_wp_id,
        category_key,
        parameter_name_key,
        parameter_unit_key,
    )
    matching_se_ids = catalog.se_ids_by_key.get(business_key, ())
    notes: list[str] = []
    if len(matching_wp_ids) > 1:
        notes.append(
            "WP matched IDs "
            f"{', '.join(str(value) for value in matching_wp_ids)}; "
            f"selected smallest wp.id {selected_wp_id}."
        )
    if not matching_se_ids:
        notes.append(
            "No se_parameter matched the four input values under "
            f"wp.id {selected_wp_id}."
        )
        return LookupResult("NOT_FOUND", None, " ".join(notes))

    selected_se_id = matching_se_ids[0]
    if len(matching_se_ids) > 1:
        notes.append(
            "se_parameter matched IDs "
            f"{', '.join(str(value) for value in matching_se_ids)}; "
            f"selected smallest ID {selected_se_id}."
        )
    return LookupResult(
        "SELECTED_MIN" if notes else "FOUND",
        selected_se_id,
        " ".join(notes) or "Exact match found.",
    )


def _set_result_cell(cell: Any, result: LookupResult) -> None:
    cell.value = result.se_parameter_id
    cell.comment = Comment(result.message, "SPDM lookup")
    cell.number_format = "0"
    if result.status == "FOUND":
        cell.fill = copy(FOUND_FILL)
    elif result.status == "SELECTED_MIN":
        cell.fill = copy(SELECTED_MIN_FILL)
    elif result.status in {"NOT_FOUND", "INVALID"}:
        cell.fill = copy(NOT_FOUND_FILL)
    else:
        cell.fill = copy(SKIPPED_FILL)


def enrich_workbook(
    input_path: Path | str,
    *,
    sheet_name: str,
    hierarchy_id: int,
    output_path: Path | str | None = None,
    db_connection=connection,
) -> RunResult:
    resolved_input = Path(input_path).expanduser().resolve()
    if not resolved_input.is_file():
        raise FileNotFoundError(f"Excel file was not found: {resolved_input}")
    if resolved_input.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Input Excel must use the .xlsx or .xlsm extension.")
    resolved_output = make_output_path(resolved_input, output_path)

    keep_vba = resolved_input.suffix.lower() == ".xlsm"
    workbook = load_workbook(resolved_input, data_only=False, keep_vba=keep_vba)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet {sheet_name!r} was not found; available={workbook.sheetnames}"
            )
        worksheet = workbook[sheet_name]
        header_columns = find_header_columns(worksheet)
        output_column = find_or_add_output_column(worksheet)

        with db_connection.cursor() as cursor:
            catalog = load_reference_catalog(cursor, hierarchy_id=hierarchy_id)

        counts: Counter[str] = Counter()
        for row_number in range(2, worksheet.max_row + 1):
            source_cells = {
                header: worksheet.cell(row_number, column)
                for header, column in header_columns.items()
            }
            formula_headers = [
                header for header, cell in source_cells.items() if cell.data_type == "f"
            ]
            if formula_headers:
                result = LookupResult(
                    "INVALID",
                    None,
                    "Formula cells are not supported for lookup fields: "
                    + ", ".join(formula_headers)
                    + ".",
                )
            else:
                result = resolve_se_parameter(
                    catalog,
                    wp=source_cells["WP"].value,
                    category=source_cells["Category"].value,
                    parameter_name=source_cells["パラメータ名"].value,
                    parameter_unit=source_cells["単位"].value,
                )
            counts[result.status] += 1
            _set_result_cell(worksheet.cell(row_number, output_column), result)

        resolved_output.parent.mkdir(parents=True, exist_ok=True)
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{resolved_output.stem}_",
                suffix=resolved_output.suffix,
                dir=resolved_output.parent,
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            temporary_path.replace(resolved_output)
        finally:
            if temporary_path is not None and temporary_path.exists():
                temporary_path.unlink()
    finally:
        workbook.close()

    database, host, port = database_identity(db_connection)
    print("SE parameter lookup completed")
    print(f"Database (read-only): {database}")
    print(f"Host: {host}:{port}")
    print(f"Hierarchy ID: {hierarchy_id}")
    print(f"Input Excel: {resolved_input}")
    print(f"Input sheet: {sheet_name}")
    print(f"Output Excel: {resolved_output}")
    print("Results:")
    for status in ("FOUND", "SELECTED_MIN", "NOT_FOUND", "INVALID", "SKIPPED"):
        print(f"- {status}: {counts[status]}")
    return RunResult(resolved_output, counts)


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Read WP/Category/パラメータ名/単位 from Excel and write the matching "
            "se_parameter.id values to a separate output workbook."
        )
    )
    parser.add_argument("--excel", type=Path, default=RUN_EXCEL_PATH)
    parser.add_argument("--sheet", default=RUN_SHEET_NAME)
    parser.add_argument("--hierarchy-id", type=int, default=RUN_HIERARCHY_ID)
    parser.add_argument("--output", type=Path, default=RUN_OUTPUT_PATH)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        enrich_workbook(
            args.excel,
            sheet_name=args.sheet,
            hierarchy_id=args.hierarchy_id,
            output_path=args.output,
        )
    except Exception as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
