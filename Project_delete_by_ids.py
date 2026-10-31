"""Xem trước hoặc xóa dữ liệu của từng cặp project-phase an toàn.

Để chạy bằng PyCharm, sửa RUN_INPUT_FILE/RUN_EXECUTE rồi bấm Run.
Mặc định RUN_EXECUTE=False nên chương trình chỉ xem trước, không xóa dữ liệu.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PROJECT_ROOT = Path(__file__).resolve().parents[2]

# Khi chạy trực tiếp, phải khởi tạo Django trước khi import database connection.
# Cách này giống utils/update_member_list.py và các script khác trong project.
if __name__ == "__main__":
    import os

    import django

    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "backend.settings")
    django.setup()

from django.db import connection, transaction  # noqa: E402

PROJECT_TABLE = "project_info"
PROJECT_KEY = "id"
NOTIFICATION_TABLE = "project_se_notification"
NOTIFICATION_TARGET_TABLE = "project_se_notification_target"

# Cấu hình khi chạy trực tiếp bằng nút Run của PyCharm.
# File phải có hai cột project_id, phase_id và có thể có match_status.
# Quy trình an toàn: chạy False trước, kiểm tra kết quả, rồi mới đổi thành True.
RUN_INPUT_FILE = (
    PROJECT_ROOT
    / "outputs"
    / "clw_project_phase_ids"
    / "CLW_project_phase_ids.xlsx"
)
RUN_REPORT_DIR = PROJECT_ROOT / "outputs" / "project_delete_reports"
RUN_EXECUTE = False
RUN_ASSUME_YES = False  # Nên giữ False để xác nhận trong Run console.

# Các bảng legacy có project_id nhưng chưa khai báo FK tới project_info.id.
# Chỉ bổ sung bảng sau khi đã kiểm tra ý nghĩa của cột trong schema.
LEGACY_PROJECT_COLUMNS: dict[str, tuple[str, ...]] = {
    "detail_data": ("project_id",),
    "detail_data_rlist": ("project_id",),
    "project_parameter_attachments": ("project_id",),
    "project_r_parameter_history": ("project_id",),
    "project_r_parameter_manager_approval_history": ("project_id",),
    "project_se_confirmation_history": ("project_id",),
    "project_se_input_completion": ("project_id",),
    "project_se_notification_target": ("project_id",),
    "studies_comparison_dashboard_settings": ("project_id",),
    "temp_project_parameter": ("project_id",),
}

IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z0-9_$]+$")


def quote_identifier(identifier: str) -> str:
    """Kiểm tra rồi quote tên bảng/cột lấy từ metadata."""
    if not IDENTIFIER_PATTERN.fullmatch(identifier):
        raise ValueError(f"Tên bảng/cột không hợp lệ: {identifier!r}")
    return f"`{identifier}`"


def make_placeholders(values: Sequence[object]) -> str:
    if not values:
        raise ValueError("Không thể tạo SQL placeholders cho danh sách rỗng.")
    return ", ".join(["%s"] * len(values))


def normalize_positive_id(raw_id: int, *, label: str) -> int:
    if isinstance(raw_id, bool):
        raise ValueError(f"{label} phải là số nguyên dương, không phải boolean.")
    try:
        normalized_id = int(raw_id)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} không hợp lệ: {raw_id!r}") from exc
    if normalized_id <= 0:
        raise ValueError(f"{label} phải lớn hơn 0: {normalized_id}")
    return normalized_id


def normalize_targets(
    targets: Iterable[Sequence[int]],
) -> list[tuple[int, int]]:
    """Chuẩn hóa các cặp (project_id, phase_id), không trùng lặp."""
    normalized: list[tuple[int, int]] = []
    seen: set[tuple[int, int]] = set()

    for target in targets:
        try:
            project_raw, phase_raw = target
        except (TypeError, ValueError) as exc:
            raise ValueError(
                "Mỗi target phải có đúng dạng (project_id, phase_id)."
            ) from exc
        pair = (
            normalize_positive_id(project_raw, label="project_id"),
            normalize_positive_id(phase_raw, label="phase_id"),
        )
        if pair not in seen:
            normalized.append(pair)
            seen.add(pair)

    if not normalized:
        raise ValueError("Danh sách project-phase targets rỗng.")
    return normalized


def load_targets_from_excel(input_path: Path) -> list[tuple[int, int]]:
    """Đọc các cặp project_id, phase_id từ workbook output đã đối chiếu."""
    if not input_path.is_file():
        raise FileNotFoundError(f"Không tìm thấy file input: {input_path}")

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration as exc:
        workbook.close()
        raise ValueError(f"File Excel rỗng: {input_path}") from exc

    positions = {
        str(value).strip(): index
        for index, value in enumerate(headers)
        if value is not None and str(value).strip()
    }
    required = ("project_id", "phase_id")
    missing_headers = [header for header in required if header not in positions]
    if missing_headers:
        workbook.close()
        raise ValueError(
            "File Excel thiếu cột bắt buộc: " + ", ".join(missing_headers)
        )

    raw_targets: list[tuple[object, object]] = []
    for row_number, row in enumerate(rows, start=2):
        project_id = row[positions["project_id"]]
        phase_id = row[positions["phase_id"]]
        if "match_status" in positions:
            status = str(row[positions["match_status"]] or "").strip()
            if status and status != "MATCHED":
                workbook.close()
                raise ValueError(
                    f"Dòng {row_number} có match_status={status!r}; "
                    "chỉ chấp nhận MATCHED."
                )
        if project_id in (None, "") and phase_id in (None, ""):
            continue
        raw_targets.append((project_id, phase_id))

    workbook.close()
    return normalize_targets(raw_targets)


def get_existing_legacy_columns(cursor) -> dict[str, list[str]]:
    """Lấy cột allowlist ở bảng có phase_id trong database hiện tại."""
    if not LEGACY_PROJECT_COLUMNS:
        return {}

    table_names = sorted(LEGACY_PROJECT_COLUMNS)
    placeholders = make_placeholders(table_names)
    cursor.execute(
        f"""
        SELECT project_column.TABLE_NAME, project_column.COLUMN_NAME
        FROM information_schema.COLUMNS AS project_column
        JOIN information_schema.COLUMNS AS phase_column
          ON phase_column.TABLE_SCHEMA = project_column.TABLE_SCHEMA
         AND phase_column.TABLE_NAME = project_column.TABLE_NAME
         AND phase_column.COLUMN_NAME = 'phase_id'
        JOIN information_schema.TABLES AS table_metadata
          ON table_metadata.TABLE_SCHEMA = project_column.TABLE_SCHEMA
         AND table_metadata.TABLE_NAME = project_column.TABLE_NAME
         AND table_metadata.TABLE_TYPE = 'BASE TABLE'
        WHERE project_column.TABLE_SCHEMA = DATABASE()
          AND project_column.TABLE_NAME IN ({placeholders})
        """,
        table_names,
    )

    existing: dict[str, list[str]] = defaultdict(list)
    for table_name, column_name in cursor.fetchall():
        if column_name in LEGACY_PROJECT_COLUMNS[table_name]:
            existing[table_name].append(column_name)
    return dict(existing)


def get_phase_related_tables(cursor) -> list[dict[str, object]]:
    """Lấy bảng có cả cột project liên quan và phase_id."""
    cursor.execute(
        """
        SELECT DISTINCT
            kcu.TABLE_NAME,
            kcu.COLUMN_NAME
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        JOIN information_schema.COLUMNS AS phase_column
          ON phase_column.TABLE_SCHEMA = kcu.TABLE_SCHEMA
         AND phase_column.TABLE_NAME = kcu.TABLE_NAME
         AND phase_column.COLUMN_NAME = 'phase_id'
        JOIN information_schema.TABLES AS t
          ON t.TABLE_SCHEMA = kcu.TABLE_SCHEMA
         AND t.TABLE_NAME = kcu.TABLE_NAME
        WHERE kcu.TABLE_SCHEMA = DATABASE()
          AND t.TABLE_TYPE = 'BASE TABLE'
          AND kcu.TABLE_NAME <> %s
          AND kcu.REFERENCED_TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_NAME = %s
          AND kcu.REFERENCED_COLUMN_NAME = %s
        ORDER BY kcu.TABLE_NAME, kcu.COLUMN_NAME
        """,
        (PROJECT_TABLE, PROJECT_TABLE, PROJECT_KEY),
    )

    table_columns: dict[str, set[str]] = defaultdict(set)
    relation_types: dict[str, set[str]] = defaultdict(set)
    for table_name, column_name in cursor.fetchall():
        table_columns[table_name].add(column_name)
        relation_types[table_name].add("foreign_key")

    for table_name, columns in get_existing_legacy_columns(cursor).items():
        table_columns[table_name].update(columns)
        relation_types[table_name].add("legacy_allowlist")

    return [
        {
            "table_name": table_name,
            "project_columns": sorted(columns),
            "phase_column": "phase_id",
            "relation_types": sorted(relation_types[table_name]),
        }
        for table_name, columns in sorted(table_columns.items())
    ]


def get_project_usage_tables(cursor) -> list[dict[str, object]]:
    """Lấy mọi bảng base có thể chặn việc xóa project_info.

    Cột legacy không FK chỉ được dùng để kiểm tra còn tham chiếu; dữ liệu ở các
    bảng đó không bị xóa tự động nếu không thuộc allowlist phase hiện có.
    """
    cursor.execute(
        """
        SELECT DISTINCT kcu.TABLE_NAME, kcu.COLUMN_NAME
        FROM information_schema.KEY_COLUMN_USAGE AS kcu
        JOIN information_schema.TABLES AS table_metadata
          ON table_metadata.TABLE_SCHEMA = kcu.TABLE_SCHEMA
         AND table_metadata.TABLE_NAME = kcu.TABLE_NAME
         AND table_metadata.TABLE_TYPE = 'BASE TABLE'
        WHERE kcu.TABLE_SCHEMA = DATABASE()
          AND kcu.TABLE_NAME <> %s
          AND kcu.REFERENCED_TABLE_SCHEMA = DATABASE()
          AND kcu.REFERENCED_TABLE_NAME = %s
          AND kcu.REFERENCED_COLUMN_NAME = %s
        ORDER BY kcu.TABLE_NAME, kcu.COLUMN_NAME
        """,
        (PROJECT_TABLE, PROJECT_TABLE, PROJECT_KEY),
    )

    table_columns: dict[str, set[str]] = defaultdict(set)
    relation_types: dict[str, set[str]] = defaultdict(set)
    for table_name, column_name in cursor.fetchall():
        table_columns[table_name].add(column_name)
        relation_types[table_name].add("foreign_key")

    cursor.execute(
        """
        SELECT column_metadata.TABLE_NAME, column_metadata.COLUMN_NAME
        FROM information_schema.COLUMNS AS column_metadata
        JOIN information_schema.TABLES AS table_metadata
          ON table_metadata.TABLE_SCHEMA = column_metadata.TABLE_SCHEMA
         AND table_metadata.TABLE_NAME = column_metadata.TABLE_NAME
         AND table_metadata.TABLE_TYPE = 'BASE TABLE'
        WHERE column_metadata.TABLE_SCHEMA = DATABASE()
          AND column_metadata.TABLE_NAME <> %s
          AND column_metadata.COLUMN_NAME IN ('project_id', 'project_info_id')
        ORDER BY column_metadata.TABLE_NAME, column_metadata.COLUMN_NAME
        """,
        (PROJECT_TABLE,),
    )
    for table_name, column_name in cursor.fetchall():
        if column_name not in table_columns[table_name]:
            relation_types[table_name].add("unconstrained_blocker")
        table_columns[table_name].add(column_name)

    return [
        {
            "table_name": table_name,
            "project_columns": sorted(columns),
            "relation_types": sorted(relation_types[table_name]),
        }
        for table_name, columns in sorted(table_columns.items())
    ]


def order_tables_child_first(
    cursor,
    related_tables: Sequence[dict[str, object]],
) -> list[dict[str, object]]:
    """Sắp xếp để bảng con được xóa trước bảng cha."""
    by_name = {str(item["table_name"]): item for item in related_tables}
    if not by_name:
        return []

    table_names = sorted(by_name)
    placeholders = make_placeholders(table_names)
    params = [*table_names, *table_names]
    cursor.execute(
        f"""
        SELECT DISTINCT TABLE_NAME, REFERENCED_TABLE_NAME
        FROM information_schema.KEY_COLUMN_USAGE
        WHERE TABLE_SCHEMA = DATABASE()
          AND REFERENCED_TABLE_SCHEMA = DATABASE()
          AND TABLE_NAME IN ({placeholders})
          AND REFERENCED_TABLE_NAME IN ({placeholders})
        """,
        params,
    )

    children: dict[str, set[str]] = defaultdict(set)
    for child_table, parent_table in cursor.fetchall():
        if child_table != parent_table:
            children[parent_table].add(child_table)

    ordered_names: list[str] = []
    visiting: set[str] = set()
    visited: set[str] = set()

    def visit(table_name: str) -> None:
        if table_name in visited:
            return
        if table_name in visiting:
            raise RuntimeError(
                f"Phát hiện vòng lặp FK giữa các bảng liên quan tại {table_name}."
            )
        visiting.add(table_name)
        for child in sorted(children.get(table_name, ())):
            visit(child)
        visiting.remove(table_name)
        visited.add(table_name)
        ordered_names.append(table_name)

    for table_name in table_names:
        visit(table_name)
    return [by_name[name] for name in ordered_names]


def build_target_filter(
    project_columns: Sequence[str],
    phase_column: str,
    targets: Sequence[tuple[int, int]],
) -> tuple[str, tuple[int, ...]]:
    """Tạo điều kiện chỉ khớp chính xác từng cặp project-phase."""
    if not project_columns:
        raise ValueError("Không có cột project để tạo điều kiện SQL.")
    if not targets:
        raise ValueError("Không có project-phase target để tạo điều kiện SQL.")

    phase_identifier = quote_identifier(phase_column)
    target_clauses: list[str] = []
    params: list[int] = []
    for project_id, phase_id in targets:
        project_clauses = [
            f"{quote_identifier(column)} = %s" for column in project_columns
        ]
        target_clauses.append(
            f"(({' OR '.join(project_clauses)}) AND {phase_identifier} = %s)"
        )
        params.extend([project_id] * len(project_columns))
        params.append(phase_id)
    return " OR ".join(target_clauses), tuple(params)


def build_project_filter(
    project_columns: Sequence[str], project_id: int
) -> tuple[str, tuple[int, ...]]:
    if not project_columns:
        raise ValueError("Không có cột project để kiểm tra tham chiếu.")
    clauses = [
        f"{quote_identifier(column)} = %s" for column in project_columns
    ]
    return " OR ".join(clauses), tuple([project_id] * len(project_columns))


def find_projects(cursor, project_ids: Sequence[int]) -> list[dict[str, object]]:
    placeholders = make_placeholders(project_ids)
    cursor.execute(
        f"""
        SELECT id, project_code
        FROM {quote_identifier(PROJECT_TABLE)}
        WHERE {quote_identifier(PROJECT_KEY)} IN ({placeholders})
        ORDER BY id
        """,
        project_ids,
    )
    return [
        {"id": project_id, "project_code": project_code}
        for project_id, project_code in cursor.fetchall()
    ]


def find_phases(cursor, phase_ids: Sequence[int]) -> list[dict[str, object]]:
    placeholders = make_placeholders(phase_ids)
    cursor.execute(
        f"SELECT id, phase FROM phase WHERE id IN ({placeholders}) ORDER BY id",
        phase_ids,
    )
    return [
        {"id": phase_id, "phase": phase_name}
        for phase_id, phase_name in cursor.fetchall()
    ]


def count_related_rows(
    cursor,
    related_tables: Sequence[dict[str, object]],
    targets: Sequence[tuple[int, int]],
) -> dict[str, int]:
    counts: dict[str, int] = {}
    for item in related_tables:
        table_name = str(item["table_name"])
        project_columns = list(item["project_columns"])
        phase_column = str(item["phase_column"])
        where_sql, params = build_target_filter(
            project_columns, phase_column, targets
        )
        cursor.execute(
            f"SELECT COUNT(*) FROM {quote_identifier(table_name)} "
            f"WHERE {where_sql}",
            params,
        )
        counts[table_name] = int(cursor.fetchone()[0])
    return counts


def count_project_references(
    cursor,
    usage_tables: Sequence[dict[str, object]],
    project_ids: Sequence[int],
) -> dict[int, dict[str, int]]:
    counts = {project_id: {} for project_id in project_ids}
    for item in usage_tables:
        table_name = str(item["table_name"])
        project_columns = list(item["project_columns"])
        for project_id in project_ids:
            where_sql, params = build_project_filter(
                project_columns, project_id
            )
            cursor.execute(
                f"SELECT COUNT(*) FROM {quote_identifier(table_name)} "
                f"WHERE {where_sql}",
                params,
            )
            count = int(cursor.fetchone()[0])
            if count:
                counts[project_id][table_name] = count
    return counts


def count_references_after_phase_delete(
    cursor,
    usage_tables: Sequence[dict[str, object]],
    phase_tables: Sequence[dict[str, object]],
    targets: Sequence[tuple[int, int]],
) -> dict[int, dict[str, int]]:
    """Dự đoán tham chiếu còn lại sau khi xóa các target phase."""
    project_ids = list(dict.fromkeys(project_id for project_id, _ in targets))
    phase_tables_by_name = {
        str(item["table_name"]): item for item in phase_tables
    }
    counts = {project_id: {} for project_id in project_ids}

    for usage_item in usage_tables:
        table_name = str(usage_item["table_name"])
        project_columns = list(usage_item["project_columns"])
        phase_item = phase_tables_by_name.get(table_name)
        for project_id in project_ids:
            usage_sql, params = build_project_filter(
                project_columns, project_id
            )
            sql = (
                f"SELECT COUNT(*) FROM {quote_identifier(table_name)} "
                f"WHERE ({usage_sql})"
            )
            query_params: list[int] = list(params)
            if phase_item is not None:
                target_sql, target_params = build_target_filter(
                    list(phase_item["project_columns"]),
                    str(phase_item["phase_column"]),
                    targets,
                )
                sql += f" AND NOT ({target_sql})"
                query_params.extend(target_params)
            cursor.execute(sql, query_params)
            count = int(cursor.fetchone()[0])
            if count:
                counts[project_id][table_name] = count
    return counts


def count_notification_cleanup_preview(
    cursor,
    targets: Sequence[tuple[int, int]],
) -> tuple[int, dict[int, int]]:
    """Đếm notification header sẽ xóa và còn lại sau cleanup target."""
    project_ids = list(dict.fromkeys(project_id for project_id, _ in targets))
    target_sql, target_params = build_target_filter(
        ["project_id"], "phase_id", targets
    )
    orphan_count = 0
    remaining_counts: dict[int, int] = {}

    for project_id in project_ids:
        common_sql = f"""
            FROM {quote_identifier(NOTIFICATION_TABLE)} AS notification
            WHERE notification.project_id = %s
        """
        remaining_target_sql = f"""
            SELECT 1
            FROM {quote_identifier(NOTIFICATION_TARGET_TABLE)} AS target
            WHERE target.notification_id = notification.id
              AND NOT ({target_sql})
        """
        params = [project_id, *target_params]

        cursor.execute(
            f"SELECT COUNT(*) {common_sql} "
            f"AND NOT EXISTS ({remaining_target_sql})",
            params,
        )
        orphan_count += int(cursor.fetchone()[0])

        cursor.execute(
            f"SELECT COUNT(*) {common_sql} "
            f"AND EXISTS ({remaining_target_sql})",
            params,
        )
        remaining_count = int(cursor.fetchone()[0])
        if remaining_count:
            remaining_counts[project_id] = remaining_count

    return orphan_count, remaining_counts


def delete_orphan_notifications(cursor, project_ids: Sequence[int]) -> int:
    """Xóa notification header khi không còn target nào sử dụng."""
    placeholders = make_placeholders(project_ids)
    cursor.execute(
        f"""
        SELECT notification.id
        FROM {quote_identifier(NOTIFICATION_TABLE)} AS notification
        WHERE notification.project_id IN ({placeholders})
          AND NOT EXISTS (
              SELECT 1
              FROM {quote_identifier(NOTIFICATION_TARGET_TABLE)} AS target
              WHERE target.notification_id = notification.id
          )
        FOR UPDATE
        """,
        project_ids,
    )
    notification_ids = [int(row[0]) for row in cursor.fetchall()]
    if not notification_ids:
        return 0

    notification_placeholders = make_placeholders(notification_ids)
    cursor.execute(
        f"DELETE FROM {quote_identifier(NOTIFICATION_TABLE)} "
        f"WHERE id IN ({notification_placeholders})",
        notification_ids,
    )
    deleted_count = int(cursor.rowcount)
    if deleted_count != len(notification_ids):
        raise RuntimeError(
            "Số project_se_notification xóa không đúng; transaction sẽ "
            f"rollback: expected={len(notification_ids)}, actual={deleted_count}"
        )
    return deleted_count


def make_report_path() -> Path:
    timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
    return RUN_REPORT_DIR / f"project_delete_report_{timestamp}.xlsx"


def write_deletion_report(
    report_path: Path,
    *,
    db_connection,
    input_path: Path | None,
    mode: str,
    status: str,
    targets: Sequence[tuple[int, int]],
    projects: Sequence[dict[str, object]],
    phases: Sequence[dict[str, object]],
    related_tables: Sequence[dict[str, object]],
    planned_counts: dict[str, int],
    deleted_counts: dict[str, int] | None,
    remaining_project_references: dict[int, dict[str, int]],
    deleted_project_ids: Sequence[int] = (),
    error_message: str = "",
) -> None:
    """Ghi báo cáo audit Excel; không chứa credential hoặc connection URL."""
    report_path.parent.mkdir(parents=True, exist_ok=True)
    deleted_project_id_set = set(deleted_project_ids)
    predicted_parent_count = sum(
        not table_counts
        for table_counts in remaining_project_references.values()
    )
    planned_with_parent = dict(planned_counts)
    planned_with_parent[PROJECT_TABLE] = predicted_parent_count

    committed_counts = deleted_counts if status == "COMPLETED" else {}
    committed_total = sum(committed_counts.values())
    run_at = datetime.now().astimezone()
    run_at_local = run_at.replace(tzinfo=None)
    settings = db_connection.settings_dict

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    targets_sheet = workbook.create_sheet("Targets")
    tables_sheet = workbook.create_sheet("Table details")

    title_font = Font(name="Arial", size=14, bold=True, color="1F2937")
    body_font = Font(name="Arial", size=10, color="1F2937")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    rule = Side(style="thin", color="9CA3AF")

    summary["A2"] = "SPDM project deletion report"
    summary["A2"].font = title_font
    summary["A3"].border = Border(bottom=rule)
    summary["B3"].border = Border(bottom=rule)
    summary_rows = [
        ("Run time", run_at_local),
        ("Time zone", f"UTC{run_at.strftime('%z')[:3]}:{run_at.strftime('%z')[3:]}"),
        ("Mode", mode),
        ("Status", status),
        ("Database", settings.get("NAME") or ""),
        ("Host", settings.get("HOST") or ""),
        ("Input file", input_path.name if input_path else "CLI targets"),
        ("Target pairs", len(targets)),
        ("Planned rows", sum(planned_with_parent.values())),
        ("Committed deleted rows", committed_total),
        ("Message", error_message),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=5):
        summary.cell(row_index, 1, label)
        summary.cell(row_index, 2, value)
    summary["B5"].number_format = "yyyy-mm-dd hh:mm:ss"
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 50
    summary["B15"].alignment = Alignment(vertical="top", wrap_text=True)

    project_by_id = {
        int(project["id"]): str(project["project_code"])
        for project in projects
    }
    phase_by_id = {
        int(phase["id"]): str(phase["phase"])
        for phase in phases
    }
    target_headers = (
        "project_id",
        "phase_id",
        "project_code",
        "phase",
        "project_info_action",
        "remaining_reference_rows",
        "remaining_tables",
    )
    targets_sheet.append(target_headers)
    for project_id, phase_id in targets:
        table_counts = remaining_project_references.get(project_id, {})
        if status == "COMPLETED":
            action = (
                "DELETED"
                if project_id in deleted_project_id_set
                else "KEPT"
            )
        elif status == "ROLLED_BACK":
            action = "ROLLED_BACK"
        elif status == "CANCELLED":
            action = "NOT_EXECUTED"
        else:
            action = "WOULD_KEEP" if table_counts else "WOULD_DELETE"
        targets_sheet.append(
            (
                project_id,
                phase_id,
                project_by_id.get(project_id, ""),
                phase_by_id.get(phase_id, ""),
                action,
                sum(table_counts.values()),
                ", ".join(sorted(table_counts)),
            )
        )

    relation_sources = {
        str(item["table_name"]): ", ".join(item["relation_types"])
        for item in related_tables
    }
    relation_sources[NOTIFICATION_TABLE] = "custom orphan cleanup"
    relation_sources[PROJECT_TABLE] = "parent record"
    table_headers = (
        "table_name",
        "relation_source",
        "planned_rows",
        "committed_deleted_rows",
        "status",
    )
    tables_sheet.append(table_headers)
    ordered_tables = [
        str(item["table_name"]) for item in related_tables
    ] + [NOTIFICATION_TABLE, PROJECT_TABLE]
    for table_name in ordered_tables:
        planned = int(planned_with_parent.get(table_name, 0))
        if status == "COMPLETED":
            committed = int(committed_counts.get(table_name, 0))
            table_status = "DELETED" if committed else "NO_ROWS"
        elif status == "ROLLED_BACK":
            committed = 0
            table_status = "ROLLED_BACK"
        elif status == "CANCELLED":
            committed = None
            table_status = "NOT_EXECUTED"
        else:
            committed = None
            table_status = "PLANNED" if planned else "NO_ROWS"
        tables_sheet.append(
            (
                table_name,
                relation_sources.get(table_name, ""),
                planned,
                committed,
                table_status,
            )
        )

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center")
    summary["A2"].font = title_font
    for row_index in range(5, 5 + len(summary_rows)):
        summary.cell(row_index, 1).font = Font(
            name="Arial", size=10, bold=True, color="374151"
        )

    for sheet in (targets_sheet, tables_sheet):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(right=Side(style="thin", color="FFFFFF"))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    targets_sheet.column_dimensions["A"].width = 14
    targets_sheet.column_dimensions["B"].width = 12
    targets_sheet.column_dimensions["C"].width = 32
    targets_sheet.column_dimensions["D"].width = 22
    targets_sheet.column_dimensions["E"].width = 22
    targets_sheet.column_dimensions["F"].width = 26
    targets_sheet.column_dimensions["G"].width = 70
    tables_sheet.column_dimensions["A"].width = 46
    tables_sheet.column_dimensions["B"].width = 28
    tables_sheet.column_dimensions["C"].width = 16
    tables_sheet.column_dimensions["D"].width = 25
    tables_sheet.column_dimensions["E"].width = 18
    workbook.save(report_path)


def print_plan(
    db_connection,
    projects: Sequence[dict[str, object]],
    phases: Sequence[dict[str, object]],
    targets: Sequence[tuple[int, int]],
    related_tables: Sequence[dict[str, object]],
    counts: dict[str, int],
    remaining_project_references: dict[int, dict[str, int]],
) -> None:
    settings = db_connection.settings_dict
    print(f"Database: {settings.get('NAME')}")
    print(f"Host: {settings.get('HOST')}:{settings.get('PORT') or 3306}")
    print("Projects tìm thấy:")
    for project in projects:
        print(f"- {project['id']}: {project['project_code']}")

    print("Phases tìm thấy:")
    for phase in phases:
        print(f"- {phase['id']}: {phase['phase']}")

    print("Targets sẽ xóa:")
    for project_id, phase_id in targets:
        print(f"- project_id={project_id}, phase_id={phase_id}")

    print("\nCác bảng sẽ được xử lý (con trước, cha sau):")
    for item in related_tables:
        table_name = str(item["table_name"])
        columns = ", ".join(
            f"`{column}`" for column in item["project_columns"]
        )
        relation_types = ", ".join(item["relation_types"])
        print(
            f"- {table_name}: {counts[table_name]} rows; "
            f"project_columns={columns}; phase_column=`{item['phase_column']}`; "
            f"source={relation_types}"
        )
    print(
        f"- {NOTIFICATION_TABLE}: {counts[NOTIFICATION_TABLE]} rows; "
        "source=custom orphan cleanup"
    )
    print(f"Tổng số dòng liên quan: {sum(counts.values())}")

    print("\nDự kiến xử lý project_info sau khi xóa phase:")
    for project_id, table_counts in remaining_project_references.items():
        remaining_count = sum(table_counts.values())
        if remaining_count:
            print(
                f"- project_id={project_id}: giữ lại; còn "
                f"{remaining_count} tham chiếu trong "
                f"{', '.join(sorted(table_counts))}"
            )
        else:
            print(f"- project_id={project_id}: có thể xóa project_info")


def delete_project_phase_data(
    targets: Iterable[Sequence[int]],
    db_connection=connection,
    *,
    execute: bool = False,
    assume_yes: bool = False,
    report_path: Path | None = None,
    input_path: Path | None = None,
) -> dict[str, int]:
    """Xóa đúng cặp project-phase và project_info khi không còn được dùng."""
    normalized_targets = normalize_targets(targets)
    project_ids = list(dict.fromkeys(item[0] for item in normalized_targets))
    phase_ids = list(dict.fromkeys(item[1] for item in normalized_targets))
    report_path = report_path or make_report_path()

    with db_connection.cursor() as cursor:
        projects = find_projects(cursor, project_ids)
        found_ids = {int(project["id"]) for project in projects}
        missing_ids = [
            project_id
            for project_id in project_ids
            if project_id not in found_ids
        ]
        if missing_ids:
            raise ValueError(f"Không tìm thấy project IDs: {missing_ids}")

        phases = find_phases(cursor, phase_ids)
        found_phase_ids = {int(phase["id"]) for phase in phases}
        missing_phase_ids = [
            phase_id for phase_id in phase_ids if phase_id not in found_phase_ids
        ]
        if missing_phase_ids:
            raise ValueError(f"Không tìm thấy phase IDs: {missing_phase_ids}")

        related_tables = order_tables_child_first(
            cursor, get_phase_related_tables(cursor)
        )
        counts = count_related_rows(cursor, related_tables, normalized_targets)
        usage_tables = get_project_usage_tables(cursor)
        remaining_project_references = count_references_after_phase_delete(
            cursor,
            usage_tables,
            related_tables,
            normalized_targets,
        )
        orphan_notification_count, remaining_notification_counts = (
            count_notification_cleanup_preview(cursor, normalized_targets)
        )
        counts[NOTIFICATION_TABLE] = orphan_notification_count
        for project_id in project_ids:
            remaining_project_references[project_id].pop(
                NOTIFICATION_TABLE, None
            )
            remaining_count = remaining_notification_counts.get(project_id, 0)
            if remaining_count:
                remaining_project_references[project_id][
                    NOTIFICATION_TABLE
                ] = remaining_count
        print_plan(
            db_connection,
            projects,
            phases,
            normalized_targets,
            related_tables,
            counts,
            remaining_project_references,
        )
        if not any(counts.values()):
            raise ValueError(
                "Không tìm thấy dữ liệu nào thuộc các cặp project-phase đã chọn."
            )

    preview_references = {
        project_id: dict(table_counts)
        for project_id, table_counts in remaining_project_references.items()
    }

    if not execute:
        write_deletion_report(
            report_path,
            db_connection=db_connection,
            input_path=input_path,
            mode="DRY_RUN",
            status="DRY_RUN",
            targets=normalized_targets,
            projects=projects,
            phases=phases,
            related_tables=related_tables,
            planned_counts=counts,
            deleted_counts=None,
            remaining_project_references=preview_references,
        )
        print("\nDRY RUN: chưa có dữ liệu nào bị xóa. Đặt RUN_EXECUTE=True để xóa.")
        print(f"Báo cáo Excel: {report_path.resolve()}")
        return counts

    # Bắt buộc tạo được báo cáo kế hoạch trước khi cho phép mutation.
    write_deletion_report(
        report_path,
        db_connection=db_connection,
        input_path=input_path,
        mode="EXECUTE",
        status="READY_TO_EXECUTE",
        targets=normalized_targets,
        projects=projects,
        phases=phases,
        related_tables=related_tables,
        planned_counts=counts,
        deleted_counts=None,
        remaining_project_references=preview_references,
    )

    database_name = db_connection.settings_dict.get("NAME") or "<unknown>"
    target_label = ",".join(
        f"{project_id}:{phase_id}"
        for project_id, phase_id in normalized_targets
    )
    confirmation = f"DELETE_PHASE_DATA {database_name}:{target_label}"
    if not assume_yes:
        answer = input(f"\nNhập chính xác '{confirmation}' để tiếp tục: ").strip()
        if answer != confirmation:
            write_deletion_report(
                report_path,
                db_connection=db_connection,
                input_path=input_path,
                mode="EXECUTE",
                status="CANCELLED",
                targets=normalized_targets,
                projects=projects,
                phases=phases,
                related_tables=related_tables,
                planned_counts=counts,
                deleted_counts=None,
                remaining_project_references=preview_references,
                error_message="Người dùng không xác nhận chuỗi DELETE_PHASE_DATA.",
            )
            print("Đã hủy. Không có dữ liệu nào bị xóa.")
            print(f"Báo cáo Excel: {report_path.resolve()}")
            return counts

    deleted_counts: dict[str, int] = {}
    deletable_project_ids: list[int] = []
    try:
        with transaction.atomic(using=db_connection.alias):
            with db_connection.cursor() as cursor:
                project_placeholders = make_placeholders(project_ids)
                cursor.execute(
                    f"SELECT id FROM {quote_identifier(PROJECT_TABLE)} "
                    f"WHERE id IN ({project_placeholders}) FOR UPDATE",
                    project_ids,
                )
                locked_project_ids = {int(row[0]) for row in cursor.fetchall()}
                if locked_project_ids != set(project_ids):
                    raise RuntimeError(
                        "Project thay đổi sau preview; transaction sẽ rollback."
                    )

                # Đọc lại metadata để hạn chế schema đổi giữa preview và xóa.
                related_tables = order_tables_child_first(
                    cursor, get_phase_related_tables(cursor)
                )
                for item in related_tables:
                    table_name = str(item["table_name"])
                    project_columns = list(item["project_columns"])
                    phase_column = str(item["phase_column"])
                    where_sql, params = build_target_filter(
                        project_columns, phase_column, normalized_targets
                    )
                    cursor.execute(
                        f"SELECT COUNT(*) FROM {quote_identifier(table_name)} "
                        f"WHERE {where_sql} FOR UPDATE",
                        params,
                    )
                    expected = int(cursor.fetchone()[0])
                    cursor.execute(
                        f"DELETE FROM {quote_identifier(table_name)} "
                        f"WHERE {where_sql}",
                        params,
                    )
                    actual = int(cursor.rowcount)
                    if actual != expected:
                        raise RuntimeError(
                            f"Số dòng xóa ở {table_name} không đúng; "
                            "transaction sẽ rollback: "
                            f"expected={expected}, actual={actual}"
                        )
                    deleted_counts[table_name] = actual
                    print(f"{table_name}: deleted {actual} rows")

                remaining = count_related_rows(
                    cursor, related_tables, normalized_targets
                )
                not_empty = {
                    table: count for table, count in remaining.items() if count
                }
                if not_empty:
                    raise RuntimeError(
                        "Vẫn còn dữ liệu tham chiếu trực tiếp; transaction "
                        f"sẽ rollback: {not_empty}"
                    )

                deleted_notification_count = delete_orphan_notifications(
                    cursor, project_ids
                )
                deleted_counts[NOTIFICATION_TABLE] = deleted_notification_count
                print(
                    f"{NOTIFICATION_TABLE}: deleted "
                    f"{deleted_notification_count} orphan rows"
                )

                usage_tables = get_project_usage_tables(cursor)
                remaining_project_references = count_project_references(
                    cursor, usage_tables, project_ids
                )
                deletable_project_ids = [
                    project_id
                    for project_id, table_counts in (
                        remaining_project_references.items()
                    )
                    if not table_counts
                ]
                retained_project_ids = [
                    project_id
                    for project_id, table_counts in (
                        remaining_project_references.items()
                    )
                    if table_counts
                ]

                deleted_project_count = 0
                if deletable_project_ids:
                    placeholders = make_placeholders(deletable_project_ids)
                    cursor.execute(
                        f"DELETE FROM {quote_identifier(PROJECT_TABLE)} "
                        f"WHERE id IN ({placeholders})",
                        deletable_project_ids,
                    )
                    deleted_project_count = int(cursor.rowcount)
                    if deleted_project_count != len(deletable_project_ids):
                        raise RuntimeError(
                            "Số project_info xóa không đúng; transaction sẽ "
                            "rollback: "
                            f"expected={len(deletable_project_ids)}, "
                            f"actual={deleted_project_count}"
                        )
                deleted_counts[PROJECT_TABLE] = deleted_project_count

                if retained_project_ids:
                    print(
                        "project_info giữ lại vì còn tham chiếu: "
                        + ", ".join(map(str, retained_project_ids))
                    )
    except Exception as exc:
        write_deletion_report(
            report_path,
            db_connection=db_connection,
            input_path=input_path,
            mode="EXECUTE",
            status="ROLLED_BACK",
            targets=normalized_targets,
            projects=projects,
            phases=phases,
            related_tables=related_tables,
            planned_counts=counts,
            deleted_counts={},
            remaining_project_references=preview_references,
            error_message=str(exc),
        )
        print(f"Báo cáo Excel: {report_path.resolve()}")
        raise

    write_deletion_report(
        report_path,
        db_connection=db_connection,
        input_path=input_path,
        mode="EXECUTE",
        status="COMPLETED",
        targets=normalized_targets,
        projects=projects,
        phases=phases,
        related_tables=related_tables,
        planned_counts=counts,
        deleted_counts=deleted_counts,
        remaining_project_references=remaining_project_references,
        deleted_project_ids=deletable_project_ids,
    )

    print(
        "\nXóa dữ liệu project-phase thành công; "
        f"đã xóa {deleted_counts[PROJECT_TABLE]} project_info không còn sử dụng."
    )
    print(f"Báo cáo Excel: {report_path.resolve()}")
    return deleted_counts


def parse_target_arg(value: str) -> tuple[int, int]:
    try:
        project_text, phase_text = value.split(":", maxsplit=1)
        return (
            normalize_positive_id(project_text, label="project_id"),
            normalize_positive_id(phase_text, label="phase_id"),
        )
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"Target phải có dạng PROJECT_ID:PHASE_ID, nhận được {value!r}."
        ) from exc


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Preview hoặc xóa đúng dữ liệu của từng cặp project-phase."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=RUN_INPUT_FILE,
        help="File Excel chứa cột project_id và phase_id.",
    )
    parser.add_argument(
        "--report",
        type=Path,
        help="Đường dẫn file báo cáo .xlsx; mặc định dùng tên có timestamp.",
    )
    parser.add_argument(
        "--targets",
        nargs="+",
        type=parse_target_arg,
        metavar="PROJECT_ID:PHASE_ID",
        help="Một hoặc nhiều cặp project-phase; nếu có sẽ ghi đè file input.",
    )
    parser.add_argument(
        "--execute",
        action="store_true",
        help="Thực hiện DELETE; nếu bỏ qua thì chương trình chỉ dry-run.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        help="Bỏ qua xác nhận tương tác; chỉ dùng cho automation đã kiểm soát.",
    )
    return parser.parse_args(argv)


def main() -> None:
    args = parse_args()
    targets = (
        args.targets
        if args.targets is not None
        else load_targets_from_excel(args.input)
    )
    execute = args.execute or RUN_EXECUTE
    assume_yes = args.yes or RUN_ASSUME_YES

    if assume_yes and not execute:
        raise ValueError("RUN_ASSUME_YES/--yes chỉ hợp lệ khi bật chế độ xóa.")

    delete_project_phase_data(
        targets,
        execute=execute,
        assume_yes=assume_yes,
        report_path=args.report,
        input_path=args.input if args.targets is None else None,
    )


if __name__ == "__main__":
    main()
